"""
Unit tests for PR 4: Scope Dependency Analysis, Historical Error Isolation, and Delivery Gate.
"""

import unittest
import tempfile
import shutil
import hashlib
from decimal import Decimal
from pathlib import Path
import openpyxl

from dsh_modules.spreadsheet_update.contract import SpreadsheetUpdateReceipt
from dsh_modules.spreadsheet_update.capabilities import BaselineErrorFingerprint
from dsh_modules.spreadsheet_update.validation import (
    validate_scope_and_historical_errors,
)
from dsh_modules.spreadsheet_update.delivery import (
    validate_update_delivery,
    register_verified_update,
    is_verified_update_file,
)


class TestSpreadsheetUpdateValidation(unittest.TestCase):

    def setUp(self):
        self.test_dir = Path(tempfile.mkdtemp(prefix="test_update_val_"))

    def tearDown(self):
        shutil.rmtree(self.test_dir, ignore_errors=True)

    def test_target_sheet_error_strictly_rejected(self):
        # Target sheet itself contains #DIV/0!
        wb_path = self.test_dir / "target_has_error.xlsx"
        wb = openpyxl.Workbook()
        ws_old = wb.active
        ws_old.title = "1019"
        ws_old["A1"] = "OK"

        ws_new = wb.create_sheet("1023")
        ws_new["A1"] = 100
        ws_new["A2"] = 0
        ws_new["A3"] = "#DIV/0!"  # cell error
        ws_new["A3"].data_type = 'e'
        wb.save(wb_path)
        wb.close()

        verdict = validate_scope_and_historical_errors(
            candidate_path=wb_path,
            target_sheet_name="1023",
            baseline_errors=[]
        )
        self.assertEqual(verdict.verdict, "rejected")
        self.assertFalse(verdict.is_deliverable)
        self.assertGreater(verdict.target_sheet_errors_count, 0)
        self.assertTrue(any("Target sheet '1023' has formula error" in iss for iss in verdict.issues))

    def test_isolated_historical_error_passed_with_warnings(self):
        # Historical sheet 1019 has #DIV/0!, but target 1023 only reads valid cells (A1)
        wb_path = self.test_dir / "historical_isolated.xlsx"
        wb = openpyxl.Workbook()
        ws_old = wb.active
        ws_old.title = "1019"
        ws_old["A1"] = 100  # valid cell
        ws_old["B9"] = "#DIV/0!"  # error cell in historical sheet
        ws_old["B9"].data_type = 'e'

        ws_new = wb.create_sheet("1023")
        ws_new["A1"] = 120
        # Formula references valid cell in 1019, completely ignoring B9
        ws_new["A2"] = "=A1-'1019'!A1"
        wb.save(wb_path)
        wb.close()

        # Inject evaluated cache for formula cell A2 so it does not fail the uncalculated formula gate
        import zipfile
        tmp_unzip = self.test_dir / "unzip_isolated"
        with zipfile.ZipFile(wb_path, "r") as zin:
            zin.extractall(tmp_unzip)
        sheet2_xml_p = tmp_unzip / "xl" / "worksheets" / "sheet2.xml"
        s2_content = sheet2_xml_p.read_text(encoding="utf-8")
        s2_content = s2_content.replace("<v></v>", "<v>20</v>")
        sheet2_xml_p.write_text(s2_content, encoding="utf-8")
        with zipfile.ZipFile(wb_path, "w", zipfile.ZIP_DEFLATED) as zout:
            for p in tmp_unzip.rglob("*"):
                if p.is_file():
                    zout.write(p, p.relative_to(tmp_unzip))

        baseline = [
            BaselineErrorFingerprint(sheet_name="1019", coordinate="B9", error_type="#DIV/0!")
        ]

        verdict = validate_scope_and_historical_errors(
            candidate_path=wb_path,
            target_sheet_name="1023",
            baseline_errors=baseline
        )
        self.assertEqual(verdict.verdict, "passed_with_warnings")
        self.assertTrue(verdict.is_deliverable)
        self.assertEqual(verdict.target_sheet_errors_count, 0)
        self.assertEqual(verdict.historical_isolated_errors_count, 1)
        self.assertEqual(len(verdict.issues), 0)
        self.assertTrue(any("is not referenced by target deliverable" in w for w in verdict.warnings))

    def test_target_depending_on_historical_error_rejected(self):
        # Target 1023 references the historical error cell B9
        wb_path = self.test_dir / "historical_leaked.xlsx"
        wb = openpyxl.Workbook()
        ws_old = wb.active
        ws_old.title = "1019"
        ws_old["B9"] = "#DIV/0!"
        ws_old["B9"].data_type = 'e'

        ws_new = wb.create_sheet("1023")
        ws_new["A1"] = 120
        # Target formula directly depends on the broken cell in 1019!
        ws_new["A2"] = "='1019'!B9+10"
        wb.save(wb_path)
        wb.close()

        baseline = [
            BaselineErrorFingerprint(sheet_name="1019", coordinate="B9", error_type="#DIV/0!")
        ]

        verdict = validate_scope_and_historical_errors(
            candidate_path=wb_path,
            target_sheet_name="1023",
            baseline_errors=baseline
        )
        self.assertEqual(verdict.verdict, "rejected")
        self.assertFalse(verdict.is_deliverable)
        self.assertGreater(verdict.unresolved_dependencies_count, 0)
        self.assertTrue(any("Target sheet depends on historical error cell 1019!B9" in iss for iss in verdict.issues))

    def test_validate_update_delivery_hash_drift_and_tampering(self):
        f = self.test_dir / "delivered.xlsx"
        f.write_bytes(b"initial_bytes_content_12345")
        orig_hash = hashlib.sha256(b"initial_bytes_content_12345").hexdigest()

        receipt = SpreadsheetUpdateReceipt(
            schema_version="SpreadsheetUpdateReceipt/v1",
            run_id="run_1",
            plan_id="plan_1",
            plan_hash="p_hash",
            request_hash="r_hash",
            target_id="tgt_1",
            target_input_hash="in_hash",
            output_hash=orig_hash,
            output_path=str(f),
            status="published",
            validation_verdict="passed"
        )

        # 1. Valid receipt check
        is_valid, issues = validate_update_delivery(f, receipt=receipt)
        self.assertTrue(is_valid)
        self.assertTrue(is_verified_update_file(f))

        # 2. Simulate post-validation file tampering
        f.write_bytes(b"tampered_malicious_bytes_67890")
        is_valid_tampered, issues_tampered = validate_update_delivery(f, receipt=receipt)
        self.assertFalse(is_valid_tampered)
        self.assertTrue(any("hash drift detected" in iss for iss in issues_tampered))


if __name__ == "__main__":
    unittest.main()
