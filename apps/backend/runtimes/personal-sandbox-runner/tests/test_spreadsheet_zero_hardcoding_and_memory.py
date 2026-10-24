"""Comprehensive verification for zero-hardcoding, tri-evidence grounding, active clarification and memory lifecycle.

Tests:
1. Absence of code-level aliases triggers structured clarification (status: needs_input) with standalone title candidates.
2. User confirmation grounded in actual prompt successfully binds and auto-persists to domain memory.
3. Subsequent turns automatically load domain memory without prompting the user again.
4. Independent boundary scanning captures leading target-only keys (e.g. FC48005, FC34032) in unmatched warnings.
5. Deterministic reconciliation calculates exact variance between source actuals, preserved old forecasts, and delivered totals.
"""
import copy
import json
import os
import tempfile
import unittest
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch

import openpyxl

from dsh_modules.task_intent import validate_intent
from dsh_modules.domain_memory import (
    clear_domain_memory,
    load_domain_knowledge,
    resolve_section_mapping,
    save_section_mapping
)
from dsh_modules.spreadsheet_update.layout_binding import BindingNeedsInput, locate
from dsh_modules.spreadsheet_update.reconciliation import (
    compute_reconciliation,
    render_reconciliation_markdown
)
from dsh_modules.spreadsheet_update.result_catalog import (
    RESULTS, fingerprint, reset_results, harvest_analytical_facts
)
from dsh_modules.spreadsheet_update.semantic_plan import build_semantic_plan
from dsh_modules.spreadsheet_update.tools import (
    reset_turn_results,
    tool_plan_spreadsheet_update
)
from dsh_modules.user_declarations import (
    bind_user_declarations,
    section_confirmation,
    unit_confirmation
)


class ZeroHardcodingAndMemoryLifecycleTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.ws = Path(self.temp.name).resolve()
        self.target = self.ws / '工数3Q预实管理.xlsx'
        self.source = self.ws / '工数实绩明细.xlsx'

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = '1019'
        ws['F1'] = '单位：K'

        # Section 1: 有偿
        ws['A1'] = '有偿'
        ws['A2'] = '部门'
        ws['C3'] = '202610'
        ws['C4'] = '预测'
        ws['B5'] = 'FC45031'; ws['C5'] = 100
        ws['A10'] = '总计'; ws['C10'] = '=C5'

        # Section 2: Non-billable with realistic legacy title
        # Note: contains leading departments (FC48005, FC34032) before FC46005
        ws['A22'] = '*P/*Y/EX（不含*L/*B）'
        ws['A23'] = '部门'
        ws['C24'] = '202610'
        ws['C25'] = '预测'
        ws['B26'] = 'FC48005'; ws['C26'] = 0
        ws['B27'] = 'FC34032'; ws['C27'] = 0
        ws['B28'] = 'Subtotal'; ws['C28'] = '=C26+C27'
        ws['B29'] = 'FC46005'; ws['C29'] = 53.8353
        ws['B30'] = 'FC34022'; ws['C30'] = 2.87059  # Has legacy forecast value!
        ws['A35'] = '总计'; ws['C35'] = '=SUM(C26:C30)'

        wb.save(self.target)
        wb.close()

        # Build source detail with only FC46005 for unpaid
        wb_src = openpyxl.Workbook()
        ws_src = wb_src.active
        ws_src.title = '工数实绩明细'
        ws_src.append(['工数年月', '部门编号', '有偿无偿', '金额(K)'])
        ws_src.append(['202610', 'FC45031', '有偿', 570.75345])
        ws_src.append(['202610', 'FC46005', '无偿', 14.3])
        wb_src.save(self.source)
        wb_src.close()

        prov = {'file_path': str(self.source), 'workbook_id': fingerprint(self.source)}
        receipt_paid = {
            'kind': 'spreadsheet_analysis', 'evidence_id': 'res_paid',
            'plan': {'where': {'and': [{'field': '工数年月', 'op': 'eq', 'value': '202610'},
                                      {'field': '有偿无偿', 'op': 'eq', 'value': '有偿'}]},
                    'group_by': ['部门编号'],
                    'metrics': [{'id': 'amt', 'field': '金额(K)', 'op': 'sum', 'unit': 'source'}]},
            'facts': [{'id': 'amt__FC45031', 'group_fields': {'部门编号': 'FC45031'}, 'value': '570.75345', 'unit': 'K'}],
            'coverage': [{'coverage_complete': True}]
        }
        receipt_unpaid = {
            'kind': 'spreadsheet_analysis', 'evidence_id': 'res_unpaid',
            'plan': {'where': {'and': [{'field': '工数年月', 'op': 'eq', 'value': '202610'},
                                      {'field': '有偿无偿', 'op': 'eq', 'value': '无偿'}]},
                    'group_by': ['部门编号'],
                    'metrics': [{'id': 'amt', 'field': '金额(K)', 'op': 'sum', 'unit': 'source'}]},
            'facts': [{'id': 'amt__FC46005', 'group_fields': {'部门编号': 'FC46005'}, 'value': '14.3', 'unit': 'K'}],
            'coverage': [{'coverage_complete': True}]
        }
        reset_turn_results()
        harvest_analytical_facts(receipt_paid, provenance=prov)
        harvest_analytical_facts(receipt_unpaid, provenance=prov)

        self.patch_env = patch.dict(os.environ, {'WORKSPACE_DIR': str(self.ws)})
        self.patch_env.start()
        self.patch_config = patch('dsh_modules.config.WORKSPACE_DIR', str(self.ws))
        self.patch_config.start()

        clear_domain_memory(workspace=str(self.ws))
        bind_user_declarations({})

    def tearDown(self):
        clear_domain_memory(workspace=str(self.ws))
        bind_user_declarations({})
        self.patch_config.stop()
        self.patch_env.stop()
        reset_turn_results()
        self.temp.cleanup()

    def test_missing_alias_triggers_active_clarification_with_candidates(self):
        """Zero hardcoded aliases: unknown '无偿' must ask user with discovered candidates."""
        req = {
            'target_file': '工数3Q预实管理.xlsx',
            'base_sheet': '1019',
            'new_sheet': '1023',
            'period': '202610',
            'period_field': '工数年月',
            'target_label': '预测',
            'target_kind': 'forecast',
            'target_unit': 'K',
            'unit_cell': 'F1',
            'sections': [
                {'section': '有偿', 'result_id': 'res_paid'},
                {'section': '无偿', 'result_id': 'res_unpaid'}
            ]
        }
        res = tool_plan_spreadsheet_update(req)
        self.assertTrue(res.is_partial, str(res))
        self.assertEqual(res.data.get('status'), 'needs_input')
        question = res.data.get('question', '')
        self.assertIn('未找到名为【无偿】的区块', question)
        self.assertIn('*P/*Y/EX（不含*L/*B）', question)

    def test_user_confirmation_persists_to_memory_and_enables_planning(self):
        """User clarification saves mapping to workspace memory and unblocks execution."""
        req = {
            'target_file': '工数3Q预实管理.xlsx',
            'base_sheet': '1019',
            'new_sheet': '1023',
            'period': '202610',
            'period_field': '工数年月',
            'target_label': '预测',
            'target_kind': 'forecast',
            'target_unit': 'K',
            'unit_cell': 'F1',
            'sections': [
                {'section': '有偿', 'result_id': 'res_paid'},
                {'section': '无偿', 'result_id': 'res_unpaid'}
            ]
        }
        # Simulate user replying with clarification in chat
        user_reply = "确认无偿对应*P/*Y/EX那块"
        bind_user_declarations({
            'section_confirmation': {'section': '无偿', 'target_title': '*P/*Y/EX（不含*L/*B）', 'quote': user_reply},
            'user_statements': [user_reply],
            'target_file_hash': fingerprint(self.target)
        })

        # Planning now succeeds
        res = tool_plan_spreadsheet_update(req)
        self.assertTrue(res.is_success, str(res))
        self.assertEqual(res.data.get('status'), 'planned')

        # Verify mapping was written into domain memory
        remembered = resolve_section_mapping('无偿', file_hash=fingerprint(self.target), workspace=str(self.ws))
        self.assertEqual(remembered, '*P/*Y/EX（不含*L/*B）')

    def test_domain_memory_auto_loads_without_asking_user_again(self):
        """Subsequent turns automatically reload domain memory, no re-prompting needed."""
        # Pre-populate domain memory (simulating a prior confirmed session)
        save_section_mapping(
            section_name='无偿',
            target_title='*P/*Y/EX（不含*L/*B）',
            file_hash=fingerprint(self.target),
            quote='用户已在历史会话确认',
            workspace=str(self.ws)
        )

        req = {
            'target_file': '工数3Q预实管理.xlsx',
            'base_sheet': '1019',
            'new_sheet': '1023',
            'period': '202610',
            'period_field': '工数年月',
            'target_label': '预测',
            'target_kind': 'forecast',
            'target_unit': 'K',
            'unit_cell': 'F1',
            'sections': [
                {'section': '有偿', 'result_id': 'res_paid'},
                {'section': '无偿', 'result_id': 'res_unpaid'}
            ]
        }
        # Notice: no user declarations bound for this turn!
        plan, mapped, payload, preview, warnings = build_semantic_plan(req, self.ws)
        self.assertEqual(len(preview), 2)  # 1 paid + 1 unpaid write
        self.assertEqual(preview[0]['coordinate'], 'C5')
        self.assertEqual(preview[1]['coordinate'], 'C29')

    def test_independent_boundary_scan_captures_leading_unmatched_departments(self):
        """FC48005 and FC34032 (rows 26 & 27) are captured despite FC46005 (row 29) being the first source match."""
        save_section_mapping('无偿', '*P/*Y/EX（不含*L/*B）', file_hash=fingerprint(self.target), workspace=str(self.ws))
        req = {
            'target_file': '工数3Q预实管理.xlsx',
            'base_sheet': '1019',
            'new_sheet': '1023',
            'period': '202610',
            'period_field': '工数年月',
            'target_label': '预测',
            'target_kind': 'forecast',
            'target_unit': 'K',
            'unit_cell': 'F1',
            'sections': [
                {'section': '无偿', 'result_id': 'res_unpaid'}
            ]
        }
        plan, mapped, payload, preview, warnings = build_semantic_plan(req, self.ws)
        unpaid_spec = plan['semantic_bindings']['region_specs'][0]
        dept_rows = unpaid_spec['department_rows']

        # Both leading departments MUST be present in target keys
        self.assertIn('FC48005', dept_rows)
        self.assertEqual(dept_rows['FC48005'], 26)
        self.assertIn('FC34032', dept_rows)
        self.assertEqual(dept_rows['FC34032'], 27)
        self.assertIn('FC34022', dept_rows)
        self.assertEqual(dept_rows['FC34022'], 30)

        # Warnings must specifically flag the unmatched count and non-zero legacy forecast
        warn_str = " ".join(warnings)
        self.assertIn('FC34022', warn_str)
        self.assertIn('FC48005', warn_str)
        self.assertIn('FC34032', warn_str)

    def test_deterministic_reconciliation_balances_and_discloses_variance(self):
        """Reconciliation discloses exact legacy forecast retention (FC34022 = 2.87059 K)."""
        save_section_mapping('无偿', '*P/*Y/EX（不含*L/*B）', file_hash=fingerprint(self.target), workspace=str(self.ws))
        req = {
            'target_file': '工数3Q预实管理.xlsx',
            'base_sheet': '1019',
            'new_sheet': '1023',
            'period': '202610',
            'period_field': '工数年月',
            'target_label': '预测',
            'target_kind': 'forecast',
            'target_unit': 'K',
            'unit_cell': 'F1',
            'sections': [
                {'section': '无偿', 'result_id': 'res_unpaid'}
            ]
        }
        plan, mapped, payload, preview, warnings = build_semantic_plan(req, self.ws)
        reconciliations = plan['semantic_bindings']['reconciliations']
        self.assertEqual(len(reconciliations), 1)

        recon = reconciliations[0]
        # Source actual: 14.3 K
        self.assertEqual(recon['source_actual_sum'], '14.3')
        # Preserved old forecast: 2.87059 K (from FC34022)
        self.assertEqual(recon['preserved_old_forecast_sum'], '2.87059')
        # Delivered forecast total: 14.3 + 2.87059 = 17.17059 K
        self.assertEqual(recon['delivered_forecast_sum'], '17.17059')
        self.assertTrue(recon['is_balanced'])

        # Render markdown and verify transparency
        md = render_reconciliation_markdown(recon)
        self.assertIn('FC34022', md)
        self.assertIn('14.3 + 2.87059 = 17.17059 K', md)

    @patch('dsh_modules.spreadsheet_update.executor.recalculate_candidate')
    def test_two_turn_clarification_end_to_end(self, mock_recalc):
        """End-to-end multi-turn flow: Ask -> User Clarifies -> Followup Locked -> File Delivered."""
        from test_spreadsheet_update_multi_section_e2e import MultiSectionSpreadsheetUpdateTests
        mock_recalc.side_effect = MultiSectionSpreadsheetUpdateTests._simulate_recalculate_candidate
        from dsh_modules.task_intent import validate_intent
        from dsh_modules.spreadsheet_update.routing import is_workbook_update
        from dsh_modules.spreadsheet_update.tools import tool_execute_spreadsheet_update

        # --- Turn 1: Initial request triggers clarification ---
        turn1_prompt = "需要更新工数管理，里面有计划和预测，先复制最新的sheet 生成今天 1023的新sheet，用 部门编码 统计 部门的 有偿和无偿 部分的金额 更新到 这个月的预测部分，以K为单位"
        req_turn1 = {
            'target_file': '工数3Q预实管理.xlsx',
            'base_sheet': '1019',
            'new_sheet': '1023',
            'period': '202610',
            'period_field': '工数年月',
            'target_label': '预测',
            'target_kind': 'forecast',
            'target_unit': 'K',
            'unit_cell': 'F1',
            'sections': [
                {'section': '有偿', 'result_id': 'res_paid'},
                {'section': '无偿', 'result_id': 'res_unpaid'}
            ]
        }
        res_t1 = tool_plan_spreadsheet_update(req_turn1)
        self.assertTrue(res_t1.is_partial)
        self.assertEqual(res_t1.data.get('status'), 'needs_input')
        assistant_question = res_t1.data.get('question', '')
        self.assertIn('未找到名为【无偿】的区块', assistant_question)

        # --- Turn 2: User answers clarification: "【*P/Y/EX（不含L/*B）】是无偿" ---
        turn2_user_msg = "【*P/Y/EX（不含L/*B）】是无偿"
        full_history = [
            {'role': 'user', 'content': turn1_prompt},
            {'role': 'assistant', 'content': assistant_question}
        ]

        # 1. Verify TaskIntent accurately locks relation=followup, task_mode=execute, output_mode=artifact
        dummy_model_result = {
            'instruction_relation': 'standalone',  # Model mistakenly outputted standalone
            'task_mode': 'consult',               # Model mistakenly outputted consult
            'output_mode': 'conversation',        # Model mistakenly outputted conversation
            'artifact_types': [],
            'skill_invocation': {'mode': 'none', 'name': '', 'replace_defaults': False},
            'requirements': [{'id': '1', 'quote': turn2_user_msg}]
        }
        corrected_intent = validate_intent(
            dummy_model_result,
            turn2_user_msg,
            [turn1_prompt],
            skill_names=[],
            allow_sanitize=True,
            history=full_history
        )
        self.assertEqual(corrected_intent['instruction_relation'], 'followup')
        self.assertEqual(corrected_intent['task_mode'], 'execute')
        self.assertEqual(corrected_intent['output_mode'], 'artifact')
        self.assertIn('.xlsx', corrected_intent['artifact_types'])
        self.assertIsNotNone(corrected_intent.get('section_confirmation'))
        self.assertIn('无偿', corrected_intent['section_confirmation'])
        self.assertEqual(corrected_intent['section_confirmation']['无偿']['section'], '无偿')

        # 2. Verify routing selects native workbook update
        is_wb = is_workbook_update(turn2_user_msg, admission=corrected_intent, history=full_history)
        self.assertTrue(is_wb)

        # 3. Plan and execute spreadsheet update now that declaration is bound
        res_t2 = tool_plan_spreadsheet_update(req_turn1)
        self.assertTrue(res_t2.is_success, str(res_t2))
        plan_id = res_t2.data['plan_id']

        exec_res = tool_execute_spreadsheet_update(plan_id)
        self.assertTrue(exec_res.is_success, str(exec_res))
        self.assertEqual(exec_res.data['status'], 'published')
        out_path = Path(exec_res.data['published_file'])
        self.assertTrue(out_path.exists())

        # 4. Verify delivered workbook has sheet 1023 and filled values
        wb_out = openpyxl.load_workbook(out_path, data_only=True)
        self.assertIn('1023', wb_out.sheetnames)
        ws_1023 = wb_out['1023']
        # Paid: FC45031 = 570.75345
        self.assertAlmostEqual(float(ws_1023['C5'].value), 570.75345, places=3)
        # Unpaid: FC46005 = 14.3
        self.assertAlmostEqual(float(ws_1023['C29'].value), 14.3, places=3)
        # Unmatched FC34022 preserved legacy forecast = 2.87059
        self.assertAlmostEqual(float(ws_1023['C30'].value), 2.87059, places=3)

    def test_natural_language_clarification_no_regex_trap(self):
        """Verify that natural phrasing '【*P/Y/EX（不含L/*B）】 就是无偿' never degrades to '就'."""
        assistant_question = (
            "目标工作表中未找到名为【无偿】的区块。检测到以下候选区块标题：【有偿】、【*P/*Y/EX（不含*L/*B）】。"
            "请确认【无偿】是否对应上述某个候选区块？本次尚未生成更新文件。"
        )
        user_prompt = "【*P/Y/EX（不含L/*B）】 就是无偿"
        history = [
            {'role': 'user', 'content': '更新工数管理'},
            {'role': 'assistant', 'content': assistant_question}
        ]

        # 1. Fallback candidate matcher test
        from dsh_modules.user_declarations import direct_section_reply
        matched = direct_section_reply(user_prompt, assistant_question)
        self.assertIsNotNone(matched)
        self.assertEqual(matched['section'], '无偿')
        self.assertEqual(matched['target_title'], '*P/*Y/EX（不含*L/*B）')
        self.assertNotEqual(matched['target_title'], '就')

        # 2. End-to-end validate_intent with model proposal containing abbreviation
        model_proposal = {
            'instruction_relation': 'followup',
            'task_mode': 'execute',
            'output_mode': 'artifact',
            'artifact_types': ['.xlsx'],
            'skill_invocation': {'mode': 'none', 'name': '', 'replace_defaults': False},
            'unit_confirmation': {'unit': 'K', 'quote': '以K为单位'},
            'section_confirmation': {'section': '无偿', 'target_title': '*P/Y/EX', 'quote': user_prompt},
            'requirements': [{'id': '1', 'quote': user_prompt}]
        }
        validated = validate_intent(
            model_proposal,
            user_prompt,
            ['更新工数管理以K为单位'],
            skill_names=[],
            allow_sanitize=True,
            history=history
        )
        self.assertEqual(validated['instruction_relation'], 'followup')
        self.assertEqual(validated['task_mode'], 'execute')
        self.assertEqual(validated['output_mode'], 'artifact')
        sec_conf = validated.get('section_confirmation')
        self.assertIsNotNone(sec_conf)
        self.assertIn('无偿', sec_conf)
        # Ensure abbreviation *P/Y/EX was resolved to full workbook title *P/*Y/EX（不含*L/*B）
        self.assertEqual(sec_conf['无偿']['target_title'], '*P/*Y/EX（不含*L/*B）')
        self.assertNotEqual(sec_conf['无偿']['target_title'], '就')


if __name__ == '__main__':
    unittest.main()
