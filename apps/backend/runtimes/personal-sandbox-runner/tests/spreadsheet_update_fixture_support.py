"""Explicit metadata for the small, known internal-engine test fixtures (never runtime inference)."""
import openpyxl


def bind_fixture_regions(path,sheet_name,specs):
    wb=openpyxl.load_workbook(path)
    try:
        ws=wb[sheet_name]
        for spec in specs:
            spec['value_kind']='actual'
            spec['period_columns']={str(ws.cell(spec['header_row'],c).value):c
                                    for c in range(spec['min_col'],spec['max_col']+1)
                                    if c!=spec['key_col'] and ws.cell(spec['header_row'],c).value is not None}
            spec['department_rows']={str(ws.cell(r,spec['key_col']).value):r
                                     for r in range(spec['header_row']+1,spec['max_row']+1)
                                     if ws.cell(r,spec['key_col']).value is not None}
        return specs
    finally:
        wb.close()
