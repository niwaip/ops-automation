"""Reusable comparisons, execution evidence and the model-tool-model lifecycle."""

import hashlib
import json
import sys
import tempfile
import time
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src"))
from dsh_modules.analysis_contract import build_analysis_contract
from dsh_modules.analysis_validation import validate_reply_against_contract
from dsh_modules.table_comparison import calculate_column_comparison
from dsh_modules.tools import execute_tool, get_sandbox_tools
from dsh_modules.agent_loop import run_agent_loop
from dsh_modules.runtime_policy import RuntimePolicy
from dsh_modules.runner import select_active_tools
from dsh_modules.skill_router import SkillRouter
from dsh_modules.tool_result import ToolResult
from dsh_modules.prompt_builder import build_system_prompt, build_user_turn
from dsh_modules.comparison_evidence import validate_execution_output
from dsh_modules.comparison_delivery import compile_comparison_response
from dsh_modules.comparison_context import discover_comparison_candidates, build_comparison_source_context


QUERY = "找出收入与净利润偏差最大的月份。"
ANSWER = "营业收入12月相对预算偏差+12千元；净利润7月相对预算偏差+50千元。"


class TableComparisonTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.path = Path(self.temp.name) / "arbitrary.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "任意名称"
        ws.append(["测试数据"])
        ws.append(["月份", "收入预算", "收入实际", "净利润预算", "净利润实际"])
        for month in range(1, 13):
            ws.append([f"{month}月", 100, 100+month, 30, 80 if month == 7 else 20])
        wb.save(self.path)
        wb.close()
        self.params = {"file_path": str(self.path), "sheet": "任意名称", "data_range": "A3:E14",
                       "header_row": 2, "group_column": "A", "basis": "budget",
                       "rank_by": "absolute_difference", "extreme": "max",
                       "expected_groups": [f"{m}月" for m in range(1, 13)],
                       "comparisons": [{"label": "营业收入", "left_column": "C", "right_column": "B"},
                                       {"label": "净利润", "left_column": "E", "right_column": "D"}]}
        self.contract = build_analysis_contract(QUERY)

    def tearDown(self):
        self.temp.cleanup()

    def analyze(self, **changes):
        arguments = {k: v for k, v in {**self.params, **changes}.items() if k != "file_path"}
        return calculate_column_comparison(self.path, **arguments)

    def edit(self, action):
        wb = openpyxl.load_workbook(self.path)
        action(wb)
        wb.save(self.path)
        wb.close()

    def validate(self, text, evidence):
        return validate_reply_against_contract(self.contract, text, 0, 4, execution_evidence=evidence)

    def call(self, params=None):
        with patch("dsh_modules.table_comparison.resolve_sandboxed_path", return_value=(self.path, None)):
            return execute_tool("compare_spreadsheet_columns", params or self.params)

    def test_separate_extrema_and_original_protection(self):
        before = self.path.read_bytes()
        result = self.analyze()
        self.assertTrue(result.is_success)
        self.assertEqual([p["matches"][0]["group"] for p in result.data["results"]], ["12月", "7月"])
        self.assertEqual(result.data["results"][1]["matches"][0]["source"]["left"], "任意名称!E9")
        self.assertEqual(result.provenance["workbook_id"], hashlib.sha256(before).hexdigest())
        self.assertEqual(self.path.read_bytes(), before)

    def test_source_schema_proposes_pairs_without_answering(self):
        before = self.path.read_bytes()
        candidates = discover_comparison_candidates(self.path, self.contract)
        self.assertEqual(len(candidates), 1)
        self.assertEqual(candidates[0]["sheet"], "任意名称")
        self.assertEqual(candidates[0]["comparisons"], self.params["comparisons"])
        self.assertEqual(candidates[0]["expected_groups"], self.params["expected_groups"])
        self.assertNotIn("difference", json.dumps(candidates))
        self.assertEqual(self.path.read_bytes(), before)

    def test_source_schema_never_selects_between_multiple_tables(self):
        self.edit(lambda w: setattr(w.copy_worksheet(w.active), "title", "另一份计划"))
        self.assertEqual(len(discover_comparison_candidates(self.path, self.contract)), 2)

    def test_source_schema_follows_direct_references_for_metric_identity(self):
        def change(w):
            source = w.create_sheet("原始指标")
            source.append(["月份", "净利润"])
            for m in range(1, 13):
                source.append([f"{m}月", 80 if m == 7 else 20])
            w.active["E2"] = "利润实际"
            w.active["D2"] = "利润预算"
            for r in range(3, 15):
                w.active.cell(r, 1, f"=原始指标!A{r-1}")
                w.active.cell(r, 5, f"=原始指标!B{r-1}")
        self.edit(change)
        candidates = discover_comparison_candidates(self.path, self.contract)
        self.assertEqual(candidates[0]["comparisons"], self.params["comparisons"])
        self.assertEqual(candidates[0]["expected_groups"], self.params["expected_groups"])

    def test_source_schema_rejects_ambiguous_or_duplicate_groups(self):
        self.edit(lambda w: setattr(w.active["A4"], "value", "1月"))
        self.assertEqual(discover_comparison_candidates(self.path, self.contract), [])
        self.assertEqual(build_comparison_source_context(self.path, QUERY), "")

    def test_source_context_uses_overview_without_calculation(self):
        with patch("dsh_modules.comparison_context.resolve_sandboxed_path", return_value=(self.path, None)), \
             patch("dsh_modules.table_comparison.get_or_create_recalculated_workbook") as recalc:
            context = build_comparison_source_context(self.path, QUERY)
        recalc.assert_not_called()
        self.assertIn('"mode": "CATALOG"', context)
        self.assertIn("Column Comparison Candidates", context)
        self.assertNotIn('"matches"', context)

    def test_source_context_respects_file_boundary(self):
        with patch("dsh_modules.comparison_context.resolve_sandboxed_path", return_value=(None, "越界")), \
             patch("dsh_modules.comparison_context.discover_comparison_candidates") as discover:
            self.assertEqual(build_comparison_source_context(self.path, QUERY), "")
        discover.assert_not_called()

    def test_any_metric_and_nonmonthly_population(self):
        def change(w):
            w.active["A2"] = "部门"
            w.active["B2"] = "成本目标"
            w.active["C2"] = "成本实际"
            for i in range(3, 15):
                w.active.cell(i, 1, f"部门{i-2}")
        self.edit(change)
        result = self.analyze(comparisons=[{"label": "成本", "left_column": "C", "right_column": "B"}],
                              expected_groups=[f"部门{m}" for m in range(1, 13)], basis="custom")
        self.assertEqual(result.data["results"][0]["matches"][0]["group"], "部门12")

    def test_partial_scope_is_explicit_and_supported(self):
        result = self.analyze(data_range="A3:E8", expected_groups=[f"{m}月" for m in range(1, 7)])
        self.assertEqual(result.data["scope"]["row_count"], 6)
        self.assertEqual(result.data["results"][0]["matches"][0]["group"], "6月")

    def test_missing_group_is_not_a_full_population(self):
        self.assertEqual(self.analyze(data_range="A3:E8").error_code, "incomplete_group_coverage")

    def test_duplicate_and_blank_groups_fail(self):
        for value in ["1月", None]:
            original = self.path.read_bytes()
            self.edit(lambda w: setattr(w.active["A14"], "value", value))
            with self.assertRaises(ValueError):
                self.analyze()
            self.path.write_bytes(original)

    def test_all_ties_are_retained(self):
        self.edit(lambda w: setattr(w.active["C13"], "value", 112))
        result = self.analyze()
        self.assertEqual([r["group"] for r in result.data["results"][0]["matches"]], ["11月", "12月"])
        self.assertFalse(self.validate(ANSWER, [result]).is_pass)

    def test_relative_and_amount_rankings_differ(self):
        def change(w):
            w.active["B12"] = 1
            w.active["C12"] = 5
        self.edit(change)
        result = self.analyze(rank_by="absolute_relative_difference")
        self.assertEqual(result.data["results"][0]["matches"][0]["group"], "10月")

    def test_zero_baseline_blocks_relative_but_not_amount(self):
        self.edit(lambda w: setattr(w.active["B3"], "value", 0))
        self.assertTrue(self.analyze().is_success)
        self.assertEqual(self.analyze(rank_by="absolute_relative_difference").error_code, "undefined_relative_difference")

    def test_signed_minimum(self):
        result = self.analyze(rank_by="difference", extreme="min")
        self.assertEqual(result.data["results"][0]["matches"][0]["group"], "1月")

    def test_cross_metric_comparison_is_explicit(self):
        result = self.analyze(basis="cross_metric", comparisons=[{"label": "收入减净利润", "left_column": "C", "right_column": "E"}])
        contract = build_analysis_contract("找出收入减净利润差额最大的月份。")
        reply = "12月收入112千元，净利润20千元，收入减净利润差额92千元。"
        self.assertTrue(validate_reply_against_contract(contract, reply, 0, 4, execution_evidence=[result]).is_pass)
        other = build_analysis_contract("找出营业利润减毛利润差额最大的月份。")
        self.assertEqual(other.comparison_basis, "cross_metric")

    def test_wrong_plan_fails_at_tool_boundary(self):
        params = {**self.params, "basis": "custom", "comparisons": [
            {"label": "营业收入", "left_column": "C", "right_column": "E"},
            {"label": "净利润", "left_column": "E", "right_column": "C"}]}
        rejected = validate_execution_output(QUERY, self.call(params))
        self.assertEqual(rejected.tool_result.error_code, "comparison_plan_mismatch")

    def test_invalid_parameter_and_numeric_types_fail(self):
        for changes in [{"group_column": "Z"}, {"data_range": "A:E"}, {"header_row": 3},
                        {"comparisons": []}, {"rank_by": "unknown"}]:
            with self.assertRaises(ValueError):
                self.analyze(**changes)
        self.edit(lambda w: setattr(w.active["C3"], "value", True))
        self.assertTrue(self.call().tool_result.is_error)

    def test_nonexistent_sheet_does_not_guess(self):
        self.assertEqual(self.analyze(sheet="不存在").error_code, "sheet_not_found")

    def test_unverified_calculation_fails(self):
        self.edit(lambda w: setattr(w.active["C3"], "value", "=101"))
        with patch("dsh_modules.table_comparison.get_or_create_recalculated_workbook",
                   return_value=(self.path, {"status": "failed"})):
            self.assertEqual(self.analyze().error_code, "calculation_unverified")

    def test_mismatched_calculation_credentials_fail(self):
        self.edit(lambda w: setattr(w.active["C3"], "value", "=101"))
        with patch("dsh_modules.table_comparison.get_or_create_recalculated_workbook",
                   return_value=(self.path, {"status": "recalculated_verified", "provenance": {}})):
            self.assertEqual(self.analyze().error_code, "calculation_provenance_mismatch")

    def test_formula_groups_use_verified_values(self):
        cached = self.path.with_name("verified.xlsx")
        cached.write_bytes(self.path.read_bytes())
        self.edit(lambda w: setattr(w.active["A3"], "value", '= "1月"'))
        metadata = {"status": "recalculated_verified", "provenance": {
            "workbook_id": hashlib.sha256(self.path.read_bytes()).hexdigest(),
            "output_hash": hashlib.sha256(cached.read_bytes()).hexdigest()}}
        with patch("dsh_modules.table_comparison.get_or_create_recalculated_workbook", return_value=(cached, metadata)):
            self.assertTrue(self.analyze().is_success)

    def test_nonfinite_derived_values_fail(self):
        def change(w):
            w.active["C3"] = 1e308
            w.active["B3"] = -1e308
        self.edit(change)
        with self.assertRaises(ValueError):
            self.analyze()

    def test_deadline_propagates(self):
        with self.assertRaises(TimeoutError):
            self.analyze(deadline=time.monotonic()-1)

    def test_proof_survives_context_compaction(self):
        output = self.call()
        compact = json.loads(output)
        self.assertNotIn("evaluated", compact["data"])
        self.assertIn("evaluated", output.tool_result.data)
        self.assertTrue(self.validate(ANSWER, [output.tool_result]).is_pass)

    def test_numbers_or_saved_history_are_not_execution_proof(self):
        self.assertFalse(self.validate(ANSWER, []).is_pass)
        self.assertFalse(validate_reply_against_contract(self.contract, ANSWER, 0, 4,
                         messages=[{"role": "tool", "content": self.call()}], execution_evidence=[]).is_pass)

    def test_incorrect_group_value_and_rank_are_rejected(self):
        result = self.analyze()
        for answer in [ANSWER.replace("7月", "12月"), ANSWER.replace("+50", "+10")]:
            self.assertFalse(self.validate(answer, [result]).is_pass)
        self.assertFalse(self.validate(ANSWER, [self.analyze(extreme="min")]).is_pass)

    def test_markdown_metric_sections_do_not_force_recalculation(self):
        report = "### 营业收入\n最大偏差月份：12月\n差额：+12千元\n### 净利润\n最大偏差月份：7月\n差额：+50千元"
        self.assertTrue(self.validate(report, [self.analyze()]).is_pass)
        self.assertFalse(self.validate(report.replace("+50", "+12"), [self.analyze()]).is_pass)

    def test_source_unit_and_magnitude_conversion_are_verified(self):
        self.edit(lambda w: setattr(w.active["A1"], "value", "测试表（单位：人民币千元）"))
        result = self.analyze()
        self.assertEqual(result.data["value_unit"], "千元")
        self.assertEqual(result.provenance["unit_sources"], ["任意名称!A1"])
        self.assertTrue(self.validate(ANSWER, [result]).is_pass)
        converted = "营业收入12月偏差+1.2万元；净利润7月偏差+5万元。"
        self.assertTrue(self.validate(converted, [result]).is_pass)
        self.assertFalse(self.validate(ANSWER.replace("千元", "万元"), [result]).is_pass)
        self.assertFalse(self.validate(ANSWER.replace("千元", ""), [result]).is_pass)

    def test_wrong_narrative_unit_cannot_hide_behind_a_correct_table(self):
        self.edit(lambda w: setattr(w.active["A1"], "value", "单位：千元"))
        report = ANSWER+"\n营业收入12月比预算高12万元。"
        self.assertFalse(self.validate(report, [self.analyze()]).is_pass)

    def test_conflicting_source_units_are_left_unresolved(self):
        def change(w):
            w.active["A1"] = "单位：千元"
            w.active["B1"] = "单位：万元"
        self.edit(change)
        self.assertIsNone(self.analyze().data["value_unit"])

    def test_unit_dictionary_field_can_be_on_a_separate_sheet(self):
        def change(w):
            notes = w.create_sheet("说明页")
            notes["A6"] = "币种及单位"
            notes["B6"] = "人民币，千元"
        self.edit(change)
        result = self.analyze()
        self.assertEqual(result.data["value_unit"], "千元")
        self.assertEqual(result.provenance["unit_sources"], ["说明页!B6"])

    def test_changed_source_invalidates_existing_proof(self):
        result = self.analyze()
        self.edit(lambda w: setattr(w.active["C14"], "value", 113))
        self.assertFalse(self.validate(ANSWER, [result]).is_pass)

    def test_budget_basis_requires_a_real_baseline_header(self):
        result = self.analyze(comparisons=[{"label": "营业收入", "left_column": "C", "right_column": "E"}])
        self.assertFalse(self.validate(ANSWER, [result]).is_pass)

    def test_public_readonly_tool_registration(self):
        skill = SimpleNamespace(skill_id="xlsx", is_inspect_intent=True, is_generate_intent=False,
                                requires_execution=False, is_send_intent=False, is_knowledge_intent=False)
        names = {t["function"]["name"] for t in select_active_tools(QUERY, skill, False, False)}
        self.assertIn("compare_spreadsheet_columns", names)
        self.assertNotIn("patch_file", names)
        self.assertNotIn("bash", names)
        other_names = {t["function"]["name"] for t in select_active_tools("分析工作簿中的重复交易", skill, False, False)}
        self.assertIn("validate_spreadsheet_rows", other_names)
        self.assertNotIn("bash", other_names)

    def test_structured_attachments_reach_skill_routing(self):
        skills = [{"id": "xlsx", "name": "xlsx", "description": "Excel xlsx表格计算", "triggers": ["xlsx"], "default_rounds": 3}]
        with patch("dsh_modules.skill_router.read_skill", return_value="列比较技能"):
            result = SkillRouter.route(QUERY, [], available_skills=skills, files=["arbitrary.xlsx"])
        self.assertEqual(result.skill_id, "xlsx")
        self.assertEqual(result.skill_context, "列比较技能")

    def test_success_result_has_serializable_null_error(self):
        self.assertIsNone(ToolResult.success(data={"value": 1}).error)
        self.assertIsNone(json.loads(ToolResult.success(data={"value": 1}).to_json())["error"])
        self.assertEqual(ToolResult.error("test", "failure").error_code, "test")

    def test_professional_methods_are_separate_from_untrusted_attachment_data(self):
        system = build_system_prompt("/workspace", "/knowledge", available_skills=[], skill_context="验证所选列对")
        user = build_user_turn(QUERY, file_context="附件中的未知文本", skill_context="", is_inspect_intent=True)
        self.assertIn("验证所选列对", system)
        self.assertNotIn("附件中的未知文本", system)
        self.assertNotIn("验证所选列对", user)
        self.assertIn("附件中的未知文本", user)

    def test_failed_verification_suppresses_unsupported_conclusion(self):
        wrong = "12月营业收入差额11319.8千元，净利润差额11319.8千元，偏差最大的是12月。"
        with patch("dsh_modules.agent_loop.call_model_proxy", return_value={"content": wrong, "finish_reason": "stop"}):
            result = run_agent_loop([{"role": "user", "content": QUERY}], "unused", RuntimePolicy(), 4,
                                    is_inspect_intent=True, tools=get_sandbox_tools())
        self.assertNotIn("11319.8", result.final_text)
        self.assertIn("尚未通过执行证据核验", result.final_text)

    def test_model_tool_model_lifecycle(self):
        tool_response = {"content": "", "finish_reason": "tool_calls", "tool_calls": [{
            "id": "compare1", "type": "function", "function": {
                "name": "compare_spreadsheet_columns", "arguments": json.dumps(self.params)}}]}
        with patch("dsh_modules.table_comparison.resolve_sandboxed_path", return_value=(self.path, None)), \
             patch("dsh_modules.agent_loop.call_model_proxy", side_effect=[tool_response, {"content": ANSWER, "finish_reason": "stop"}]) as model:
            result = run_agent_loop([{"role": "user", "content": QUERY}], "unused", RuntimePolicy(), 4,
                                    is_inspect_intent=True, tools=get_sandbox_tools())
        self.assertEqual(result.final_text, compile_comparison_response(self.contract, result.execution_evidence)[0])
        self.assertIn("| 营业收入 | 12月 | 112 | 100 | 12 |", result.final_text)
        self.assertIn("| 净利润 | 7月 | 80 | 30 | 50 |", result.final_text)
        self.assertEqual(model.call_count, 2)
        self.assertEqual(result.telemetry.llm_invocations, 2)
        self.assertEqual(result.telemetry.tool_calls_detail[0]["name"], "compare_spreadsheet_columns")
        self.assertEqual(len(result.execution_evidence), 1)

    def test_separate_tool_calls_accumulate_metric_evidence(self):
        outputs = []
        calls = []
        for i, pair in enumerate(self.params["comparisons"]):
            params = {**self.params, "comparisons": [pair]}
            output = validate_execution_output(QUERY, self.call(params))
            self.assertTrue(output.tool_result.is_success)
            outputs.append(output.tool_result)
            calls.append({"id": f"metric{i}", "type": "function", "function": {
                "name": "compare_spreadsheet_columns", "arguments": json.dumps(params)}})
        self.assertFalse(self.validate(ANSWER, [outputs[0]]).is_pass)
        self.assertTrue(self.validate(ANSWER, outputs).is_pass)
        with patch("dsh_modules.table_comparison.resolve_sandboxed_path", return_value=(self.path, None)), \
             patch("dsh_modules.agent_loop.call_model_proxy", side_effect=[
                 {"content": "", "finish_reason": "tool_calls", "tool_calls": calls},
                 {"content": ANSWER, "finish_reason": "stop"}]):
            result = run_agent_loop([{"role": "user", "content": QUERY}], "unused", RuntimePolicy(), 4,
                                    is_inspect_intent=True, tools=get_sandbox_tools())
        self.assertEqual(result.final_text, compile_comparison_response(self.contract, result.execution_evidence)[0])
        self.assertIn("| 营业收入 | 12月 | 112 | 100 | 12 |", result.final_text)
        self.assertIn("| 净利润 | 7月 | 80 | 30 | 50 |", result.final_text)
        self.assertEqual(len(result.execution_evidence), 2)

    def test_valid_partial_work_cannot_hide_an_invalid_extra_pair(self):
        pairs = [self.params["comparisons"][0],
                 {"label": "净利润", "left_column": "E", "right_column": "B"}]
        output = validate_execution_output(QUERY, self.call({**self.params, "comparisons": pairs}))
        self.assertTrue(output.tool_result.is_error)

    def test_unavailable_tool_is_rejected_without_execution(self):
        forbidden = {"content": "", "finish_reason": "tool_calls", "tool_calls": [{
            "id": "unavailable", "type": "function", "function": {
                "name": "bash", "arguments": json.dumps({"cmd": "echo bypass"})}}]}
        selected = [t for t in get_sandbox_tools() if t["function"]["name"] == "read_file"]
        with patch("dsh_modules.agent_loop.execute_tool") as execute, \
             patch("dsh_modules.agent_loop.call_model_proxy", side_effect=[forbidden,
                   {"content": "工具不可用。", "finish_reason": "stop"}]):
            messages = [{"role": "user", "content": "读取文件"}]
            result = run_agent_loop(messages, "unused", RuntimePolicy(), 4,
                                    is_inspect_intent=True, tools=selected)
        execute.assert_not_called()
        self.assertEqual(result.telemetry.tool_calls_detail[0]["status"], "blocked")
        response = next(m for m in messages if m.get("role") == "tool")
        self.assertEqual(response["tool_call_id"], "unavailable")
        self.assertEqual(json.loads(response["content"])["error"]["code"], "tool_not_available")

    def test_unsupported_answer_is_corrected_through_model_not_auto_return(self):
        responses = [{"content": ANSWER, "finish_reason": "stop"},
                     {"content": "", "finish_reason": "tool_calls", "tool_calls": [{
                         "id": "compare1", "type": "function", "function": {
                             "name": "compare_spreadsheet_columns", "arguments": json.dumps(self.params)}}]},
                     {"content": ANSWER, "finish_reason": "stop"}]
        with patch("dsh_modules.table_comparison.resolve_sandboxed_path", return_value=(self.path, None)), \
             patch("dsh_modules.agent_loop.call_model_proxy", side_effect=responses) as model:
            result = run_agent_loop([{"role": "user", "content": QUERY}], "unused", RuntimePolicy(), 4,
                                    is_inspect_intent=True, tools=get_sandbox_tools())
        self.assertEqual(model.call_count, 3)
        self.assertEqual(result.final_text, compile_comparison_response(self.contract, result.execution_evidence)[0])
        self.assertIn("| 营业收入 | 12月 | 112 | 100 | 12 |", result.final_text)
        self.assertIn("| 净利润 | 7月 | 80 | 30 | 50 |", result.final_text)

    def test_remaining_round_can_repair_a_wrong_tool_plan(self):
        wrong_params = {**self.params, "rank_by": "absolute_relative_difference"}
        def tool(params, identity):
            return {"content": "", "finish_reason": "tool_calls", "tool_calls": [{
                "id": identity, "type": "function", "function": {
                    "name": "compare_spreadsheet_columns", "arguments": json.dumps(params)}}]}
        responses = [{"content": ANSWER, "finish_reason": "stop"}, tool(wrong_params, "wrong"),
                     {"content": ANSWER, "finish_reason": "stop"}, tool(self.params, "fixed"),
                     {"content": ANSWER, "finish_reason": "stop"}]
        with patch("dsh_modules.table_comparison.resolve_sandboxed_path", return_value=(self.path, None)), \
             patch("dsh_modules.agent_loop.call_model_proxy", side_effect=responses) as model:
            result = run_agent_loop([{"role": "user", "content": QUERY}], "unused", RuntimePolicy(), 4,
                                    is_inspect_intent=True, tools=get_sandbox_tools())
        self.assertEqual(model.call_count, 4)
        self.assertEqual(result.final_text, compile_comparison_response(self.contract, result.execution_evidence)[0])
        self.assertIn("| 营业收入 | 12月 | 112 | 100 | 12 |", result.final_text)
        self.assertIn("| 净利润 | 7月 | 80 | 30 | 50 |", result.final_text)
        self.assertEqual(len(result.execution_evidence), 2)

    def test_model_extra_ranks_signs_units_and_causes_never_enter_report(self):
        tool_response = {'content':'','finish_reason':'tool_calls','tool_calls':[{
            'id':'verified','type':'function','function':{
                'name':'compare_spreadsheet_columns','arguments':json.dumps(self.params)}}]}
        reports=[]
        for prose in [ANSWER+'\n第二名是10月，负绝对偏差-9999万元。原因已经确认。',
                      '已核算999999千元，所有结果正常。',
                      '{"evidence_ids":["历史ID"],"hypotheses":[]}']:
            with patch('dsh_modules.table_comparison.resolve_sandboxed_path',return_value=(self.path,None)), \
                 patch('dsh_modules.agent_loop.call_model_proxy',side_effect=[tool_response,{'content':prose,'finish_reason':'stop'}]) as model:
                result=run_agent_loop([{'role':'user','content':QUERY}],'test',RuntimePolicy(),4,
                    is_inspect_intent=True,tools=get_sandbox_tools())
            reports.append(result.final_text)
            self.assertNotIn('9999',result.final_text)
            self.assertNotIn('第二名',result.final_text)
            self.assertNotIn('原因已经确认',result.final_text)
            self.assertTrue(all(call.kwargs['stream_deltas'] is False for call in model.call_args_list))
            self.assertEqual(model.call_args_list[0].kwargs['tool_choice'],'required')
        self.assertEqual(len(set(reports)),1)


if __name__ == "__main__":
    unittest.main()
