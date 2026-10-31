"""
Tests for guard relaxation on read/inspect tasks, URL markdown artifact disambiguation,
and ScriptLeakJudge protection of substantive answers.
"""

import unittest
from unittest.mock import patch, MagicMock
from pathlib import Path

from dsh_modules.deliverable_contract import (
    requests_markdown_artifact,
    resolve_markdown_filename,
    is_substantive_final_content
)
from dsh_modules.llm import detect_unexecuted_script_leak
from dsh_modules.completion_evaluator import evaluate_no_tool_turn
from dsh_modules.final_answer import finalize_agent_text, evaluate_answer_state
from dsh_modules.telemetry import TelemetryStats
from dsh_modules.runtime_policy import RuntimePolicy
from dsh_modules.script_leak_judge import ScriptLeakJudge, judge_script_leak


class TestInspectScriptLeakAndGuardRelaxation(unittest.TestCase):

    def test_github_url_ending_in_md_not_treated_as_markdown_deliverable_request(self):
        prompt = "打开https://github.com/anthropics/skills/blob/main/skills/docx/SKILL.md 查看主要内容"
        self.assertFalse(requests_markdown_artifact(prompt))
        filename = resolve_markdown_filename(prompt)
        self.assertNotEqual(filename, "SKILL.md")
        self.assertEqual(filename, "result.md")

    def test_explicit_markdown_request_still_works(self):
        prompt = "请帮我整理一份分析总结，生成并保存为 summary.md 文件"
        self.assertTrue(requests_markdown_artifact(prompt))
        self.assertEqual(resolve_markdown_filename(prompt), "summary.md")

    def test_bash_commands_with_docx_in_args_not_flagged_as_script_leak(self):
        # Snippets from anthropics docx/SKILL.md explaining CLI tools
        bash_example = """
## 四、验证输出
```bash
python scripts/office/soffice.py --headless --convert-to pdf output.docx
pdftoppm -jpeg -r 100 output.pdf page
ls page-*.jpg
```

## 七、批注
```bash
python scripts/comment.py unpacked/ "Fees cap is too low"
python scripts/comment.py contract.docx "This cap is too low" -o annotated.docx
```
"""
        self.assertFalse(detect_unexecuted_script_leak(bash_example))

    def test_actual_python_docx_pptx_fpdf_generation_code_still_flagged(self):
        pptx_leak = '```python\nfrom pptx import Presentation\nprs = Presentation("AIGC.pptx")\nprint(prs)\n```'
        self.assertTrue(detect_unexecuted_script_leak(pptx_leak))

        fpdf_leak = '```python\nfrom fpdf import FPDF\npdf = FPDF()\npdf.output("/workspace/test.pdf")\n```'
        self.assertTrue(detect_unexecuted_script_leak(fpdf_leak))

    def test_evaluate_no_tool_turn_exempts_inspect_tasks(self):
        prompt = "打开https://github.com/anthropics/skills/blob/main/skills/docx/SKILL.md 查看主要内容"
        reply = """
以下是该 SKILL.md 的主要内容：
```bash
python scripts/office/validate.py out.docx --original doc.docx
```
该规范强调了对 docx 文件的验证与编辑流程。
"""
        dec = evaluate_no_tool_turn(
            reply_text=reply,
            last_user_prompt=prompt,
            round_idx=4,
            max_rounds=10,
            start_ts=0,
            is_guide_intent=False,
            expected_deliverables=[],
            is_generate_intent=False,
            is_inspect_intent=True,
            tools=[{"name": "fetch_page"}, {"name": "read_file"}]
        )
        self.assertEqual(dec.action, "pass")

    def test_evaluate_no_tool_turn_skips_script_leak_when_no_bash(self):
        prompt = "如何生成 Word 批注？"
        reply = """
```python
from pptx import Presentation
prs = Presentation()
prs.save('/workspace/out.pptx')
```
"""
        dec = evaluate_no_tool_turn(
            reply_text=reply,
            last_user_prompt=prompt,
            round_idx=2,
            max_rounds=10,
            start_ts=0,
            is_guide_intent=False,
            expected_deliverables=[],
            is_generate_intent=False,
            is_inspect_intent=False,
            tools=[{"name": "read_file"}]  # no bash!
        )
        self.assertEqual(dec.action, "pass")

    def test_finalize_agent_text_preserves_substantive_answer_on_inspect_task(self):
        prompt = "打开https://github.com/anthropics/skills/blob/main/skills/docx/SKILL.md 查看主要内容"
        reply = """
## 📄 anthropics/skills 仓库 — `docx/SKILL.md` 主要内容

该文件是 Anthropics 官方 Skill 仓库中 docx 技能的定义文档。
核心内容包含：
1. 元信息：定义了 Word 技能触发条件。
2. 工作策略：创建新文档写 docx 脚本，读取内容使用 pandoc。
3. 验证方式：
```bash
python scripts/office/soffice.py --headless --convert-to pdf output.docx
```
完整总结如上。
"""
        final_text, healed = finalize_agent_text(
            reply_text=reply,
            messages=[],
            model="default",
            policy=RuntimePolicy.from_env(),
            deadline=None,
            is_guide_intent=False,
            is_inspect_intent=True,
            is_generate_intent=False,
            is_search_intent=False,
            has_bash=False,
            executed_calls_history=["fetch_page:...", "read_file:..."],
            was_token_truncated=False,
            telemetry=TelemetryStats(),
            last_user_prompt=prompt,
            expected_deliverables=[],
            execution_evidence=[],
            call_model=MagicMock(),
            workspace_dir="/tmp",
            max_agent_rounds=10,
            max_spreadsheet_rounds=10
        )
        self.assertIn("## 📄 anthropics/skills 仓库", final_text)
        self.assertNotIn("⚠️ 沙箱已成功调用工具", final_text)

    def test_script_leak_judge_fallback_with_substantive_text(self):
        prompt = "解释一下 docx 库用法"
        reply = """
这里为您详细解释 docx 库的使用方法：
首先安装 python-docx，然后通过 Document 对象创建文档，添加段落与表格。
例如：
```python
import docx
doc = docx.Document()
doc.add_heading("标题")
doc.save("/workspace/test.docx")
```
以上就是基本的快速入门流程，注意处理表格列宽时使用 DXA 单位。
"""
        verdict = ScriptLeakJudge._fallback_verdict(reply)
        self.assertFalse(verdict["is_unexecuted_script_leak"])
        self.assertIn("substantive text", verdict["reason"])

    def test_fake_tool_claim_interception(self):
        prompt = "打开https://github.com/anthropics/skills/blob/main/skills/docx/SKILL.md 查看主要内容"
        reply = "调用 `fetch_page` 获取 GitHub 页面信息失败。\n\n调用 `web_search` 搜索相关信息失败。"
        dec = evaluate_no_tool_turn(
            reply_text=reply,
            last_user_prompt=prompt,
            round_idx=0,
            max_rounds=5,
            start_ts=0,
            is_guide_intent=False,
            is_inspect_intent=True,
            is_generate_intent=False,
            executed_calls_history=[],
            tools=[{"name": "fetch_page"}, {"name": "web_search"}]
        )
        self.assertEqual(dec.action, "continue")
        self.assertEqual(dec.category, "fake_tool_claim")
        self.assertIn("未发起任何实际工具调用", dec.guard_msg)

    def test_unexecuted_url_fetch_interception(self):
        prompt = "打开https://github.com/anthropics/skills/blob/main/skills/docx/SKILL.md 查看主要内容"
        reply = "无法打开该链接。"
        dec = evaluate_no_tool_turn(
            reply_text=reply,
            last_user_prompt=prompt,
            round_idx=0,
            max_rounds=5,
            start_ts=0,
            is_guide_intent=False,
            is_inspect_intent=True,
            is_generate_intent=False,
            executed_calls_history=[],
            tools=[{"name": "fetch_page"}, {"name": "web_search"}]
        )
        self.assertEqual(dec.action, "continue")
        self.assertEqual(dec.category, "unexecuted_url_fetch")
        self.assertIn("fetch_page", dec.guard_msg)

    def test_skill_router_url_does_not_hijack_docx_skill(self):
        from dsh_modules.skill_router import SkillRouter
        prompt = "打开https://github.com/anthropics/skills/blob/main/skills/docx/SKILL.md 查看主要内容"
        res = SkillRouter.route(prompt)
        self.assertNotEqual(res.skill_id, "docx")
        self.assertFalse(res.is_docx_intent)

    def test_prompt_builder_no_contradictory_instructions_on_url(self):
        from dsh_modules.prompt_builder import build_user_turn
        prompt = "打开https://github.com/anthropics/skills/blob/main/skills/docx/SKILL.md 查看主要内容"
        tools = ["fetch_page", "read_file", "scan_knowledge", "vision_inspect", "web_search"]
        turn = build_user_turn(
            prompt=prompt,
            is_docx_intent=False,
            is_inspect_intent=True,
            is_generate_intent=False,
            available_tool_names=tools
        )
        self.assertNotIn("【Word/合同审阅与批注硬性执行要求】", turn)
        self.assertNotIn("【Markdown Deliverable】", turn)
        self.assertIn("【数据分析与只读查阅准则】", turn)
        self.assertIn("[Available Capabilities]: fetch_page", turn)

    def test_extract_source_and_target_artifacts_ignores_urls(self):
        from dsh_modules.intent_contract import extract_source_and_target_artifacts
        prompt = "打开https://github.com/anthropics/skills/blob/main/skills/docx/SKILL.md 查看主要内容"
        src, tgt_fmt, tgt_files = extract_source_and_target_artifacts(prompt)
        self.assertEqual(src, [])
        self.assertEqual(tgt_fmt, [])
        self.assertEqual(tgt_files, [])

    def test_prompt_builder_does_not_inject_search_when_search_intent_false(self):
        from dsh_modules.prompt_builder import build_user_turn
        prompt = "基于附件生成 2026 年最新产品销售预测 PPT"
        turn = build_user_turn(
            prompt=prompt,
            is_search_intent=False,
            is_ppt_intent=True,
            available_tool_names=["bash"]
        )
        self.assertNotIn("[Live Information Retrieval Directive]", turn)

    def test_prompt_builder_weather_gated_on_tool(self):
        from dsh_modules.prompt_builder import build_user_turn
        prompt = "今天上海天气怎么样"
        # Tool list does NOT contain weather
        turn = build_user_turn(
            prompt=prompt,
            is_weather_intent=True,
            available_tool_names=["read_file", "bash"]
        )
        self.assertNotIn("【Weather】", turn)

        # Tool list contains weather
        turn_with_tool = build_user_turn(
            prompt=prompt,
            is_weather_intent=True,
            available_tool_names=["weather", "read_file"]
        )
        self.assertIn("【Weather】", turn_with_tool)

    def test_prompt_builder_reminder_uses_action_protocol(self):
        from dsh_modules.prompt_builder import build_user_turn
        # Mention of reminder in report/explanation is not a command
        prompt_explanation = "制作一份关于团队效率的报告，提醒大家珍惜时间"
        turn_exp = build_user_turn(
            prompt=prompt_explanation,
            available_tool_names=["create_reminders", "bash"]
        )
        self.assertNotIn("【Reminder】", turn_exp)

        # Real command with time
        prompt_real = "提醒我明天下午3点准时开会"
        turn_real = build_user_turn(
            prompt=prompt_real,
            available_tool_names=["create_reminders"]
        )
        self.assertIn("【Reminder】按下方当前时间把用户给出的时间换算为 ISO 8601", turn_real)

    def test_prompt_builder_send_file_gated_on_tool(self):
        from dsh_modules.prompt_builder import build_user_turn
        prompt = "把文件发我"
        turn_no_tool = build_user_turn(
            prompt=prompt,
            is_send_intent=True,
            available_tool_names=["read_file"]
        )
        self.assertNotIn("[Delivery Intent]", turn_no_tool)

        turn_with_tool = build_user_turn(
            prompt=prompt,
            is_send_intent=True,
            available_tool_names=["send_file", "read_file"]
        )
        self.assertIn("[Delivery Intent]", turn_with_tool)


if __name__ == '__main__':
    unittest.main()
