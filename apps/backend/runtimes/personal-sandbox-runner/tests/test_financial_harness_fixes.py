"""Unit tests validating harness fixes for finance workbook evaluation (F01-F06)."""

import json
import unittest
from pathlib import Path
from tempfile import NamedTemporaryFile
import openpyxl

from dsh_modules.analysis_contract import build_analysis_contract
from dsh_modules.analysis_task_state import analysis_task_state
from dsh_modules.query_executor import execute_query_plan
from dsh_modules.query_result import QueryResult
from dsh_modules.tool_result import ToolResult
from dsh_modules.evidence_store import persist_turn_evidence


class TestFinancialHarnessFixes(unittest.TestCase):
    def test_f01_formula_calculation_description_not_misidentified_as_grouping(self):
        """F01: Prompts describing formula definitions ('按全年利润除以全年收入计算') must not be misidentified as group dimension."""
        prompt = (
            "核算2026年全年营业收入、毛利、毛利率、营业利润、净利润以及净利率六项指标，给出金额单位和计算结果。"
            "说明毛利率与净利率的计算口径或来源（例如全年利润率按全年利润除以全年收入计算，还是12个月简单平均？）。"
            "核对各项指标在不同表格中是否一致。只在聊天回复，不生成文件，不联网。"
        )
        contract = build_analysis_contract(prompt)
        self.assertFalse(contract.is_grouped, "F01 calculation description must not be treated as grouped query")
        self.assertNotEqual(contract.group_dimension, "全年利润除以全年收入")

        # Explicit groupings should still work
        grouped_prompt = "按部门统计各部门营业利润与毛利"
        grouped_contract = build_analysis_contract(grouped_prompt)
        self.assertTrue(grouped_contract.is_grouped)
        self.assertEqual(grouped_contract.group_dimension, "部门")

    def test_f02_ratio_to_percentage_scaling_and_gate_compliance(self):
        """F02: Metrics with unit=ratio and output_unit=% must be multiplied by 100 and marked with unit=%."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "月度经营"
        ws.cell(3, 1, "月份")
        ws.cell(3, 2, "营业收入")
        ws.cell(3, 4, "毛利")
        ws.cell(3, 5, "毛利率")
        ws.cell(16, 1, "全年合计/平均")
        ws.cell(16, 2, 124020)
        ws.cell(16, 4, 53443.25)
        ws.cell(16, 5, 0.430924447669731)

        q_spec = {
            "source": {"sheet": "月度经营", "header_row": 3, "data_range": "A4:E16"},
            "where": {"field": "月份", "op": "eq", "value": "全年合计/平均"},
            "metrics": [
                {"id": "revenue", "label": "营业收入", "op": "sum", "field": "营业收入", "unit": "amount", "output_unit": "千元"},
                {"id": "gross_margin", "label": "毛利率", "op": "avg", "field": "毛利率", "unit": "ratio", "output_unit": "%"},
            ]
        }

        with NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
            wb.save(tmp.name)
            tmp_path = tmp.name

        try:
            wb_data = openpyxl.load_workbook(tmp_path, data_only=True)
            res = execute_query_plan(wb_data, q_spec, check_deadline=lambda: None, default_amount_unit="千元")
            self.assertIsInstance(res, QueryResult)

            gm_fact = next(f for f in res.facts if f["id"] == "gross_margin")
            self.assertEqual(gm_fact["unit"], "%", "output_unit='%' must convert final unit to '%'")
            self.assertTrue(gm_fact["value"].startswith("43.092444"), f"Ratio must be multiplied by 100: {gm_fact['value']}")

            # Validate that analysis_task_state accepts this fact for percentage requirement
            tool_res = res.to_tool_result(
                source=Path(tmp_path),
                source_hash="dummy_hash",
                requested_plan=q_spec,
                effective_plan=q_spec,
                plan_hash="plan_hash_1",
                amount_unit="千元",
                semantic_domain={"id": "finance", "version": "1"}
            )
            state = analysis_task_state(
                [tool_res],
                prompt="只核算2026年营业收入、毛利率、营业利润和净利润，提供金额单位和结果。"
            )
            missing_names = [m["name"] for m in state.get("missing_metrics", [])]
            self.assertNotIn("毛利率", missing_names, "Scaled percentage fact must satisfy percentage requirement")
        finally:
            Path(tmp_path).unlink(missing_ok=True)

    def test_f04_null_predicate_and_delimiters_not_misidentified_as_entities(self):
        """F04: Null predicates ('部门为空') and multi-clause statements with semicolons must not extract false entities."""
        prompt = (
            "接下来检查同一个附件的“交易明细”第4至83行。分别列出：凭证号重复的全部记录；部门为空的记录；"
            "金额绝对值大于10000的记录；日期不在2026年内的记录；审批状态不是已审批的记录。"
            "每类给出条数、Excel行号和关键字段。备注为空不算异常，费用负数不算异常。只在聊天回复，不修改原件，不联网。"
        )
        contract = build_analysis_contract(prompt)
        self.assertNotIn("为空", contract.entity_predicates, "'为空' is a filter predicate, not a business entity name")
        self.assertFalse(any("证号重复" in e for e in contract.entity_predicates), "Regex must not swallow adjacent clauses across semicolons")

        # Legitimate entities should still be captured
        legit_prompt = "查询华东MIS部门杨烨在2026年的报销记录"
        legit_contract = build_analysis_contract(legit_prompt)
        self.assertTrue(any("杨烨" in e for e in legit_contract.entity_predicates))

    def test_f06_artifact_delivery_turn_verified(self):
        """F06: Turns delivering files without calculation gate failures should be marked verified."""
        with NamedTemporaryFile(suffix=".json", delete=False) as tmp:
            history_file = tmp.name
        try:
            guard_decisions = [
                {"guard": "task_intent_review", "action": "accept"},
                {"guard": "task_intent", "action": "accept", "reason": {"output_mode": "artifact"}},
                {"guard": "skill_admission", "action": "use"},
                {"guard": "user_unit_confirmation", "action": "accept"},
            ]
            receipt = ToolResult.success(
                data="<<<DSH_OUTBOUND_FILE:len=50:{\"filePath\": \"/workspace/out.xlsx\"}>>>",
                provenance={"tool": "send_file"}
            )
            res = persist_turn_evidence(
                history_file=history_file,
                source_hashes={"test.xlsx": "hash123"},
                query="生成可下载文件",
                receipts=[receipt],
                guard_decisions=guard_decisions
            )
            self.assertTrue(res["verified"], "Outbound delivered artifact without rejection should be verified")
        finally:
            Path(history_file).unlink(missing_ok=True)
            Path(history_file).with_suffix('.evidence.json').unlink(missing_ok=True)

    def test_llm_judge_dismisses_gate_false_positive_required_metrics(self):
        """LLM as a judge: Gate review must honor judge dismissal of required_metrics_missing when facts exist."""
        from dsh_modules.analysis_gate_review import revalidate_review

        issues = [{
            "code": "required_metrics_missing",
            "metrics": [{"name": "毛利率", "metric_id": "gross_margin"}],
            "message": "分析计划遗漏所需指标计算：毛利率"
        }]
        verdict = {
            "action": "repair",
            "category": "gate_false_positive",
            "complete": False,
            "dismissed_issue_codes": ["required_metrics_missing"],
            "supporting_fact_ids": ["gross_margin"],
            "reason": "系统误报：毛利率实际已完成计算并存在于事实中。"
        }
        facts = [
            {"id": "gross_margin", "label": "毛利率(ratio)", "value": "0.4309", "unit": "ratio"}
        ]
        receipts = [
            ToolResult.success(data={"facts": facts, "coverage": [{"coverage_complete": True}]})
        ]

        remaining, updated_verdict = revalidate_review(issues, verdict, receipts, facts)
        self.assertEqual(len(remaining), 0, "LLM Judge false-positive dismissal must remove the issue")
        self.assertIn("required_metrics_missing", updated_verdict["dismissed_issue_codes"])
        self.assertTrue(updated_verdict["complete"])
        self.assertEqual(updated_verdict["action"], "recheck")

    def test_llm_judge_dismisses_unsolicited_filter_applied(self):
        """LLM as a judge: Gate review must honor judge dismissal of unsolicited_filter_applied."""
        from dsh_modules.analysis_gate_review import revalidate_review

        issues = [{
            "code": "unsolicited_filter_applied",
            "message": "执行计划包含了未经请求的明细行过滤条件"
        }]
        verdict = {
            "action": "recheck",
            "category": "gate_false_positive",
            "complete": False,
            "dismissed_issue_codes": ["unsolicited_filter_applied"],
            "supporting_fact_ids": [],
            "reason": "用户指令明确要求筛选'部门为空'，过滤器符合用户真实意图。"
        }
        facts = [{"id": "count", "label": "条数", "value": "1", "unit": "number"}]
        receipts = [ToolResult.success(data={"facts": facts, "coverage": [{"coverage_complete": True}]})]

        remaining, updated_verdict = revalidate_review(issues, verdict, receipts, facts)
        self.assertEqual(len(remaining), 0)
        self.assertIn("unsolicited_filter_applied", updated_verdict["dismissed_issue_codes"])
        self.assertTrue(updated_verdict["complete"])

    def test_row_filter_not_applied_cannot_be_dismissed_by_group_by(self):
        """Hard red line: group_by across full coverage must NEVER dismiss row_filter_not_applied."""
        from dsh_modules.analysis_gate_review import revalidate_review

        issues = [{
            "code": "row_filter_not_applied",
            "message": "请求包含明细行筛选，但执行范围未带行筛选条件；分组不能代替明细筛选。"
        }]
        verdict = {
            "action": "recheck",
            "category": "gate_false_positive",
            "complete": False,
            "dismissed_issue_codes": ["row_filter_not_applied"],
            "supporting_fact_ids": [],
            "reason": "已执行全表分组"
        }
        # Receipt only did grouping on voucher number, NO row filters!
        receipts = [ToolResult.success(data={
            "plan": {"group_by": ["凭证号"], "metrics": [{"field": "凭证号", "op": "count"}]},
            "coverage": [{"id": "t1", "coverage_complete": True, "filters": []}],
            "facts": [{"id": "cnt", "label": "凭证号计数", "value": "2"}]
        })]
        facts = [{"id": "cnt", "label": "凭证号计数", "value": "2"}]

        remaining, updated_verdict = revalidate_review(issues, verdict, receipts, facts)
        self.assertEqual(len(remaining), 1, "row_filter_not_applied MUST NOT be dismissed by group_by!")
        self.assertNotIn("row_filter_not_applied", updated_verdict["dismissed_issue_codes"])
        self.assertFalse(updated_verdict["complete"])

    def test_required_metrics_missing_cannot_be_dismissed_by_partial_match(self):
        """Gate review must require ALL missing metrics to be covered, never allowing 1 metric to excuse 3."""
        from dsh_modules.analysis_gate_review import revalidate_review

        issues = [{
            "code": "required_metrics_missing",
            "metrics": [
                {"name": "营业收入", "metric_id": "revenue"},
                {"name": "净利润", "metric_id": "net_profit"},
                {"name": "净利率", "metric_id": "net_profit_margin"}
            ],
            "message": "分析计划遗漏所需指标计算"
        }]
        verdict = {
            "action": "recheck",
            "category": "gate_false_positive",
            "complete": False,
            "dismissed_issue_codes": ["required_metrics_missing"],
            "supporting_fact_ids": [],
            "reason": "已算出营业收入"
        }
        # Only revenue exists; net profit and net profit margin are missing!
        facts = [{"id": "revenue", "label": "营业收入", "value": "124020", "unit": "千元"}]
        receipts = [ToolResult.success(data={"facts": facts, "coverage": [{"coverage_complete": True}]})]

        remaining, updated_verdict = revalidate_review(issues, verdict, receipts, facts)
        self.assertEqual(len(remaining), 1, "Partial match cannot dismiss required_metrics_missing for all metrics")
        self.assertFalse(updated_verdict["complete"])

    def test_audit_checks_incomplete_detected_when_multi_rules_unexecuted(self):
        """Audit intent with 5 numbered checks must detect incomplete execution when only 1 check was run."""
        prompt = (
            "请对附件“AI能力测试_财务报表.xlsx”中的“交易明细”表进行以下五项合规性检查：1. 查找重复的凭证号；2. 筛选出部门为空的记录；"
            "3. 找出金额绝对值大于10000的交易；4. 检查日期不在2026年度的跨期凭证；5. 检查审批状态非“已审批”的记录。请逐项输出检查结果与对应的行号。"
        )
        receipt_data = {
            "kind": "spreadsheet_analysis",
            "plan": {
                "source": {"sheet": "交易明细", "header_row": 3, "data_range": "A4:K83"},
                "group_by": ["凭证号"],
                "metrics": [{"id": "cnt", "label": "凭证号出现次数", "op": "count", "field": "凭证号", "unit": "number"}],
                "having": {"and": [{"metric": "cnt", "op": "gt", "value": 1}]}
            },
            "coverage": [{
                "id": "t1", "sheet": "交易明细", "coverage_complete": True, "selected_rows": 80,
                "headers": {"A": "凭证号", "B": "日期", "C": "类型", "D": "部门", "E": "金额", "K": "审批状态"}
            }],
            "facts": [{"id": "f1", "label": "重复凭证", "value": "2"}],
            "evidence_id": "eid1"
        }
        import hashlib
        with NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
            tmp.write(b"mock_content")
            tmp_path = tmp.name
        wb_id = hashlib.sha256(b"mock_content").hexdigest()
        try:
            r1 = ToolResult.success(data=receipt_data, text=json.dumps(receipt_data), provenance={"tool": "query_spreadsheet", "file_path": tmp_path, "workbook_id": wb_id})

            state = analysis_task_state([r1], prompt)
            self.assertFalse(state["complete"])
            issue_codes = [i["code"] for i in state["issues"]]
            self.assertIn("audit_checks_incomplete", issue_codes)
            audit_issue = next(i for i in state["issues"] if i["code"] == "audit_checks_incomplete")
            self.assertIn("筛选出部门为空的记录", audit_issue["missing_checks"])
            self.assertIn("找出金额绝对值大于10000的交易", audit_issue["missing_checks"])
        finally:
            Path(tmp_path).unlink(missing_ok=True)

    def test_parenthesized_unit_header_metric_stem_matching_and_revalidation(self):
        """Parenthesized column metrics like 金额(K) must match custom aggregated labels like 有偿金额合计."""
        from dsh_modules.analysis_contract import _contract_from_judge_verdict
        from dsh_modules.analysis_task_state import analysis_task_state, task_feedback
        from dsh_modules.analysis_gate_review import revalidate_review

        # 1. Contract extraction matches stem aliases
        verdict = {'task_type': 'tabular_compute', 'metrics': ['金额(K)'], 'is_grouped': True, 'group_dimension': '部门名称'}
        contract = _contract_from_judge_verdict('查看 各个部门的 有偿 金额的合计', verdict, 'generic')
        metric = contract.metrics[0]
        self.assertEqual(metric.name, '金额(K)')
        import re
        self.assertTrue(any(re.search(alias, '華東MIS 有偿金额合计') for alias in metric.aliases))

        # 2. Fact with field and custom label is matched by analysis_task_state
        facts = [{
            'id': "total_amount__(('str', '華東MIS'),)",
            'label': '華東MIS 有偿金额合计',
            'group_key': '華東MIS',
            'value': '570.75345',
            'unit': 'K',
            'field': '金额(K)',
            'metric_label': '有偿金额合计'
        }]
        receipt_data = {
            'kind': 'spreadsheet_analysis',
            'plan': {
                'source': {'sheet': '明细', 'header_row': 4, 'data_range': 'A5:R278'},
                'group_by': ['部门名称'],
                'metrics': [{'id': 'total_amount', 'label': '有偿金额合计', 'op': 'sum', 'field': '金额(K)'}],
                'where': {'field': '有偿无偿', 'op': 'eq', 'value': '有偿'}
            },
            'coverage': [{'id': 't1', 'sheet': '明细', 'coverage_complete': True, 'selected_rows': 259, 'headers': {'C': '部门名称', 'P': '金额(K)'}}],
            'facts': facts,
            'evidence_id': 'eid_test_123'
        }
        r = ToolResult.success(data=receipt_data, text=json.dumps(receipt_data), provenance={'tool': 'aggregate_spreadsheet', 'file_path': '/workspace/test.xlsx', 'workbook_id': 'wb123'})
        state = analysis_task_state([r], '查看 各个部门的 有偿 金额的合计')
        # Crucial assertion: static gate must NOT fail with required_metrics_missing
        missing_metric_names = [m.get('name') for m in state.get('missing_metrics', [])]
        self.assertNotIn('金额(K)', missing_metric_names)

        # 3. Gate review revalidation handles stem matching and semantic code correlation
        issues = [{
            'code': 'required_metrics_missing',
            'metrics': [{'name': '金额(K)'}],
            'message': '分析计划遗漏所需指标计算：金额(K)'
        }]
        review_verdict = {
            'action': 'repair',
            'category': 'gate_false_positive',
            'complete': False,
            'dismissed_issue_codes': ['per_group_metrics_missing'],
            'supporting_fact_ids': ["total_amount__(('str', '華東MIS'),)"],
            'repair_instruction': '无需补做，直接基于facts输出。'
        }
        remaining, updated_review = revalidate_review(issues, review_verdict, [r], facts)
        self.assertEqual(len(remaining), 0, "Gate review must correlate per_group_metrics_missing to required_metrics_missing on single issue")
        self.assertTrue(updated_review['complete'])

        # 4. task_feedback includes prompt reminder suppressing internal terminology
        feedback_text = task_feedback([r], '查看 各个部门的 有偿 金额的合计')
        self.assertIn('切勿在正文中提及“门禁系统”、“误报”', feedback_text)

    def test_physical_column_lineage_and_fact_retention_for_timesheet_and_amount(self):
        """Physical column lineage (field) and synonyms must preserve facts and pass gate without false alarm."""
        from dsh_modules.analysis_contract import AnalysisContract, MetricRequirement
        from dsh_modules.analysis_task_state import fact_matches_metric, analysis_task_state
        from dsh_modules.analysis_fact_selection import resolve_conforming_facts
        from dsh_modules.analysis_gate_review import revalidate_review

        # 1. fact_matches_metric matches '工数(小时)' and '工时' via field and synonym stems
        m_hours = MetricRequirement(name='工数(小时)', aliases=['工数\\(小时\\)', '工时\\(小时\\)'])
        f_hours = {
            'id': 'hours',
            'label': '工时合计(h)',
            'value': '170.59',
            'unit': 'h',
            'field': '工数(小时)',
            'metric_label': '工时合计',
            'matched_records': 6,
            'source_cells': ['工数实绩明细!O176']
        }
        self.assertTrue(fact_matches_metric(f_hours, m_hours))

        f_amount = {
            'id': 'amount',
            'label': '金额合计',
            'value': '45.15617',
            'unit': 'K',
            'field': '金额(K)',
            'metric_label': '金额合计',
            'matched_records': 6,
            'source_cells': ['工数实绩明细!P176']
        }
        m_amount = MetricRequirement(name='金额(K)', aliases=['金额\\(K\\)', '金额'])
        self.assertTrue(fact_matches_metric(f_amount, m_amount))

        # 2. resolve_conforming_facts must NOT purge hours when both hours and amount are computed
        receipt_data = {
            'kind': 'spreadsheet_analysis',
            'plan': {
                'source': {'sheet': '工数实绩明细', 'data_range': 'A5:R278', 'header_row': 4},
                'where': {'and': [{'field': '员工编号', 'op': 'eq', 'value': 'S22014'}]},
                'metrics': [
                    {'id': 'hours', 'label': '工时合计', 'field': '工数(小时)', 'op': 'sum'},
                    {'id': 'amount', 'label': '金额合计', 'field': '金额(K)', 'op': 'sum'}
                ]
            },
            'facts': [f_hours, f_amount],
            'coverage': [{'id': 'detail', 'sheet': '工数实绩明细', 'coverage_complete': True, 'headers': {'E': '员工编号', 'O': '工数(小时)', 'P': '金额(K)'}}],
            'evidence_id': 'eid_timesheet_1'
        }
        r = ToolResult.success(receipt_data, provenance={'workbook_id': 'wb_ts_1', 'file_path': '/workspace/test.xlsx', 'tool': 'aggregate_spreadsheet'})
        prompt = '统计 员工  s22014 的 工时 和金额的 合计'
        resolved = resolve_conforming_facts([r], prompt)
        resolved_ids = [f['id'] for f in resolved]
        self.assertIn('hours', resolved_ids, "Physical field fact 'hours' must not be purged by fact selection")
        self.assertIn('amount', resolved_ids, "Physical field fact 'amount' must be retained")

        # 3. revalidate_review correctly covers metrics using all_facts and physical column fields
        issues = [{
            'code': 'required_metrics_missing',
            'metrics': [{'name': '工数(小时)'}],
            'message': '指标已有执行结果，但尚未通过来源或口径核验：工数(小时)。'
        }]
        verdict = {
            'action': 'recheck',
            'category': 'gate_false_positive',
            'complete': False,
            'dismissed_issue_codes': ['required_metrics_missing'],
            'supporting_fact_ids': ['amount'],
            'reason': '工数(小时)字段已被聚合计算'
        }
        remaining, updated_v = revalidate_review(issues, verdict, [r], [f_amount])
        self.assertEqual(len(remaining), 0, "Receipt physical field lineage must satisfy revalidation")
        self.assertTrue(updated_v['complete'])


if __name__ == "__main__":
    unittest.main()
