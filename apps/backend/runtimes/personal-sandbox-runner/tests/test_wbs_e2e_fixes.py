"""
Unit tests for standardized spreadsheet analysis and personal mode fixes:
1. Generic metric extraction in analysis_contract (W01, W03, W04, W05)
2. Amount unit detection in spreadsheet_units (e.g. 金额(K) -> 'K')
3. Trailing note/footer pruning in spreadsheet_context (e.g. A5:R278)
4. Extended filter operations in spreadsheet_analysis (prefixes, exclude_prefixes, op)
5. Empty set handling for zero-match filters (W05 count=0, sum=0)
6. Deliverable filename extraction without stripping leading '名'
"""

import sys
import re
import unittest
import openpyxl
from decimal import Decimal
from pathlib import Path

MODULES_DIR = Path(__file__).resolve().parent.parent / "src"
if str(MODULES_DIR) not in sys.path:
    sys.path.insert(0, str(MODULES_DIR))

from dsh_modules.analysis_contract import build_analysis_contract
from dsh_modules.spreadsheet_units import declared_amount_unit, AMOUNT_UNIT_FACTORS
from dsh_modules.spreadsheet_expression import AMOUNT_SCALE, ExpressionEvaluator, Quantity
from dsh_modules.spreadsheet_analysis import prepare_tables, calculate_spreadsheet_analysis
from dsh_modules.spreadsheet_analysis_evidence import describe_expression
from dsh_modules.spreadsheet_context import spreadsheet_structure
from dsh_modules.intent_contract import extract_source_and_target_artifacts


class TestWbsStandardizedFixes(unittest.TestCase):
    def test_generic_metric_extraction(self):
        """1. 验证通用数值度量提取：剥离动作/连词前缀，支持复合指标"""
        # W01
        c1 = build_analysis_contract("帮我统计下附件中按 WBS负责成本中心 统计 金额 的数据")
        metric_names1 = [m.name for m in c1.metrics]
        self.assertIn("金额", metric_names1)
        self.assertTrue(any(any(re.search(a, "金额(K)") for a in m.aliases) for m in c1.metrics))
        self.assertTrue(any(any(re.search(a, "FC34000金额") for a in m.aliases) for m in c1.metrics))

        # W03: 剥离 "每组明细条数和"
        c3 = build_analysis_contract("按WBS编号前两个字母分组，统计每组明细条数和金额(K)，并核对合计")
        metric_names3 = [m.name for m in c3.metrics]
        self.assertIn("金额", metric_names3)
        self.assertNotIn("每组明细条数和金额", metric_names3)

        # W04: 复合指标提取 "金额" 和 "工数"
        c4 = build_analysis_contract("按WBS负责成本中心汇总金额(K)和工数(小时)，附总计")
        metric_names4 = [m.name for m in c4.metrics]
        self.assertIn("金额", metric_names4)
        self.assertIn("工数", metric_names4)

        # W05: 剥离 "条数及"
        c5 = build_analysis_contract("仅保留WBS编号以SY、IB或SZ开头的明细，统计条数及金额(K)")
        metric_names5 = [m.name for m in c5.metrics]
        self.assertIn("金额", metric_names5)
        self.assertNotIn("条数及金额", metric_names5)

    def test_amount_unit_detection(self):
        """2. 验证表头金额单位识别（如 金额(K) -> 'K'）及量纲系数支持"""
        self.assertIn("K", AMOUNT_UNIT_FACTORS)
        self.assertEqual(AMOUNT_UNIT_FACTORS["K"], 1000)
        self.assertIn("M", AMOUNT_UNIT_FACTORS)
        self.assertEqual(AMOUNT_UNIT_FACTORS["M"], 1000000)

        self.assertIn("K", AMOUNT_SCALE)
        self.assertEqual(AMOUNT_SCALE["K"], Decimal(1000))

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "工数明细"
        ws["A1"] = "WBS编号"
        ws["B1"] = "金额(K)"
        ws["A2"] = "CK001"
        ws["B2"] = 12.5

        unit, locs = declared_amount_unit(wb, "工数明细", 1)
        self.assertEqual(unit, "K")
        self.assertIn("工数明细!B1", locs)

    def test_spreadsheet_context_footer_pruning(self):
        """3. 验证表尾查询条件、报表说明等非业务行自动修剪"""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "明细表"
        # Header at row 2
        headers = ["WBS编号", "成本中心", "工数(小时)", "金额(K)"]
        for col_idx, h in enumerate(headers, 1):
            ws.cell(row=2, column=col_idx, value=h)
        # Data rows 3-10
        for r in range(3, 11):
            ws.cell(row=r, column=1, value=f"WBS{r}")
            ws.cell(row=r, column=2, value="FC34000")
            ws.cell(row=r, column=3, value=10.0)
            ws.cell(row=r, column=4, value=100.0)
        # Footers at rows 11-13
        ws.cell(row=11, column=1, value="查询条件：")
        ws.cell(row=12, column=1, value="*会计年度:2026")
        ws.cell(row=13, column=1, value="报表说明：内部工数实绩明细")

        test_file = Path("/tmp/test_footer_pruning.xlsx")
        wb.save(test_file)

        res = spreadsheet_structure(str(test_file))
        self.assertTrue(res.is_success)
        sheet_data = res.data["sheets"][0]
        candidates = sheet_data["data_region_candidates"]
        self.assertTrue(len(candidates) > 0)
        # Candidate data range should end at row 10, not 13
        self.assertEqual(candidates[0]["data_range"], "A3:D10")
        self.assertEqual(candidates[0]["header_row"], 2)

        test_file.unlink(missing_ok=True)

    def test_spreadsheet_analysis_filter_extensions(self):
        """4. 验证过滤条件扩展：exclude_prefixes、prefixes、大小写不敏感"""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "测试表"
        ws["A1"] = "WBS编号"
        ws["B1"] = "金额"
        rows_data = [
            ("EX001", 10),
            ("ex002", 20),
            ("IP001", 30),
            ("CK001", 100),
            ("CK002", 200),
            ("SK001", 300),
        ]
        for idx, (code, amt) in enumerate(rows_data, 2):
            ws[f"A{idx}"] = code
            ws[f"B{idx}"] = amt

        # 4.1 exclude_prefixes
        table_plan = [{
            "id": "t1",
            "sheet": "测试表",
            "data_range": "A2:B7",
            "header_row": 1,
            "filters": [{
                "column": "A",
                "exclude_prefixes": ["ex", "ip"]
            }]
        }]
        tables = prepare_tables(table_plan, wb, lambda: None)
        # Should keep only CK001, CK002, SK001 (rows 5, 6, 7)
        self.assertEqual(tables["t1"]["selected_rows"], [5, 6, 7])

        # 4.2 prefixes
        table_plan2 = [{
            "id": "t2",
            "sheet": "测试表",
            "data_range": "A2:B7",
            "header_row": 1,
            "filters": [{
                "column": "WBS编号",
                "prefixes": ["ck"]
            }]
        }]
        tables2 = prepare_tables(table_plan2, wb, lambda: None)
        self.assertEqual(tables2["t2"]["selected_rows"], [5, 6])

    def test_empty_filter_results_handling(self):
        """5. 验证过滤结果为空集时正常返回 0 条 / 金额 0，而非崩溃"""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "测试表"
        ws["A1"] = "WBS编号"
        ws["B1"] = "金额(K)"
        ws["A2"] = "CK001"
        ws["B2"] = 100.0

        # Filter for non-existent prefixes: SY, IB, SZ
        table_plan = [{
            "id": "t_empty",
            "sheet": "测试表",
            "data_range": "A2:B2",
            "header_row": 1,
            "filters": [{
                "column": "A",
                "prefixes": ["SY", "IB", "SZ"]
            }]
        }]
        tables = prepare_tables(table_plan, wb, lambda: None)
        self.assertEqual(tables["t_empty"]["selected_rows"], [])
        self.assertEqual(tables["t_empty"]["evaluated_records"], 0)

        # Expression evaluation on empty table
        evaluator = ExpressionEvaluator(wb, tables, "K", lambda: None)
        # count -> 0
        q_count = evaluator.evaluate({
            "aggregate": "count",
            "table": "t_empty",
            "column": "A",
            "unit": "number"
        })
        self.assertEqual(q_count.value, Decimal("0"))
        self.assertEqual(q_count.unit, "number")

        # sum -> 0
        q_sum = evaluator.evaluate({
            "aggregate": "sum",
            "table": "t_empty",
            "column": "B",
            "unit": "amount"
        })
        self.assertEqual(q_sum.value, Decimal("0"))
        self.assertEqual(q_sum.unit, "K")

        # describe_expression should return '无匹配记录'
        expr_data = {"coverage": [{"id": "t_empty", "sheet": "测试表", "headers": {"B": "金额(K)"}, "selected_rows": []}]}
        desc = describe_expression({"aggregate": "sum", "table": "t_empty", "column": "B"}, expr_data)
        self.assertIn("无匹配记录", desc)

    def test_filename_extraction_leading_ming(self):
        """6. 验证文件名抽取不会误截掉开头的'名'字符，并覆盖名单、名册、名称对照表"""
        cases = [
            ("生成名单.xlsx", "名单.xlsx"),
            ("生成名册.xlsx", "名册.xlsx"),
            ("生成名称对照表.xlsx", "名称对照表.xlsx"),
            ("生成可下载Excel，文件名WBS成本中心汇总.xlsx", "WBS成本中心汇总.xlsx"),
            ("导出文件名为明细表.xlsx", "明细表.xlsx"),
        ]
        for prompt, exp in cases:
            sources, formats, targets = extract_source_and_target_artifacts(prompt)
            self.assertIn(exp, targets, f"Failed for prompt: {prompt}")

    def test_uncalculated_plan_truthful_attribution(self):
        """7. P1: 验证计划未计算指标时真实归因，绝不误报为原件缺少列"""
        from dsh_modules.analysis_task_state import analysis_task_state
        from dsh_modules.report_presentation import render_unverified_guidance
        from dsh_modules.tool_result import ToolResult

        # Mock evidence: structure found '金额(K)' column, but validate_spreadsheet_rows ran without aggregate_spreadsheet
        structure_receipt = ToolResult.success(data={
            "kind": "spreadsheet_structure",
            "sheets": [{
                "sheet": "工数实绩明细",
                "header_candidates": [{
                    "row": 4,
                    "columns": {"I": "WBS编号", "J": "WBS负责成本中心", "P": "金额(K)"}
                }]
            }]
        })
        validate_receipt = ToolResult.success(data={
            "kind": "spreadsheet_analysis",
            "evidence_id": "ev_001",
            "facts": [],  # No aggregation facts computed!
            "rules": [{"label": "WBS非空", "evaluated_records": 274, "violation_count": 0}],
            "coverage": [{"id": "t1", "sheet": "工数实绩明细", "headers": {"I": "WBS编号", "P": "金额(K)"}, "coverage_complete": True}]
        })

        prompt = "按WBS编号前两个字母分组，统计每组明细条数和金额(K)，并核对合计"
        state = analysis_task_state([structure_receipt, validate_receipt], prompt)

        self.assertFalse(state["complete"])
        missing = state["missing_metrics"]
        self.assertTrue(len(missing) > 0)
        amount_metric = next(m for m in missing if m["name"] == "金额")
        # Critical assertion: system correctly recognizes that the source column DOES exist!
        self.assertTrue(amount_metric["source_column_exists"])
        self.assertIn("金额(K)", amount_metric["matched_source_columns"])

        # Guidance assertion: MUST NOT claim table lacks column
        guidance = render_unverified_guidance(state, prompt)
        self.assertIn("表格中已存在对应列", guidance)
        self.assertIn("金额(K)", guidance)
        self.assertIn("但本轮分析计划尚未执行对", guidance)
        self.assertIn("金额", guidance)
        self.assertNotIn("表格中缺少计算 【金额】 所需的原始列", guidance)

    def test_prepare_tables_prunes_trailing_footers(self):
        """8. P2: 验证prepare_tables在执行层自动剪枝表尾说明行，阻断A5:R281产生的伪缺失"""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "工数明细"
        # Row 4: headers
        ws.cell(row=4, column=1, value="WBS编号")
        ws.cell(row=4, column=2, value="成本中心")
        ws.cell(row=4, column=3, value="金额(K)")
        # Rows 5-14: 10 data rows
        for r in range(5, 15):
            ws.cell(row=r, column=1, value=f"CK00{r}")
            ws.cell(row=r, column=2, value="FC34000")
            ws.cell(row=r, column=3, value=100.0)
        # Rows 15-17: Trailing query notes
        ws.cell(row=15, column=1, value="查询条件：")
        ws.cell(row=16, column=1, value="*会计年度:2026")
        ws.cell(row=17, column=1, value="报表说明：实绩明细")

        # Even if plan passed A5:C17 (including trailing footers)
        plan = [{
            "id": "t_raw",
            "sheet": "工数明细",
            "data_range": "A5:C17",
            "header_row": 4
        }]
        tables = prepare_tables(plan, wb, lambda: None)
        t = tables["t_raw"]
        # Must be trimmed to A5:C14 (10 rows), excluding rows 15-17
        self.assertEqual(t["data_range"], "A5:C14")
        self.assertEqual(t["selected_rows"], list(range(5, 15)))
        self.assertEqual(t["evaluated_records"], 10)
        self.assertEqual(t["input_records"], 10)

    def test_zero_record_group_omitted_in_presentation(self):
        """9. 验证全被过滤的空分组项（如FC32003）在展示时被自动过滤，不污染正常分组表"""
        from dsh_modules.report_presentation import analysis_overview
        from dsh_modules.tool_result import ToolResult

        receipt = ToolResult.success(data={
            "kind": "spreadsheet_analysis",
            "evidence_id": "ev_002",
            "plan": {
                "calculations": [
                    {"id": "f1", "label": "FC34000金额", "expression": {"aggregate": "sum", "table": "t1", "column": "C", "unit": "amount"}},
                    {"id": "f2", "label": "FC45031金额", "expression": {"aggregate": "sum", "table": "t2", "column": "C", "unit": "amount"}},
                    {"id": "f3", "label": "FC32003金额", "expression": {"aggregate": "sum", "table": "t3", "column": "C", "unit": "amount"}},
                ]
            },
            "facts": [
                {"id": "f1", "label": "FC34000金额", "value": "150.0", "unit": "K", "matched_records": 10},
                {"id": "f2", "label": "FC45031金额", "value": "500.0", "unit": "K", "matched_records": 40},
                {"id": "f3", "label": "FC32003金额", "value": "0", "unit": "K", "matched_records": 0},  # Fully filtered out!
            ],
            "coverage": [
                {"id": "t1", "sheet": "Sheet1", "data_range": "A1:C10", "selected_rows": list(range(1, 11))},
                {"id": "t2", "sheet": "Sheet1", "data_range": "A1:C40", "selected_rows": list(range(1, 41))},
                {"id": "t3", "sheet": "Sheet1", "data_range": "A1:C1", "selected_rows": []},
            ]
        }, provenance={"workbook_id": "wb_01", "file_path": "/workspace/工数实绩明细_20261020134124.xlsx"})
        overview = analysis_overview([receipt], "按成本中心汇总金额")
        self.assertIn("FC34000金额", overview)
        self.assertIn("FC45031金额", overview)
        # FC32003 had 0 matched records and should NOT be displayed
        self.assertNotIn("FC32003金额", overview)

    def test_footer_preserves_business_name_like_totalenergies(self):
        """10. 验证自动修剪表尾时，绝不误杀 TotalEnergies 等带 total 的业务实体"""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "明细"
        for row in (["客户名称", "金额"], ["Acme", 10], ["TotalEnergies", 20]):
            ws.append(row)
        table = prepare_tables([{"id": "t", "sheet": "明细", "header_row": 1, "data_range": "A2:B3"}], wb, lambda: None)["t"]
        self.assertEqual(table["selected_rows"], [2, 3])
        self.assertEqual(table["data_range"], "A2:B3")

    def test_exclude_prefix_supports_prefixes_array(self):
        """11. 验证过滤条件中 op=exclude_prefix 与 prefixes 数组组合正常排除"""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "明细"
        for row in (["WBS编号", "金额"], ["IK001", 10], ["SK001", 20]):
            ws.append(row)
        table = prepare_tables([{
            "id": "t",
            "sheet": "明细",
            "header_row": 1,
            "data_range": "A2:B3",
            "filters": [{"column": "A", "op": "exclude_prefix", "prefixes": ["IK"]}]
        }], wb, lambda: None)["t"]
        self.assertEqual(table["selected_rows"], [3])

    def test_grouping_incomplete_detected_when_single_group_computed(self):
        """12. 验证用户请求按前缀分组统计时，仅计算了单组（如IK）会被执行门禁拦截为grouping_incomplete"""
        from dsh_modules.tool_result import ToolResult
        from dsh_modules.analysis_task_state import analysis_task_state

        receipt = ToolResult.success(data={
            "kind": "spreadsheet_analysis",
            "evidence_id": "ev_003",
            "facts": [
                {"id": "f1", "label": "总明细条数", "value": "274", "unit": "number"},
                {"id": "f2", "label": "总金额", "value": "1887.73", "unit": "K"},
                {"id": "f3", "label": "IK明细条数", "value": "123", "unit": "number"},
                {"id": "f4", "label": "IK金额", "value": "561.16", "unit": "K"},
            ],
            "coverage": [
                {"id": "t_raw", "sheet": "工数实绩明细", "data_range": "A5:R278", "evaluated_records": 274, "coverage_complete": True},
                {"id": "t_ik", "sheet": "工数实绩明细", "data_range": "A5:R278", "evaluated_records": 123, "filters": [{"column": "I", "op": "prefix", "prefixes": ["IK"]}], "coverage_complete": True},
            ]
        }, provenance={"workbook_id": "wb_01", "file_path": "/workspace/test.xlsx"})

        prompt = "按WBS编号前两个字母分组，统计每组明细条数和金额(K)，并核对合计。仅统计明细，不含表尾合计和条件行。"
        state = analysis_task_state([receipt], prompt)
        self.assertFalse(state["complete"])
        issue_codes = [i["code"] for i in state["issues"]]
        self.assertIn("grouping_incomplete", issue_codes)

    def test_facts_digest_omits_zero_matched_group_when_positive_groups_exist(self):
        """13. 验证生成给模型的总结提示词中，自动剔除匹配为0的空组，防止模型在正文中输出 FC32003: 0"""
        from dsh_modules.tool_result import ToolResult
        from dsh_modules.analysis_synthesis import build_verified_facts_digest

        receipt = ToolResult.success(data={
            "kind": "spreadsheet_analysis",
            "evidence_id": "ev_004",
            "facts": [
                {"id": "f1", "label": "FC34000工数", "value": "1.2", "unit": "小时", "matched_records": 1},
                {"id": "f2", "label": "FC34000金额", "value": "0.43", "unit": "K", "matched_records": 1},
                {"id": "f3", "label": "FC32003工数", "value": "0", "unit": "小时", "matched_records": 0},
                {"id": "f4", "label": "FC32003金额", "value": "0", "unit": "K", "matched_records": 0},
                {"id": "f5", "label": "总工数", "value": "5692.78", "unit": "小时", "matched_records": 253},
                {"id": "f6", "label": "总金额", "value": "1637.42", "unit": "K", "matched_records": 253},
            ],
            "coverage": [
                {"id": "t1", "sheet": "Sheet1", "data_range": "A1:C1", "evaluated_records": 1},
                {"id": "t2", "sheet": "Sheet1", "data_range": "A1:C1", "evaluated_records": 0},
            ]
        }, provenance={"workbook_id": "wb_01", "file_path": "/workspace/test.xlsx"})

        prompt = "按 WBS 负责成本中心汇总的工数与金额"
        digest = build_verified_facts_digest([receipt], prompt)
        self.assertIn("FC34000工数", digest)
        self.assertIn("FC34000金额", digest)
        self.assertIn("总金额", digest)
        self.assertNotIn("FC32003", digest)

    def test_native_group_by_expression_aggregation(self):
        """14. 验证原生 aggregate 表达式支持 group_by 字段，自动按唯一取值展开多维聚合事实"""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "明细"
        ws["A1"] = "WBS编号"; ws["B1"] = "成本中心"; ws["C1"] = "金额(K)"
        test_data = [
            ("CK001", "FC34000", 10.0),
            ("CK002", "FC34000", 25.0),
            ("CK003", "FC45031", 50.0),
            ("CK004", "FC45031", 15.0),
            ("CK005", "FC49000", 100.0),
        ]
        for r, (code, cc, amt) in enumerate(test_data, 2):
            ws[f"A{r}"] = code; ws[f"B{r}"] = cc; ws[f"C{r}"] = amt
        test_path = Path("/tmp/test_native_groupby.xlsx")
        wb.save(test_path)

        plan = {
            "tables": [{
                "id": "t_data", "sheet": "明细", "header_row": 1, "data_range": "A2:C6"
            }],
            "calculations": [{
                "id": "amt_by_cc",
                "label": "金额",
                "expression": {
                    "aggregate": "sum",
                    "table": "t_data",
                    "column": "C",
                    "group_by": "B",
                    "unit": "amount"
                },
                "output_unit": "元"
            }]
        }
        res = calculate_spreadsheet_analysis(str(test_path), plan)
        self.assertTrue(res.is_success, res.error)
        facts = res.data["facts"]
        self.assertEqual(len(facts), 3)
        by_key = {f["group_key"]: Decimal(f["value"]) for f in facts}
        # FC34000: 35 K = 35000 元
        self.assertEqual(by_key["FC34000"], Decimal("35000"))
        # FC45031: 65 K = 65000 元
        self.assertEqual(by_key["FC45031"], Decimal("65000"))
        # FC49000: 100 K = 100000 元
        self.assertEqual(by_key["FC49000"], Decimal("100000"))
        test_path.unlink(missing_ok=True)

    def test_grouping_gate_blocks_zero_group_scalar_only_totals(self):
        """15. 验证执行门禁拦截仅输出了全表总计（金额总和+记录数）而缺少分组计算的计划"""
        from dsh_modules.tool_result import ToolResult
        from dsh_modules.analysis_task_state import analysis_task_state

        receipt = ToolResult.success(data={
            "kind": "spreadsheet_analysis",
            "evidence_id": "ev_005",
            "facts": [
                {"id": "total_amount", "label": "金额总和", "value": "1637419.9", "unit": "元"},
                {"id": "count_records", "label": "记录数", "value": "253", "unit": "number"},
            ],
            "coverage": [{
                "id": "filtered_data", "sheet": "工数实绩明细", "data_range": "A5:R278",
                "evaluated_records": 253, "coverage_complete": True,
                "filters": [{"column": "I", "op": "exclude_prefix", "prefixes": ["ex", "sy"]}]
            }]
        }, provenance={"workbook_id": "wb_01", "file_path": "/workspace/test.xlsx"})

        prompt = "用户指令：按 “WBS负责成本中心” 进行分组，对每一组内的 “金额” 字段进行 求和 计算"
        state = analysis_task_state([receipt], prompt)
        self.assertFalse(state["complete"])
        issue_codes = [i["code"] for i in state["issues"]]
        self.assertIn("grouping_incomplete", issue_codes)
        grouping_issue = next(i for i in state["issues"] if i["code"] == "grouping_incomplete")
        self.assertIn("WBS负责成本中心", grouping_issue["message"])

    def test_presentation_converges_to_filtered_contract_and_rejects_unfiltered_pollution(self):
        """16. P1: 验证呈现层在存在过滤要求时，严格收敛至过滤后契约回执，绝不混合未过滤全表事实造成口径污染 (W01)"""
        from dsh_modules.tool_result import ToolResult
        from dsh_modules.report_presentation import analysis_overview, resolve_conforming_facts

        # 模拟第1轮工具调用：未经过滤的原始全表试算
        receipt_unfiltered = ToolResult.success(data={
            "kind": "spreadsheet_analysis",
            "evidence_id": "ev_unfiltered",
            "tables": [{"id": "raw", "sheet": "明细", "header_row": 4, "data_range": "A5:R278", "filters": []}],
            "calculations": [
                {"id": "c_tot", "label": "总金额", "expression": {"aggregate": "sum", "table": "raw", "column": "P", "unit": "amount"}},
                {"id": "c_fc44004", "label": "FC44004金额", "expression": {"aggregate": "sum", "table": "raw", "column": "P", "unit": "amount"}}
            ],
            "facts": [
                {"id": "c_tot", "label": "总金额", "value": "1887.73", "unit": "K", "matched_records": 274},
                {"id": "c_fc44004", "label": "FC44004金额", "value": "328.36", "unit": "K", "matched_records": 21, "group_key": "FC44004"}
            ]
        }, provenance={"workbook_id": "wb_01", "file_path": "/workspace/test.xlsx"})

        # 模拟第2轮工具调用：模型纠偏后按排除前缀条件计算的准确事实
        receipt_filtered = ToolResult.success(data={
            "kind": "spreadsheet_analysis",
            "evidence_id": "ev_filtered",
            "tables": [{
                "id": "filtered", "sheet": "明细", "header_row": 4, "data_range": "A5:R278",
                "filters": [{"column": "I", "op": "exclude_prefix", "prefixes": ["ex", "sy", "sp", "ip", "ib", "sz"]}]
            }],
            "calculations": [
                {"id": "f_tot", "label": "总金额", "expression": {"aggregate": "sum", "table": "filtered", "column": "P", "unit": "amount"}},
                {"id": "f_fc44004", "label": "FC44004金额", "expression": {"aggregate": "sum", "table": "filtered", "column": "P", "unit": "amount"}}
            ],
            "facts": [
                {"id": "f_tot", "label": "总金额", "value": "1637.42", "unit": "K", "matched_records": 253},
                {"id": "f_fc44004", "label": "FC44004金额", "value": "231.13", "unit": "K", "matched_records": 16, "group_key": "FC44004"}
            ]
        }, provenance={"workbook_id": "wb_01", "file_path": "/workspace/test.xlsx"})

        prompt = "帮我统计下附件中按 WBS负责成本中心 统计 金额 的数据，不包含WBS编号 中ex，sy，sp，ip，ib，sz开头的数据"
        receipts = [receipt_unfiltered, receipt_filtered]

        # 1. 结构化事实收敛校验
        conforming = resolve_conforming_facts(receipts, prompt)
        values = {f["label"]: f["value"] for f in conforming}
        self.assertEqual(values.get("总金额"), "1637.42")
        self.assertEqual(values.get("FC44004金额"), "231.13")
        self.assertNotIn("1887.73", [f["value"] for f in conforming])
        self.assertNotIn("328.36", [f["value"] for f in conforming])

        # 2. 呈现层展示校验
        overview = analysis_overview(receipts, prompt)
        self.assertIn("1,637.42 K", overview)
        self.assertIn("231.13 K", overview)
        self.assertNotIn("1,887.73", overview)
        self.assertNotIn("328.36", overview)

    def test_group_validation_multi_metric_conservation_and_small_group_interception(self):
        """17. P1: 验证多指标量纲分区守恒、逐组覆盖矩阵及小额漏组拦截 (复审缺陷闭环)"""
        import tempfile
        from pathlib import Path
        import hashlib
        from dsh_modules.tool_result import ToolResult
        from dsh_modules.analysis_task_state import analysis_task_state

        prompt_multi = "按WBS编号分组，统计每组条数和金额(K)，附总金额及总条数。"

        with tempfile.TemporaryDirectory() as td:
            tf = Path(td) / "dummy.xlsx"
            tf.write_bytes(b"dummy workbook content")
            wb_hash = hashlib.sha256(tf.read_bytes()).hexdigest()
            prov = {"workbook_id": wb_hash, "file_path": str(tf)}

            # 场景 A: 多指标全量正确（金额+条数），必须守恒通过，绝不能跨量纲混加误拒 (6+3!=9)
            facts_valid = [
                {"id": "total_amount", "label": "总金额", "value": "6", "unit": "K"},
                {"id": "all_count", "label": "总条数", "value": "3", "unit": "number"},
                {"id": "c1", "label": "CK 分组金额", "group_key": "CK", "value": "2", "unit": "K"},
                {"id": "c2", "label": "IK 分组金额", "group_key": "IK", "value": "2", "unit": "K"},
                {"id": "c3", "label": "SK 分组金额", "group_key": "SK", "value": "2", "unit": "K"},
                {"id": "n1", "label": "CK 分组条数", "group_key": "CK", "value": "1", "unit": "number"},
                {"id": "n2", "label": "IK 分组条数", "group_key": "IK", "value": "1", "unit": "number"},
                {"id": "n3", "label": "SK 分组条数", "group_key": "SK", "value": "1", "unit": "number"},
            ]
            r_valid = ToolResult.success(data={"kind": "spreadsheet_analysis", "evidence_id": "ev_v", "facts": facts_valid},
                                         provenance=prov)
            state_valid = analysis_task_state([r_valid], prompt_multi)
            self.assertTrue(state_valid["complete"])
            self.assertEqual(state_valid["issues"], [])

            # 场景 B: 缺少各组条数（仅有全表总条数），逐组覆盖矩阵必须精准拦截
            facts_missing_counts = [
                {"id": "total_amount", "label": "总金额", "value": "6", "unit": "K"},
                {"id": "all_count", "label": "总条数", "value": "3", "unit": "number"},
                {"id": "c1", "label": "CK 分组金额", "group_key": "CK", "value": "2", "unit": "K"},
                {"id": "c2", "label": "IK 分组金额", "group_key": "IK", "value": "2", "unit": "K"},
                {"id": "c3", "label": "SK 分组金额", "group_key": "SK", "value": "2", "unit": "K"},
            ]
            r_missing_counts = ToolResult.success(data={"kind": "spreadsheet_analysis", "evidence_id": "ev_mc", "facts": facts_missing_counts},
                                                  provenance=prov)
            state_mc = analysis_task_state([r_missing_counts], prompt_multi)
            self.assertFalse(state_mc["complete"])
            self.assertTrue(any(i["code"] == "per_group_metrics_missing" for i in state_mc["issues"]))

            # 场景 C: 小额漏组拦截（IK=100, SK=100, 漏CK=0.1，总额200.1），精确守恒拦截
            prompt_small = "按WBS编号分组汇总金额(K)，列出各组和总计。"
            facts_small = [
                {"id": "total_amount", "label": "总金额", "value": "200.1", "unit": "K"},
                {"id": "IK_sum", "label": "IK金额", "value": "100", "unit": "K"},
                {"id": "SK_sum", "label": "SK金额", "value": "100", "unit": "K"},
            ]
            r_small = ToolResult.success(data={"kind": "spreadsheet_analysis", "evidence_id": "ev_s", "facts": facts_small},
                                         provenance=prov)
            state_small = analysis_task_state([r_small], prompt_small)
            self.assertFalse(state_small["complete"])
            self.assertTrue(any(i["code"] == "grouping_coverage_incomplete" for i in state_small["issues"]))

    def test_group_matrix_partial_coverage_and_scope_closure(self):
        """18. P1: 验证逐组覆盖矩阵（部分组缺少条数拦截）与筛选范围闭包（历史未过滤试算隔离）"""
        import tempfile
        from pathlib import Path
        import hashlib
        from dsh_modules.tool_result import ToolResult
        from dsh_modules.analysis_task_state import analysis_task_state

        prompt = "按WBS编号分组，统计每组条数和金额(K)，附总金额。"

        with tempfile.TemporaryDirectory() as td:
            tf = Path(td) / "dummy.xlsx"
            tf.write_bytes(b"dummy workbook content")
            wb_hash = hashlib.sha256(tf.read_bytes()).hexdigest()
            prov = {"workbook_id": wb_hash, "file_path": str(tf)}

            # 1. 验证逐组覆盖矩阵：仅 IK 算条数，CK 与 SK 缺少条数，必须拦截并指出缺失组
            facts_partial = [
                {"id": "total_amount", "label": "总金额", "value": "6", "unit": "K"},
                {"id": "c1", "label": "CK 分组金额", "group_key": "CK", "value": "2", "unit": "K"},
                {"id": "c2", "label": "IK 分组金额", "group_key": "IK", "value": "2", "unit": "K"},
                {"id": "c3", "label": "SK 分组金额", "group_key": "SK", "value": "2", "unit": "K"},
                {"id": "n2", "label": "IK 分组条数", "group_key": "IK", "value": "1", "unit": "number"},
            ]
            r_partial = ToolResult.success(data={"kind": "spreadsheet_analysis", "evidence_id": "ev_p", "facts": facts_partial},
                                           provenance=prov)
            state_partial = analysis_task_state([r_partial], prompt)
            self.assertFalse(state_partial["complete"])
            missing_issue = next((i for i in state_partial["issues"] if i.get("code") == "per_group_metrics_missing"), None)
            self.assertIsNotNone(missing_issue)
            self.assertIn("CK", missing_issue.get("missing_groups", []))
            self.assertIn("SK", missing_issue.get("missing_groups", []))

            # 2. 验证筛选范围闭包：未过滤试算 (总额60, AA30/BB30) 与排除后正式计算 (总额50, AA20/BB30)
            # 门禁必须通过权威契约闭包隔离未过滤试算，绝不能混加出 110 K 报出虚假差额
            prompt_filter = "排除EX前缀，按WBS编号分组统计金额(K)，附总金额"
            r_unfiltered = ToolResult.success(
                data={
                    "kind": "spreadsheet_analysis",
                    "evidence_id": "ev_unfiltered",
                    "facts": [
                        {"id": "total_60", "label": "总金额", "value": "60", "unit": "K"},
                        {"id": "grp_AA_30", "label": "AA 分组金额", "group_key": "AA", "value": "30", "unit": "K"},
                        {"id": "grp_BB_30", "label": "BB 分组金额", "group_key": "BB", "value": "30", "unit": "K"},
                    ]
                },
                provenance=prov
            )
            r_filtered = ToolResult.success(
                data={
                    "kind": "spreadsheet_analysis",
                    "evidence_id": "ev_filtered",
                    "facts": [
                        {"id": "total_50", "label": "总金额", "value": "50", "unit": "K"},
                        {"id": "grp_AA_20", "label": "AA 分组金额", "group_key": "AA", "value": "20", "unit": "K"},
                        {"id": "grp_BB_30", "label": "BB 分组金额", "group_key": "BB", "value": "30", "unit": "K"},
                    ]
                },
                provenance=prov
            )
            state_combined = analysis_task_state([r_unfiltered, r_filtered], prompt_filter)
            # Neither mock contains an executed filter/table binding. Do not trust call order.
            self.assertFalse(state_combined["complete"])
            self.assertTrue(state_combined["issues"])

    def test_delivery_reconciliation_total_and_exclusion(self):
        """19. 验证交付层事实对齐：自动校准模型心算错误总金额与排除条数"""
        import tempfile
        from pathlib import Path
        import openpyxl
        from dsh_modules.delivery_reconciliation import (
            extract_markdown_table_totals,
            reconcile_delivery_presentation,
        )

        sample_reply = (
            "WBS成本中心汇总（金额单位：千元/K）\n\n"
            "| 序号 | WBS负责成本中心 | 金额(K) | 条数 |\n"
            "|:---:|:---|:---:|:---:|\n"
            "| 1 | FC45031 | 513.091 | 79 |\n"
            "| 2 | FC46005 | 260.648 | 40 |\n"
            "| 3 | FC44004 | 231.129 | 17 |\n"
            "| 4 | FC44003 | 188.777 | 21 |\n"
            "| 5 | FC44006 | 159.881 | 15 |\n"
            "| 6 | FC34022 | 152.155 | 61 |\n"
            "| 7 | FC48005 | 67.035 | 6 |\n"
            "| 8 | FC34001 | 53.329 | 5 |\n"
            "| 9 | FC49000 | 10.944 | 8 |\n"
            "| 10 | FC34000 | 0.431 | 1 |\n\n"
            "- **覆盖成本中心数**：10 个\n"
            "- **总条数**：253 条\n"
            "- **总金额(K)**：1,436.420（Excel 已配置 `SUM` 公式自动重算）\n\n"
            "### 处理说明\n"
            "- **排除规则**：已过滤 WBS 编号以 ex/sy/sp/ip/ib/sz 开头的明细（不区分大小写），共排除 56 条。\n"
        )

        # 1. 验证从 Markdown 表格精确抽取列闭包
        totals = extract_markdown_table_totals(sample_reply)
        self.assertIsNotNone(totals)
        self.assertEqual(totals["count"], 253)
        self.assertEqual(totals["amount"], Decimal("1637.420"))

        with tempfile.TemporaryDirectory() as td:
            ws_dir = Path(td)
            # 创建模拟生成的 Excel 交付物
            gen_wb = openpyxl.Workbook()
            gen_ws = gen_wb.active
            gen_ws.title = "汇总"
            gen_ws["A1"] = "WBS负责成本中心"
            gen_ws["B1"] = "明细条数"
            gen_ws["C1"] = "汇总金额(K)"
            gen_ws["A15"] = "总计 (Total)"
            gen_ws["B15"] = 253
            gen_ws["C15"] = 1637.41990
            gen_wb.save(ws_dir / "WBS成本中心汇总.xlsx")

            # 创建模拟受保护的原始输入附件（共 274 条明细，包含表尾说明）
            src_wb = openpyxl.Workbook()
            src_ws = src_wb.active
            src_ws.title = "工数实绩明细"
            # 前3行说明，第4行表头
            src_ws.cell(row=4, column=1, value="WBS编号")
            src_ws.cell(row=4, column=2, value="金额(K)")
            # 274 条有效数据行 (row 5 to 278)
            for r in range(5, 279):
                src_ws.cell(row=r, column=1, value=f"WBS-{r}")
                src_ws.cell(row=r, column=2, value=1.0)
            # 表尾查询条件说明 (row 279 to 283, 仅1列非空)
            for r in range(279, 284):
                src_ws.cell(row=r, column=1, value="查询说明")
            src_file = ws_dir / "工数实绩明细_20261020134124.xlsx"
            src_wb.save(src_file)

            # 设置环境变量声明受保护文件
            import os
            old_env = os.environ.get("DSH_SESSION_ATTACHMENTS")
            os.environ["DSH_SESSION_ATTACHMENTS"] = "工数实绩明细_20261020134124.xlsx"
            try:
                fixed_reply = reconcile_delivery_presentation(
                    sample_reply,
                    str(ws_dir),
                    prompt="帮我统计下附件中按 WBS负责成本中心 统计 金额 的数据，不包含WBS编号 中ex，sy，sp，ip，ib，sz开头的数据，导出Excel"
                )
                # 必须将 1,436.420 自动校准为 1,637.420
                self.assertIn("1,637.420", fixed_reply)
                self.assertNotIn("1,436.420", fixed_reply)
                # 必须将 56 条自动校准为 274 - 253 = 21 条
                self.assertIn("共排除 21 条", fixed_reply)
                self.assertNotIn("共排除 56 条", fixed_reply)
            finally:
                if old_env is not None:
                    os.environ["DSH_SESSION_ATTACHMENTS"] = old_env
                else:
                    os.environ.pop("DSH_SESSION_ATTACHMENTS", None)


if __name__ == "__main__":
    unittest.main()
