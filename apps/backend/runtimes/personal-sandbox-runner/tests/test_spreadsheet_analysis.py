"""Source-driven aggregation, coverage, failed receipts and final acceptance."""

import datetime as dt
import io
import json
import os
import sys
import tempfile
import unittest
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch
from contextlib import redirect_stdout
import openpyxl

sys.path.insert(0,str(Path(__file__).resolve().parent.parent/"src"))
from dsh_modules.spreadsheet_analysis import calculate_spreadsheet_analysis
from dsh_modules.spreadsheet_analysis_evidence import evidence_error, render_analysis_response, requires_spreadsheet_analysis
from dsh_modules.spreadsheet_context import build_spreadsheet_source_context
from dsh_modules.sampling_guard import sampling_feedback
from dsh_modules.agent_loop import run_agent_loop, _finalize_agent_text
from dsh_modules.runtime_policy import RuntimePolicy
from dsh_modules.telemetry import TelemetryStats
from dsh_modules.tool_result import ToolResult
from dsh_modules.runner import select_active_tools
from dsh_modules.skill_router import SkillRouter


def aggregate(table,column,unit="amount",method="sum"):
    return {"aggregate":method,"table":table,"column":column,"unit":unit}


class SpreadsheetAnalysisTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.path=Path(self.tmp.name)/"source.xlsx"
        w=openpyxl.Workbook();s=w.active;s.title="任意业务名称"
        s.append(["单位","千元"]);s.append(["业务记录"]);s.append([]);s.append([])
        s.append(["凭证","日期","金额","审批状态","部门"])
        for i in range(80):s.append([f"K{i}",dt.date(2026,1,1),100+i,"已审批","业务部"])
        s['C45']=98000;s['B60']=dt.date(2027,1,3);s['D75']='待审批';s['E70']=None
        s['A9']='K1'
        m=w.create_sheet("月度来源改名");m.append(["单位","千元"]);m.append(["月份","收入","毛利","净利润"])
        for i in range(12):m.append([f"{i+1}月",100+i*100,40+i*30,10+i*5])
        b=w.create_sheet("任意余额表");b.append(["单位","千元"]);b.append(["项目","金额"])
        b.append(["资产",100]);b.append(["负债及权益",110])
        w.save(self.path);w.close()
        self.table={"id":"tx","sheet":"任意业务名称","header_row":5,"data_range":"A6:E85"}

    def tearDown(self):self.tmp.cleanup()

    def compute(self,**plan):return calculate_spreadsheet_analysis(self.path,plan,domain="finance")

    def rule(self,id,col,kind,**kwargs):
        return {"id":id,"label":id,"table":"tx","column":col,"kind":kind,
                "assumptions":"按来源声明的报告期和待核查审批规则检查",**kwargs}

    def response(self,*receipts,hypotheses=None):
        return json.dumps({"evidence_ids":[r.data['evidence_id'] for r in receipts],"hypotheses":hypotheses or []})

    def test_tail_anomalies_full_records_and_profiles(self):
        r=self.compute(tables=[self.table],rules=[self.rule('dates','B','date_range',min='2026-01-01',max='2026-12-31'),
            self.rule('approvals','D','allowed_values',allowed=['已审批']),self.rule('amount','C','number_range',min=0,max=1000),
            self.rule('unique_ids','A','unique'),self.rule('department','E','not_empty')])
        self.assertTrue(r.is_success,r.error)
        rules={v['id']:v for v in r.data['rules']}
        self.assertEqual(r.data['coverage'][0]['evaluated_records'],80)
        for name,row in [('dates',60),('approvals',75),('amount',45),('department',70)]:
            self.assertEqual(rules[name]['examples'][0]['row'],row)
        amount=next(p for p in r.data['profiles'] if p['column']=='C')
        self.assertEqual(amount['numeric']['max'],98000)
        approvals=next(p for p in r.data['profiles'] if p['column']=='D')
        self.assertEqual(approvals['value_counts'],{'已审批':79,'待审批':1})
        self.assertEqual(approvals['value_examples']['待审批'],['任意业务名称!D75'])
        text=render_analysis_response(self.response(r),[r])
        self.assertIn('业务行80条',text);self.assertIn('98000',text);self.assertIn('2027-01-03',text)
        self.assertIn('待审批',text)
        self.assertIn('任意业务名称!E70',text)

    def test_weighted_ratio_and_source_unit_conversion(self):
        table={'id':'m','sheet':'月度来源改名','header_row':2,'data_range':'A3:D14',
               'filters':[{'column':'A','values':['7月','8月','9月','10月','11月','12月']}]}
        expr={'op':'divide','args':[aggregate('m','C'),aggregate('m','B')]}
        r=self.compute(tables=[table],calculations=[{'id':'rev','label':'下半年收入','expression':aggregate('m','B'),'output_unit':'亿元'},
             {'id':'ratio','label':'加权毛利率','expression':expr}])
        self.assertTrue(r.is_success,r.error)
        self.assertEqual(r.data['coverage'][0]['evaluated_records'],6)
        self.assertEqual(Decimal(r.data['facts'][0]['value']),Decimal('0.057'))
        expected=sum(Decimal(40+i*30) for i in range(6,12))/sum(Decimal(100+i*100) for i in range(6,12))*100
        self.assertEqual(Decimal(r.data['facts'][1]['value']),expected)
        self.assertEqual(r.data['facts'][1]['unit'],'%')

    def test_verified_ratios_are_displayed_without_unexecuted_delta(self):
        first={'id':'h1','sheet':'月度来源改名','header_row':2,'data_range':'A3:D8'}
        second={'id':'h2','sheet':'月度来源改名','header_row':2,'data_range':'A9:D14'}
        calculations=[{'id':t['id']+'_ratio','label':caption+'毛利率',
            'expression':{'op':'divide','args':[aggregate(t['id'],'C'),aggregate(t['id'],'B')]}}
            for t,caption in [(first,'上半年'),(second,'下半年')]]
        r=self.compute(tables=[first,second],calculations=calculations)
        self.assertTrue(r.is_success,r.error)
        text=render_analysis_response(self.response(r),[r],'解释毛利率变化')
        self.assertNotIn('下降1.8045个百分点',text)
        self.assertIn('上半年毛利率',text)
        self.assertIn('下半年毛利率',text)

    def test_equation_residual_and_assumptions_are_reported(self):
        cell=lambda addr:{'cell':addr,'sheet':'任意余额表','unit':'amount'}
        r=self.compute(checks=[{'id':'balance','label':'资产与负债权益','left':cell('B3'),'right':cell('B4'),
                              'tolerance':0.01,'assumptions':'同一主体同一日期，不存在范围调整'}])
        self.assertTrue(r.is_success,r.error)
        self.assertFalse(r.data['checks'][0]['passed']);self.assertEqual(r.data['checks'][0]['residual'],'-10')
        text=render_analysis_response(self.response(r),[r]);self.assertIn('不一致',text);self.assertIn('原因',text)

    def receipt(self):
        return self.compute(tables=[self.table],calculations=[{'id':'n','label':'业务记录数','expression':aggregate('tx','A','number','count')}])

    def test_failed_partial_old_source_and_omitted_evidence_fail_closed(self):
        r=self.receipt();resp=self.response(r)
        self.assertIsNone(evidence_error(resp,[r]))
        self.assertIsNotNone(evidence_error(resp,[ToolResult.error('failed','工具失败')]))
        self.assertIsNotNone(evidence_error(resp,[ToolResult.partial(data=r.data,provenance=r.provenance)]))
        self.assertIsNotNone(evidence_error(self.response(r),[r,self.compute(tables=[self.table],calculations=[{'id':'min','label':'最小金额','expression':aggregate('tx','C',method='min')}])]))
        self.path.write_bytes(b'changed')
        self.assertIsNotNone(evidence_error(resp,[r]))

    def test_plain_numerical_answer_and_invalid_envelope_are_rejected(self):
        r=self.receipt()
        for response in ['经Python全表扫描，记录82条。','{"evidence_ids":[{}],"hypotheses":[]}',
                         self.response(r,hypotheses=['净利润9961.35千元'])]:
            self.assertIsNotNone(evidence_error(response,[r]))

    def test_successful_results_do_not_hide_a_failed_followup_plan(self):
        r=self.receipt()
        failed=ToolResult.error('invalid_analysis_plan','后续检查缺少来源',
            data={'requested_plan':{'checks':[{'id':'extra','label':'补充检查'}]}})
        text=render_analysis_response(self.response(r),[r,failed])
        self.assertIn('业务行80条',text)
        self.assertIn('后续计划未执行成功',text)
        self.assertIn('补充检查',text)

    def test_unknown_operation_missing_range_and_zero_denominator_are_errors(self):
        for plan in [{},{'tables':[dict(self.table,header_row=10)],'calculations':[{'id':'x','label':'x','expression':aggregate('tx','C')}]},
          {'calculations':[{'id':'x','label':'x','expression':{'op':'divide','args':[{'constant':1,'unit':'number'},{'constant':0,'unit':'number'}]}}]},
          {'calculations':[{'id':'x','label':'x','expression':{'eval':'__import__("os")'}}]}]:
            self.assertTrue(calculate_spreadsheet_analysis(self.path,plan).is_error)

    def test_json_encoded_array_is_decoded_without_changing_business_scope(self):
        from dsh_modules.spreadsheet_analysis import execute_spreadsheet_analysis
        with patch('dsh_modules.spreadsheet_analysis.resolve_sandboxed_path',return_value=(self.path,None)):
            output=execute_spreadsheet_analysis({'file_path':str(self.path),'tables':json.dumps([self.table]),
                'calculations':json.dumps([{'id':'n','label':'记录数','expression':aggregate('tx','A','number','count')}])})
        self.assertTrue(output.tool_result.is_success,output)
        self.assertEqual(output.tool_result.data['coverage'][0]['evaluated_records'],80)

    def test_native_tool_cannot_select_a_different_session_file(self):
        from dsh_modules.spreadsheet_analysis import execute_spreadsheet_analysis
        with patch.dict(os.environ,{'DSH_SESSION_ATTACHMENTS':'different.xlsx'}),\
             patch('dsh_modules.spreadsheet_analysis.resolve_sandboxed_path',return_value=(self.path,None)),\
             patch('dsh_modules.config.WORKSPACE_DIR',self.tmp.name):
            output=execute_spreadsheet_analysis({'file_path':str(self.path),'tables':[self.table],'profile_tables':True})
        self.assertEqual(output.tool_result.error_code,'source_outside_session_scope')

    def test_named_expression_references_and_display_unit_compile_without_guessing(self):
        r=self.compute(tables=[self.table],calculations=[{'id':'total','label':'总金额','expression':aggregate('tx','C')},
            {'id':'twice','label':'金额两倍','expression':{'op':'add','args':[{'id':'total'},{'ref':'total'}],'output_unit':'万元'}}])
        self.assertTrue(r.is_success,r.error)
        self.assertEqual(Decimal(r.data['facts'][1]['value']),Decimal(r.data['facts'][0]['value'])/5)
        self.assertTrue(r.data['normalizations'])
        cycle=self.compute(calculations=[{'id':'a','label':'循环','expression':{'ref':'a'}}])
        self.assertTrue(cycle.is_error)
        self.assertIn('循环',cycle.error_message)

    def test_explicit_cell_arithmetic_and_table_aliases_compile_safely(self):
        cell={'sheet':'任意余额表','cell':'B4-B3','label':'元数据'}
        r=self.compute(calculations=[{'id':'delta','label':'余额差额','expression':cell}])
        self.assertTrue(r.is_success,r.error)
        self.assertEqual(r.data['facts'][0]['value'],'10')
        r=self.compute(tables=[self.table],rules=[{**self.rule('min','C','number_range',min='0'),'table':'任意业务名称'}])
        self.assertTrue(r.is_success,r.error)
        self.assertEqual(r.data['plan']['rules'][0]['table'],'tx')

    def test_explicit_table_row_column_address_resolves_only_actual_headers(self):
        r=self.compute(tables=[self.table],calculations=[{'id':'maxrow','label':'指定记录金额',
            'expression':{'table_id':'tx','column':'金额','row':45}}])
        self.assertTrue(r.is_success,r.error)
        self.assertEqual(r.data['facts'][0]['value'],'98000')
        r=self.compute(tables=[self.table],calculations=[{'id':'bad','label':'越界',
            'expression':{'table_id':'tx','column':'金额','row':86}}])
        self.assertTrue(r.is_error)
        filtered={**self.table,'filters':[{'column':'D','values':['待审批']}]}
        r=self.compute(tables=[filtered],calculations=[{'id':'wrongscope','label':'指定记录金额',
            'expression':{'table_id':'tx','column':'C','row':45}}])
        self.assertTrue(r.is_error)

    def test_composite_caption_cannot_relabel_one_expense_as_three(self):
        w=openpyxl.load_workbook(self.path);w['月度来源改名']['D2']='销售费用';w.save(self.path);w.close()
        table={'id':'m','sheet':'月度来源改名','header_row':2,'data_range':'A3:D14'}
        r=self.compute(tables=[table],calculations=[{'id':'bad','label':'三费总额','expression':aggregate('m','D')}])
        self.assertTrue(r.is_error)
        self.assertIn('实际引用1',r.error_message)

    def test_missing_rule_fields_are_reported_together_with_precise_paths(self):
        r=self.compute(tables=[self.table],rules=[self.rule('dates','B','date_range'),self.rule('status','D','allowed_values')])
        self.assertTrue(r.is_error)
        self.assertIn('rules.dates',r.error_message)
        self.assertIn('rules.status',r.error_message)

    def test_constant_business_fact_is_rejected(self):
        r=self.compute(calculations=[{'id':'fake','label':'收入','expression':{'constant':10000,'unit':'number'}}])
        self.assertTrue(r.is_error)

    def test_exact_header_boundary_and_unambiguous_header_name_are_normalized(self):
        table={**self.table,'data_range':'A5:E85','filters':[{'column':'审批状态','values':['待审批']}]}
        r=self.compute(tables=[table],calculations=[{'id':'amount','label':'待审批金额',
            'expression':aggregate('tx','金额')}])
        self.assertTrue(r.is_success,r.error)
        self.assertEqual(r.data['coverage'][0]['excluded_header_rows'],[5])
        self.assertEqual(r.data['coverage'][0]['input_records'],80)
        self.assertEqual(r.data['coverage'][0]['evaluated_records'],1)

    def test_unicode_identifiers_and_profile_only_quality_capability(self):
        from dsh_modules.tools import execute_tool
        with patch('dsh_modules.spreadsheet_analysis.resolve_sandboxed_path',return_value=(self.path,None)):
            r=execute_tool('validate_spreadsheet_rows',{'file_path':str(self.path),'tables':[{**self.table,'id':'交易明细'}]}).tool_result
        self.assertTrue(r.is_success,r.error)
        self.assertEqual(r.data['coverage'][0]['evaluated_records'],80)
        self.assertIn('待审批',render_analysis_response(self.response(r),[r]))

    def test_change_request_requires_distinct_source_periods(self):
        r=self.receipt()
        self.assertIsNotNone(evidence_error(self.response(r),[r],'解释变化趋势'))

    def test_profitability_concept_requires_ratios_not_only_amounts(self):
        from dsh_modules.analysis_contract import build_analysis_contract
        contract=build_analysis_contract('查看盈利质量')
        self.assertEqual(contract.concepts,['profitability'])
        self.assertIn('净利率',contract.get_metric_names())
        r=self.receipt()
        self.assertIn('毛利率',evidence_error(self.response(r),[r],'解释盈利能力变化'))

    def test_ratio_caption_cannot_use_wrong_numerator_or_average_percentages(self):
        table={'id':'m','sheet':'月度来源改名','header_row':2,'data_range':'A3:D14'}
        r=self.compute(tables=[table],calculations=[{'id':'bad','label':'净利率','expression':
            {'op':'divide','args':[aggregate('m','C'),aggregate('m','B')]}}])
        self.assertTrue(r.is_error)
        self.assertIn('分子/分母',r.error_message)

    def test_one_sided_bounds_preserve_only_explicit_constraints(self):
        r=self.compute(tables=[self.table],rules=[self.rule('lower','C','number_range',min=0),
            self.rule('upper_date','B','date_range',max='2026-12-31')])
        self.assertTrue(r.is_success,r.error)
        self.assertEqual(r.data['rules'][0]['criteria'],{'min':0})
        self.assertEqual(r.data['rules'][1]['examples'][0]['row'],60)

    def test_period_caption_cannot_publish_annual_data_as_half_year(self):
        table={'id':'m','sheet':'月度来源改名','header_row':2,'data_range':'A3:D14'}
        plan={'tables':[table],'calculations':[{'id':'bad','label':'下半年收入','expression':aggregate('m','B')}]}
        r=self.compute(**plan)
        self.assertTrue(r.is_error)
        self.assertIn('期间与来源不一致',r.error_message)
        self.assertEqual(r.data['requested_plan'],plan)
        self.assertIsNotNone(evidence_error('{"evidence_ids":[],"hypotheses":[]}',[r]))
        for caption in ['H1营业收入','Q3营业收入']:
            r=self.compute(tables=[table],calculations=[{'id':'bad','label':caption,'expression':aggregate('m','B')}])
            self.assertTrue(r.is_error)

    def test_direct_metric_caption_must_match_actual_column(self):
        table={'id':'m','sheet':'月度来源改名','header_row':2,'data_range':'A3:D14'}
        r=self.compute(tables=[table],calculations=[{'id':'bad','label':'净利润','expression':aggregate('m','B')}])
        self.assertTrue(r.is_error)
        self.assertIn('来源字段不一致',r.error_message)

    def test_approval_enum_pass_is_not_business_compliance(self):
        r=self.compute(tables=[self.table],rules=[self.rule('审批格式','D','allowed_values',allowed=['已审批','待审批'])])
        self.assertTrue(r.data['rules'][0]['passed'])
        text=render_analysis_response(self.response(r),[r])
        self.assertIn('业务适用性尚未独立验证',text)
        self.assertIn('任意业务名称!D75',text)

    def test_additive_source_identity_rejects_component_as_total(self):
        from dsh_modules.spreadsheet_checks import validate_source_identity
        w=openpyxl.Workbook();w.active.title='自由表名'
        w.active['D9']='=D7+D8'
        cell=lambda c:{'cell':c,'sheet':'自由表名'}
        check={'left':cell('D7'),'right':cell('D9')}
        with self.assertRaisesRegex(ValueError,'分项不能'):
            validate_source_identity(check,w)
        validate_source_identity({**check,'relation_type':'component_comparison'},w)
        validate_source_identity({'left':{'op':'add','args':[cell('D7'),cell('D8')]},'right':cell('D9')},w)
        w.close()

    def test_explicit_private_stream_never_flushes_unverified_draft(self):
        from dsh_modules.llm import call_model_proxy
        captured={}
        class Response:
            def __enter__(self):
                return iter([b'data: {"choices":[{"delta":{"content":"unverified draft"}}]}\n',b'data: [DONE]\n'])
            def __exit__(self,*args):pass
        def open_request(request,**kwargs):
            captured.update(json.loads(request.data));return Response()
        output=io.StringIO()
        with redirect_stdout(output),patch('urllib.request.urlopen',side_effect=open_request):
            r=call_model_proxy([{'role':'user','content':'test'}],tools=[{'type':'function','function':{'name':'aggregate_spreadsheet'}}],
                               stream_deltas=False,tool_choice='required')
        self.assertEqual(r['content'],'unverified draft')
        self.assertNotIn('DSH_DELTA:',output.getvalue())
        self.assertEqual(captured['tool_choice'],'required')

    def test_dimensioned_optional_op_unit_and_nary_add(self):
        expr={'op':'add','unit':'amount','args':[aggregate('tx','C'),aggregate('tx','C'),aggregate('tx','C')]}
        r=self.compute(tables=[self.table],calculations=[{'id':'three','label':'三项相同金额之和','expression':expr}])
        self.assertTrue(r.is_success,r.error)
        self.assertEqual(Decimal(r.data['facts'][0]['value']),sum(Decimal(100+i) for i in range(80))*3+(Decimal(98000)-Decimal(139))*3)

    def test_overview_preserves_positions_and_does_not_claim_full_coverage(self):
        overview=build_spreadsheet_source_context(self.path)
        self.assertIn('SCHEMA_SAMPLE',overview);self.assertIn('"coverage_complete": false',overview)
        self.assertIn('A5',overview)
        self.assertNotIn('98000',overview)

    def test_failed_attempt_history_cannot_suppress_final_coverage_warning(self):
        t=TelemetryStats();t.record_tool_call('bash',{'cmd':'python3 read_excel.py'},'error')
        text='经核对，全表不存在重复。'+('已完成所有检查。'*20)
        final,_=_finalize_agent_text(text,[{'role':'user','content':'检验数据\n工作表覆盖度提醒: 抽样/截断'}],
            'test',RuntimePolicy(),None,False,['bash:python3 read_excel.py'],False,t,last_user_prompt='检验数据')
        self.assertIn('数据覆盖度与真实性提示',final)

    def test_unknown_wording_selects_readonly_native_tools_and_requires_plan(self):
        for q in ['解释下半年盈利能力变化。','识别金额、日期和审批状态异常。']:
            skill=SkillRouter.route(q,[],files=['any.xlsx'])
            self.assertTrue(requires_spreadsheet_analysis(q,skill))
            names={t['function']['name'] for t in select_active_tools(q,skill,False,False)}
            self.assertIn('aggregate_spreadsheet',names);self.assertIn('validate_spreadsheet_rows',names)
            self.assertNotIn('bash',names);self.assertNotIn('patch_file',names)
        skill=SkillRouter.route('列出工作簿中的Sheet，并说明用途。',[],files=['any.xlsx'])
        self.assertFalse(requires_spreadsheet_analysis('列出工作簿中的Sheet，并说明用途。',skill))

    def test_advisory_consulting_bypasses_spreadsheet_analysis_and_restricts_tools(self):
        advisory_queries = [
            '提出3条可执行的财务管理建议。',
            '结合报表数据，给出经营管理建议。',
            '如何提升毛利率？',
            '针对销售费用提出优化建议',
            '分析公司的经营风险并给出对策',
        ]
        for q in advisory_queries:
            skill = SkillRouter.route(q, [], files=['any.xlsx'])
            self.assertFalse(requires_spreadsheet_analysis(q, skill), f"Should not require spreadsheet analysis for: {q}")
            names = {t['function']['name'] for t in select_active_tools(q, skill, False, False)}
            self.assertIn('read_file', names)
            self.assertNotIn('aggregate_spreadsheet', names)
            self.assertNotIn('bash', names)
            self.assertNotIn('patch_file', names)

        # 显式计算带有建议时，仍应先执行计算
        hybrid_q = '计算2026年毛利率并提出3条改善建议'
        skill = SkillRouter.route(hybrid_q, [], files=['any.xlsx'])
        self.assertTrue(requires_spreadsheet_analysis(hybrid_q, skill))

    def test_loop_rejects_unexecuted_claim_even_when_budget_exhausted(self):
        with patch('dsh_modules.agent_loop.call_model_proxy',return_value={'content':'通过Python检查，全表全部正常。','finish_reason':'stop'}):
            result=run_agent_loop([{'role':'system','content':'test'},{'role':'user','content':'识别异常'}],'test',RuntimePolicy(),
                1,tools=[],spreadsheet_analysis_required=True)
        self.assertIn('尚未通过',result.final_text);self.assertNotIn('全表全部正常',result.final_text)
        self.assertEqual(result.telemetry.guard_decisions[-1]['action'],'reject')

    def test_loop_tool_then_structured_final_uses_real_receipt(self):
        params={'file_path':str(self.path),'tables':[self.table],
            'calculations':[{'id':'n','label':'业务记录数','expression':aggregate('tx','A','number','count')}]}
        receipt=calculate_spreadsheet_analysis(self.path,{k:v for k,v in params.items() if k!='file_path'})
        replies=[{'content':'','finish_reason':'tool_calls','tool_calls':[{'id':'call_a','type':'function',
            'function':{'name':'analyze_spreadsheet','arguments':json.dumps(params)}}]},
            {'content':self.response(receipt),'finish_reason':'stop'}]
        from dsh_modules.tools import get_sandbox_tools
        with patch('dsh_modules.agent_loop.call_model_proxy',side_effect=replies),patch('dsh_modules.spreadsheet_analysis.resolve_sandboxed_path',return_value=(self.path,None)):
            result=run_agent_loop([{'role':'system','content':'test'},{'role':'user','content':'统计记录'}],'test',RuntimePolicy(),
                3,tools=get_sandbox_tools({'analyze_spreadsheet'}),spreadsheet_analysis_required=True)
        self.assertIn('业务行80条',result.final_text);self.assertEqual(result.telemetry.guard_decisions[-1]['action'],'accept')

    def test_loop_allows_advisory_consulting_response_without_spreadsheet_rejection(self):
        advisory_reply = (
            "基于2026年财务经营数据，提出以下3条可执行的财务管理建议：\n\n"
            "1. **加强销售费用管控与投产比考核**：针对下半年销售费用上升，设定渠道ROI预警线；\n"
            "2. **建立负数金额交易复核机制**：对异常交易实行双人审批制；\n"
            "3. **推进预算动态管控与滚动预测**：确保经营目标按月追踪闭环。"
        )
        captured_messages = []
        def mock_proxy(messages, *args, **kwargs):
            captured_messages.extend(messages)
            return {'content': advisory_reply, 'finish_reason': 'stop'}

        skill = SkillRouter.route('提出3条可执行的财务管理建议。', [], files=['any.xlsx'])
        req_analysis = requires_spreadsheet_analysis('提出3条可执行的财务管理建议。', skill)
        self.assertFalse(req_analysis)

        with patch('dsh_modules.agent_loop.call_model_proxy', side_effect=mock_proxy):
            result = run_agent_loop(
                [{'role': 'system', 'content': 'You are an assistant.'},
                 {'role': 'user', 'content': '【当前会话有效附件清单】: any.xlsx\n用户指令：提出3条可执行的财务管理建议。'}],
                'test',
                RuntimePolicy(),
                3,
                tools=[],
                spreadsheet_analysis_required=req_analysis
            )

        self.assertIn('加强销售费用管控', result.final_text)
        self.assertIn('建立负数金额交易复核机制', result.final_text)
        self.assertNotIn('尚未通过执行证据核验', result.final_text)
        self.assertNotIn('缺少分析执行证据', result.final_text)
        # 验证注入了咨询规范指南
        sys_msgs = [m['content'] for m in captured_messages if m.get('role') == 'system']
        self.assertTrue(any('【财务与经营管理咨询规范】' in s for s in sys_msgs))

    def test_task_contract_keeps_tools_available_and_allows_clarification_after_insufficient_success(self):
        from dsh_modules.tools import get_sandbox_tools
        params={'file_path':str(self.path),'tables':[self.table],
                'calculations':[{'id':'n','label':'记录数','expression':aggregate('tx','A','number','count')}]}
        tool_reply={'content':'','finish_reason':'tool_calls','tool_calls':[{'id':'c','type':'function',
            'function':{'name':'analyze_spreadsheet','arguments':json.dumps(params)}}]}
        calls=[]
        def model(*args,**kwargs):
            calls.append(kwargs)
            return tool_reply if len(calls)==1 else {'content':'已完成','finish_reason':'stop'}
        with patch('dsh_modules.agent_loop.call_model_proxy',side_effect=model),\
             patch('dsh_modules.spreadsheet_analysis.resolve_sandboxed_path',return_value=(self.path,None)):
            result=run_agent_loop([{'role':'system','content':'test'},{'role':'user','content':'汇总营业收入'}],
                'test',RuntimePolicy(),2,tools=get_sandbox_tools({'analyze_spreadsheet'}),spreadsheet_analysis_required=True)
        self.assertIsNone(calls[1]['tool_choice'])
        self.assertTrue(calls[1]['tools'])
        self.assertIn('尚未通过',result.final_text)

    def test_browse_and_summary_intent_bypasses_spreadsheet_analysis(self):
        browse_queries = [
            '查看和归纳内容',
            '查看表格内容并进行归纳',
            '请归纳一下表格的主要内容',
            '总结这个表格讲了什么',
            '概括表格的主要信息',
            '浏览表格数据并总结',
            '梳理一下表格内容',
            '根据要求查看和归纳内容',
        ]
        for q in browse_queries:
            skill = SkillRouter.route(q, [], files=['NAS Operation Result.xlsx'])
            self.assertFalse(requires_spreadsheet_analysis(q, skill), f"Should not require spreadsheet analysis for: {q}")
            tools = select_active_tools(q, skill, False, False)
            names = {t['function']['name'] for t in tools}
            self.assertIn('read_file', names)
            self.assertNotIn('aggregate_spreadsheet', names)

        # 包含显式计算或明确指标时，仍必须强制执行表格分析
        calc_queries = [
            '计算2026年毛利率并总结',
            '汇总营业收入',
            '求和2026年销售费用',
            '识别审批状态异常',
        ]
        for q in calc_queries:
            skill = SkillRouter.route(q, [], files=['NAS Operation Result.xlsx'])
            self.assertTrue(requires_spreadsheet_analysis(q, skill), f"Should require spreadsheet analysis for: {q}")

    def test_report_informational_spreadsheet_tool(self):
        from dsh_modules.spreadsheet_informational import execute_informational_spreadsheet
        nas_path = Path(self.tmp.name) / "NAS Operation Result.xlsx"
        w = openpyxl.Workbook()
        s = w.active
        s.title = "SyncResult"
        s.append(["Host", "IP", "Status", "Timestamp"])
        s.append(["nas-01", "192.168.1.10", "Synchronization Successful", "2026-10-10 12:00:00"])
        s.append(["nas-02", "192.168.1.11", "Synchronization Successful", "2026-10-10 12:01:00"])
        w.save(nas_path)

        with patch('dsh_modules.spreadsheet_informational.resolve_sandboxed_path', return_value=(nas_path, None)):
            # 1. 正常非计算任务下的降级调用成功
            os.environ["DSH_LAST_USER_PROMPT"] = "查看和归纳内容"
            res = execute_informational_spreadsheet({
                'file_path': str(nas_path),
                'reason': '表格为运维同步日志，无数值计算指标'
            })
            tr = getattr(res, "tool_result", None)
            self.assertTrue(tr.is_success)
            self.assertEqual(tr.data.get('kind'), 'informational_spreadsheet')
            self.assertTrue(tr.data.get('is_informational'))

            # 2. 真实财务/计算任务下调用被严格阻断
            os.environ["DSH_LAST_USER_PROMPT"] = "计算2026年营业利润"
            res_calc = execute_informational_spreadsheet({
                'file_path': str(nas_path),
                'reason': '试图跳过计算'
            })
            tr_calc = getattr(res_calc, "tool_result", None)
            self.assertTrue(tr_calc.is_error)
            self.assertEqual(tr_calc.error_code, 'escape_rejected')

    def test_loop_informational_downgrade_delivers_clean_text(self):
        from dsh_modules.tools import get_sandbox_tools
        nas_path = Path(self.tmp.name) / "NAS Operation Result.xlsx"
        w = openpyxl.Workbook()
        s = w.active
        s.title = "SyncResult"
        s.append(["Host", "IP", "Status"])
        s.append(["nas-01", "192.168.1.10", "Success"])
        w.save(nas_path)

        tool_params = {
            'file_path': str(nas_path),
            'reason': 'NAS 运维结果日志，仅包含文本状态信息，无需数值计算'
        }
        tool_reply = {
            'content': '',
            'finish_reason': 'tool_calls',
            'tool_calls': [{
                'id': 'call_1',
                'type': 'function',
                'function': {'name': 'report_informational_spreadsheet', 'arguments': json.dumps(tool_params)}
            }]
        }
        summary_text = '本次 NAS 运维同步共有 1 台节点，状态全部为 Success，同步顺利完成。'
        calls = []
        def model(*args, **kwargs):
            calls.append(kwargs)
            return tool_reply if len(calls) == 1 else {'content': summary_text, 'finish_reason': 'stop'}

        with patch('dsh_modules.agent_loop.call_model_proxy', side_effect=model), \
             patch('dsh_modules.spreadsheet_informational.resolve_sandboxed_path', return_value=(nas_path, None)):
            result = run_agent_loop(
                [{'role': 'system', 'content': 'test'},
                 {'role': 'user', 'content': 'NAS Operation Result.xlsx 讲了什么'}],
                'test',
                RuntimePolicy(),
                3,
                tools=get_sandbox_tools({'report_informational_spreadsheet'}),
                spreadsheet_analysis_required=True
            )

        self.assertIn('本次 NAS 运维同步共有 1 台节点', result.final_text)
        self.assertNotIn('尚未通过执行证据核验', result.final_text)
        self.assertNotIn('无法计算的原因说明', result.final_text)
        self.assertEqual(result.telemetry.guard_decisions[-1]['action'], 'accept')

    def test_cjk_variant_fuzzy_matching_and_sheet_alignment(self):
        """Verifies zero-hardcoding CJK variant alignment for file_path and sheet names."""
        from dsh_modules.spreadsheet_context import execute_spreadsheet_structure
        from dsh_modules.spreadsheet_analysis import execute_spreadsheet_analysis

        real_filename = "工数实绩明细_20261020134124.xlsx"
        real_path = Path(self.tmp.name) / real_filename
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "工数实绩明细"
        ws.append(["部门", "WBS编号", "金额(元)"])
        ws.append(["FC01", "IK-100", 100])
        ws.append(["FC02", "EX-200", 200])
        wb.save(real_path)
        wb.close()

        # 模型产生异体字幻觉：把 "实" 写成日文汉字 "実"
        variant_file_path = f"/workspace/工数実绩明细_20261020134124.xlsx"
        variant_sheet = "工数実绩明细"

        with patch.dict(os.environ, {"DSH_SESSION_ATTACHMENTS": real_filename}), \
             patch("dsh_modules.config.WORKSPACE_DIR", self.tmp.name):
            # 1. 结构探测工具应自愈并成功解析工作簿与工作表结构
            struct_output = execute_spreadsheet_structure({
                "file_path": variant_file_path,
                "sheets": [variant_sheet]
            })
            self.assertTrue(struct_output.tool_result.is_success, str(struct_output))
            self.assertEqual(struct_output.tool_result.data["sheets"][0]["sheet"], "工数实绩明细")

            # 2. 原生聚合分析工具应自愈并成功执行计算
            analysis_output = execute_spreadsheet_analysis({
                "file_path": variant_file_path,
                "tables": [{
                    "id": "t1",
                    "sheet": variant_sheet,
                    "header_row": 1,
                    "data_range": "A2:C3",
                    "filters": [{"column": "B", "op": "exclude_prefix", "exclude_prefixes": ["EX"]}]
                }],
                "calculations": [{
                    "id": "c1",
                    "label": "金额合计",
                    "expression": {"aggregate": "sum", "table": "t1", "column": "C", "unit": "amount"}
                }]
            })
            self.assertTrue(analysis_output.tool_result.is_success, str(analysis_output))
            self.assertEqual(analysis_output.tool_result.data["facts"][0]["value"], "100")

            # 3. 传入完全不相关的别处文件时，必须严格拦截并返回包含可用附件列表的自愈提示
            irrelevant_output = execute_spreadsheet_analysis({
                "file_path": "/workspace/totally_different.xlsx",
                "tables": []
            })
            self.assertEqual(irrelevant_output.tool_result.error_code, "source_outside_session_scope")
            self.assertIn(real_filename, str(irrelevant_output))

    def test_predicate_dominance_and_targeted_entity_routing(self):
        """Standardized routing test: Entity predicates and targeted queries must NOT be collapsed into broad browsing."""
        from dsh_modules.analysis_contract import build_analysis_contract
        targeted_queries = [
            '查看员工编号 S22014的工数数据',
            '查看员工编号S22014',
            '看下员工杨烨的工数明细',
            '调出华东MIS部门的数据',
            '查找WBS编号 IW-S62662-00-001 的明细',
        ]
        for q in targeted_queries:
            skill = SkillRouter.route(q, [], files=['工数实绩明细_20261020134124.xlsx'])
            contract = build_analysis_contract(q)
            self.assertFalse(contract.is_browse_or_summary, f"Query '{q}' should NOT be browse_or_summary")
            self.assertTrue(requires_spreadsheet_analysis(q, skill), f"Query '{q}' MUST require spreadsheet analysis")
            self.assertTrue(len(contract.entity_predicates) > 0 or len(contract.attribute_targets) > 0)

    def test_context_budget_compacts_assistant_evidence_blocks(self):
        """Verifies that ContextBudget folds historical assistant evidence details to prevent context bloat and hallucination."""
        from dsh_modules.context_budget import ContextBudget
        hist = [
            {"role": "user", "content": "统计前序数据"},
            {"role": "assistant", "content": "这是前序结论。\n<details><summary>查看来源与核算明细</summary>\n来源：data.xlsx；证据 `fakehash12345678901234`。\n| a | 1 |\n</details>"}
        ]
        budgeted, _ = ContextBudget.budget_history(hist, max_total_chars=10000)
        assistant_msg = next(m for m in budgeted if m["role"] == "assistant")
        self.assertNotIn("<details>", assistant_msg["content"])
        self.assertNotIn("fakehash12345678901234", assistant_msg["content"])
        self.assertIn("前序轮次已核验明细", assistant_msg["content"])

    def test_preflight_profiling_and_multi_column_group_by(self):
        """Verifies Phase 1 Preflight profiling and multi-column group_by execution."""
        from dsh_modules.spreadsheet_preflight import profile_spreadsheet_for_planning
        from dsh_modules.spreadsheet_expression import ExpressionEvaluator

        # 1. Preflight on self.path
        res = profile_spreadsheet_for_planning(self.path, prompt="按部门和审批状态统计金额")
        self.assertIn("Pre-flight Sheet Profile", res["compact_markdown"])
        self.assertIn("通用业务表格", res["compact_markdown"])
        self.assertEqual(res["header_row"], 5)
        self.assertEqual(res["data_range"], "A6:E85")

        # 2. Multi-column group_by via compute
        r = self.compute(
            tables=[self.table],
            calculations=[{
                "id": "dept_summary",
                "label": "部门状态金额",
                "expression": {
                    "aggregate": "sum",
                    "table": "tx",
                    "column": "C",
                    "group_by": ["E", "D"],  # 部门, 审批状态
                    "unit": "amount"
                }
            }]
        )
        self.assertTrue(r.is_success, getattr(r, "error", None))
        self.assertTrue(len(r.data["facts"]) >= 1)
        for f in r.data["facts"]:
            self.assertIn(" - ", f["group_key"])
            self.assertEqual(f["unit"], "千元")


if __name__=='__main__':unittest.main()
