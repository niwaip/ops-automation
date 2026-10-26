"""Unit tests for standard Excel Query capabilities (P0 and P1).

Covers:
- query_contract: validation and refusal probes
- query_predicates: row condition trees, null semantics, scalar and text comparisons
- query_executor: deterministic 5-stage pipeline, grouping, having, order by, limit
"""

from decimal import Decimal
from pathlib import Path
import sys
import unittest
import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src"))

from dsh_modules.query_contract import (
    QueryContractError,
    validate_predicate_node,
    validate_query_plan,
)
from dsh_modules.query_predicates import (
    eval_row_predicate,
    eval_having_predicate,
    to_decimal,
)
from dsh_modules.query_executor import execute_query_plan


class QueryCapabilitiesTests(unittest.TestCase):
    def test_contract_refuses_unknown_operator(self):
        with self.assertRaises(QueryContractError) as ctx:
            validate_predicate_node(
                {"op": "typo_op", "field": "A", "value": 1},
                {"A"}
            )
        self.assertIn("不支持的操作符", str(ctx.exception))

    def test_contract_refuses_non_list_values_for_in_and_not_in(self):
        with self.assertRaises(QueryContractError):
            validate_predicate_node({"op": "in", "field": "A", "values": "scalar"}, {"A"})
        with self.assertRaises(QueryContractError):
            validate_predicate_node({"op": "not_in", "field": "A", "values": "scalar"}, {"A"})

    def test_contract_refuses_having_referencing_raw_column(self):
        with self.assertRaises(QueryContractError) as ctx:
            validate_query_plan(
                {
                    "source": {"table_id": "t1"},
                    "metrics": [{"id": "amount_sum", "op": "sum", "field": "金额"}],
                    "having": {"op": "gt", "field": "金额", "value": 100},
                },
                {"金额"}
            )
        self.assertIn("HAVING 条件引用了未定义的聚合度量", str(ctx.exception))

    def test_predicate_boolean_tree(self):
        # (billable == '无偿' OR hours > 20) AND grade != 'MIS10'
        tree = {
            "and": [
                {
                    "or": [
                        {"op": "eq", "field": "billable", "value": "无偿"},
                        {"op": "gt", "field": "hours", "value": 20},
                    ]
                },
                {"op": "ne", "field": "grade", "value": "MIS10"},
            ]
        }
        # Row 1: 有偿, 25h, BAS10 -> matches (False OR True) AND True -> True
        self.assertTrue(eval_row_predicate(tree, {"billable": "有偿", "hours": 25, "grade": "BAS10"}))
        # Row 2: 有偿, 25h, MIS10 -> False (grade is MIS10)
        self.assertFalse(eval_row_predicate(tree, {"billable": "有偿", "hours": 25, "grade": "MIS10"}))
        # Row 3: 无偿, 5h, BAS10 -> matches (True OR False) AND True -> True
        self.assertTrue(eval_row_predicate(tree, {"billable": "无偿", "hours": 5, "grade": "BAS10"}))
        # Row 4: 有偿, 15h, BAS10 -> False (False OR False)
        self.assertFalse(eval_row_predicate(tree, {"billable": "有偿", "hours": 15, "grade": "BAS10"}))

    def test_predicate_contains_and_literal_percent(self):
        node = {"op": "contains", "field": "name", "value": "%"}
        self.assertTrue(eval_row_predicate(node, {"name": "100% 完成"}))
        self.assertFalse(eval_row_predicate(node, {"name": "全部完成"}))

    def test_predicate_null_semantics(self):
        # SQL three-valued logic: comparison with NULL is False
        self.assertFalse(eval_row_predicate({"op": "eq", "field": "cust", "value": "ACME"}, {"cust": None}))
        self.assertFalse(eval_row_predicate({"op": "ne", "field": "cust", "value": "ACME"}, {"cust": None}))
        self.assertFalse(eval_row_predicate({"op": "gt", "field": "cust", "value": 10}, {"cust": None}))
        self.assertTrue(eval_row_predicate({"op": "is_null", "field": "cust"}, {"cust": None}))
        self.assertFalse(eval_row_predicate({"op": "is_not_null", "field": "cust"}, {"cust": None}))

    def test_executor_with_memory_workbook(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "TestSheet"
        ws.append(["部门", "类别", "金额", "工时"])
        ws.append(["D01", "有偿", 100, 10])
        ws.append(["D01", "有偿", 50, 15])
        ws.append(["D01", "无偿", 20, 5])
        ws.append(["D02", "有偿", 200, 30])
        ws.append(["D02", "无偿", 10, 2])
        # Summary footer row
        ws.append(["合计", "", 380, 62])

        # Query: Group by 部门, HAVING sum(金额) > 100, ORDER BY sum(金额) DESC
        plan = {
            "source": {"sheet": "TestSheet", "header_row": 1, "data_range": "A2:D6"},
            "where": {"op": "eq", "field": "类别", "value": "有偿"},
            "group_by": [{"field": "部门"}],
            "metrics": [
                {"id": "amt_sum", "op": "sum", "field": "金额"},
                {"id": "hrs_avg", "op": "avg", "field": "工时"},
                {"id": "cnt", "op": "count_rows"},
            ],
            "having": {"op": "gte", "metric": "amt_sum", "value": 150},
            "order_by": [{"ref": "amt_sum", "direction": "desc"}],
        }

        res = execute_query_plan(wb, plan)
        self.assertEqual(res.where_row_count, 3)  # 3 '有偿' rows
        self.assertEqual(res.final_row_count, 2)  # D02 amt=200, D01 amt=150 (both >= 150)
        self.assertEqual(res.rows[0]["group_key"], "D02")
        self.assertEqual(res.rows[0]["amt_sum"], "200")
        self.assertEqual(res.rows[0]["hrs_avg"], "30")
        self.assertEqual(res.rows[1]["group_key"], "D01")
        self.assertEqual(res.rows[1]["amt_sum"], "150")
        self.assertEqual(res.rows[1]["hrs_avg"], "12.5")

    def test_query_executor_preserves_source_unit_when_label_declares_source_unit(self):
        """测试单位自洽防线：当度量 label 显式声明了 (K)，但 req_unit 误填为 '元' 时，
        必须优先遵循声明的 (K) 原表单位，严禁执行 1000 倍放大！"""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        ws.append(["部门", "金额(K)"])
        ws.append(["華東MIS", 677.58544])
        ws.append(["華東BAS", 328.78323])

        plan = {
            "source": {"sheet": "Sheet1", "header_row": 1, "data_range": "A2:B3"},
            "group_by": ["部门"],
            "metrics": [
                {"id": "total_amount", "label": "合计金额(K)", "op": "sum", "field": "金额(K)", "unit": "元"}
            ],
        }

        res = execute_query_plan(wb, plan)
        # 验证数值未被乘以 1000
        self.assertEqual(res.rows[0]["group_key"], "華東MIS")
        self.assertEqual(res.rows[0]["total_amount"], "677.58544")
        mis_fact = next(f for f in res.facts if f["group_key"] == "華東MIS")
        self.assertEqual(mis_fact["value"], "677.58544")
        self.assertEqual(mis_fact["unit"], "K")
        self.assertIn("合计金额(K)", mis_fact["label"])

    def test_query_executor_synchronizes_label_when_unit_is_scaled(self):
        """测试度量标签对齐：若确实将原表 K 换算为 元（label 未写死 (K)），
        则回传 fact 的 label 必须自动同步为 (元)，杜绝标签 (K) 配 元数值 的矛盾事实！"""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        ws.append(["部门", "金额(K)"])
        ws.append(["華東MIS", 2.5])

        plan = {
            "source": {"sheet": "Sheet1", "header_row": 1, "data_range": "A2:B2"},
            "group_by": ["部门"],
            "metrics": [
                {"id": "total_amount", "label": "合计金额", "op": "sum", "field": "金额(K)", "unit": "元"}
            ],
        }

        res = execute_query_plan(wb, plan)
        self.assertEqual(res.rows[0]["total_amount"], "2500")
        mis_fact = next(f for f in res.facts if f["group_key"] == "華東MIS")
        self.assertEqual(mis_fact["value"], "2500")
        self.assertEqual(mis_fact["unit"], "元")
        self.assertIn("合计金额(元)", mis_fact["label"])

    def test_query_prefix_and_exclude_prefix_aliases_and_evaluation(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        ws.append(["部门", "WBS编号", "金额(K)"])
        ws.append(["華東MIS", "ex1001", 10.0])
        ws.append(["華東MIS", "sy2002", 20.0])
        ws.append(["華東MIS", "normal01", 30.0])
        ws.append(["華東BAS", "sp3003", 40.0])
        ws.append(["華東BAS", "normal02", 50.0])

        plan = {
            "source": {"sheet": "Sheet1", "header_row": 1, "data_range": "A2:C6"},
            "where": {
                "field": "WBS编号",
                "op": "exclude_prefix",
                "values": ["ex", "sy", "sp"]
            },
            "group_by": ["部门"],
            "metrics": [
                {"id": "amount", "label": "金额合计(K)", "op": "sum", "field": "金额(K)"}
            ]
        }

        res = execute_query_plan(wb, plan)
        # 仅 normal01 (30.0) 和 normal02 (50.0) 被保留
        self.assertEqual(len(res.rows), 2)
        mis_row = next(r for r in res.rows if r["group_key"] == "華東MIS")
        bas_row = next(r for r in res.rows if r["group_key"] == "華東BAS")
        self.assertEqual(mis_row["amount"], "30")
        self.assertEqual(bas_row["amount"], "50")

    def test_query_like_and_not_like_patterns(self):
        # 1. not_like with prefix pattern 'ex%'
        node1 = {"field": "code", "op": "not_like", "value": "ex%"}
        validate_predicate_node(node1, {"code"})
        self.assertEqual(node1["op"], "not_starts_with")
        self.assertEqual(node1["values"], ["ex"])
        self.assertFalse(eval_row_predicate(node1, {"code": "ex999"}))
        self.assertTrue(eval_row_predicate(node1, {"code": "ab999"}))

        # 2. like with suffix pattern '%txt'
        node2 = {"field": "file", "op": "like", "value": "%txt"}
        validate_predicate_node(node2, {"file"})
        self.assertEqual(node2["op"], "ends_with")
        self.assertTrue(eval_row_predicate(node2, {"file": "doc.txt"}))
        self.assertFalse(eval_row_predicate(node2, {"file": "doc.pdf"}))

        # 3. not_like with contains pattern '%temp%'
        node3 = {"field": "tag", "op": "not_like", "value": "%temp%"}
        validate_predicate_node(node3, {"tag"})
        self.assertEqual(node3["op"], "not_contains")
        self.assertFalse(eval_row_predicate(node3, {"tag": "this_temp_item"}))
        self.assertTrue(eval_row_predicate(node3, {"tag": "this_perm_item"}))


if __name__ == "__main__":
    unittest.main()
