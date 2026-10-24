"""
Unit tests for DeliveryClaimJudge (LLM-as-a-Judge for deliverable claim verification).
"""

import json
import sys
import tempfile
import time
import unittest
from pathlib import Path
from unittest.mock import patch, MagicMock

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src"))

from dsh_modules.delivery_claim_judge import DeliveryClaimJudge, judge_delivery_claims


class DeliveryClaimJudgeTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.cache_dir = self.root / "cache"
        self.cache_dir.mkdir(parents=True, exist_ok=True)
        DeliveryClaimJudge._CACHE_DIR = self.cache_dir
        self.start = time.time()

    def test_judge_clears_generic_referential_terms_with_mocked_llm(self):
        prompt = "需要更新工数管理，复制最新的sheet生成今天1023的新sheet，统计金额更新到预测部分"
        reply = "已生成《工数3Q预实管理_1023_6.xlsx》，新建了1023工作表，并将数据回填至目标工作表。"
        candidate_missing = ["新工作表_目标.xlsx"]

        # 磁盘上实际存在该工作簿和 1023 工作表
        import openpyxl
        wb = openpyxl.Workbook()
        wb.create_sheet("1023")
        wb.save(self.root / "工数3Q预实管理_1023_6.xlsx")

        mock_llm_reply = {
            "content": json.dumps({
                "phantom_claim": False,
                "missing_items": [],
                "reason": "已真实生成工数3Q预实管理_1023_6.xlsx且包含1023工作表，'目标工作表'为自然语言代词指代。"
            })
        }

        with patch("dsh_modules.analysis_synthesis._can_reach_proxy", return_value=True), \
             patch("dsh_modules.llm.call_model_proxy", return_value=mock_llm_reply) as mock_proxy:
            verdict = judge_delivery_claims(
                prompt=prompt,
                reply_text=reply,
                candidate_missing=candidate_missing,
                workspace_dir=str(self.root),
                turn_start_time=self.start - 5.0
            )
            self.assertFalse(verdict["phantom_claim"])
            self.assertEqual(verdict["missing_items"], [])
            self.assertIn("自然语言", verdict["reason"])
            self.assertEqual(mock_proxy.call_count, 1)

            # 第二次调用命中磁盘缓存，不再触发 LLM 调用
            verdict2 = judge_delivery_claims(
                prompt=prompt,
                reply_text=reply,
                candidate_missing=candidate_missing,
                workspace_dir=str(self.root),
                turn_start_time=self.start - 5.0
            )
            self.assertFalse(verdict2["phantom_claim"])
            self.assertEqual(mock_proxy.call_count, 1)

    def test_judge_confirms_truly_missing_deliverable(self):
        prompt = "统计部门金额并生成新的汇总表"
        reply = "已为您生成并导出了《部门别汇总报告.xlsx》，请下载查收。"
        candidate_missing = ["部门别汇总报告.xlsx"]

        mock_llm_reply = {
            "content": json.dumps({
                "phantom_claim": True,
                "missing_items": ["部门别汇总报告.xlsx"],
                "reason": "工作区中物理不存在部门别汇总报告.xlsx，模型口头声称已导出但实际未落盘。"
            })
        }

        with patch("dsh_modules.analysis_synthesis._can_reach_proxy", return_value=True), \
             patch("dsh_modules.llm.call_model_proxy", return_value=mock_llm_reply):
            verdict = judge_delivery_claims(
                prompt=prompt,
                reply_text=reply,
                candidate_missing=candidate_missing,
                workspace_dir=str(self.root),
                turn_start_time=self.start - 5.0
            )
            self.assertTrue(verdict["phantom_claim"])
            self.assertIn("部门别汇总报告.xlsx", verdict["missing_items"])

    def test_judge_offline_deterministic_fallback(self):
        prompt = "生成报告"
        reply = "已生成文件"
        # 包含一个合成的虚拟 Sheet 名和一个真实缺失的文件
        candidate_missing = ["新工作表_目标.xlsx", "真正的缺失报告.xlsx"]

        with patch("dsh_modules.analysis_synthesis._can_reach_proxy", return_value=False):
            verdict = judge_delivery_claims(
                prompt=prompt,
                reply_text=reply,
                candidate_missing=candidate_missing,
                workspace_dir=str(self.root),
                turn_start_time=self.start - 5.0
            )
            self.assertTrue(verdict["phantom_claim"])
            # 确定性降级自动过滤掉虚拟的 "新工作表_*.xlsx"
            self.assertNotIn("新工作表_目标.xlsx", verdict["missing_items"])
            self.assertIn("真正的缺失报告.xlsx", verdict["missing_items"])

    def test_agent_loop_dismisses_false_alarm_with_judge(self):
        from dsh_modules import agent_loop
        from dsh_modules.runtime_policy import RuntimePolicy

        import openpyxl
        wb = openpyxl.Workbook()
        wb.create_sheet("1023")
        wb_file = self.root / "工数3Q预实管理_1023_6.xlsx"
        wb.save(wb_file)

        detailed_summary = (
            "本次工数更新已完成！\n\n"
            "已在《工数3Q预实管理_1023_6.xlsx》中新建 1023 工作表，并将数据回填至目标工作表。\n\n"
            "| 部门 | 有偿(K) | 无偿(K) |\n"
            "|---|---|---|\n"
            "| FC33001 | 67.04 | 14.30 |\n"
            "公式重算验证全部通过。"
        )

        mock_model_reply = {"content": detailed_summary, "finish_reason": "stop"}

        # 模拟 Judge 裁决为 False（代词误报，已交付真实文件）
        mock_judge_verdict = {
            "phantom_claim": False,
            "missing_items": [],
            "reason": "'目标工作表'是指代已生成的1023工作表，产物完备有效。"
        }

        with patch.object(agent_loop, "WORKSPACE_DIR", str(self.root)), \
             patch.object(agent_loop, "KNOWLEDGE_DIR", str(self.root / "knowledge")), \
             patch.object(agent_loop, "call_model_proxy", return_value=mock_model_reply), \
             patch.object(agent_loop, "parse_tool_calls", return_value=[]), \
             patch.object(agent_loop, "_check_no_tool_assertion_guard", return_value=("pass", None, 1)), \
             patch.object(agent_loop, "_finalize_agent_text", return_value=(detailed_summary, [])), \
             patch("dsh_modules.delivery_claim_judge.judge_delivery_claims", return_value=mock_judge_verdict):

            result = agent_loop.run_agent_loop(
                [{"role": "user", "content": "复制最新的sheet生成1023的新sheet更新预测"}],
                model="test",
                policy=RuntimePolicy(),
                tools=[],
                max_rounds=1,
                is_generate_intent=True,
                turn_start_time=self.start - 5.0
            )

            # 验证长篇总结完整保留，没有被“本次文件生成未完成”冲刷覆盖
            self.assertIn("本次工数更新已完成！", result.final_text)
            self.assertIn("FC33001", result.final_text)
            self.assertNotIn("本次文件生成未完成", result.final_text)


if __name__ == "__main__":
    unittest.main()
