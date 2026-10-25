"""Regression tests for source cell provenance binding (T02) and generated workbook source grounding (T09)."""

import json
import os
import shutil
import sys
import tempfile
import time
import unittest
from pathlib import Path
from unittest.mock import patch
import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src"))

from dsh_modules.analysis_contract import build_analysis_contract
from dsh_modules.analysis_synthesis import (
    _is_synthesis_grounded,
    build_verified_facts_digest,
)
from dsh_modules.analysis_task_state import analysis_task_state
from dsh_modules.evidence_store import persist_turn_evidence
from dsh_modules.query_result import QueryResult
from dsh_modules.report_presentation import analysis_overview, query_overview
from dsh_modules.spreadsheet_context import execute_spreadsheet_structure
from dsh_modules.tool_result import ToolResult
from dsh_modules.workbook_delivery import (
    inspect_generated_workbook,
    quarantine_invalid_deliverables,
)
from dsh_modules.workbook_source_grounding import validate_generated_workbook_against_sources


class SourceGroundingAndCellProvenanceTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)

    def _build_financial_source_workbook(self, filename: str = "AI能力测试_财务报表.xlsx") -> Path:
        path = self.root / filename
        wb = openpyxl.Workbook()

        # Sheet 1: 月度经营
        ws_m = wb.active
        ws_m.title = "月度经营"
        ws_m["A1"] = "2026年月度经营明细"
        ws_m.append(["单位：人民币千元"])
        ws_m.append(["月份", "营业收入", "营业成本", "毛利", "毛利率", "销售费用", "管理费用", "研发费用", "营业利润", "营业利润率", "其他收益", "税前利润", "所得税", "净利润"])
        months_data = [
            ("1月", 8420, 0.405, 0.118, 0.092, 510, 20),
            ("2月", 8150, 0.398, 0.121, 0.091, 500, -35),
            ("3月", 9100, 0.421, 0.119, 0.089, 525, 10),
            ("4月", 9450, 0.430, 0.116, 0.088, 540, 25),
            ("5月", 9800, 0.417, 0.120, 0.090, 560, 15),
            ("6月", 10200, 0.425, 0.117, 0.087, 570, 30),
            ("7月", 9750, 0.409, 0.123, 0.089, 555, -120),
            ("8月", 10600, 0.438, 0.115, 0.086, 590, 40),
            ("9月", 11150, 0.442, 0.114, 0.085, 610, 35),
            ("10月", 11700, 0.447, 0.113, 0.084, 625, 50),
            ("11月", 12300, 0.451, 0.112, 0.083, 650, 45),
            ("12月", 13400, 0.458, 0.118, 0.086, 690, 60),
        ]
        for idx, (m_label, rev, gm_rate, s_rate, g_rate, rd, other) in enumerate(months_data, start=4):
            ws_m.append([
                m_label,
                rev,
                f"=B{idx}*(1-{gm_rate})",
                f"=B{idx}-C{idx}",
                f"=D{idx}/B{idx}",
                f"=B{idx}*{s_rate}",
                f"=B{idx}*{g_rate}",
                rd,
                f"=D{idx}-SUM(F{idx}:H{idx})",
                f"=I{idx}/B{idx}",
                other,
                f"=I{idx}+K{idx}",
                f"=MAX(0,L{idx}*参数!$B$8)",
                f"=L{idx}-M{idx}",
            ])
        ws_m.append([
            "全年合计/平均",
            "=SUM(B4:B15)",
            "=SUM(C4:C15)",
            "=SUM(D4:D15)",
            "=D16/B16",
            "=SUM(F4:F15)",
            "=SUM(G4:G15)",
            "=SUM(H4:H15)",
            "=SUM(I4:I15)",
            "=I16/B16",
            "=SUM(K4:K15)",
            "=SUM(L4:L15)",
            "=SUM(M4:M15)",
            "=SUM(N4:N15)",
        ])

        # Sheet 2: 利润表
        ws_p = wb.create_sheet("利润表")
        ws_p["A1"] = "2026年度利润表"
        ws_p.append([])
        ws_p.append(["项目", "2026实际", "2026预算", "差异"])
        ws_p.append(["营业收入", "=月度经营!B16", 120000, "=B4-C4"])
        ws_p.append(["营业成本", "=月度经营!C16", 70000, "=B5-C5"])
        ws_p.append(["毛利", "=月度经营!D16", 50000, "=B6-C6"])
        ws_p.append(["销售费用", "=月度经营!F16", 14000, "=B7-C7"])
        ws_p.append(["管理费用", "=月度经营!G16", 10800, "=B8-C8"])
        ws_p.append(["研发费用", "=月度经营!H16", 6900, "=B9-C9"])
        ws_p.append(["营业利润", "=月度经营!I16", 18300, "=B10-C10"])
        ws_p.append(["其他收益", "=月度经营!K16", 300, "=B11-C11"])
        ws_p.append(["税前利润", "=月度经营!L16", 18600, "=B12-C12"])
        ws_p.append(["所得税", "=月度经营!M16", 4650, "=B13-C13"])
        ws_p.append(["净利润", "=月度经营!N16", 13950, "=B14-C14"])

        # Sheet 3: 参数
        ws_param = wb.create_sheet("参数")
        ws_param["A8"] = "所得税率"
        ws_param["B8"] = 0.25

        wb.save(path)
        wb.close()
        return path

    def test_t02_wrong_source_cell_binding_rejected_and_true_cells_required(self):
        """Reproduces T02 probe: B10/B15 must be rejected when receipt provenance is 月度经营!B16 and 月度经营!N16."""
        prompt = "只汇总2026年全年营业收入和净利润两项金额，并说明它们各自来自哪个工作表和单元格。"
        receipt = ToolResult.success(
            {
                "kind": "spreadsheet_analysis",
                "evidence_id": "ev-source-probe",
                "coverage": [{"id": "t", "sheet": "月度经营", "header_row": 3, "data_range": "A4:N16", "filters": []}],
                "facts": [
                    {
                        "id": "rev",
                        "table": "t",
                        "label": "全年营业收入",
                        "metric_id": "营业收入",
                        "value": "124020",
                        "unit": "千元",
                        "source_cells": ["月度经营!B16"],
                    },
                    {
                        "id": "np",
                        "table": "t",
                        "label": "全年净利润",
                        "metric_id": "净利润",
                        "value": "16038.3",
                        "unit": "千元",
                        "source_cells": ["月度经营!N16"],
                    },
                ],
                "checks": [],
                "results": [],
                "audit": {},
            },
            provenance={"tool": "analyze_spreadsheet"},
        )
        state = analysis_task_state([receipt], prompt)

        # 1. Candidate citing wrong cells B10 and B15 MUST be rejected
        wrong_candidate = "全年营业收入为 124,020 千元（月度经营 B10），全年净利润为 16,038.3 千元（月度经营 B15）。"
        self.assertFalse(
            _is_synthesis_grounded(wrong_candidate, [receipt], prompt=prompt),
            "Synthesis citing fabricated cells B10/B15 must be rejected!",
        )

        # 2. Candidate omitting cell coordinates when prompt explicitly asked for cells MUST be rejected
        missing_cell_candidate = "全年营业收入为 124,020 千元（来自月度经营工作表），全年净利润为 16,038.3 千元（来自月度经营工作表）。"
        self.assertFalse(
            _is_synthesis_grounded(missing_cell_candidate, [receipt], prompt=prompt),
            "Synthesis omitting source cell coordinates when requested by prompt must be rejected!",
        )

        # 3. Candidate citing true cells B16 and N16 MUST be accepted
        good_candidate = "全年营业收入为 124,020 千元（来自月度经营!B16），全年净利润为 16,038.3 千元（来自月度经营!N16）。"
        self.assertTrue(
            _is_synthesis_grounded(good_candidate, [receipt], prompt=prompt),
            "Synthesis citing verified cells 月度经营!B16 and 月度经营!N16 must be accepted!",
        )

        # 4. Digest and deterministic overviews must surface source_cells
        digest = build_verified_facts_digest([receipt], prompt=prompt)
        self.assertIn("月度经营!B16", digest)
        self.assertIn("月度经营!N16", digest)
        overview = analysis_overview([receipt], prompt=prompt)
        self.assertIn("月度经营!B16", overview)
        self.assertIn("月度经营!N16", overview)

    def test_query_result_preserves_source_cells_in_tool_view_and_overview(self):
        """Verifies QueryResult.to_tool_result keeps source_cells in LLM-facing view."""
        qr = QueryResult(
            rows=[{"项目": "营业收入", "rev": "124020"}],
            selected_row_ids=[4],
            total_source_rows=11,
            where_row_count=1,
            group_count=0,
            having_group_count=0,
            final_row_count=1,
            facts=[
                {
                    "id": "rev",
                    "label": "营业收入",
                    "op": "value",
                    "value": "124020",
                    "unit": "千元",
                    "column": "B",
                    "source_cells": ["利润表!B4"],
                }
            ],
            coverage=[{"sheet": "利润表", "header_row": 3, "data_range": "A4:D14", "filters": []}],
            visited_cells=["利润表!B4"],
        )
        tr = qr.to_tool_result(
            source=Path("/workspace/AI能力测试_财务报表.xlsx"),
            source_hash="hash_wb1",
            requested_plan={},
            effective_plan={"metrics": [{"id": "rev", "label": "营业收入", "unit": "千元"}]},
            plan_hash="ph1",
        )
        rendered_view = json.loads(tr.render_text())
        self.assertEqual(rendered_view["facts"][0].get("source_cells"), ["利润表!B4"])
        self.assertIn("利润表!B4", query_overview([tr], build_analysis_contract("汇总营业收入")))

    def test_t09_generated_workbook_with_wrong_hardcoded_financials_is_quarantined(self):
        """Reproduces T09: generated Excel with 毛利=50120 and 营业利润=21438.3 is rejected and quarantined."""
        src = self._build_financial_source_workbook()

        # Case 1: With explicit source cell column ('B4', 'B6', 'B10', 'B14')
        gen_bad = self.root / "E2E_finance_interactive_20261025.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "2026财务指标"
        ws.append(["指标", "数值", "单位", "来源单元格"])
        ws.append(["营业收入", 124020, "千元", "B4"])
        ws.append(["毛利", 50120, "千元", "B6"])
        ws.append(["营业利润", 21438.3, "千元", "B10"])
        ws.append(["净利润", 16038.3, "千元", "B14"])
        wb.save(gen_bad)
        wb.close()

        # Case 2: Without source cell column (pure metric label + hardcoded wrong value)
        gen_no_ref = self.root / "E2E_finance_no_ref.xlsx"
        wb2 = openpyxl.Workbook()
        ws2 = wb2.active
        ws2.title = "核心指标"
        ws2.append(["指标", "全年数值", "单位"])
        ws2.append(["营业收入", 124020, "千元"])
        ws2.append(["毛利", 50120, "千元"])
        ws2.append(["营业利润", 21438.3, "千元"])
        ws2.append(["净利润", 16038.3, "千元"])
        wb2.save(gen_no_ref)
        wb2.close()

        # Case 3: Accurate generated workbook
        gen_good = self.root / "E2E_finance_good.xlsx"
        wb3 = openpyxl.Workbook()
        ws3 = wb3.active
        ws3.title = "核心指标"
        ws3.append(["指标", "数值", "单位", "来源单元格"])
        ws3.append(["营业收入", 124020, "千元", "利润表!B4"])
        ws3.append(["毛利", 53443.25, "千元", "利润表!B6"])
        ws3.append(["营业利润", 21209.4, "千元", "利润表!B10"])
        ws3.append(["净利润", 16038.3, "千元", "利润表!B14"])
        wb3.save(gen_good)
        wb3.close()

        with patch.dict(os.environ, {"DSH_SESSION_ATTACHMENTS": src.name}):
            # Validate Case 1
            res_bad = inspect_generated_workbook(gen_bad)
            self.assertFalse(res_bad["valid"])
            joined_issues = " | ".join(res_bad["issues"])
            self.assertIn("50120", joined_issues)
            self.assertIn("53443.25", joined_issues)
            self.assertIn("21438.3", joined_issues)
            self.assertIn("21209.4", joined_issues)

            # Validate Case 2 (no cell ref column, caught by Check B label grounding)
            issues_no_ref = validate_generated_workbook_against_sources(gen_no_ref)
            self.assertEqual(len(issues_no_ref), 2)
            self.assertIn("50120", " | ".join(issues_no_ref))
            self.assertIn("21438.3", " | ".join(issues_no_ref))
            gen_no_ref.unlink()

            # Validate Case 3 (accurate numbers pass)
            res_good = inspect_generated_workbook(gen_good)
            self.assertTrue(res_good["valid"], f"Valid workbook should pass: {res_good['issues']}")
            gen_good.unlink()

            # Quarantine bad deliverable and verify false-success text is replaced
            false_success_text = (
                "已成功生成 `E2E_finance_interactive_20261025.xlsx`。\n"
                "| 指标 | 数值 |\n|---|---|\n| 毛利 | 50,120.00 |\n| 营业利润 | 21,438.30 |\n"
                "公式校验：verified (0 errors)"
            )
            outbounds = [json.dumps({"filePath": str(gen_bad), "fileName": gen_bad.name})]
            clean_text, kept_outbounds, rejected = quarantine_invalid_deliverables(
                self.root, time.time() - 10, false_success_text, outbounds
            )
            self.assertEqual(len(rejected), 1)
            self.assertEqual(kept_outbounds, [])
            self.assertFalse(gen_bad.exists())
            self.assertNotIn("已成功生成", clean_text)
            self.assertNotIn("公式校验：verified", clean_text)
            self.assertIn("Excel 交付校验未通过", clean_text)
            self.assertIn("53443.25", clean_text)
            self.assertIn("21209.4", clean_text)

            # Verify evidence store records verified=False when workbook_delivery rejects
            history_file = self.root / "chat.json"
            ev_meta = persist_turn_evidence(
                history_file=history_file,
                source_hashes={src.name: "hash123"},
                query="生成核心财务指标Excel",
                receipts=[],
                guard_decisions=[
                    {"guard": "task_intent", "action": "accept", "reason": {"output_mode": "artifact"}},
                    {"guard": "workbook_delivery", "action": "reject", "reason": rejected},
                ],
            )
            self.assertFalse(ev_meta["verified"])

    def test_structure_tool_includes_all_rows_for_small_sheet_and_warns_on_formulas(self):
        """Verifies inspect_spreadsheet_structure includes all rows (<=32) and warns against guessing formulas."""
        src = self._build_financial_source_workbook()
        with patch.dict(os.environ, {"DSH_SESSION_ATTACHMENTS": src.name}), \
             patch("dsh_modules.config.WORKSPACE_DIR", str(self.root)):
            out = execute_spreadsheet_structure({"file_path": src.name, "sheets": ["利润表"]})
            self.assertTrue(out.tool_result.is_success)
            sample_coords = {
                k for row_map in out.tool_result.data["sheets"][0]["sample_cells"] for k in row_map.keys()
            }
            # Row 10 (营业利润) must be visible in sample_cells!
            self.assertIn("A10", sample_coords)
            self.assertIn("B10", sample_coords)
            self.assertIn("严禁据此猜测或硬编码数值", out.render_text())

    def test_multisheet_preflight_exposes_vertical_row_map_and_prunes_summary_row(self):
        """Verifies preflight profiles all sheets even when sheet[0] is a cover page, maps vertical rows, and prunes '全年合计/平均'."""
        from dsh_modules.spreadsheet_context import build_spreadsheet_source_context

        src = self._build_financial_source_workbook("multi_sheet_fin.xlsx")
        wb = openpyxl.load_workbook(src)
        ws_cover = wb.create_sheet("使用说明", 0)
        ws_cover.append(["项目", "说明"])
        ws_cover.append(["币种及单位", "人民币千元"])
        ws_cover.append(["建议任务1", "经营总览"])
        wb.save(src)
        wb.close()

        ctx = build_spreadsheet_source_context(
            src,
            prompt="请根据工作簿计算2026年全年营业收入、毛利、营业利润、净利润",
        )
        # Must include vertical table row map for 利润表 even though 使用说明 is sheet[0]
        self.assertIn("工作表 [利润表] [纵表/科目行结构]", ctx)
        self.assertIn("第4行(A4):'营业收入'", ctx)
        self.assertIn("第10行(A10):'营业利润'", ctx)
        self.assertIn("第14行(A14):'净利润'", ctx)
        self.assertIn('"group_by": ["项目"]', ctx)
        # Must prune Row 16 ('全年合计/平均') from 月度经营 so data_range is A4:N15
        self.assertIn("有效明细范围 A4:N15", ctx)
        self.assertIn("已自动排除表尾汇总行: 第16行 '全年合计/平均'", ctx)

    def test_t02_r2_ungrouped_vertical_multi_row_sum_and_cloned_metrics_blocked(self):
        """Reproduces R2 T02: ungrouped where Project in ['营业收入', '净利润'] summing to 140058.3 is blocked."""
        from dsh_modules.query_contract import QueryContractError, validate_query_plan
        from dsh_modules.query_executor import execute_query_plan

        src = self._build_financial_source_workbook()
        prompt = "只汇总2026年全年营业收入和净利润两项金额，并说明它们各自来自哪个工作表和单元格。"

        # 1. Exact T02 R2 tool call #2: duplicate sum(2026实际) metrics without group_by -> QueryContractError
        bad_plan_duplicate_metrics = {
            "source": {"sheet": "利润表", "header_row": 3, "data_range": "A4:D14"},
            "where": {"field": "项目", "op": "in", "values": ["营业收入", "净利润"]},
            "metrics": [
                {"id": "revenue", "label": "营业收入", "op": "sum", "field": "2026实际", "unit": "千元"},
                {"id": "net_profit", "label": "净利润", "op": "sum", "field": "2026实际", "unit": "千元"},
            ],
        }
        with self.assertRaises(QueryContractError):
            validate_query_plan(bad_plan_duplicate_metrics, {"项目", "2026实际", "2026预算", "差异"})

        # 2. Single metric claiming '营业收入' while filtering ['营业收入', '净利润'] without group_by -> QueryContractError
        bad_plan_single_claim = {
            "source": {"sheet": "利润表", "header_row": 3, "data_range": "A4:D14"},
            "where": {"field": "项目", "op": "in", "values": ["营业收入", "净利润"]},
            "metrics": [
                {"id": "rev", "label": "营业收入", "op": "sum", "field": "2026实际", "unit": "千元"},
            ],
        }
        wb_q = openpyxl.load_workbook(src, data_only=False)
        self.addCleanup(wb_q.close)
        with self.assertRaises(QueryContractError):
            execute_query_plan(wb_q, bad_plan_single_claim)

        # 3. Synthetic ungrouped conflated receipt (140058.3 for both metrics) must be rejected by analysis_task_state
        conflated_receipt = ToolResult.success(
            {
                "kind": "spreadsheet_analysis",
                "evidence_id": "ev-conflated-t02",
                "coverage": [
                    {
                        "id": "query_source",
                        "sheet": "利润表",
                        "header_row": 3,
                        "data_range": "A4:D14",
                        "coverage_complete": True,
                        "filters": [{"field": "项目", "op": "in", "values": ["营业收入", "净利润"]}],
                    }
                ],
                "facts": [
                    {
                        "id": "revenue",
                        "label": "营业收入",
                        "op": "sum",
                        "value": "140058.3",
                        "unit": "千元",
                        "source_cells": ["利润表!B4", "利润表!B14"],
                    },
                    {
                        "id": "net_profit",
                        "label": "净利润",
                        "op": "sum",
                        "value": "140058.3",
                        "unit": "千元",
                        "source_cells": ["利润表!B4", "利润表!B14"],
                    },
                ],
                "rows": [{"revenue": "140058.3", "net_profit": "140058.3"}],
            },
            provenance={"tool": "aggregate_spreadsheet"},
        )
        state_bad = analysis_task_state([conflated_receipt], prompt)
        self.assertFalse(state_bad["complete"], "Conflated 140058.3 receipt must NOT pass task gate!")
        self.assertTrue(any(i.get("code") == "conflated_multi_row_metrics" for i in state_bad.get("issues", [])))
        bad_text = "2026年全年营业收入为 140,058.3 千元（利润表!B4），全年净利润为 140,058.3 千元（利润表!B14）。"
        self.assertFalse(_is_synthesis_grounded(bad_text, [conflated_receipt], prompt=prompt))

        # 4. With group_by: ["项目"], query succeeds and produces exact 124020 (利润表!B4) and 16038.3 (利润表!B14)
        from dsh_modules.spreadsheet_analysis import calculate_spreadsheet_analysis

        good_plan = {
            "tables": [{"id": "p", "sheet": "利润表", "header_row": 3, "data_range": "A4:D14"}],
            "query": {
                "source": {"sheet": "利润表", "header_row": 3, "data_range": "A4:D14"},
                "where": {"field": "项目", "op": "in", "values": ["营业收入", "净利润"]},
                "group_by": ["项目"],
                "metrics": [
                    {"id": "actual", "label": "2026实际", "op": "sum", "field": "2026实际", "unit": "千元"},
                ],
            },
        }
        tr_good = calculate_spreadsheet_analysis(src, good_plan, domain="finance")
        self.assertTrue(tr_good.is_success, getattr(tr_good, "error", None))
        facts_by_key = {f["group_key"]: f for f in tr_good.data["facts"]}
        self.assertEqual(facts_by_key["营业收入"]["value"], "124020")
        self.assertEqual(facts_by_key["营业收入"]["source_cells"], ["利润表!B4"])
        self.assertEqual(facts_by_key["净利润"]["value"], "16038.3")
        self.assertEqual(facts_by_key["净利润"]["source_cells"], ["利润表!B14"])
        state_good = analysis_task_state([tr_good], prompt)
        self.assertTrue(state_good["complete"], f"Grouped vertical query should complete task: {state_good}")

    def test_t01_r2_monthly_summary_row_pruned_and_horizontal_receipt_preserved(self):
        """Reproduces R2 T01: 月度经营 A4:N16 auto-prunes Row 16 ('全年合计/平均') and is not wiped out by rejected 利润表 receipt."""
        from decimal import Decimal
        from dsh_modules.spreadsheet_analysis import calculate_spreadsheet_analysis

        src = self._build_financial_source_workbook()
        prompt = "请根据工作簿计算2026年全年营业收入、毛利、营业利润、净利润、加权毛利率和加权净利率（净利率按净利润/营业收入计算）。"

        # Earlier rejected vertical-table receipt (like T01 tool call #3)
        bad_vertical_receipt = ToolResult.success(
            {
                "kind": "spreadsheet_analysis",
                "evidence_id": "ev-bad-vertical",
                "coverage": [
                    {
                        "id": "query_source",
                        "sheet": "利润表",
                        "header_row": 3,
                        "data_range": "A4:D14",
                        "coverage_complete": True,
                        "filters": [{"field": "项目", "op": "in", "values": ["营业收入", "毛利", "营业利润", "净利润"]}],
                    }
                ],
                "facts": [
                    {
                        "id": "sum_all",
                        "label": "2026实际汇总",
                        "op": "sum",
                        "value": "214710.95",
                        "unit": "千元",
                        "source_cells": ["利润表!B4", "利润表!B6", "利润表!B10", "利润表!B14"],
                    }
                ],
                "rows": [{"sum_all": "214710.95"}],
            },
            provenance={"tool": "aggregate_spreadsheet"},
        )

        # T01 tool call #4 on 月度经营 with data_range A4:N16 (including Row 16 '全年合计/平均')
        monthly_plan = {
            "tables": [{"id": "m", "sheet": "月度经营", "header_row": 3, "data_range": "A4:N16"}],
            "calculations": [
                {"id": "rev", "label": "全年营业收入", "expression": {"aggregate": "sum", "table": "m", "column": "营业收入", "unit": "amount"}},
                {"id": "gp", "label": "全年毛利", "expression": {"aggregate": "sum", "table": "m", "column": "毛利", "unit": "amount"}},
                {"id": "op", "label": "全年营业利润", "expression": {"aggregate": "sum", "table": "m", "column": "营业利润", "unit": "amount"}},
                {"id": "np", "label": "全年净利润", "expression": {"aggregate": "sum", "table": "m", "column": "净利润", "unit": "amount"}},
                {
                    "id": "gm_rate",
                    "label": "加权毛利率",
                    "expression": {
                        "op": "divide",
                        "args": [
                            {"aggregate": "sum", "table": "m", "column": "毛利", "unit": "amount"},
                            {"aggregate": "sum", "table": "m", "column": "营业收入", "unit": "amount"},
                        ],
                    },
                },
                {
                    "id": "np_rate",
                    "label": "加权净利率",
                    "expression": {
                        "op": "divide",
                        "args": [
                            {"aggregate": "sum", "table": "m", "column": "净利润", "unit": "amount"},
                            {"aggregate": "sum", "table": "m", "column": "营业收入", "unit": "amount"},
                        ],
                    },
                },
            ],
        }
        good_monthly_receipt = calculate_spreadsheet_analysis(src, monthly_plan, domain="finance")
        self.assertTrue(good_monthly_receipt.is_success, getattr(good_monthly_receipt, "error", None))
        facts_map = {f["id"]: Decimal(f["value"]) for f in good_monthly_receipt.data["facts"]}
        # Row 16 ('全年合计/平均') must be pruned so revenue is 124020 (NOT double-counted 248040!)
        self.assertEqual(facts_map["rev"], Decimal("124020"))
        self.assertEqual(facts_map["gp"], Decimal("53443.25"))
        self.assertEqual(facts_map["op"], Decimal("21209.4"))
        self.assertEqual(facts_map["np"], Decimal("16038.3"))

        # Combining bad_vertical_receipt + good_monthly_receipt must NOT wipe out good_monthly_receipt!
        bad_vertical_receipt = ToolResult.success(
            bad_vertical_receipt.data,
            provenance=dict(good_monthly_receipt.provenance or {}),
        )
        state = analysis_task_state([bad_vertical_receipt, good_monthly_receipt], prompt)
        self.assertTrue(state["complete"], f"Monthly receipt must not be wiped out by rejected vertical receipt: {state}")

    def test_t06_r2_clarification_turn_overwrites_stale_evidence_sidecar(self):
        """Verifies persist_turn_evidence overwrites stale verified=True when a clarification turn has no tool receipts."""
        history_file = self.root / "session_history.json"
        # Turn T05: verified=True
        ev_t05 = persist_turn_evidence(
            history_file=history_file,
            source_hashes={"AI能力测试_财务报表.xlsx": "hash1"},
            query="T05: 检查报表勾稽等式",
            receipts=[],
            guard_decisions=[{"guard": "spreadsheet_evidence", "action": "accept", "reason": "ok"}],
        )
        self.assertTrue(ev_t05["verified"])

        # Turn T06: no receipts, guard_decisions has reject + clarify
        ev_t06 = persist_turn_evidence(
            history_file=history_file,
            source_hashes={"AI能力测试_财务报表.xlsx": "hash1"},
            query="T06: 帮我看一下利润表里的偏差情况，为什么利润和预期差这么多？",
            receipts=[],
            guard_decisions=[
                {"guard": "spreadsheet_evidence", "action": "reject", "reason": "missing_tool_receipt"},
                {"guard": "analysis_gate_review", "action": "clarify", "reason": "scope_ambiguous"},
            ],
        )
        self.assertFalse(ev_t06["verified"])
        sidecar = history_file.with_suffix(".evidence.json")
        loaded = json.loads(sidecar.read_text(encoding="utf-8"))
        self.assertFalse(loaded["verified"])
        self.assertIn("T06", loaded["query"])


if __name__ == "__main__":
    unittest.main()
