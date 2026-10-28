"""Business calculations require model tool calls, not runtime or presentation fallback."""
import copy
import io
import json
import os
import sys
import tempfile
import unittest
from contextlib import redirect_stdout
from pathlib import Path
from unittest.mock import patch

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'src'))
from dsh_modules.agent_loop import run_agent_loop
from dsh_modules.analysis_task_state import analysis_task_state, task_feedback
from dsh_modules.analysis_synthesis import _is_synthesis_grounded
from dsh_modules.bounded_summary_recovery import recover_summary, SummaryRecovery
from dsh_modules.deliverable_contract import materialize_script_deliverable
from dsh_modules.registered_metric_completion import RegisteredMetricCompletion
from dsh_modules.report_presentation import analysis_overview
from dsh_modules.runtime_policy import RuntimePolicy
from dsh_modules.spreadsheet_analysis import calculate_spreadsheet_analysis
from dsh_modules.spreadsheet_trends import render_ratio_trends
from dsh_modules.telemetry import TelemetryStats
from dsh_modules.tool_result import ToolResult
from dsh_modules.tools import get_sandbox_tools


class ExplicitModelComputationTests(unittest.TestCase):
    def setUp(self):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        self.folder = Path(temp.name)
        self.source = self.folder / 'source.xlsx'
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = '利润表'
        sheet.append(['2026年度利润表，单位：千元'])
        sheet.append(['项目', '2026实际', '2026预算', '差异'])
        for row in [('营业收入', 124020, 120000, 4020), ('毛利', 53443.25, 50000, 3443.25),
                    ('营业利润', 21209.4, 18300, 2909.4), ('净利润', 16038.3, 13950, 2088.3)]:
            sheet.append(row)
        workbook.save(self.source)
        workbook.close()
        self.prompt = '总结2026年收入、毛利率、营业利润和净利润。'
        self.query = {'query': {'source': {'sheet': '利润表', 'header_row': 2, 'data_range': 'A3:D6'},
            'group_by': ['项目'], 'metrics': [
                {'id': name, 'label': field, 'field': field, 'op': 'sum', 'unit': 'source'}
                for name, field in [('actual', '2026实际'), ('budget', '2026预算'), ('diff', '差异')]]}}
        self.margin = {'tables': [{'id': 't', 'sheet': '利润表', 'header_row': 2, 'data_range': 'A3:D6'}],
            'calculations': [{'id': 'margin', 'label': '2026实际毛利率', 'metric_id': 'finance.gross_margin',
                'output_unit': '%', 'expression': {'op': 'divide', 'args': [
                    {'sheet': '利润表', 'cell': 'B4', 'unit': 'amount'},
                    {'sheet': '利润表', 'cell': 'B3', 'unit': 'amount'}]}}]}
        for target, value in [('dsh_modules.analysis_synthesis._can_reach_proxy', False),
                              ('dsh_modules.semantic_filter_judge.judge_filter_conformance', True),
                              ('dsh_modules.task_completeness_judge.judge_task_completeness', None)]:
            patcher = patch(target, return_value=value)
            patcher.start()
            self.addCleanup(patcher.stop)

    def execute(self, plan):
        result = calculate_spreadsheet_analysis(self.source, plan, domain='finance')
        self.assertTrue(result.is_success, result.error)
        return result

    def call(self, call_id, plan):
        return {'content': '', 'finish_reason': 'tool_calls', 'tool_calls': [
            {'id': call_id, 'type': 'function', 'function': {'name': 'aggregate_spreadsheet',
             'arguments': json.dumps({'file_path': str(self.source), **plan})}}]}

    def run_loop(self, replies):
        with patch('dsh_modules.agent_loop.call_model_proxy', side_effect=replies) as model, \
             patch('dsh_modules.spreadsheet_analysis.resolve_sandboxed_path', return_value=(self.source, None)), \
             patch.dict(os.environ, {'DSH_SESSION_ATTACHMENTS': ''}), redirect_stdout(io.StringIO()):
            result = run_agent_loop([{'role': 'system', 'content': 'test'},
                                     {'role': 'user', 'content': self.prompt}],
                'test', RuntimePolicy(), 4, tools=get_sandbox_tools({'aggregate_spreadsheet'}),
                spreadsheet_analysis_required=True)
        return result, model

    def test_missing_margin_is_repaired_by_a_second_model_tool_call(self):
        result, model = self.run_loop([self.call('read-actual', self.query),
                                       self.call('calculate-margin', self.margin)])
        self.assertEqual(model.call_count, 2)
        self.assertEqual(result.telemetry.tool_invocations, 2)
        self.assertEqual([r.provenance['tool_call_id'] for r in result.execution_evidence],
                         ['read-actual', 'calculate-margin'])
        self.assertTrue(all(r.provenance['execution_origin'] == 'model_tool_call'
                            for r in result.execution_evidence))
        self.assertIn('43.09%', result.final_text)
        self.assertNotIn('85.65%', result.final_text)
        self.assertNotIn('注册关系补算', result.final_text)
        feedback = [m['content'] for m in result.messages if m.get('role') == 'user']
        self.assertTrue(any('显式调用' in text and 'gross_margin' in text for text in feedback))

    def test_missing_metric_cannot_be_filled_by_the_legacy_executor(self):
        evidence = [self.execute(self.query)]
        before = copy.deepcopy(evidence[0].to_dict())
        telemetry = TelemetryStats()
        self.assertEqual(RegisteredMetricCompletion().complete(
            evidence, self.prompt, telemetry, ['aggregate_spreadsheet']), [])
        self.assertEqual(len(evidence), 1)
        self.assertEqual(evidence[0].to_dict(), before)
        self.assertEqual(telemetry.tool_invocations, 0)
        self.assertFalse(analysis_task_state(evidence, self.prompt)['complete'])
        self.assertIn('显式调用', task_feedback(evidence, self.prompt))

    def test_prose_cannot_trigger_summary_calculation_or_export(self):
        before = set(self.folder.iterdir())
        evidence, outbound = [], []
        self.assertIsNone(recover_summary('按项目汇总金额，生成Excel。', self.folder))
        self.assertFalse(SummaryRecovery('统计合计', self.folder).apply(
            evidence, TelemetryStats(), outbound))
        self.assertEqual(set(self.folder.iterdir()), before)
        self.assertEqual(evidence, [])
        self.assertEqual(outbound, [])

    def test_presentation_never_adds_totals_or_mixes_metrics(self):
        for facts, unwanted in [
            ([{'id': 'a', 'label': '甲毛利率', 'value': '40', 'unit': '%', 'group_key': '甲'},
              {'id': 'b', 'label': '乙毛利率', 'value': '60', 'unit': '%', 'group_key': '乙'}], '100%'),
            ([{'id': 'a', 'label': '甲收入', 'value': '1000', 'unit': '千元', 'group_key': '甲'},
              {'id': 'b', 'label': '甲净利润', 'value': '100', 'unit': '千元', 'group_key': '甲'},
              {'id': 'c', 'label': '乙收入', 'value': '2000', 'unit': '千元', 'group_key': '乙'},
              {'id': 'd', 'label': '乙净利润', 'value': '200', 'unit': '千元', 'group_key': '乙'}], '3,300')]:
            receipt = ToolResult.success(data={'kind': 'spreadsheet_analysis', 'facts': facts})
            with patch('dsh_modules.report_presentation.resolve_conforming_facts', return_value=facts):
                rendered = analysis_overview([receipt], '各组和总计')
            self.assertNotIn('**各组总计**', rendered)
            self.assertNotIn(unwanted, rendered)

    def test_trends_hook_does_not_derive_new_values(self):
        self.assertEqual(render_ratio_trends([self.execute(self.margin)], str), [])

    def test_explicit_difference_plan_returns_percentage_points(self):
        plan = copy.deepcopy(self.margin)
        actual = plan['calculations'][0]
        budget = copy.deepcopy(actual)
        budget.update(id='budget_margin', label='2026预算毛利率')
        for arg in budget['expression']['args']:
            arg['cell'] = arg['cell'].replace('B', 'C')
        plan['calculations'].extend([budget, {
            'id': 'margin_change', 'label': '毛利率变化', 'output_unit': '%',
            'expression': {'op': 'subtract', 'args': [{'ref': 'margin'}, {'ref': 'budget_margin'}]}}])
        receipt = self.execute(plan)
        change = receipt.data['facts'][-1]
        self.assertEqual(change['unit'], '百分点')
        self.assertAlmostEqual(float(change['value']), 1.4257781)
        rendered = analysis_overview([receipt], '比较毛利率')
        self.assertIn('1.43 百分点', rendered)

    def test_missing_requested_total_is_feedback_until_an_explicit_total_plan(self):
        source = self.folder / 'groups.xlsx'
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = '明细'
        for row in [('中心', '金额'), ('甲', 7), ('乙', 9)]:
            sheet.append(row)
        book.save(source)
        book.close()
        group = calculate_spreadsheet_analysis(source, {'query': {
            'source': {'sheet': '明细', 'header_row': 1, 'data_range': 'A2:B3'},
            'group_by': ['中心'], 'metrics': [{'id': 'amount', 'label': '金额',
                'field': '金额', 'op': 'sum', 'unit': 'source'}]}})
        prompt = '按中心汇总金额，附总计。'
        self.assertTrue(group.is_success, group.error)
        state = analysis_task_state([group], prompt)
        self.assertIn('required_totals_missing', [i['code'] for i in state['issues']])
        total = calculate_spreadsheet_analysis(source, {'tables': [{
            'id': 't', 'sheet': '明细', 'header_row': 1, 'data_range': 'A2:B3'}],
            'calculations': [{'id': 'total', 'label': '总金额', 'expression': {
                'aggregate': 'sum', 'table': 't', 'column': 'B', 'unit': 'number'}}]})
        self.assertTrue(total.is_success, total.error)
        self.assertTrue(analysis_task_state([group, total], prompt)['complete'])

    def test_unexecuted_ratio_in_synthesis_is_rejected_and_executed_ratio_is_allowed(self):
        raw = self.execute(self.query)
        summary = ('净利润实际16038.3、预算13950、差额2088.3；营业利润实际21209.4、预算18300、'
                   '差额2909.4；营业收入实际124020、预算120000、差额4020。')
        self.assertFalse(_is_synthesis_grounded(summary + '全年毛利率85.65%。',
                                               [raw], prompt=self.prompt))
        self.assertFalse(_is_synthesis_grounded(summary + '实际毛利率43.09%。',
                                               [raw], prompt=self.prompt))
        self.assertTrue(_is_synthesis_grounded(summary + '实际毛利率43.09%。',
                                              [raw, self.execute(self.margin)], prompt=self.prompt))

    def test_retired_implicit_receipts_cannot_satisfy_the_gate(self):
        for origin in ['registered_derivation', 'bounded_summary_recovery']:
            receipt = self.execute(self.margin)
            receipt.provenance['execution_origin'] = origin
            state = analysis_task_state([self.execute(self.query), receipt], self.prompt)
            self.assertFalse(state['complete'])
            self.assertIn('毛利率', [m['name'] for m in state['missing_metrics']])

    def test_prose_python_is_not_executed(self):
        script = "```python\nimport openpyxl\nbook = openpyxl.Workbook()\nbook.save('unexpected.xlsx')\n```"
        with patch('dsh_modules.deliverable_contract.subprocess.run') as execute:
            output, feedback = materialize_script_deliverable(script, str(self.folder), ['.xlsx'])
        execute.assert_not_called()
        self.assertIsNone(output)
        self.assertIn('显式调用', feedback)
        self.assertFalse((self.folder / 'unexpected.xlsx').exists())

    def test_grounded_missing_column_explanation_is_delivered_with_guidance(self):
        from dsh_modules.analysis_turn import AnalysisTurn
        from dsh_modules.tool_result import ToolResult
        turn = AnalysisTurn('总结毛利率', model='test')
        model_reply = (
            '依据当前表格，营业收入事实总额为 124,020 千元，业务整体规模保持稳定。'
            '但由于原表中未包含计算毛利率所需的营业成本列，当前无法直接得出毛利率数值。'
            '建议您补充上传包含成本明细的完整利润表，以便为您继续核算毛利率。'
        )
        receipt = ToolResult.success(data={'kind': 'spreadsheet_analysis', 'facts': [
            {'label': '营业收入', 'value': '124020', 'unit': '千元'}
        ]})
        rendered = turn.render([receipt], model_reply, TelemetryStats())
        self.assertIn('营业收入事实总额为 124,020 千元', rendered)
        self.assertIn('未包含计算毛利率所需的营业成本列', rendered)
        self.assertIn('建议您补充上传包含成本明细', rendered)
        self.assertIn('数据核查与未完成指标说明', rendered)
        self.assertIn('下一步处理建议', rendered)

    def test_zero_receipt_missing_data_explanation_is_delivered_with_notice(self):
        from dsh_modules.analysis_turn import AnalysisTurn
        turn = AnalysisTurn('总结毛利率', model='test')
        model_reply = (
            '已全面核对您上传的表格。原表格中仅有销售额，未包含计算毛利率所需的营业成本列，'
            '当前无法直接核算毛利率。建议您补充提供成本数据，以便完成完整核算。'
        )
        rendered = turn.render([], model_reply, TelemetryStats())
        self.assertIn('未包含计算毛利率所需的营业成本列', rendered)
        self.assertIn('建议您补充提供成本数据', rendered)

    def test_pure_clarification_is_delivered_directly(self):
        from dsh_modules.analysis_turn import AnalysisTurn
        turn = AnalysisTurn('总结毛利率', model='test')
        clarify_reply = (
            '请确认您所需的毛利率计算口径：表中未找到明确的成本列，'
            '请选择是否按默认的综合毛利率标准继续测算，或请补充提供成本明细表。'
        )
        rendered = turn.render([], clarify_reply, TelemetryStats())
        self.assertIn('请确认您所需的毛利率计算口径', rendered)
        self.assertIn('请选择是否按默认的综合毛利率标准继续测算', rendered)

    def test_fabricated_number_in_prose_is_still_rejected(self):
        from dsh_modules.analysis_turn import AnalysisTurn
        turn = AnalysisTurn('总结毛利率', model='test')
        fake_reply = '虽然没有计算工具，但我估算实际毛利率为 45.67%，属于行业上游。'
        rendered = turn.render([], fake_reply, TelemetryStats())
        self.assertNotIn('45.67%', rendered)
        self.assertIn('本次分析尚未通过执行证据核验', rendered)

    def test_post_query_calculation_with_annotated_ref_succeeds(self):
        plan = {
            'query': {
                'source': {'sheet': '利润表', 'header_row': 2, 'data_range': 'A3:D6'},
                'group_by': ['项目'],
                'metrics': [{'field': '2026实际', 'id': 'val', 'label': '2026实际', 'op': 'sum', 'unit': 'amount'}]
            },
            'calculations': [
                {
                    'id': 'gross_margin',
                    'label': '毛利率',
                    'output_unit': '%',
                    'expression': {
                        'op': 'divide',
                        'unit': 'ratio',
                        'args': [
                            {'ref': "val__(('str', '毛利'),)", 'unit': 'amount'},
                            {'ref': "val__(('str', '营业收入'),)", 'unit': 'amount'}
                        ]
                    }
                }
            ]
        }
        receipt = self.execute(plan)
        facts = receipt.data['facts']
        margin_fact = next((f for f in facts if f['id'] == 'gross_margin'), None)
        self.assertIsNotNone(margin_fact)
        self.assertEqual(margin_fact['unit'], '%')
        # 53443.25 / 124020 ≈ 43.09%
        self.assertAlmostEqual(float(margin_fact['value']), 43.0924447669)

    def test_unresolved_ref_provides_clear_diagnostic_error(self):
        plan = {
            'query': {
                'source': {'sheet': '利润表', 'header_row': 2, 'data_range': 'A3:D6'},
                'group_by': ['项目'],
                'metrics': [{'field': '2026实际', 'id': 'val', 'label': '2026实际', 'op': 'sum', 'unit': 'amount'}]
            },
            'calculations': [
                {
                    'id': 'gross_margin',
                    'label': '毛利率',
                    'output_unit': '%',
                    'expression': {
                        'op': 'divide',
                        'unit': 'ratio',
                        'args': [
                            {'ref': 'nonexistent_fact_id', 'unit': 'amount'},
                            {'ref': "val__(('str', '营业收入'),)", 'unit': 'amount'}
                        ]
                    }
                }
            ]
        }
        res = calculate_spreadsheet_analysis(self.source, plan, domain='finance')
        self.assertTrue(res.is_error)
        self.assertIn('未知计算引用 [nonexistent_fact_id]', res.error_message)


if __name__ == '__main__':
    unittest.main()
