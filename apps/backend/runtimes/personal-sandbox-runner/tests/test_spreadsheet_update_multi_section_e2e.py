"""
End-to-end verification for multi-section spreadsheet update and title alias resolution:
1. Multi-section planning via sections list or comma-separated string.
2. Title alias resolution (e.g. '无偿' mapping to '*P/*Y/EX（不含*L/*B）').
3. Boundary detection stopping cleanly before subsequent tables.
4. Full execution delivering verified workbook with both sections updated.
"""

import copy
import json
import os
import shutil
import tempfile
import unittest
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch

import openpyxl

from dsh_modules.spreadsheet_update.result_catalog import (
    RESULTS, fingerprint, reset_results, harvest_analytical_facts
)
from dsh_modules.spreadsheet_update.tools import (
    reset_turn_results,
    tool_plan_spreadsheet_update,
    tool_execute_spreadsheet_update
)
from dsh_modules.spreadsheet_update.semantic_plan import build_semantic_plan
from dsh_modules.user_declarations import bind_user_declarations
import zipfile
import re


class MultiSectionSpreadsheetUpdateTests(unittest.TestCase):
    @staticmethod
    def _simulate_recalculate_candidate(file_path, output_path=None, force=False):
        dest = Path(output_path).resolve() if output_path else Path(file_path).resolve()
        if str(dest) != str(Path(file_path).resolve()):
            dest.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(file_path, dest)

        tmp_unzip = dest.parent / f"tmp_recalc_{dest.stem}"
        tmp_unzip.mkdir(parents=True, exist_ok=True)
        try:
            with zipfile.ZipFile(dest, "r") as zin:
                zin.extractall(tmp_unzip)
            for s_xml in (tmp_unzip / "xl" / "worksheets").glob("sheet*.xml"):
                content = s_xml.read_text("utf-8")
                content = re.sub(r'(<c [^>]*t="e"[^>]*><f>[^<]+</f><v>)(</v>)', r'\1#DIV/0!\2', content)
                content = content.replace("<v></v>", "<v>100</v>")
                content = re.sub(r'(<f>[^<]+</f>)(?!<v>)', r'\1<v>100</v>', content)
                s_xml.write_text(content, "utf-8")
            with zipfile.ZipFile(dest, "w", zipfile.ZIP_DEFLATED) as zout:
                for p in tmp_unzip.rglob("*"):
                    if p.is_file():
                        zout.write(p, p.relative_to(tmp_unzip))
        finally:
            shutil.rmtree(tmp_unzip, ignore_errors=True)

        from dsh_modules.workbook_calculation import check_workbook_formulas, RECALC_CACHE_VERSION
        f_check = check_workbook_formulas(dest, is_post_recalc=True)
        meta = {
            "schema_version": 1,
            "status": "verified" if f_check.get("total_errors", 0) == 0 else "partial",
            "engine_execution": "success",
            "provenance": {"workbook_id": "mock", "engine": "libreoffice_calc_headless", "validator_version": RECALC_CACHE_VERSION},
            **f_check
        }
        return dest, meta
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.ws = Path(self.temp.name).resolve()
        self.target = self.ws / '工数3Q预实管理.xlsx'
        self.source = self.ws / '工数实绩明细.xlsx'

        # Build realistic workbook mimicking user's actual sheet
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = '1019'

        # Section 1: 有偿
        ws['A1'] = '有偿'
        ws['A2'] = '部门'
        ws['C3'] = '202610'
        ws['C4'] = '预测'
        ws['D4'] = '计划'
        ws['B5'] = 'FC48005'; ws['C5'] = 0; ws['D5'] = 100
        ws['B6'] = 'FC34032'; ws['C6'] = 0; ws['D6'] = 50
        ws['B7'] = 'Subtotal'; ws['C7'] = '=SUM(C5:C6)'
        ws['B8'] = 'FC46005'; ws['C8'] = 0; ws['D8'] = 200
        ws['B9'] = 'Subtotal'; ws['C9'] = '=C8'
        ws['A10'] = '总计'; ws['C10'] = '=C7+C9'

        # Section 2: Non-billable (*P/*Y/EX（不含*L/*B）)
        ws['A12'] = '*P/*Y/EX（不含*L/*B）'
        ws['A13'] = '部门'
        ws['C14'] = '202610'
        ws['C15'] = '预测'
        ws['D15'] = '计划'
        ws['B16'] = 'FC48005'; ws['C16'] = 0; ws['D16'] = 30
        ws['B17'] = 'FC46005'; ws['C17'] = 0; ws['D17'] = 40
        ws['B18'] = 'Subtotal'; ws['C18'] = '=SUM(C16:C17)'
        ws['A19'] = '总计'; ws['C19'] = '=C18'

        # Subsequent separate table (e.g. utilization rates)
        ws['B22'] = 'FC48005'; ws['C22'] = 0.8
        ws['B23'] = 'FC46005'; ws['C23'] = 0.9

        wb.save(self.target)
        wb.close()

        # Build source detail workbook
        wb_src = openpyxl.Workbook()
        ws_src = wb_src.active
        ws_src.title = '工数实绩明细'
        ws_src.append(['工数年月', '部门编号', '有偿无偿', '金额(K)'])
        ws_src.append(['202610', 'FC48005', '有偿', 67.0353])
        ws_src.append(['202610', 'FC46005', '有偿', 314.48323])
        ws_src.append(['202610', 'FC48005', '无偿', 15.5])
        ws_src.append(['202610', 'FC46005', '无偿', 14.3])
        wb_src.save(self.source)
        wb_src.close()

        reset_turn_results()
        bind_user_declarations({'unit_confirmation': {'unit': 'K', 'quote': '以K为单位'},
                                'user_statements': ['以K为单位']})

        # Register analytical facts for paid
        receipt_paid = {
            'kind': 'spreadsheet_analysis', 'evidence_id': 'result_paid',
            'plan': {
                'where': {'and': [{'field': '工数年月', 'op': 'eq', 'value': '202610'},
                                  {'field': '有偿无偿', 'op': 'eq', 'value': '有偿'}]},
                'group_by': ['部门编号'],
                'metrics': [{'id': 'amt', 'field': '金额(K)', 'op': 'sum', 'unit': 'source'}]
            },
            'facts': [
                {'id': 'amt__FC48005', 'group_fields': {'部门编号': 'FC48005'}, 'value': '67.0353', 'unit': 'K'},
                {'id': 'amt__FC46005', 'group_fields': {'部门编号': 'FC46005'}, 'value': '314.48323', 'unit': 'K'}
            ],
            'coverage': [{'coverage_complete': True}]
        }
        # Register analytical facts for unpaid
        receipt_unpaid = {
            'kind': 'spreadsheet_analysis', 'evidence_id': 'result_unpaid',
            'plan': {
                'where': {'and': [{'field': '工数年月', 'op': 'eq', 'value': '202610'},
                                  {'field': '有偿无偿', 'op': 'eq', 'value': '无偿'}]},
                'group_by': ['部门编号'],
                'metrics': [{'id': 'amt', 'field': '金额(K)', 'op': 'sum', 'unit': 'source'}]
            },
            'facts': [
                {'id': 'amt__FC48005', 'group_fields': {'部门编号': 'FC48005'}, 'value': '15.5', 'unit': 'K'},
                {'id': 'amt__FC46005', 'group_fields': {'部门编号': 'FC46005'}, 'value': '14.3', 'unit': 'K'}
            ],
            'coverage': [{'coverage_complete': True}]
        }
        prov = {'file_path': str(self.source), 'workbook_id': fingerprint(self.source)}
        harvest_analytical_facts(receipt_paid, provenance=prov)
        harvest_analytical_facts(receipt_unpaid, provenance=prov)

        self.patch_env = patch.dict(os.environ, {'WORKSPACE_DIR': str(self.ws)})
        self.patch_env.start()
        self.patch_config = patch('dsh_modules.config.WORKSPACE_DIR', str(self.ws))
        self.patch_config.start()

    def tearDown(self):
        self.patch_config.stop()
        self.patch_env.stop()
        reset_turn_results()
        self.temp.cleanup()

    def test_multi_section_planning_with_alias_resolution(self):
        req = {
            'target_file': '工数3Q预实管理.xlsx',
            'base_sheet': '1019',
            'new_sheet': '1023',
            'period': '202610',
            'period_field': '工数年月',
            'target_label': '预测',
            'target_kind': 'forecast',
            'target_unit': 'K',
            'sections': [
                {'section': '有偿', 'result_id': 'result_paid'},
                {'section': '无偿', 'result_id': 'result_unpaid'}
            ]
        }
        plan, mapped, payload, preview, warnings = build_semantic_plan(req, self.ws)
        # 2 paid writes + 2 unpaid writes = 4 writes
        self.assertEqual(len(preview), 4)
        preview_map = {p['coordinate']: p['value'] for p in preview}
        # Paid block
        self.assertEqual(preview_map['C5'], '67.0353')
        self.assertEqual(preview_map['C8'], '314.48323')
        # Unpaid block (mapped to row 16 & 17 via alias resolution)
        self.assertEqual(preview_map['C16'], '15.5')
        self.assertEqual(preview_map['C17'], '14.3')
        self.assertEqual(len(plan['semantic_bindings']['region_specs']), 2)

    def test_multi_section_auto_matching_with_comma_separated_sections(self):
        req = {
            'target_file': '工数3Q预实管理.xlsx',
            'base_sheet': '1019',
            'new_sheet': '1023',
            'period': '202610',
            'period_field': '工数年月',
            'target_label': '预测',
            'target_kind': 'forecast',
            'target_unit': 'K',
            'section': '有偿, 无偿'
        }
        plan, mapped, payload, preview, warnings = build_semantic_plan(req, self.ws)
        self.assertEqual(len(preview), 4)
        preview_map = {p['coordinate']: p['value'] for p in preview}
        self.assertEqual(preview_map['C5'], '67.0353')
        self.assertEqual(preview_map['C8'], '314.48323')
        self.assertEqual(preview_map['C16'], '15.5')
        self.assertEqual(preview_map['C17'], '14.3')

    @patch('dsh_modules.spreadsheet_update.executor.recalculate_candidate')
    def test_multi_section_end_to_end_execute(self, mock_recalc):
        mock_recalc.side_effect = self._simulate_recalculate_candidate

        req = {
            'target_file': '工数3Q预实管理.xlsx',
            'base_sheet': '1019',
            'new_sheet': '1023',
            'period': '202610',
            'period_field': '工数年月',
            'target_label': '预测',
            'target_kind': 'forecast',
            'target_unit': 'K',
            'sections': [
                {'section': '有偿', 'result_id': 'result_paid'},
                {'section': '无偿', 'result_id': 'result_unpaid'}
            ]
        }
        plan_res = tool_plan_spreadsheet_update(req)
        self.assertTrue(plan_res.is_success, str(plan_res))
        plan_id = plan_res.data['plan_id']

        exec_res = tool_execute_spreadsheet_update(plan_id)
        self.assertTrue(exec_res.is_success, str(exec_res))
        self.assertEqual(exec_res.data['status'], 'published')
        out_file = Path(exec_res.data['published_file'])
        self.assertTrue(out_file.exists())

        # Inspect resulting workbook
        wb_out = openpyxl.load_workbook(out_file, data_only=True)
        self.assertIn('1023', wb_out.sheetnames)
        ws_1023 = wb_out['1023']
        # Check both sections are updated
        self.assertAlmostEqual(float(ws_1023['C5'].value), 67.0353, places=3)
        self.assertAlmostEqual(float(ws_1023['C8'].value), 314.48323, places=3)
        self.assertAlmostEqual(float(ws_1023['C16'].value), 15.5, places=3)
        self.assertAlmostEqual(float(ws_1023['C17'].value), 14.3, places=3)
        wb_out.close()


if __name__ == '__main__':
    unittest.main()
