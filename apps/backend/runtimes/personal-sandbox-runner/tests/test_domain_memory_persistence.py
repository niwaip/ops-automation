import os
import tempfile
import unittest
from pathlib import Path

from dsh_modules.domain_memory import (
    save_section_mapping,
    resolve_section_mapping,
    clear_domain_memory,
    _get_memory_path,
    _IN_MEMORY_STORE
)
from dsh_modules.spreadsheet_update.semantic_plan import _extract_where_values
from dsh_modules.spreadsheet_update.result_catalog import RESULTS


class TestDomainMemoryPersistence(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.workspace = self.temp_dir.name
        self.orig_ws = os.environ.get("WORKSPACE")
        self.orig_ws_dir = os.environ.get("WORKSPACE_DIR")
        os.environ["WORKSPACE"] = self.workspace
        if "WORKSPACE_DIR" in os.environ:
            del os.environ["WORKSPACE_DIR"]
        clear_domain_memory(self.workspace)

    def tearDown(self):
        clear_domain_memory(self.workspace)
        if self.orig_ws is not None:
            os.environ["WORKSPACE"] = self.orig_ws
        elif "WORKSPACE" in os.environ:
            del os.environ["WORKSPACE"]
        if self.orig_ws_dir is not None:
            os.environ["WORKSPACE_DIR"] = self.orig_ws_dir
        self.temp_dir.cleanup()

    def test_memory_path_with_workspace_env_only(self):
        path = _get_memory_path()
        self.assertIsNotNone(path)
        expected = Path(self.workspace).resolve() / ".dsh" / "memory" / "domain_knowledge.json"
        self.assertEqual(path.resolve(), expected)

    def test_persistence_across_in_memory_resets(self):
        # Save mapping
        save_section_mapping("无偿", "*P/*Y/EX（不含*L/*B）", quote="【*P/Y/EX（不含L/*B）】 就是无偿")
        file_path = _get_memory_path()
        self.assertTrue(file_path.exists())

        # Clear in-memory dictionary explicitly to simulate fresh process
        _IN_MEMORY_STORE.clear()

        # Should load from persisted file
        resolved = resolve_section_mapping("无偿")
        self.assertEqual(resolved, "*P/*Y/EX（不含*L/*B）")

    def test_result_disambiguation_with_metric_labels(self):
        # Setup mock results
        r_paid = {
            "data": {
                "plan": {
                    "where": {"field": "工数年月", "op": "eq", "value": "202610"},
                    "metrics": [{"id": "m1", "label": "有偿金额(K)"}]
                }
            }
        }
        r_unpaid = {
            "data": {
                "plan": {
                    "where": {"field": "工数年月", "op": "eq", "value": "202610"},
                    "metrics": [{"id": "m1", "label": "无偿金额(K)"}]
                }
            }
        }
        RESULTS["res_paid"] = r_paid
        RESULTS["res_unpaid"] = r_unpaid

        # Verify descriptors matching
        def match(sec_name):
            candidate_names = {sec_name}
            matched = None
            for r_id, r in RESULTS.items():
                plan_obj = r.get("data", {}).get("plan", {})
                plan_where = plan_obj.get("where", {})
                where_vals = _extract_where_values(plan_where)
                metric_labels = {
                    str(m.get("label", "")).strip()
                    for m in plan_obj.get("metrics", [])
                    if isinstance(m, dict)
                }
                for f in r.get("data", {}).get("filters", []):
                    if isinstance(f, dict) and "value" in f:
                        where_vals.add(str(f["value"]).strip())
                all_descriptors = where_vals | metric_labels
                if any(any(c in desc or desc in c for desc in all_descriptors if desc) for c in candidate_names):
                    matched = r_id
                    break
            return matched

        self.assertEqual(match("有偿"), "res_paid")
        self.assertEqual(match("无偿"), "res_unpaid")

    def test_memory_fallback_to_general_when_file_hash_unregistered(self):
        # 验证用户确认的记忆在遇到全新未注册的 file_hash 时，能够平滑从 general 回退命中，绝不会失忆
        save_section_mapping("无偿", "*P/*Y/EX（不含*L/*B）", quote="【*P/Y/EX（不含L/*B）】")
        # Querying with an unregistered new file_hash should fallback to general
        resolved = resolve_section_mapping("无偿", file_hash="new_unseen_file_sha256")
        self.assertEqual(resolved, "*P/*Y/EX（不含*L/*B）")

    def test_monthly_excel_inherits_template_mapping_despite_different_hash(self):
        # 需求1：每月 Excel 哈希不同但格式相同，通过模版特征自动继承记忆
        # 10月在工数3Q预实管理_1024.xlsx (hash10) 下保存确认
        save_section_mapping(
            section_name="无偿",
            target_title="*P/*Y/EX（不含*L/*B）",
            file_name="工数3Q预实管理_1024.xlsx",
            file_hash="sha256_october_data",
            quote="【*P/Y/EX（不含L/*B）】"
        )
        # 11月上传了全新数据文件 工数3Q预实管理_1120.xlsx (hash11)，哈希完全变了
        resolved_nov = resolve_section_mapping(
            section_name="无偿",
            file_name="工数3Q预实管理_1120.xlsx",
            file_hash="sha256_november_different_data"
        )
        self.assertEqual(resolved_nov, "*P/*Y/EX（不含*L/*B）")

    def test_scenario_isolation_between_different_templates(self):
        # 需求2：记忆必须绑定特定场景/文件，不同业务模版互不污染
        save_section_mapping(
            section_name="无偿",
            target_title="*P/*Y/EX（不含*L/*B）",
            template_name="工数3Q预实管理",
            quote="工数专属"
        )
        # 在完全不相干的财务模版中查询，不能随意套用工数模版的专属记忆
        from dsh_modules.domain_memory import load_domain_knowledge
        fin_knowledge = load_domain_knowledge(file_name="财务成本损益表_202611.xlsx")
        # 财务模版的专属映射中不应有该工数特定绑定
        self.assertNotIn("无偿", fin_knowledge.get("section_mappings", {}))

    def test_strict_exact_match_gate_in_locate(self):
        # 需求4：记忆只是补足信息，必须在当前 sheet 完全匹配且唯一才作为入力，否则抛出确认
        import openpyxl
        from dsh_modules.spreadsheet_update.layout_binding import locate, BindingNeedsInput
        
        # 1. 创建模拟工作簿
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "1019"
        
        # 放入表头和键列，但目标区块标题故意不是记忆里的 "*P/*Y/EX（不含*L/*B）"，而是全新未见的 "未知新区块"
        ws["A1"] = "工数预实"
        ws["A3"] = "未知新区块"
        ws["A4"] = "部门编号"
        ws["B4"] = "202610"
        ws["B5"] = "预测"
        ws["A6"] = "FC001"
        ws["B6"] = 100
        
        # 记忆中存有 "无偿 -> *P/*Y/EX（不含*L/*B）"
        save_section_mapping("无偿", "*P/*Y/EX（不含*L/*B）", quote="旧记忆")
        
        # 尝试 locate "无偿"：因为当前 Sheet 中根本不存在 "*P/*Y/EX（不含*L/*B）"，记忆绝对不能强行生效！
        # 必须抛出 BindingNeedsInput 找用户确认！
        with self.assertRaises(BindingNeedsInput) as ctx:
            locate(
                ws, section="无偿", period="202610", target_label="预测",
                source_keys={"FC001"}, file_name="工数3Q预实管理_1120.xlsx"
            )
        self.assertIn("未知新区块", str(ctx.exception))




if __name__ == "__main__":
    unittest.main()
