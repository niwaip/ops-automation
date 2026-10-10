"""Engineering boundaries independent of workbook names and financial answers."""

import hashlib
import io
import json
import os
import sys
import tempfile
import time
import unittest
from contextlib import redirect_stdout
from pathlib import Path
from unittest.mock import patch

import openpyxl

sys.path.insert(0,str(Path(__file__).resolve().parent.parent/'src'))
from dsh_modules.workbook_reader import inspect_workbook
from dsh_modules.worksheet_metadata import worksheet_dimensions
from dsh_modules.spreadsheet_context import (
    spreadsheet_structure, build_spreadsheet_source_context, execute_spreadsheet_structure, STRUCTURE_TOOL_NAME)
from dsh_modules.spreadsheet_analysis import calculate_spreadsheet_analysis
from dsh_modules.spreadsheet_analysis_evidence import render_analysis_response, verified_response
from dsh_modules.evidence_store import persist_turn_evidence, load_task_context
from dsh_modules.prompt_builder import build_user_turn
from dsh_modules.table_comparison import calculate_column_comparison
from dsh_modules.comparison_delivery import compile_comparison_response
from dsh_modules.analysis_contract import build_analysis_contract
from dsh_modules.agent_loop import run_agent_loop
from dsh_modules.runtime_policy import RuntimePolicy
from dsh_modules.tool_result import ToolResult
from dsh_modules.tools import get_sandbox_tools


class GeneralEngineeringTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory()
        self.path=Path(self.temp.name)/'experiments.xlsx'
        self.history=Path(self.temp.name)/'session.json'
        workbook=openpyxl.Workbook(); sheet=workbook.active; sheet.title='实验结果'
        sheet.append(['批次',None,'结果','状态'])
        sheet.append(['a',None,2,'OK']); sheet.append(['b',None,3,'OK'])
        sheet.append(['b',None,5,'REVIEW'])
        table=workbook.create_sheet('产量对照')
        table.append(['分组','产量实际','产量目标'])
        table.append(['样本甲',10,12]); table.append(['样本乙',30,35]); table.append(['样本丙',40,35])
        workbook.save(self.path); workbook.close()
        self.table={'id':'results','sheet':'实验结果','header_row':1,'data_range':'A2:D4'}

    def tearDown(self):
        self.temp.cleanup()

    def hash(self):
        return hashlib.sha256(self.path.read_bytes()).hexdigest()

    def compute(self, table=None, label='观测合计'):
        selected=table or self.table
        return calculate_spreadsheet_analysis(self.path,{'tables':[selected],'calculations':[
            {'id':'sum','label':label,'expression':{
                'aggregate':'sum','table':selected['id'],'column':'C','unit':'number'}}]})

    def persist(self, receipts, query='观测结果', verified=True):
        return persist_turn_evidence(self.history,{self.path.name:self.hash()},query,receipts,
            [{'guard':'spreadsheet_evidence','action':'accept' if verified else 'reject'}])

    def comparison(self, **changes):
        params={'sheet':'产量对照','data_range':'A2:C4','header_row':1,'group_column':'A',
                'comparisons':[{'label':'产量偏差','left_column':'B','right_column':'C'}],
                'basis':'custom','rank_by':'absolute_difference','extreme':'max',
                'expected_groups':['样本甲','样本乙','样本丙']}
        return calculate_column_comparison(self.path,**{**params,**changes})

    def test_streaming_dimensions_and_sparse_header_positions(self):
        manifest=inspect_workbook(self.path)
        meta=manifest.sheets['实验结果']
        self.assertEqual(meta.dimensions,'A1:D4')
        self.assertEqual(meta.headers,['批次','','结果','状态'])
        self.assertEqual(meta.sample_rows[0],['a','','2','OK'])
        workbook=openpyxl.load_workbook(self.path,read_only=True)
        try:self.assertEqual(worksheet_dimensions(workbook.active),'A1:D4')
        finally:workbook.close()

    def test_missing_producer_dimension_is_computed(self):
        workbook=openpyxl.load_workbook(self.path,read_only=True)
        try:
            workbook.active.reset_dimensions()
            self.assertEqual(worksheet_dimensions(workbook.active),'A1:D4')
        finally:workbook.close()

    def test_structure_expansion_is_targeted_and_not_execution_proof(self):
        compact=build_spreadsheet_source_context(self.path)
        expanded=spreadsheet_structure(self.path,['实验结果'])
        self.assertEqual([s['sheet'] for s in expanded.data['sheets']],['实验结果'])
        self.assertEqual(expanded.data['available_sheets'],['实验结果','产量对照'])
        self.assertFalse(expanded.data['coverage_complete'])
        self.assertIsNone(verified_response([expanded]))
        self.assertNotIn('nonempty_rows_including_titles_and_headers',compact)
        self.assertGreater(len(expanded.data['sheets'][0]['sample_cells']),1)

    def test_structure_rejects_guessed_sheet_and_expired_deadline(self):
        self.assertEqual(spreadsheet_structure(self.path,['汇总表']).error_code,'invalid_sheet_selection')
        self.assertEqual(spreadsheet_structure(self.path,[]).error_code,'invalid_sheet_selection')
        with self.assertRaises(TimeoutError):spreadsheet_structure(self.path,deadline=time.monotonic()-1)

    def test_structure_adapter_keeps_attachment_scope(self):
        with patch('dsh_modules.spreadsheet_context.resolve_sandboxed_path',return_value=(self.path,None)), \
             patch.dict(os.environ,{'DSH_SESSION_ATTACHMENTS':'different.xlsx'}), \
             patch('dsh_modules.config.WORKSPACE_DIR',self.temp.name):
            result=execute_spreadsheet_structure({'file_path':str(self.path),'sheets':['实验结果']})
        self.assertEqual(result.tool_result.error_code,'source_outside_session_scope')

    def test_prompt_uses_actual_capabilities(self):
        readonly=build_user_turn('核算记录',is_inspect_intent=True,available_tool_names=['aggregate_spreadsheet'])
        executable=build_user_turn('核算记录',is_inspect_intent=True,available_tool_names=['bash'])
        self.assertNotIn('本轮允许 bash',readonly)
        self.assertIn('不得调用未提供的工具',readonly)
        self.assertIn('[Available Capabilities]: aggregate_spreadsheet',readonly)
        self.assertIn('本轮允许 bash',executable)

    def test_archive_retains_prior_turn_and_failures_after_history_replacement(self):
        success=self.compute(); first=self.persist([success],query='首次核算')
        error=ToolResult.error('invalid_analysis_plan','字段无法定位',data={'requested_plan':{
            'tables':[self.table],'calculations':[{'id':'bad','label':'未知字段核算'}]}})
        second=self.persist([error],query='继续核算',verified=False)
        self.history.write_text('[{"role":"assistant","content":"不可信的历史数字123456"}]')
        paths=list(self.history.with_suffix('.evidence').glob('*.json'))
        self.assertEqual(len(paths),2)
        self.assertNotEqual(first['turn_id'],second['turn_id'])
        self.assertEqual(json.loads(self.history.with_suffix('.evidence.json').read_text())['turn_id'],second['turn_id'])
        context=load_task_context(self.history,{self.path.name:self.hash()})
        self.assertIn('首次核算',context);self.assertIn('未知字段核算',context)
        self.assertIn('invalid_analysis_plan',context);self.assertNotIn('123456',context)
        self.assertIn('不构成本轮完成证据',context)
        self.assertIn('"value": "10"',context)

    def test_task_context_invalidates_replaced_source(self):
        self.persist([self.compute()])
        workbook=openpyxl.load_workbook(self.path);workbook.active['C2']=999
        workbook.save(self.path);workbook.close()
        self.assertEqual(load_task_context(self.history,{self.path.name:self.hash()}),'')

    def test_task_context_also_checks_actual_receipt_source(self):
        old=self.hash();self.persist([self.compute()])
        self.path.write_bytes(b'replaced')
        self.assertEqual(load_task_context(self.history,{self.path.name:old}),'')

    def test_task_state_is_bounded_without_truncated_json(self):
        for i in range(6):self.persist([self.compute()],query=f'任务{i}')
        context=load_task_context(self.history,{self.path.name:self.hash()},max_chars=1300,max_turns=2)
        self.assertLessEqual(len(context),1300)
        states=[json.loads(row) for row in context.splitlines()[1:]]
        self.assertTrue(states);self.assertLessEqual(len(states),2)
        self.assertTrue(all(state['verified'] for state in states))

    def test_last_rejection_does_not_inherit_prior_acceptance(self):
        receipt=self.compute()
        state=persist_turn_evidence(self.history,{self.path.name:self.hash()},'未完成',[receipt],[
            {'guard':'spreadsheet_evidence','action':'accept'},
            {'guard':'spreadsheet_evidence','action':'reject'}])
        self.assertFalse(state['verified'])
        context=load_task_context(self.history,{self.path.name:self.hash()})
        self.assertNotIn('"value": "10"',context)

    def test_persistence_error_is_visible_and_temporary_file_removed(self):
        with patch('dsh_modules.evidence_store.os.replace',side_effect=OSError('disk failure')):
            with self.assertRaises(OSError):self.persist([self.compute()])
        self.assertEqual(list(self.history.with_suffix('.evidence').glob('*.tmp')),[])

    def test_fact_deduplication_preserves_all_receipt_ids(self):
        first=self.compute();second=self.compute({**self.table,'id':'renamed'},label='重试观测合计')
        response=verified_response([first,second])
        final=render_analysis_response(response,[first,second])
        self.assertIn(first.data['evidence_id'],final);self.assertIn(second.data['evidence_id'],final)
        main, details = final.split('<details>', 1)
        self.assertEqual(main.count('| 观测合计 |'),1)
        self.assertEqual(details.count('| 观测合计 |'),1)
        self.assertNotIn('| 重试观测合计 |',final)
        self.assertEqual(len(json.loads(response)['evidence_ids']),2)

    def test_different_filters_are_not_deduplicated(self):
        first=self.compute();second=self.compute({**self.table,'filters':[{'column':'A','values':['b']}]},label='筛选合计')
        final=render_analysis_response(verified_response([first,second]),[first,second])
        self.assertIn('| 观测合计 | 10 |',final);self.assertIn('| 筛选合计 | 8 |',final)

    def test_rule_report_scopes_profiles_to_requested_field(self):
        receipt=calculate_spreadsheet_analysis(self.path,{'tables':[self.table],'rules':[
            {'id':'unique','label':'批次唯一','table':'results','column':'批次','kind':'unique','assumptions':'声明批次应唯一'}]})
        self.assertTrue(receipt.is_success,receipt.error)
        final=render_analysis_response(verified_response([receipt]),[receipt])
        self.assertIn('results.A',final);self.assertNotIn('results.C',final)
        self.assertGreater(len(receipt.data['profiles']),1)

    def test_generic_comparison_signs_ties_and_exact_values_are_compiled(self):
        receipt=self.comparison()
        rendered,error=compile_comparison_response(build_analysis_contract('比较实验结果'),[receipt])
        self.assertIsNone(error)
        self.assertIn('| 样本乙 | 30 | 35 | -5 | 5 |',rendered)
        self.assertIn('| 样本丙 | 40 | 35 | 5 | 5 |',rendered)
        self.assertEqual(receipt.data['results'][0]['matches'][0]['exact']['difference'],'-5')
        self.assertNotIn('千元',rendered)
        self.assertTrue(receipt.data['evidence_id']);self.assertTrue(receipt.data['plan_hash'])

    def test_comparison_changed_source_and_partial_result_fail_closed(self):
        receipt=self.comparison();contract=build_analysis_contract('比较实验结果')
        self.assertIsNotNone(compile_comparison_response(contract,[ToolResult.partial(data=receipt.data)])[1])
        self.path.write_bytes(b'changed')
        self.assertIn('失效',compile_comparison_response(contract,[receipt])[1])

    def test_comparison_render_deduplicates_same_receipt(self):
        receipt=self.comparison()
        rendered,error=compile_comparison_response(build_analysis_contract('比较实验结果'),[receipt,receipt])
        self.assertIsNone(error)
        main, details = rendered.split('<details>', 1)
        self.assertEqual(main.count('| 产量偏差 | 样本乙 |'),1)
        self.assertEqual(details.count('| 产量偏差 | 样本乙 |'),1)

    def test_comparison_zero_baseline_amount_report_has_undefined_rate(self):
        workbook=openpyxl.load_workbook(self.path);workbook['产量对照']['C4']=0
        workbook.save(self.path);workbook.close()
        rendered,error=compile_comparison_response(build_analysis_contract('比较实验结果'),[self.comparison()])
        self.assertIsNone(error);self.assertIn('未定义（基准为零）',rendered)

    def test_failed_comparison_retains_plan_and_pending_work(self):
        receipt=self.comparison()
        from dsh_modules.table_comparison import execute_column_comparison
        params={'file_path':str(self.path),'sheet':'不存在','data_range':'A2:C4','header_row':1,'group_column':'A',
                'comparisons':[{'label':'待核项目','left_column':'B','right_column':'C'}],
                'basis':'custom','rank_by':'absolute_difference','extreme':'max'}
        with patch('dsh_modules.table_comparison.resolve_sandboxed_path',return_value=(self.path,None)):
            failed=execute_column_comparison(params).tool_result
        self.assertEqual(failed.data['requested_plan'],params)
        report,error=compile_comparison_response(build_analysis_contract('比较实验结果'),[receipt,failed])
        self.assertIsNone(error);self.assertIn('后续计划未执行成功',report)
        self.assertIn('sheet_not_found',report)
        self.persist([receipt,failed])
        context=load_task_context(self.history,{self.path.name:self.hash()})
        self.assertIn('待核项目',context)

    def test_structure_tool_remains_available_but_does_not_satisfy_analysis(self):
        selected=get_sandbox_tools(allowed_names={STRUCTURE_TOOL_NAME,'aggregate_spreadsheet'})
        structure=spreadsheet_structure(self.path,['实验结果'])
        responses=[{'content':'','finish_reason':'tool_calls','tool_calls':[{'id':'schema','type':'function',
            'function':{'name':STRUCTURE_TOOL_NAME,'arguments':json.dumps({'file_path':str(self.path),'sheets':['实验结果']})}}]},
            {'content':'已完成全部计算','finish_reason':'stop'}]
        with patch('dsh_modules.spreadsheet_context.resolve_sandboxed_path',return_value=(self.path,None)), \
             patch.dict(os.environ,{'DSH_SESSION_ATTACHMENTS':''}), \
             patch('dsh_modules.agent_loop.call_model_proxy',side_effect=responses+[responses[-1]]*3) as model,redirect_stdout(io.StringIO()):
            result=run_agent_loop([{'role':'system','content':'test'},{'role':'user','content':'核算实验结果'}],
                'test',RuntimePolicy(),2,tools=selected,spreadsheet_analysis_required=True)
        self.assertIn('尚未通过',result.final_text)
        self.assertIsNone(verified_response([structure]))
        self.assertIn(STRUCTURE_TOOL_NAME,{t['function']['name'] for t in model.call_args_list[0].kwargs['tools']})


if __name__=='__main__':unittest.main()
