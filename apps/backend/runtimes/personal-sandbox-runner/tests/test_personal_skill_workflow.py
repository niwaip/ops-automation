"""Replay saved logic on changed data; reject missing, altered and unverifiable flows."""

import copy
import io
import json
import sys
import tempfile
import time
import unittest
from contextlib import redirect_stdout
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / 'src'))
from dsh_modules.personal_skill_workflow_engine import execute_workflow, verify_receipt
from dsh_modules.personal_skill_workflow import try_run_personal_workflow
from dsh_modules.telemetry import extract_dsh_markers


class PersonalSkillWorkflowTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(); self.root = Path(self.tmp.name)
        self.source = self.root / 'new-input.xlsx'
        self.plan = {'schema_version': 1, 'engine': 'spreadsheet-group-sum-v1',
                     'group_by': '部门编号', 'measure': '金额',
                     'filters': [{'field': 'WBS编号', 'op': 'exclude_prefix', 'values': ['ex', 'sy', 'sp', 'ip', 'ib', 'sz']}],
                     'instructions': ['按部门汇总金额并排除指定WBS前缀'], 'steps': ['绑定本轮附件', '执行并验收'],
                     'verification': {'source_hash': True, 'row_accounting': True, 'group_total': True}}
        self.write_source()
        self.gate_model = patch('dsh_modules.personal_skill_gate.call_model_proxy', side_effect=self.gate_response)
        self.gate_model.start()
        self.addCleanup(self.gate_model.stop)

    def gate_response(self, messages, **kwargs):
        payload = json.loads(messages[-1]['content'])
        if 'execution' in payload:
            result = {'status': 'accept', 'evidence_id': payload['execution']['evidence_id'],
                      'covered': ['r1'], 'reason': '', 'question': ''}
        else:
            result = {'mode': 'query', 'requirements': [{'id': 'r1', 'quote': payload['request'], 'target': '/metrics'}],
                      'query': payload['saved_query_template'], 'question': ''}
        return {'content': json.dumps(result), 'finish_reason': 'stop'}

    def tearDown(self):
        self.tmp.cleanup()

    def write_source(self, amount='20.20000'):
        w = openpyxl.Workbook(); s = w.active; s.title = '任意新表名'
        s.append(['另一份业务明细']); s.append([])
        s.append(['有偿无偿', 'WBS编号', '金额(K)', '部门编号'])
        s.append(['无偿', 'ok-1', '10.10000', 'A'])
        s.append(['有偿', 'OK-2', amount, 'A'])
        s.append(['有偿', 'EX-1', 999, 'B'])
        s.append(['有偿', 'ip-1', 999, 'B'])
        s.append(['有偿', 'ok-3', '-0.30000', 'B'])
        w.save(self.source); w.close()

    def test_full_replay_new_columns_names_and_data_preserves_exact_rules(self):
        r = execute_workflow(self.source, self.plan)
        self.assertEqual(r['total'], '30.00000')
        self.assertEqual((r['source_rows'], r['included_rows'], r['excluded_rows']), (5, 3, 2))
        self.assertEqual(r['field_mapping']['金额'], 'C')
        self.write_source('40.20000')
        changed = execute_workflow(self.source, self.plan)
        self.assertEqual(changed['total'], '50.00000')
        self.assertNotEqual(changed['workbook_id'], r['workbook_id'])
        with self.assertRaisesRegex(ValueError, '摘要不一致'):
            verify_receipt(r, self.plan, self.source)

    def test_tampered_numbers_cannot_pass_receipt_verification(self):
        r = execute_workflow(self.source, self.plan)
        r['total'] = '114.87346'
        with self.assertRaisesRegex(ValueError, '总额不一致'):
            verify_receipt(r, self.plan, self.source)

    def test_explicit_paid_filter_changes_the_plan_and_result(self):
        p = copy.deepcopy(self.plan)
        p['filters'].append({'field': '有偿无偿', 'op': 'in', 'values': ['有偿']})
        self.assertEqual(execute_workflow(self.source, p)['total'], '19.90000')

    def test_missing_ambiguous_fields_and_formula_amounts_fail_closed(self):
        p = copy.deepcopy(self.plan); p['measure'] = '不存在的金额'
        with self.assertRaisesRegex(ValueError, '唯一表头'):
            execute_workflow(self.source, p)
        w = openpyxl.load_workbook(self.source); w.active['C4'] = '=1+2'; w.save(self.source); w.close()
        with self.assertRaisesRegex(ValueError, '显式重算'):
            execute_workflow(self.source, self.plan)

    def test_unknown_operations_cannot_be_silently_ignored(self):
        p = copy.deepcopy(self.plan); p['filters'][0]['op'] = 'guess'
        with self.assertRaisesRegex(ValueError, '不受支持'):
            execute_workflow(self.source, p)

    def test_trailing_query_metadata_is_excluded_but_incomplete_business_rows_are_not(self):
        w = openpyxl.load_workbook(self.source); s = w.active
        s.append(['查询条件：']); s.append(['工数年月:202610']); s.append(['*部门编号:全部'])
        w.save(self.source); w.close()
        r = execute_workflow(self.source, self.plan)
        self.assertEqual(r['source_rows'], 5)
        self.assertEqual(r['skipped_metadata_rows'], [9, 10, 11])
        w = openpyxl.load_workbook(self.source); w.active.append(['有偿', 'ok-bad', 10, None]); w.save(self.source); w.close()
        with self.assertRaisesRegex(ValueError, '混杂'):
            execute_workflow(self.source, self.plan)

    def install(self):
        import hashlib
        pkg = self.root / 'skills' / 'custom-sum'
        (pkg / 'references').mkdir(parents=True); (pkg / 'scripts').mkdir()
        (pkg / 'SKILL.md').write_text('execution_mode: verified_workflow')
        from dsh_modules import personal_skill_workflow_engine
        engine = Path(personal_skill_workflow_engine.__file__)
        (pkg / 'scripts/run.py').write_bytes(engine.read_bytes())
        plan_bytes = json.dumps(self.plan).encode()
        (pkg / 'references/workflow.json').write_bytes(plan_bytes)
        manifest = {'schema_version': 1, 'kind': 'personal_skill_workflow', 'entrypoint': 'scripts/run.py',
                    'plan': 'references/workflow.json', 'engine_sha256': hashlib.sha256(engine.read_bytes()).hexdigest(),
                    'plan_sha256': hashlib.sha256(plan_bytes).hexdigest()}
        (pkg / 'execution.json').write_text(json.dumps(manifest))
        (pkg / 'references/validation.json').write_text(json.dumps({'status': 'replayed', 'plan_sha256': manifest['plan_sha256']}))
        return pkg

    def invoke(self, files=None):
        out = io.StringIO(); history_file = self.root / 'session.json'
        with patch('dsh_modules.personal_skill_workflow.CUSTOM_SKILL_DIR', self.root / 'skills'), \
             patch('dsh_modules.personal_skill_workflow.WORKSPACE_DIR', self.root), redirect_stdout(out):
            self.assertTrue(try_run_personal_workflow(SimpleNamespace(skill_id='custom-sum', admission={'action': 'use', 'skill_id': 'custom-sum', 'use_default_rules': True}),
                files if files is not None else [self.source.name], 'custom-sum', history_file, [], time.time(), None))
        metrics = json.loads(extract_dsh_markers(out.getvalue(), 'METRICS')[0][0])
        return json.loads(history_file.read_text()), metrics

    def test_runtime_executes_only_after_admission_and_persists_reviewed_evidence(self):
        self.install(); history, metrics = self.invoke()
        self.assertTrue(history[-1]['verification']['verified'])
        self.assertEqual(metrics['toolCallsCount'], 1)
        self.assertEqual(metrics['llmInvocations'], 2)
        self.assertIn('30.00000', history[-1]['content'])

    def test_missing_flow_altered_script_and_multiple_inputs_never_return_numbers(self):
        pkg = self.install(); (pkg / 'scripts/run.py').write_text('print("invented")')
        history, metrics = self.invoke()
        self.assertFalse(history[-1]['verification']['verified'])
        self.assertEqual(metrics['guardDecisions'][-1]['action'], 'reject')
        self.assertNotIn('30.00000', history[-1]['content'])
        (pkg / 'execution.json').unlink()
        history, _ = self.invoke()
        self.assertFalse(history[-1]['verification']['verified'])

    def test_multiple_input_files_are_not_guessed(self):
        self.install(); history, _ = self.invoke([self.source.name, 'another.xlsx'])
        self.assertFalse(history[-1]['verification']['verified'])

    def test_source_changes_during_model_presentation_retract_completion(self):
        self.install(); history_file = self.root / 'presentation-change.json'
        def changed(data, table, *args):
            self.write_source('70.00000')
            return table
        with patch('dsh_modules.personal_skill_workflow.CUSTOM_SKILL_DIR', self.root / 'skills'), \
             patch('dsh_modules.personal_skill_workflow.WORKSPACE_DIR', self.root), \
             patch('dsh_modules.personal_skill_presentation.present_verified_workflow', side_effect=changed), \
             redirect_stdout(io.StringIO()):
            try_run_personal_workflow(SimpleNamespace(skill_id='custom-sum', admission={'action': 'use', 'skill_id': 'custom-sum', 'use_default_rules': True}), [self.source.name], 'custom-sum',
                                      history_file, [], time.time(), None, present=True)
        answer = json.loads(history_file.read_text())[-1]
        self.assertFalse(answer['verification']['verified'])
        self.assertNotIn('30.00000', answer['content'])

    def test_text_delivery_does_not_remove_mandatory_workflow_execution(self):
        from dsh_modules.skills import get_available_skills
        from dsh_modules.skill_router import SkillRouter
        pkg = self.install()
        (pkg / 'SKILL.md').write_text('---\nname: custom-sum\nzh_name: 部门别有偿统计\n'
            'triggers: [部门别有偿统计]\nrequires_execution: true\n'
            'execution_mode: verified_workflow\ndeliverables: []\n---\n核算流程')
        with patch('dsh_modules.skills.CUSTOM_SKILL_DIR', self.root / 'skills'), \
             patch('dsh_modules.skills.SKILL_DIR', self.root / 'empty'):
            skills = get_available_skills()
        with patch('dsh_modules.skill_router.read_skill', return_value='核算流程'):
            routed = SkillRouter.route('部门别有偿统计', [], files=[self.source.name], available_skills=skills)
        self.assertEqual(routed.skill_id, 'custom-sum')
        self.assertTrue(routed.requires_execution)
        self.assertEqual(routed.deliverables, [])


if __name__ == '__main__':
    unittest.main()
