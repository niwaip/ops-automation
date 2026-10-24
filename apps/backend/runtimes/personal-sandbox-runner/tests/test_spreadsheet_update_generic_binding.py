"""Metamorphic coverage: unrelated field names, key formats, periods and layouts."""
import copy
import random
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
import os
import openpyxl
from dsh_modules.spreadsheet_update.result_catalog import harvest_analytical_facts,fingerprint,reset_results
from dsh_modules.spreadsheet_update.semantic_plan import build_semantic_plan
from dsh_modules.spreadsheet_update.tools import tool_plan_spreadsheet_update


class GenericBindingTests(unittest.TestCase):
    def test_layout_fields_keys_and_periods_are_request_bound(self):
        generator=random.Random(83746)
        for index in range(8):
            with self.subTest(index=index),tempfile.TemporaryDirectory() as directory:
                root=Path(directory)
                key_col=generator.randint(4,10);value_col=key_col+generator.randint(2,5)
                period_row=generator.randint(15,35);label_row=period_row+generator.randint(12,23)
                first=label_row+generator.randint(3,7)
                period='期-Ω-'+str(index);section='产品组·'+str(index);label='見込額·'+str(index)
                keys=['003','组织/乙','A-9',0]
                wb=openpyxl.Workbook();ws=wb.active;ws.title='版本·α'
                ws.cell(2,1,section);ws.cell(period_row,value_col,period);ws.cell(label_row,value_col,label)
                ws.cell(1,1,'单位：元')
                for offset,key in enumerate(keys):
                    ws.cell(first+offset,key_col,key);ws.cell(first+offset,value_col,0)
                ws.cell(first+5,key_col,'没有统计的数据项');ws.cell(first+5,value_col,987)
                ws.cell(first+6,key_col,'汇总');ws.cell(first+6,value_col,f'=SUM({openpyxl.utils.get_column_letter(value_col)}{first}:{openpyxl.utils.get_column_letter(value_col)}{first+5})')
                ws.cell(first+12,value_col,period);ws.cell(first+16,value_col,label)
                ws.cell(first+20,key_col,keys[0]);ws.cell(first+20,value_col,765)
                target=root/'unrelated.xlsx';wb.save(target);wb.close()
                source=root/'income.xlsx';wb=openpyxl.Workbook();wb.active.append(['任意键','会计窗口','收入(K)']);wb.save(source);wb.close()
                data={'kind':'spreadsheet_analysis','evidence_id':'evidence_'+str(index),
                      'plan':{'group_by':['任意键'],'where':{'field':'会计窗口','op':'eq','value':period},
                              'metrics':[{'id':'revenue','field':'收入(K)','op':'sum'}]},
                      'facts':[{'id':'revenue__'+str(i),'group_fields':{'任意键':key},'value':str(i+1),'unit':'K'} for i,key in enumerate(keys)],
                      'coverage':[{'coverage_complete':True}]}
                reset_results();harvest_analytical_facts(data,provenance={'file_path':str(source),'workbook_id':fingerprint(source)})
                request={'target_file':target.name,'base_sheet':'版本·α','new_sheet':'本次·β','period':period,
                         'period_field':'会计窗口','section':section,'target_label':label,'target_kind':'forecast','target_unit':'元'}
                plan,facts,payload,preview,warnings=build_semantic_plan(request,root)
                self.assertEqual({row['coordinate'] for row in preview},
                                 {f'{openpyxl.utils.get_column_letter(value_col)}{first+i}' for i in range(len(keys))})
                self.assertEqual([row['value'] for row in preview],['1000','2000','3000','4000'])
                self.assertEqual(len(plan['operations']),5)
                self.assertEqual(len(warnings),1);self.assertIn('没有统计的数据项',warnings[0])
                self.assertFalse((root/'unrelated_本次·β.xlsx').exists())

    def test_missing_unit_requires_input_without_guessing_source_unit(self):
        from test_spreadsheet_update_semantic_api import SemanticUpdateTests
        fixture=SemanticUpdateTests();fixture.setUp()
        try:
            wb=openpyxl.load_workbook(fixture.target);wb['1019']['F1']=None;wb.save(fixture.target);wb.close()
            request={k:v for k,v in fixture.request.items() if k!='unit_cell'}
            result=tool_plan_spreadsheet_update(request)
            self.assertTrue(result.is_partial,str(result))
            self.assertEqual(result.data['status'],'needs_input')
            self.assertFalse((fixture.ws/'target_1023.xlsx').exists())
        finally:
            fixture.tearDown()

    def test_old_month_and_unit_defaults_are_rejected(self):
        from dsh_modules.spreadsheet_update.sources import ingest_facts_from_source_queries
        facts,issues=ingest_facts_from_source_queries([], [{'source_id':'unknown'}],[])
        self.assertTrue(issues);self.assertEqual(facts.all_facts(),[])
        from dsh_modules.spreadsheet_update.bindings import LogicalKey
        with self.assertRaises(TypeError):
            LogicalKey('r','k','m','period')

    def test_unfiltered_or_escape_cannot_establish_a_period(self):
        from dsh_modules.spreadsheet_update.result_catalog import _requires_equality
        eq={'field':'窗口','op':'eq','value':'P'}
        self.assertTrue(_requires_equality({'and':[eq,{'field':'类别','op':'eq','value':'收入'}]},'窗口','P'))
        self.assertFalse(_requires_equality({'or':[eq,{'field':'类别','op':'eq','value':'收入'}]},'窗口','P'))

    def test_structure_evidence_covers_deep_headers_without_numeric_unit_guessing(self):
        from dsh_modules.spreadsheet_structure_evidence import collect_label_bands
        wb=openpyxl.Workbook();ws=wb.active
        ws['F45']='产品区';ws['G81']='期间·α';ws['G99']='見込額'
        ws['F103']='键/乙';ws['G103']=123.4
        ws['Z150']='比例（A）'
        evidence=collect_label_bands(ws)
        self.assertTrue(evidence['metadata_scan_complete'])
        self.assertTrue(any('G99' in band['cells'] for band in evidence['label_bands']))
        self.assertEqual(evidence['unit_declarations'],[])
        ws['A1']='单位：万元'
        evidence=collect_label_bands(ws)
        self.assertEqual(evidence['unit_declarations'],[{'cell':'A1','text':'单位：万元'}])
        self.assertFalse(collect_label_bands(ws,cell_budget=100)['metadata_scan_complete'])
        wb.close()

    def test_base_sheet_is_target_snapshot_and_attachment_alias_is_authorized(self):
        from test_spreadsheet_update_semantic_api import SemanticUpdateTests
        from dsh_modules.spreadsheet_update.tools import tool_plan_spreadsheet_update
        fixture=SemanticUpdateTests();fixture.setUp()
        try:
            wrong=tool_plan_spreadsheet_update({**fixture.request,'base_sheet':'统计明细'})
            self.assertTrue(wrong.is_error)
            self.assertIn('1019',str(wrong))
            with patch.dict(os.environ,{'DSH_SESSION_ATTACHMENTS':fixture.target.name}):
                alias=tool_plan_spreadsheet_update({**fixture.request,'target_file':'/obsolete/'+fixture.target.name})
                self.assertTrue(alias.is_success,str(alias))
                forbidden=tool_plan_spreadsheet_update({**fixture.request,'target_file':'/obsolete/unrelated.xlsx'})
                self.assertTrue(forbidden.is_error)
        finally:
            fixture.tearDown()

    def test_unit_evidence_is_checked_before_any_missing_statistical_result(self):
        from test_spreadsheet_update_semantic_api import SemanticUpdateTests
        fixture=SemanticUpdateTests();fixture.setUp()
        try:
            reset_results()
            wb=openpyxl.load_workbook(fixture.target);wb['1019']['F1']=None;wb.save(fixture.target);wb.close()
            request={k:v for k,v in fixture.request.items() if k not in ('unit_cell','target_unit')}
            result=tool_plan_spreadsheet_update(request)
            self.assertTrue(result.is_partial,str(result))
            self.assertEqual(result.data['status'],'needs_input')
        finally:
            fixture.tearDown()

    def test_flat_aggregation_compiles_filters_and_executes_real_source(self):
        from test_spreadsheet_update_semantic_api import SemanticUpdateTests
        from dsh_modules.tools import execute_tool
        from dsh_modules.spreadsheet_update.aggregation_api import compile_aggregation
        fixture=SemanticUpdateTests();fixture.setUp()
        try:
            params={'file_path':str(fixture.source),'sheet':'Sheet','header_row':1,
                    'filters':[{'field':'工数年月','op':'eq','value':'202610'},
                               {'field':'有偿无偿','op':'eq','value':'有偿'}],
                    'group_by':['部门编号'],'value_field':'金额(K)','operation':'sum'}
            compiled=compile_aggregation(params)
            self.assertEqual(compiled['query']['where'],{'and':params['filters']})
            self.assertEqual(compiled['query']['metrics'][0]['unit'],'source')
            with patch('dsh_modules.file_tools.WORKSPACE_DIR',str(fixture.ws)):
                result=execute_tool('aggregate_spreadsheet',params)
            self.assertTrue(result.is_success,str(result))
            self.assertEqual(result.data['facts'][0]['value'],'67.0353')
            self.assertEqual(result.data['facts'][0]['unit'],'K')
            with self.assertRaises(ValueError):
                compile_aggregation({**params,'where':{'op':'and','children':[]}})
        finally:
            fixture.tearDown()
