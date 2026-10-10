"""Task-state and registered-recipe recovery from successful but incomplete plans."""

import io
import json
import os
import sys
import tempfile
import types
import unittest
from contextlib import redirect_stdout
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch
import openpyxl

sys.path.insert(0,str(Path(__file__).resolve().parent.parent/'src'))
from dsh_modules.analysis_task_state import analysis_task_state, task_feedback
from dsh_modules.analysis_turn import AnalysisTurn
from dsh_modules.analysis_semantics import DOMAIN_MODULES, derivation_recipes
from dsh_modules.spreadsheet_analysis import calculate_spreadsheet_analysis
from dsh_modules.registered_metric_completion import RegisteredMetricCompletion
from dsh_modules.spreadsheet_analysis_evidence import verified_response
from dsh_modules.telemetry import TelemetryStats
from dsh_modules.tool_result import ToolResult
from dsh_modules.agent_loop import run_agent_loop
from dsh_modules.runtime_policy import RuntimePolicy
from dsh_modules.tools import get_sandbox_tools
from dsh_modules.llm import _process_sse_data


QUERY='总结2026年收入、毛利率、营业利润和净利润。'


class AnalysisCompletionTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.source=Path(self.temp.name)/'renamed.xlsx'
        workbook=openpyxl.Workbook();sheet=workbook.active;sheet.title='来源甲'
        sheet.append(['单位','千元']);sheet.append(['月份','营业收入','毛利','营业利润','净利润','主营业务毛利'])
        for m in range(1,13):sheet.append([f'{m}月',100+m,40+m,20+m,10+m,30+m])
        workbook.save(self.source);workbook.close()
        self.table={'id':'period','sheet':'来源甲','header_row':2,'data_range':'A3:F14'}

    def tearDown(self):self.temp.cleanup()

    def receipt(self, metrics=None, table=None):
        table=table or self.table
        definitions=metrics or [('finance.revenue','营业收入','B'),('finance.gross_profit','毛利润','C'),
            ('finance.operating_profit','营业利润','D'),('finance.net_profit','净利润','E')]
        plan={'tables':[table],'calculations':[{'id':f'm{i}','metric_id':metric,'label':label,
            'expression':{'aggregate':'sum','table':table['id'],'column':column,'unit':'amount'}}
            for i,(metric,label,column) in enumerate(definitions)]}
        result=calculate_spreadsheet_analysis(self.source,plan,domain='finance')
        self.assertTrue(result.is_success,result.error)
        return result

    def complete(self,evidence,prompt=QUERY,allowed=('aggregate_spreadsheet',),completion=None):
        completion=completion or RegisteredMetricCompletion()
        return completion.complete(evidence,prompt,TelemetryStats(),allowed)

    def test_structured_task_state_detects_missing_rate_independent_of_json(self):
        state=analysis_task_state([self.receipt()],QUERY)
        self.assertFalse(state['complete'])
        self.assertEqual(state['missing_metrics'],[{'metric_id':'finance.gross_margin','name':'毛利率','expected_type':'percentage'}])
        self.assertIn('finance.gross_margin',task_feedback([self.receipt()],QUERY))

    def test_duplicate_realistic_receipts_are_completed_once_and_ids_are_runtime_owned(self):
        evidence=[self.receipt(),self.receipt(table={**self.table,'id':'retry'})]
        completed=self.complete(evidence)
        self.assertEqual(len(completed),1)
        fact=completed[0].data['facts'][0]
        self.assertEqual(fact['metric_id'],'finance.gross_margin');self.assertEqual(fact['unit'],'%')
        self.assertAlmostEqual(float(fact['value']),float(Decimal(558)/Decimal(1278)*100))
        self.assertEqual(completed[0].provenance['execution_origin'],'registered_derivation')
        response=json.loads(verified_response(evidence,QUERY,'{"evidence_ids":["wrong"],"hypotheses":["2026数字",{},"待验证原因"]}'))
        self.assertEqual(set(response['evidence_ids']),{r.data['evidence_id'] for r in evidence})
        self.assertEqual(response['hypotheses'],['待验证原因'])
        self.assertEqual(self.complete(evidence),[])

    def test_pipeline_finishes_from_executed_plan_without_a_model_protocol_round(self):
        receipt=self.receipt()
        params={'file_path':str(self.source),**receipt.data['plan']}
        tool={'content':'未验证草稿','finish_reason':'tool_calls','tool_calls':[{'id':'one','type':'function',
              'function':{'name':'aggregate_spreadsheet','arguments':json.dumps(params)}}]}
        with patch('dsh_modules.spreadsheet_analysis.resolve_sandboxed_path',return_value=(self.source,None)), \
             patch.dict(os.environ,{'DSH_SESSION_ATTACHMENTS':''}), \
             patch('dsh_modules.agent_loop.call_model_proxy',return_value=tool) as model,redirect_stdout(io.StringIO()):
            result=run_agent_loop([{'role':'system','content':'test'},{'role':'user','content':QUERY}],
                'test',RuntimePolicy(),4,tools=get_sandbox_tools({'aggregate_spreadsheet'}),spreadsheet_analysis_required=True)
        self.assertEqual(model.call_count,1)
        self.assertIn('毛利率（注册关系补算）',result.final_text)
        self.assertNotIn('未验证草稿',result.final_text)
        self.assertEqual(result.telemetry.guard_decisions[-1]['action'],'accept')
        self.assertEqual(result.telemetry.tool_invocations,2)

    def test_malformed_final_protocol_cannot_block_complete_generic_execution(self):
        source=self.receipt()
        evidence=[source];self.complete(evidence)
        turn=AnalysisTurn(QUERY)
        for text in ['bad JSON','{"evidence_ids":[],"hypotheses":"wrong"}',
                     '{"evidence_ids":["old"],"hypotheses":["2026年收入999"]}']:
            final=turn.render(evidence,text,TelemetryStats())
            self.assertNotIn('尚未通过',final);self.assertIn('毛利率（注册关系补算）',final)
            self.assertNotIn('999',final)

    def test_cross_period_operands_are_not_combined(self):
        first={**self.table,'filters':[{'column':'A','values':[f'{m}月' for m in range(1,7)]}]}
        last={**self.table,'filters':[{'column':'A','values':[f'{m}月' for m in range(7,13)]}]}
        evidence=[self.receipt([('finance.gross_profit','毛利润','C')],first),
                  self.receipt([('finance.revenue','营业收入','B')],last)]
        self.assertEqual(self.complete(evidence,'总结毛利率'),[])

    def test_registered_relationship_is_completed_for_each_matching_period(self):
        evidence=[]
        for months in (range(1,7),range(7,13)):
            evidence.append(self.receipt(table={**self.table,'filters':[{'column':'A','values':[f'{m}月' for m in months]}]}))
        completed=self.complete(evidence,'解释盈利能力变化')
        self.assertEqual(len(completed),1)
        self.assertEqual(len(completed[0].data['facts']),6)
        self.assertTrue(analysis_task_state(evidence,'解释盈利能力变化')['complete'])

    def test_ambiguous_source_fields_remain_missing(self):
        evidence=[self.receipt(),self.receipt([('finance.gross_profit','毛利润','F')])]
        self.assertEqual(self.complete(evidence),[])
        self.assertFalse(analysis_task_state(evidence,QUERY)['complete'])

    def test_no_recipe_is_applied_in_generic_or_unsupported_capability_mode(self):
        evidence=[self.receipt()]
        self.assertEqual(self.complete(evidence,allowed=['validate_spreadsheet_rows']),[])
        with patch.dict(os.environ,{'DSH_ANALYSIS_DOMAIN':'generic'}):
            self.assertEqual(self.complete(evidence),[])
        self.assertEqual(derivation_recipes('generic'),[])

    def test_zero_denominator_is_execution_failure_and_does_not_loop(self):
        workbook=openpyxl.load_workbook(self.source)
        for row in range(3,15):workbook.active.cell(row,2,0)
        workbook.save(self.source);workbook.close()
        evidence=[self.receipt()];completion=RegisteredMetricCompletion()
        result=self.complete(evidence,completion=completion)
        self.assertEqual(len(result),1);self.assertTrue(result[0].is_error)
        self.assertIsNone(verified_response(evidence,QUERY))
        self.assertEqual(self.complete(evidence,completion=completion),[])

    def test_source_change_and_partial_evidence_cannot_trigger_completion(self):
        receipt=self.receipt()
        partial=ToolResult.partial(data=receipt.data,provenance=receipt.provenance)
        self.assertEqual(self.complete([partial]),[])
        self.source.write_bytes(b'changed')
        evidence=[receipt]
        self.assertEqual(self.complete(evidence),[])
        codes={i['code'] for i in analysis_task_state(evidence,QUERY)['issues']}
        self.assertEqual(codes,{'source_changed','required_metrics_missing'})

    def test_receipt_domain_version_must_match_current_registry(self):
        receipt=self.receipt();receipt.data['semantic_domain']['version']='old'
        self.assertEqual(self.complete([receipt]),[])

    def test_unregistered_nonfinancial_ratio_uses_the_same_completion_engine_when_registered(self):
        pack=types.ModuleType('dsh_modules.analysis_domain_experiment')
        pack.VERSION='test';pack.CONCEPTS=[];pack.GUIDANCE='实验指标'
        pack.METRICS=[{'id':'experiment.input','name':'投入','patterns':['投入'],'expected_type':'number'},
            {'id':'experiment.pass','name':'合格','patterns':['合格数量'],'expected_type':'number'},
            {'id':'experiment.yield','name':'合格率','patterns':['合格率'],'expected_type':'percentage',
             'numerator_metric':'合格','denominator_metric':'投入'}]
        workbook=openpyxl.Workbook();ws=workbook.active;ws.title='随意实验表'
        ws.append(['投入','合格数量']);ws.append([10,8]);ws.append([30,24])
        workbook.save(self.source);workbook.close()
        plan={'tables':[{'id':'batch','sheet':'随意实验表','header_row':1,'data_range':'A2:B3'}],
            'calculations':[{'id':str(i),'metric_id':metric,'label':label,
                'expression':{'aggregate':'sum','table':'batch','column':column,'unit':'number'}}
                for i,(metric,label,column) in enumerate([('experiment.input','投入','A'),('experiment.pass','合格','B')])]}
        with patch.dict(DOMAIN_MODULES,{'experiment':'.analysis_domain_experiment'}), \
             patch.dict(sys.modules,{'dsh_modules.analysis_domain_experiment':pack}), \
             patch.dict(os.environ,{'DSH_ANALYSIS_DOMAIN':'experiment'}):
            evidence=[calculate_spreadsheet_analysis(self.source,plan,domain='experiment')]
            self.assertTrue(evidence[0].is_success,evidence[0].error)
            result=self.complete(evidence,'总结合格率')
            self.assertEqual(len(result),1);self.assertEqual(Decimal(result[0].data['facts'][0]['value']),80)
            self.assertTrue(analysis_task_state(evidence,'总结合格率')['complete'])

    def test_internal_stream_suppresses_both_thoughts_and_content(self):
        chunk=json.dumps({'choices':[{'delta':{'reasoning_content':'I calculated all results','content':'unverified draft'}}]})
        output=io.StringIO();chunks=[]
        with redirect_stdout(output):_process_sse_data(chunk,chunks,{},0,None,stream_deltas=False)
        self.assertEqual(chunks,['unverified draft']);self.assertEqual(output.getvalue(),'')

    def test_presentation_audit_records_shape_without_untrusted_text(self):
        evidence=[self.receipt()];self.complete(evidence)
        telemetry=TelemetryStats()
        final=AnalysisTurn(QUERY).render(evidence,'{"evidence_ids":["sensitive-unknown-text"],"hypotheses":[]}',telemetry)
        audit=telemetry.guard_decisions[0]
        self.assertEqual(audit['guard'],'analysis_presentation')
        self.assertEqual(audit['reason']['unknown_id_count'],1)
        self.assertNotIn('sensitive-unknown-text',json.dumps(telemetry.guard_decisions))
        self.assertNotIn('尚未通过',final)

    def test_completion_reexecutes_source_expressions_not_converted_cached_values(self):
        receipt=self.receipt()
        # Output formatting may express different scales; recipes use the source
        # expressions, not the already-presented numeric strings.
        for fact in receipt.data['facts']:
            if fact['metric_id']=='finance.revenue':
                fact['value']=str(Decimal(fact['value'])/10);fact['unit']='万元'
        result=self.complete([receipt])
        self.assertEqual(len(result),1)
        self.assertAlmostEqual(float(result[0].data['facts'][0]['value']),float(Decimal(558)/Decimal(1278)*100))

    def test_registered_number_cannot_be_satisfied_by_a_percentage(self):
        receipt=self.receipt()
        revenue=next(f for f in receipt.data['facts'] if f['metric_id']=='finance.revenue')
        revenue['unit']='%'
        missing=analysis_task_state([receipt],QUERY)['missing_metrics']
        self.assertIn('finance.revenue',{m['metric_id'] for m in missing})
        expr={'aggregate':'sum','table':'period','column':'B','unit':'amount'}
        result=calculate_spreadsheet_analysis(self.source,{'tables':[self.table],'calculations':[
            {'id':'wrong','label':'营业收入','metric_id':'finance.revenue',
             'expression':{'op':'divide','args':[expr,expr]}}]},domain='finance')
        self.assertTrue(result.is_error);self.assertIn('结果类型不符',result.error_message)

    def test_analysis_overview_includes_rule_violation_examples(self):
        """阶段一测试：验证数据检查规则核查在正文中透出具体违规单元格与取值"""
        from dsh_modules.report_presentation import analysis_overview
        from dsh_modules.tool_result import ToolResult
        receipt = ToolResult.success({
            "kind": "spreadsheet_analysis",
            "facts": [],
            "checks": [],
            "rules": [
                {
                    "label": "日期范围规则",
                    "evaluated_records": 80,
                    "violation_count": 1,
                    "examples": [{"row": 15, "cell": "明细表!B15", "value": "2027-01-02"}]
                },
                {
                    "label": "金额正数规则",
                    "evaluated_records": 80,
                    "violation_count": 64,
                    "examples": [
                        {"row": 5, "cell": "明细表!F5", "value": "-120.00"},
                        {"row": 6, "cell": "明细表!F6", "value": "-85.50"}
                    ]
                }
            ]
        }, provenance={"file_path": "/workspace/AI能力测试_财务报表.xlsx"})
        rendered = analysis_overview([receipt], "识别金额、日期和审批状态异常")
        self.assertIn("日期范围规则", rendered)
        self.assertIn("需核查记录具体位置与异常值", rendered)
        self.assertIn("`明细表!B15` (值: `2027-01-02`)", rendered)
        self.assertIn("`明细表!F5` (值: `-120.00`)", rendered)
        self.assertIn("等共 64 处", rendered)

    def test_render_unverified_guidance_explains_missing_content_and_guides_user(self):
        """阶段三测试：验证当数据缺失时明确回答缺少什么并引导用户输入"""
        from dsh_modules.report_presentation import render_unverified_guidance
        state = {
            "missing_metrics": [{"name": "毛利率", "metric_id": "finance.gross_margin"}],
            "issues": [
                {"code": "required_metrics_missing", "message": "分析计划遗漏所需指标或类型：毛利率。"}
            ]
        }
        guidance = render_unverified_guidance(state, "总结2026年毛利率")
        self.assertIn("本次分析尚未通过执行证据核验", guidance)
        self.assertIn("无法完成计算的原因", guidance)
        self.assertIn("缺少计算 【毛利率】 所需的原始列或对应数据", guidance)
        self.assertIn("建议您补充以下信息以继续", guidance)
        self.assertIn("请补充提供包含 【毛利率】 对应原始列的工作表", guidance)
        self.assertIn("尚未验证的项目不会被列为完成", guidance)

    def test_synthesize_analysis_summary_with_mock_llm(self):
        """阶段二测试：验证 LLM 总结生效时输出自然语言业务报告并附带折叠明细"""
        from types import SimpleNamespace
        from dsh_modules.analysis_synthesis import synthesize_analysis_summary
        from dsh_modules.tool_result import ToolResult
        receipt = ToolResult.success({
            "kind": "spreadsheet_analysis",
            "facts": [{"label": "2026年收入", "value": "124020", "unit": "千元"}],
            "rules": []
        }, provenance={"file_path": "/workspace/test.xlsx"})

        with patch("dsh_modules.analysis_synthesis._can_reach_proxy", return_value=True), \
             patch("dsh_modules.llm.call_model_proxy") as mock_proxy:
            mock_proxy.return_value = {
                "content": "2026年全年营业收入为 124,020 千元，业务规模稳健增长。"
            }
            res = synthesize_analysis_summary(
                prompt="总结2026年收入",
                receipts=[receipt],
                audit="审计明细文本",
                fallback_text="兜底文本",
                model="test-model",
                policy=SimpleNamespace(single_request_timeout=10)
            )
            self.assertIn("2026年全年营业收入为 124,020 千元", res)
            self.assertIn("查看来源与核算明细", res)
            self.assertIn("审计明细文本", res)


if __name__=='__main__':unittest.main()
