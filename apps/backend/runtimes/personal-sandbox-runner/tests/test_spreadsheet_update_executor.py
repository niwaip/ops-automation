from spreadsheet_update_fixture_support import bind_fixture_regions
"""
Unit tests for PR 3: Controlled Execution, Compilation Constraints, and Semantic Diffing.
"""

import unittest
import tempfile
import shutil
from decimal import Decimal
from pathlib import Path
import openpyxl

from dsh_modules.spreadsheet_update.schema import SCHEMA_VERSION_V1
from dsh_modules.spreadsheet_update.bindings import LogicalKey
from dsh_modules.spreadsheet_update.discovery import extract_workbook_ir
from dsh_modules.spreadsheet_update.sources import FactRecord, FactTable
from dsh_modules.spreadsheet_update.compiler import (
    compile_update_plan,
    UpdateCompilationError,
)
from dsh_modules.spreadsheet_update.executor import execute_compiled_update
from dsh_modules.spreadsheet_update.diff import compute_semantic_diff


class TestSpreadsheetUpdateExecutor(unittest.TestCase):

    def setUp(self):
        self.test_dir = Path(tempfile.mkdtemp(prefix="test_update_exec_"))
        self.target_wb_path = self.test_dir / "target_source.xlsx"
        self._build_test_workbook(self.target_wb_path)

    def tearDown(self):
        shutil.rmtree(self.test_dir, ignore_errors=True)

    def _build_test_workbook(self, path: Path):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "1019"

        # Headers
        ws["B2"] = "快照日期"
        ws["C2"] = "2026-10-19"

        ws["B4"] = "部门"
        ws["C4"] = "202610"
        ws["D4"] = "202611"
        ws["E4"] = "合计"

        # Row 5
        ws["B5"] = "FC001"
        ws["C5"] = 100
        ws["D5"] = 150
        ws["E5"] = "=SUM(C5:D5)"

        # Row 6
        ws["B6"] = "FC002"
        ws["C6"] = 200
        ws["D6"] = 250
        ws["E6"] = "=SUM(C6:D6)"

        wb.save(path)

    def _create_facts(self) -> FactTable:
        facts = FactTable()
        k1 = LogicalKey("eng_paid", "FC001", "hours", "202610", "actual", "hours")
        k2 = LogicalKey("eng_paid", "FC002", "hours", "202610", "actual", "hours")
        facts.add_fact(FactRecord(k1, Decimal("105.0"), "hours", "202610", "2026-10-23", "s1", 2, "C2", "ev_1"))
        facts.add_fact(FactRecord(k2, Decimal("210.0"), "hours", "202610", "2026-10-23", "s1", 3, "C3", "ev_2"))
        return facts

    def test_compiler_rejects_formula_overwrite(self):
        facts = self._create_facts()
        spec = [{
            "region_id": "eng_paid",
            "title": "工数",
            "min_row": 4,
            "min_col": 2,
            "max_row": 6,
            "max_col": 5,
            "header_row": 4,
            "key_col": 2,
            "metric_id": "hours",
            "unit": "hours"
        }]
        ir = extract_workbook_ir(self.target_wb_path, "1019", region_specs=bind_fixture_regions(self.target_wb_path, "1019", spec))

        plan = {
            "schema_version": SCHEMA_VERSION_V1,
            "task_contract_ref": "c1",
            "target_ref": "t1",
            "as_of": "2026-10-23",
            "snapshot": {
                "selected_sheet_id": "1019",
                "new_sheet_name": "1023",
            },
            "operations": [
                {
                    "op": "copy_sheet",
                    "op_id": "op1",
                    "source_sheet": "1019",
                    "target_sheet": "1023"
                },
                {
                    # Attempt to overwrite E5 which contains formula =SUM(C5:D5)
                    "op": "fill_from_result",
                    "op_id": "op_bad_fill",
                    "coordinate": "E5",
                    "region_id": "eng_paid",
                    "department_code": "FC001",
                    "metric_id": "hours",
                    "period": "202610",
                    "value_kind": "actual",
                    "unit": "hours",
                    "fact_ref": "ev_1"
                }
            ]
        }

        with self.assertRaises(UpdateCompilationError) as ctx:
            compile_update_plan(plan, ir, facts)
        self.assertIn("attempted to overwrite existing formula at E5", str(ctx.exception))

    def test_compiler_and_executor_end_to_end_in_sandbox(self):
        facts = self._create_facts()
        spec = [{
            "region_id": "eng_paid",
            "title": "工数",
            "min_row": 4,
            "min_col": 2,
            "max_row": 6,
            "max_col": 5,
            "header_row": 4,
            "key_col": 2,
            "metric_id": "hours",
            "unit": "hours"
        }]
        ir = extract_workbook_ir(self.target_wb_path, "1019", region_specs=bind_fixture_regions(self.target_wb_path, "1019", spec))

        plan = {
            "schema_version": SCHEMA_VERSION_V1,
            "task_contract_ref": "c1",
            "target_ref": "t1",
            "as_of": "2026-10-23",
            "snapshot": {
                "selected_sheet_id": "1019",
                "new_sheet_name": "1023",
            },
            "operations": [
                {
                    "op": "copy_sheet",
                    "op_id": "op1",
                    "source_sheet": "1019",
                    "target_sheet": "1023"
                },
                {
                    "op": "set_snapshot_date",
                    "op_id": "op_date",
                    "sheet": "1023",
                    "coordinate": "C2",
                    "date_value": "2026-10-23"
                },
                {
                    "op": "fill_from_result",
                    "op_id": "op_fill_1",
                    "coordinate": "C5",
                    "region_id": "eng_paid",
                    "department_code": "FC001",
                    "metric_id": "hours",
                    "period": "202610",
                    "value_kind": "actual",
                    "unit": "hours",
                    "fact_ref": "ev_1"
                },
                {
                    "op": "fill_from_result",
                    "op_id": "op_fill_2",
                    "coordinate": "C6",
                    "region_id": "eng_paid",
                    "department_code": "FC002",
                    "metric_id": "hours",
                    "period": "202610",
                    "value_kind": "actual",
                    "unit": "hours",
                    "fact_ref": "ev_2"
                }
            ]
        }

        compiled = compile_update_plan(plan, ir, facts)
        self.assertEqual(len(compiled.cell_writes), 3)  # Date + 2 cells
        self.assertIn("C5", compiled.allowed_write_coordinates)
        self.assertIn("C6", compiled.allowed_write_coordinates)
        self.assertIn("C2", compiled.allowed_write_coordinates)

        cand_output = self.test_dir / "candidate.xlsx"
        # Execute without full libreoffice in pure unittest runner or test recalc
        res = execute_compiled_update(self.target_wb_path, compiled, cand_output, recalculate=False)
        self.assertTrue(res.is_success)
        self.assertTrue(cand_output.exists())

        # Inspect candidate with openpyxl
        wb_cand = openpyxl.load_workbook(cand_output)
        self.assertIn("1019", wb_cand.sheetnames)
        self.assertIn("1023", wb_cand.sheetnames)

        ws_1023 = wb_cand["1023"]
        self.assertEqual(ws_1023["C2"].value, "2026-10-23")
        self.assertEqual(ws_1023["C5"].value, 105.0)
        self.assertEqual(ws_1023["C6"].value, 210.0)
        # Formula preserved!
        self.assertEqual(ws_1023["E5"].value, "=SUM(C5:D5)")
        wb_cand.close()

        # Semantic diff verification
        diff_report = compute_semantic_diff(self.target_wb_path, cand_output, compiled)
        self.assertTrue(diff_report.is_valid)
        self.assertTrue(diff_report.historical_sheets_preserved)
        self.assertTrue(diff_report.is_write_set_bounded)
        self.assertTrue(diff_report.protected_formulas_preserved)

    def test_semantic_diff_catches_historical_tampering(self):
        facts = self._create_facts()
        spec = [{
            "region_id": "eng_paid",
            "title": "工数",
            "min_row": 4,
            "min_col": 2,
            "max_row": 6,
            "max_col": 5,
            "header_row": 4,
            "key_col": 2,
            "metric_id": "hours",
            "unit": "hours"
        }]
        ir = extract_workbook_ir(self.target_wb_path, "1019", region_specs=bind_fixture_regions(self.target_wb_path, "1019", spec))
        plan = {
            "schema_version": SCHEMA_VERSION_V1,
            "snapshot": {"selected_sheet_id": "1019", "new_sheet_name": "1023"},
            "operations": [{"op": "copy_sheet", "op_id": "1", "source_sheet": "1019", "target_sheet": "1023"}]
        }
        compiled = compile_update_plan(plan, ir, facts)
        cand_output = self.test_dir / "tampered.xlsx"
        execute_compiled_update(self.target_wb_path, compiled, cand_output, recalculate=False)

        # Tamper historical sheet in candidate
        wb = openpyxl.load_workbook(cand_output)
        wb["1019"]["C5"] = 99999
        wb.save(cand_output)
        wb.close()

        diff_report = compute_semantic_diff(self.target_wb_path, cand_output, compiled)
        self.assertFalse(diff_report.is_valid)
        self.assertFalse(diff_report.historical_sheets_preserved)
        self.assertTrue(any("Historical sheet '1019'" in err for err in diff_report.unexpected_modifications))


if __name__ == "__main__":
    unittest.main()
