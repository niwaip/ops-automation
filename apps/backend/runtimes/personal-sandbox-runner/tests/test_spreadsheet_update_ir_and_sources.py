from spreadsheet_update_fixture_support import bind_fixture_regions
"""
Unit tests for PR 2: Multi-Region WorkbookIR, Typed LogicalKeys, and Source Fact Cardinality.
"""

import unittest
import tempfile
import shutil
from decimal import Decimal
from pathlib import Path
import openpyxl

from dsh_modules.spreadsheet_update.bindings import (
    LogicalKey,
    CellBinding,
    BindingRegistry,
)
from dsh_modules.spreadsheet_update.discovery import (
    BusinessRegion,
    WorkbookIR,
    extract_workbook_ir,
)
from dsh_modules.spreadsheet_update.sources import (
    FactRecord,
    FactTable,
    CardinalityReport,
)


class TestWorkbookIRAndSources(unittest.TestCase):

    def setUp(self):
        self.test_dir = Path(tempfile.mkdtemp(prefix="test_update_ir_"))

    def tearDown(self):
        shutil.rmtree(self.test_dir, ignore_errors=True)

    def test_typed_logical_key_distinguishes_regions(self):
        # Same dept code FC001 in paid vs non-paid regions
        k_paid = LogicalKey(
            region_id="paid_hours",
            department_code="FC001",
            metric_id="hours",
            period="202610",
            value_kind="actual",
            unit="hours"
        )
        k_other = LogicalKey(
            region_id="non_paid_hours",
            department_code="FC001",
            metric_id="hours",
            period="202610",
            value_kind="actual",
            unit="hours"
        )
        self.assertNotEqual(k_paid, k_other)
        self.assertNotEqual(k_paid.to_key_string(), k_other.to_key_string())

        registry = BindingRegistry()
        registry.register(CellBinding(k_paid, "1019", "E5"))
        registry.register(CellBinding(k_other, "1019", "E18"))

        self.assertEqual(registry.get_by_key(k_paid).coordinate, "E5")
        self.assertEqual(registry.get_by_key(k_other).coordinate, "E18")
        self.assertEqual(registry.get_by_coordinate("1019", "E5"), k_paid)
        self.assertEqual(registry.get_by_coordinate("1019", "E18"), k_other)

    def test_workbook_ir_extraction_and_column_reorder_resilience(self):
        wb_path = self.test_dir / "target_reordered.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "1019"

        # Headers on row 4 with reordered periods: 202612 before 202610
        ws["B4"] = "部门编码"
        ws["C4"] = "202612"
        ws["D4"] = "202610"
        ws["E4"] = "202611"
        ws["F4"] = "3Q"

        # Data rows
        ws["B5"] = "FC001"
        ws["C5"] = 120
        ws["D5"] = 100
        ws["E5"] = 110
        ws["F5"] = "=SUM(C5:E5)"

        ws["B6"] = "FC002"
        ws["C6"] = 220
        ws["D6"] = 200
        ws["E6"] = 210
        ws["F6"] = "=SUM(C6:E6)"

        wb.save(wb_path)

        spec = [{
            "region_id": "engineering_paid",
            "title": "有偿工数",
            "min_row": 4,
            "min_col": 2,
            "max_row": 6,
            "max_col": 6,
            "header_row": 4,
            "key_col": 2,
            "metric_id": "paid_hours",
            "unit": "hours"
        }]

        ir = extract_workbook_ir(wb_path, "1019", region_specs=bind_fixture_regions(wb_path, "1019", spec))
        self.assertEqual(len(ir.regions), 1)
        region = ir.regions[0]

        # Verify periods mapped to correct column even when reordered
        self.assertEqual(region.period_to_col["202610"], 4)  # Col D
        self.assertEqual(region.period_to_col["202612"], 3)  # Col C
        self.assertEqual(region.period_to_col["3Q"], 6)      # Col F

        # Verify formulas vs input cells
        self.assertIn("F5", region.formula_cells)
        self.assertIn("F6", region.formula_cells)
        self.assertIn("D5", region.input_cells)

        # Lookup binding by typed key for 202610
        target_key = LogicalKey(
            region_id="engineering_paid",
            department_code="FC001",
            metric_id="paid_hours",
            period="202610",
            value_kind="actual",
            unit="hours"
        )
        binding = ir.registry.get_by_key(target_key)
        self.assertIsNotNone(binding)
        self.assertEqual(binding.coordinate, "D5")
        self.assertFalse(binding.is_formula)

    def test_fact_table_cardinality_and_conflict_detection(self):
        facts = FactTable()

        k1 = LogicalKey("r1", "FC001", "hours", "202610", "actual", "hours")
        k2 = LogicalKey("r1", "FC002", "hours", "202610", "actual", "hours")

        facts.add_fact(FactRecord(k1, Decimal("100.5"), "hours", "202610", "2026-10-23", "src1", 2, "C2", "ev_1"))
        facts.add_fact(FactRecord(k2, Decimal("200.0"), "hours", "202610", "2026-10-23", "src1", 3, "C3", "ev_2"))

        # Normal 1:1 check against expected keys
        report = facts.check_cardinality_against_bindings([k1, k2])
        self.assertTrue(report.is_valid)
        self.assertEqual(report.matched_count, 2)
        self.assertEqual(len(report.missing_keys), 0)

        # Now simulate conflicting row for FC001 with different amount
        facts.add_fact(FactRecord(k1, Decimal("999.0"), "hours", "202610", "2026-10-23", "src2", 5, "C5", "ev_3"))

        conflict_report = facts.check_cardinality_against_bindings([k1, k2])
        self.assertFalse(conflict_report.is_valid)
        self.assertIn(k1, conflict_report.conflicting_keys)
        self.assertTrue(any("Conflicting values detected" in err for err in conflict_report.errors))


if __name__ == "__main__":
    unittest.main()
