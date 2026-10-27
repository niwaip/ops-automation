"""
Spreadsheet pre-flight profiling and domain identification module.
Provides zero-hallucination table physical schemas, distinct value samplings,
and JIT domain guidance before agent planning.
"""

from decimal import Decimal
import datetime
from pathlib import Path
import re
from typing import Dict, Any, List, Optional, Set, Tuple
import openpyxl
from openpyxl.utils import get_column_letter


SUMMARY_ROW_LABEL_RE = re.compile(
    r'^\s*(?:'
    r'(?:全年|本年|年度|全期|本期|本季|季度|本月|月度|当期|累计|全表|全部|各月|各期|总体|综合)?\s*'
    r'(?:合\s*计|总\s*计|小\s*计|汇\s*总)'
    r'(?:\s*[/／、与和及]\s*(?:平\s*均|均\s*值|合\s*计|总\s*计|小\s*计))?'
    r'|grand\s+total|subtotal|total|summary'
    r')\s*[:：]?\s*$',
    re.I,
)

COVER_SHEET_NAMES = frozenset({"使用说明", "说明", "填报说明", "编制说明", "目录", "封面", "索引", "readme", "instructions"})
CONFIG_SHEET_NAMES = frozenset({"参数", "模型参数", "假设", "情景假设", "情景参数", "参数配置", "配置", "config", "parameters", "assumptions", "dictionary", "数据字典"})
VERTICAL_SUBJECT_HEADERS = frozenset({"项目", "科目", "指标", "核心指标", "参数", "行项目", "报表项目", "指标名称", "科目名称"})


def _profile_single_worksheet(
    ws,
    max_distinct_sample: int = 15,
) -> Dict[str, Any]:
    max_r = ws.max_row or 1
    max_c = min(ws.max_column or 1, 64)

    # 1. Detect header row by scanning first 15 rows for highest density of non-formula strings
    observed_rows: Dict[int, List[Any]] = {}
    scan_limit = min(max_r, 15)
    for row in ws.iter_rows(min_row=1, max_row=scan_limit, max_col=max_c):
        filled = [c for c in row if c.value is not None]
        if filled:
            observed_rows[filled[0].row] = filled

    best_row = 1
    best_headers: Dict[str, str] = {}
    best_is_unit_banner = False
    for r_idx, filled in observed_rows.items():
        strings = {
            get_column_letter(c.column): str(c.value).strip()
            for c in filled
            if isinstance(c.value, str) and not str(c.value).startswith("=") and str(c.value).strip()
        }
        first_str = next(iter(strings.values()), "") if strings else ""
        is_unit_banner = len(strings) <= 2 and first_str.rstrip(":：") in ("单位", "币种", "金额单位", "货币单位", "报表单位")
        if len(strings) > len(best_headers) or (
            len(strings) == len(best_headers) and len(strings) >= 2 and best_is_unit_banner and not is_unit_banner
        ):
            best_headers = strings
            best_row = r_idx
            best_is_unit_banner = is_unit_banner

    # 2. Detect footer boundary from the bottom (including trailing summary/total rows)
    tail_scan_min = max(best_row + 1, max_r - 20)
    tail_rows: Dict[int, List[Any]] = {}
    for row in ws.iter_rows(min_row=tail_scan_min, max_row=max_r, max_col=max_c):
        filled = [c for c in row if c.value is not None]
        if filled:
            tail_rows[filled[0].row] = filled

    footer_prefix_patterns = (
        r'^(?:[\*＊#※]|查询条件|报表类别|填报说明|编制说明|制表人|制表时间|打印时间|导出时间|生成时间|温馨提示|特别提示|数据来源)(?:[:：\s]|$)',
        r'^(?:备注|注|说明|附注|声明)[:：]',
        r'^(?:合\s*计|总\s*计|小\s*计|grand\s+total|total|subtotal|summary)(?:[:：\s]|$)',
    )
    last_data_row = max_r
    header_col_count = len(best_headers)
    excluded_summary_rows: List[Tuple[int, str]] = []
    while last_data_row > best_row:
        if last_data_row in tail_rows:
            cells = tail_rows[last_data_row]
            first_val = str(cells[0].value or "").strip()
            is_footer = False
            if SUMMARY_ROW_LABEL_RE.match(first_val):
                is_footer = True
                excluded_summary_rows.append((last_data_row, first_val))
            elif len(cells) <= max(3, header_col_count // 3):
                if any(re.search(pat, first_val, re.I) for pat in footer_prefix_patterns):
                    is_footer = True
                elif ":" in first_val or "：" in first_val:
                    is_footer = True
            elif re.match(r'^(?:合\s*计|总\s*计|小\s*计)(?:[:：\s]|$)', first_val):
                is_footer = True
                excluded_summary_rows.append((last_data_row, first_val))

            if is_footer:
                last_data_row -= 1
                continue
        break

    data_start_row = best_row + 1
    total_data_records = max(0, last_data_row - data_start_row + 1)
    max_header_col = max(
        (openpyxl.utils.column_index_from_string(let) for let in best_headers),
        default=max_c,
    )
    effective_max_c = min(max_c, max(max_header_col, 1))
    data_range_str = f"A{data_start_row}:{get_column_letter(effective_max_c)}{last_data_row}"

    # 3. Sample column values, formulas, and types
    distincts: Dict[str, Set[str]] = {letter: set() for letter in best_headers}
    top_samples: Dict[str, List[str]] = {letter: [] for letter in best_headers}
    numeric_or_formula_counts: Dict[str, int] = {letter: 0 for letter in best_headers}
    non_empty_counts: Dict[str, int] = {letter: 0 for letter in best_headers}
    pct_format_counts: Dict[str, int] = {letter: 0 for letter in best_headers}
    date_counts: Dict[str, int] = {letter: 0 for letter in best_headers}
    first_col_letter = min(
        best_headers.keys(),
        key=lambda x: openpyxl.utils.column_index_from_string(x),
        default="A",
    )
    row_coord_pairs: List[Tuple[int, str, str]] = []

    val_sample_limit = min(last_data_row, data_start_row + 300)
    is_sampled_scan = last_data_row > val_sample_limit

    for r_offset, row in enumerate(ws.iter_rows(min_row=data_start_row, max_row=val_sample_limit, max_col=effective_max_c)):
        r_num = data_start_row + r_offset
        for c in row:
            c_letter = get_column_letter(c.column)
            if c_letter not in distincts or c.value is None:
                continue
            raw_val = c.value
            val_str = str(raw_val).strip()
            if not val_str:
                continue
            non_empty_counts[c_letter] += 1
            is_formula = (isinstance(raw_val, str) and val_str.startswith("=")) or getattr(c, "data_type", None) == "f"
            is_num = isinstance(raw_val, (int, float, Decimal)) and not isinstance(raw_val, bool)
            if is_num or is_formula:
                numeric_or_formula_counts[c_letter] += 1
            if "%" in str(getattr(c, "number_format", "") or ""):
                pct_format_counts[c_letter] += 1
            if isinstance(raw_val, (datetime.date, datetime.datetime)) or re.fullmatch(r"\d{4}[-/]\d{1,2}(?:[-/]\d{1,2})?(?:\s+.*)?", val_str):
                date_counts[c_letter] += 1

            distincts[c_letter].add(val_str)
            if r_offset < 5 and val_str not in top_samples[c_letter] and len(top_samples[c_letter]) < 3:
                top_samples[c_letter].append(val_str)
            if c_letter == first_col_letter and not is_formula and len(row_coord_pairs) <= 32:
                row_coord_pairs.append((r_num, f"{c_letter}{r_num}", val_str))

    # 4. Infer physical dimension & role using both header and actual cell types
    col_summaries: List[Dict[str, Any]] = []
    for c_letter in sorted(best_headers.keys(), key=lambda x: openpyxl.utils.column_index_from_string(x)):
        hname = best_headers[c_letter]
        vals = distincts.get(c_letter, set())
        d_count = len(vals)
        ne_count = non_empty_counts.get(c_letter, 0)
        nf_count = numeric_or_formula_counts.get(c_letter, 0)
        pct_count = pct_format_counts.get(c_letter, 0)
        dt_count = date_counts.get(c_letter, 0)
        mostly_numeric = ne_count > 0 and (nf_count / ne_count) >= 0.6

        dim = "categorical"
        if any(k in hname for k in ("率", "占比", "比重", "margin", "rate", "ratio")) or (ne_count > 0 and pct_count / ne_count >= 0.5):
            dim = "percentage" if mostly_numeric else "categorical"
        elif any(k in hname for k in ("工数", "工时", "小时", "时数", "duration")):
            dim = "duration"
        elif any(k in hname for k in ("条数", "人数", "数量", "笔数")) and mostly_numeric:
            dim = "count"
        elif any(k in hname for k in ("年月", "日期", "年度", "时间", "月份")) or (ne_count > 0 and dt_count / ne_count >= 0.6):
            dim = "datetime"
        elif any(k in hname for k in ("编号", "代码", "ID", "id", "No", "账号", "凭证号")) and not mostly_numeric:
            dim = "identifier"
        elif any(k in hname for k in ("单价", "金额", "费用", "价格", "收入", "成本", "毛利", "利润", "余额", "预算", "实际", "差异", "变动", "税额", "现金", "资产", "负债", "权益")) and mostly_numeric:
            dim = "currency"
        elif mostly_numeric and c_letter != first_col_letter:
            dim = "currency" if any(k in hname for k in ("元", "K", "k", "万", "数值")) else "numeric"
        else:
            dim = "categorical"

        summary: Dict[str, Any] = {
            "letter": c_letter,
            "name": hname,
            "dimension": dim,
            "distinct_count": d_count,
            "is_sampled": is_sampled_scan,
        }
        eff_limit = max(max_distinct_sample, 20) if c_letter == first_col_letter else max_distinct_sample
        if not is_sampled_scan and dim in ("categorical", "identifier", "datetime") and 0 < d_count <= eff_limit:
            summary["enum_values"] = sorted(vals)
        elif top_samples.get(c_letter):
            summary["sample_values"] = top_samples[c_letter]
        col_summaries.append(summary)

    # Detect vertical subject-row table layout (e.g. 利润表, 资产负债表, 现金流量表, 参数, 管理驾驶舱)
    first_col_summary = col_summaries[0] if col_summaries else None
    measure_cols = [c for c in col_summaries[1:] if c["dimension"] in ("currency", "numeric", "percentage")]
    is_vertical_metric_table = bool(
        first_col_summary
        and 2 <= total_data_records <= 40
        and first_col_summary["dimension"] == "categorical"
        and len(row_coord_pairs) == total_data_records
        and len({lbl for _, _, lbl in row_coord_pairs}) == total_data_records
        and measure_cols
        and (
            first_col_summary["name"] in VERTICAL_SUBJECT_HEADERS
            or any(
                 any(k in lbl for k in ("营业收入", "毛利", "净利润", "营业利润", "资产总计", "负债合计", "现金流"))
                 for _, _, lbl in row_coord_pairs
            )
        )
    )

    return {
        "sheet": ws.title,
        "header_row": best_row,
        "data_range": data_range_str,
        "records": total_data_records,
        "headers": best_headers,
        "columns": col_summaries,
        "is_sampled": is_sampled_scan,
        "max_column": ws.max_column or 1,
        "max_c": effective_max_c,
        "excluded_summary_rows": excluded_summary_rows,
        "row_coord_pairs": row_coord_pairs,
        "is_vertical_metric_table": is_vertical_metric_table,
    }


def _render_sheet_profile_lines(sp: Dict[str, Any]) -> List[str]:
    truncation_note = f"（注意：物理表格共有 {sp['max_column']} 列，当前扫描前 {sp['max_c']} 列）" if sp["max_column"] > sp["max_c"] else ""
    sampling_note = "（注意：数据行数超过300行，离散字段取值基于前300行采样，非全量枚举）" if sp.get("is_sampled") else ""
    summary_excl_note = ""
    if sp.get("excluded_summary_rows"):
        excl_desc = "、".join(f"第{r}行 '{lbl}'" for r, lbl in sorted(sp["excluded_summary_rows"]))
        summary_excl_note = f"（已自动排除表尾汇总行: {excl_desc}）"

    layout_tag = " [纵表/科目行结构]" if sp.get("is_vertical_metric_table") else ""
    lines = [
        f"- 工作表 [{sp['sheet']}]{layout_tag} | 结构: 表头第 {sp['header_row']} 行，有效明细范围 {sp['data_range']}，共计 {sp['records']} 条有效业务明细{summary_excl_note}{truncation_note}{sampling_note}。",
    ]
    for idx, col in enumerate(sp.get("columns", [])):
        c_str = f"  * 列 {col['letter']}: [{col['name']}] ({col['dimension']})"
        if idx == 0 and sp.get("is_vertical_metric_table") and sp.get("row_coord_pairs"):
            coord_items = [f"第{r}行({cell}):'{lbl}'" for r, cell, lbl in sp["row_coord_pairs"]]
            c_str += f" | 科目行坐标映射: [{', '.join(coord_items)}]"
        elif "enum_values" in col:
            c_str += f" | 离散取值枚举: {col['enum_values']}"
        elif "sample_values" in col:
            if col.get("is_sampled"):
                c_str += f" | 样例: {col['sample_values']} (前300行采样发现 {col['distinct_count']} 种取值，非全量枚举)"
            else:
                c_str += f" | 样例: {col['sample_values']} (共 {col['distinct_count']} 种取值)"
        lines.append(c_str)

    if sp.get("is_vertical_metric_table") and sp.get("columns"):
        first_col_name = sp["columns"][0]["name"]
        measure_names = [c["name"] for c in sp["columns"][1:] if c["dimension"] in ("currency", "numeric", "percentage")]
        first_measure = measure_names[0] if measure_names else sp["columns"][-1]["name"]
        lines.append(
            f"  ⚠️ 【纵表查询规则】: [{sp['sheet']}] 的 [{first_col_name}] 列每一行是不同财务科目/指标，共用数值列 {measure_names}。"
            f"若使用 aggregate_spreadsheet 的 where 筛选多个科目（如 {first_col_name} in [...]），必须同时声明 "
            f"\"group_by\": [\"{first_col_name}\"]（且 metrics 仅需声明单个数值列如 {{\"id\": \"val\", \"label\": \"{first_measure}\", \"op\": \"sum\", \"field\": \"{first_measure}\"}}），"
            f"严禁在未设置 group_by 的情况下对多个科目行直接求和！或在 calculations 中按上表行坐标直接引用单元格。"
        )
    return lines


def profile_spreadsheet_for_planning(
    file_path: Path | str,
    prompt: str = "",
    max_distinct_sample: int = 12
) -> Dict[str, Any]:
    """
    Fast (<50ms) deterministic inspection of spreadsheet physical structure.
    Extracts true header rows, data boundaries (excluding titles/footers/summary rows),
    column physical dimensions, vertical/horizontal layout rules, and discrete value enumerations.
    """
    p = Path(file_path)
    if not p.exists() or p.suffix.lower() not in (".xlsx", ".xlsm"):
        return {"compact_markdown": "", "domain": "generic", "headers": {}}

    try:
        wb = openpyxl.load_workbook(p, data_only=False)
    except Exception:
        return {"compact_markdown": "", "domain": "generic", "headers": {}}

    try:
        sheet_names = wb.sheetnames
        if not sheet_names:
            return {"compact_markdown": "", "domain": "generic", "headers": {}}

        sheet_profiles: List[Dict[str, Any]] = []
        for s_name in sheet_names[:10]:
            ws = wb[s_name]
            sp = _profile_single_worksheet(ws, max_distinct_sample=max_distinct_sample)
            if sp["headers"]:
                sheet_profiles.append(sp)

        if not sheet_profiles:
            return {"compact_markdown": "", "domain": "generic", "headers": {}}

        # Select primary sheet (prefer sheet mentioned in prompt, or column/row terms matched in prompt, or first non-cover business sheet)
        primary_sp = sheet_profiles[0]
        domain_target_profiles = sheet_profiles
        if len(sheet_profiles) > 1:
            prompt_text = prompt or ""
            prompt_matched = [sp for sp in sheet_profiles if sp["sheet"] in prompt_text]
            if prompt_matched:
                primary_sp = prompt_matched[0]
                domain_target_profiles = prompt_matched
            else:
                non_cover = [sp for sp in sheet_profiles if sp["sheet"].strip().lower() not in COVER_SHEET_NAMES] or sheet_profiles
                business_candidates = [sp for sp in non_cover if sp["sheet"].strip().lower() not in CONFIG_SHEET_NAMES] or non_cover
                if prompt_text:
                    scored: List[Tuple[int, int, Dict[str, Any]]] = []
                    concept_terms: Set[str] = set()
                    for kw, synonyms in (
                        ("盈利", ("利润", "毛利", "营业利润", "净利润", "利润率", "收益", "经营")),
                        ("利润", ("毛利", "营业利润", "净利润", "利润率", "营收", "营业收入")),
                        ("下半年", ("月份", "7月", "8月", "9月", "10月", "11月", "12月", "H2")),
                        ("上半年", ("月份", "1月", "2月", "3月", "4月", "5月", "6月", "H1")),
                        ("费用", ("销售费用", "管理费用", "研发费用", "成本")),
                        ("成本", ("营业成本", "人员成本", "费用")),
                    ):
                        if kw in prompt_text:
                            concept_terms.update(synonyms)

                    for sp in business_candidates:
                        terms = set(sp["headers"].values()) | {lbl for _, _, lbl in sp.get("row_coord_pairs", [])}
                        direct_hits = sum(1 for t in terms if len(t) >= 2 and t not in ("项目", "指标", "单位", "备注") and t in prompt_text)
                        concept_hits = sum(1 for t in terms if any(c in t for c in concept_terms)) if concept_terms else 0
                        total_score = direct_hits * 3 + concept_hits
                        record_bonus = 1 if sp.get("records", 0) >= 3 else 0
                        if total_score > 0:
                            scored.append((total_score + record_bonus, sp.get("records", 0), sp))
                    if scored:
                        scored.sort(key=lambda item: (item[0], item[1]), reverse=True)
                        max_score = scored[0][0]
                        primary_sp = scored[0][2]
                        domain_target_profiles = [sp for s, r, sp in scored if s == max_score]
                    else:
                        primary_sp = business_candidates[0]
                else:
                    primary_sp = business_candidates[0]

        # Domain identification across targeted sheets (or all sheets when prompt doesn't narrow to specific sheets)
        all_header_names: Set[str] = set()
        all_row_labels: Set[str] = set()
        for sp in domain_target_profiles:
            all_header_names.update(sp["headers"].values())
            for _, _, lbl in sp.get("row_coord_pairs", []):
                all_row_labels.add(lbl)
        combined_terms = all_header_names | all_row_labels

        domain_id = "generic"
        domain_title = "通用业务表格"
        domain_notes = ""

        if any("WBS" in h or "工数" in h or "职级" in h for h in all_header_names):
            domain_id = "timesheet_wbs"
            domain_title = "工数实绩与人力工时管理"
            domain_notes = (
                "- 【统计指导】: 当需要按主实体维度（如部门、项目）与业务分类维度交叉统计不同口径合计时：\n"
                "  * 推荐在 group_by 中传入多列维度组合，如 group_by: ['实体列名', '分类列名']。\n"
                "  * 或声明不同过滤条件的 table 分别汇总目标分类，切勿在单表上直接混算不同分类口径。"
            )
        elif any(any(k in t for k in ("营收", "营业收入", "净利润", "毛利", "负债", "资产", "现金流")) for t in combined_terms):
            domain_id = "finance_fpna"
            domain_title = "财务会计与经营分析"
            domain_notes = (
                "- 【字段语义与纵横表规范】: 金额类字段遵循财务报表勾稽关系。\n"
                "  * 横表（如月度明细表，各指标为独立列）：对明细月份范围直接按列聚合（表尾合计行已自动排除，无需额外添加月份 in 过滤）。\n"
                "  * 纵表（如利润表/资产负债表/现金流量表，[项目]列每行为不同指标）：若用 query.where 一次筛选多个项目，必须同时指定 group_by: ['项目']，或按明确单元格地址引用，严禁无 group_by 跨项目混加！"
            )

        lines = ["【表格预检与物理数据字典 (Pre-flight Sheet Profile)】"]
        if len(sheet_profiles) == 1:
            sp = sheet_profiles[0]
            truncation_note = f"（注意：物理表格共有 {sp['max_column']} 列，当前扫描前 {sp['max_c']} 列）" if sp["max_column"] > sp["max_c"] else ""
            sampling_note = "（注意：数据行数超过300行，离散字段取值基于前300行采样，非全量枚举）" if sp.get("is_sampled") else ""
            lines.append(
                f"- 目标工作表: {sp['sheet']} | 结构: 表头第 {sp['header_row']} 行，有效明细范围 {sp['data_range']}，共计 {sp['records']} 条有效业务明细（已排除表尾说明）{truncation_note}{sampling_note}。"
            )
            lines.append("- 物理字段清单与离散取值枚举:")
            lines.extend(_render_sheet_profile_lines(sp)[1:])
        else:
            for sp in sheet_profiles:
                # Skip verbose column listing for pure cover/instruction sheets when business sheets exist
                if sp["sheet"].strip().lower() in COVER_SHEET_NAMES and len(sheet_profiles) > 1:
                    lines.append(f"- 工作表 [{sp['sheet']}] (说明/目录页): 范围 {sp['data_range']}，字段 {list(sp['headers'].values())}。")
                    continue
                lines.extend(_render_sheet_profile_lines(sp))

        lines.append("\n【确认业务领域与口径指引】")
        lines.append(f"- 业务领域: {domain_title} ({domain_id})")
        if domain_notes:
            lines.append(domain_notes)

        compact_md = "\n".join(lines) + "\n"

        return {
            "compact_markdown": compact_md,
            "domain": domain_id,
            "sheet": primary_sp["sheet"],
            "header_row": primary_sp["header_row"],
            "data_range": primary_sp["data_range"],
            "records": primary_sp["records"],
            "headers": primary_sp["headers"],
            "columns": primary_sp["columns"],
            "sheets": sheet_profiles,
        }
    finally:
        wb.close()
