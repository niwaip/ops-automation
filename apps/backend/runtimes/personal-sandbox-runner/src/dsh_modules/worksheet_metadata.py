"""Coordinate metadata shared by normal and streaming worksheet readers."""

from openpyxl.utils import get_column_letter


def worksheet_dimensions(worksheet):
    """Never fabricate A1:A1 when a streaming sheet has no dimensions property.

    The OOXML used range is structural metadata, not proof of populated cells.
    Missing producer metadata may require openpyxl's streaming size calculation.
    """
    dimension = getattr(worksheet, "dimensions", None)
    if dimension:
        return dimension
    calculate = getattr(worksheet, "calculate_dimension", None)
    if calculate:
        try:
            return calculate()
        except ValueError:
            return calculate(force=True)
    rows, columns = worksheet.max_row, worksheet.max_column
    return f"A1:{get_column_letter(columns)}{rows}" if rows and columns else "未知范围"
