"""Read declared amount units from source notes; never infer magnitude from values."""

import re
from openpyxl.utils import get_column_letter

AMOUNT_UNIT_FACTORS = {
    "元": 1,
    "千元": 1000,
    "K": 1000,
    "k": 1000,
    "万元": 10000,
    "百万元": 1000000,
    "M": 1000000,
    "m": 1000000,
    "亿元": 100000000,
}


def normalize_unit(unit):
    """Canonicalize standard unit aliases without choosing a business default."""
    return {'千元': 'K', 'k': 'K', '百万元': 'M', 'm': 'M', 'yuan': '元',
            '小时': 'hours', 'h': 'hours'}.get(unit, unit)


def declared_amount_unit(workbook, selected_sheet, header_row):
    def collect(sheets):
        found = {}
        for sheet, last_row in sheets:
            for row in sheet.iter_rows(max_row=last_row, max_col=min(sheet.max_column, 128)):
                for cell in row:
                    value = cell.value
                    if not isinstance(value, str) or value.startswith("="):
                        continue
                    if "单位" in value:
                        for match in re.finditer(r"单位\s*[:：为]?\s*(?:人民币[，,\s]*)?([kKmM]|亿元|百万元|千元|万元|元)", value):
                            u = match.group(1)
                            if u in ('k', 'K'): u = 'K'
                            elif u in ('m', 'M'): u = 'M'
                            found.setdefault(u, []).append(f"{sheet.title}!{get_column_letter(cell.column)}{cell.row}")
                        # Common metadata layout: a unit label in one cell, its
                        # declared value in the adjacent cell. Read that field only.
                        if re.fullmatch(r"(?:币种及|金额|币种和)?单位[:：\s]*", value.strip()):
                            adjacent = sheet.cell(cell.row, cell.column+1)
                            if isinstance(adjacent.value, str):
                                for match in re.finditer(r"([kKmM]|亿元|百万元|千元|万元|元)(?=$|[\s，,。；;）)])", adjacent.value):
                                    u = match.group(1)
                                    if u in ('k', 'K'): u = 'K'
                                    elif u in ('m', 'M'): u = 'M'
                                    found.setdefault(u, []).append(
                                        f"{sheet.title}!{get_column_letter(adjacent.column)}{adjacent.row}")
                    # Also detect units in column headers, e.g. "金额(K)", "金额(万元)", "金额（千元）"
                    m_hdr = re.search(r"(?:金额|收入|成本|费用|利润|总额|工数|单价)?\s*[（(]([kKmM]|亿元|百万元|千元|万元|元)[)）]", value)
                    if m_hdr:
                        u = m_hdr.group(1)
                        if u in ('k', 'K'): u = 'K'
                        elif u in ('m', 'M'): u = 'M'
                        found.setdefault(u, []).append(f"{sheet.title}!{get_column_letter(cell.column)}{cell.row}")
        if len(found) == 1:
            unit, locations = next(iter(found.items()))
            return unit, locations
        if found:
            return None, [ref for refs in found.values() for ref in refs]
        return None, []
    unit, locations = collect([(workbook[selected_sheet], min(header_row, 20))])
    if unit or locations:
        return unit, locations
    return collect([(s, min(s.max_row, 20)) for s in workbook])
