"""Source-bound numeric column comparisons; no question routing or financial inference."""

import datetime
import hashlib
import json
import math
import re
import time
from decimal import Decimal
from pathlib import Path

import openpyxl
from openpyxl.utils import column_index_from_string, range_boundaries, get_column_letter

from .file_tools import resolve_sandboxed_path
from .tool_result import ToolResult, TextToolOutput
from .workbook_calculation import get_or_create_recalculated_workbook
from .spreadsheet_units import declared_amount_unit


def _deadline(deadline):
    if deadline is not None and time.monotonic() >= deadline:
        raise TimeoutError("Task deadline exceeded during column comparison")


def _column(value):
    if not isinstance(value, str) or not re.fullmatch(r"[A-Za-z]{1,3}", value):
        raise ValueError("Column must be an Excel column letter")
    column = column_index_from_string(value.upper())
    if column > 16384:
        raise ValueError("Column exceeds Excel limits")
    return column


def _group(value):
    if isinstance(value, (datetime.datetime, datetime.date)):
        return value.isoformat()
    return str(value).strip() if value is not None else ""


def _number(value):
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
        raise ValueError("Comparison inputs must be finite numeric values, not blanks or text")
    return Decimal(str(value))


def _source_headers(workbook, sheet, column, header_row, rows):
    """Attach actual column identity, including direct reference headers when available."""
    headers = {str(sheet.cell(header_row, column).value or "")}
    for row in rows:
        formula = str(sheet.cell(row, column).value or "")
        match = re.fullmatch(r"='?([^']+?)'?!\$?([A-Z]+)\$?\d+", formula)
        if match and match.group(1) in workbook.sheetnames:
            target = workbook[match.group(1)]
            col = column_index_from_string(match.group(2))
            # A reference's column heading may precede the data by several title rows.
            ref_row = int(re.search(r"\d+$", formula).group(0))
            for r in range(1, min(ref_row, 20)):
                value = target.cell(r, col).value
                if isinstance(value, str) and not value.startswith("="):
                    headers.add(value)
    return sorted(headers - {""})


def calculate_column_comparison(source, *, sheet, data_range, header_row, group_column,
                                comparisons, basis, rank_by, extreme, expected_groups=None,
                                deadline=None):
    """Compare explicit columns over an explicit population and retain all tied extrema."""
    _deadline(deadline)
    if rank_by not in ("absolute_difference", "difference", "absolute_relative_difference") or extreme not in ("max", "min"):
        raise ValueError("Invalid comparison ranking or extreme")
    if not isinstance(basis, str) or not basis.strip():
        raise ValueError("An explicit comparison basis is required")
    if not re.fullmatch(r"\$?[A-Za-z]+\$?\d+:\$?[A-Za-z]+\$?\d+", data_range or ""):
        raise ValueError("A finite data range such as A4:H15 is required")
    min_col, start, max_col, end = range_boundaries(data_range)
    if not 1 <= start <= end <= 1048576 or not 1 <= min_col <= max_col <= 16384 or end-start+1 > 100000:
        raise ValueError("Invalid or oversized data range")
    if not isinstance(header_row, int) or isinstance(header_row, bool) or not 1 <= header_row < start:
        raise ValueError("Header row must precede data rows")
    group_col = _column(group_column)
    if not isinstance(comparisons, list) or not 1 <= len(comparisons) <= 16:
        raise ValueError("Provide between 1 and 16 column comparisons")
    pairs = []
    for pair in comparisons:
        if not isinstance(pair, dict) or not isinstance(pair.get("label"), str) or not pair["label"].strip():
            raise ValueError("Each comparison requires a label")
        left, right = _column(pair.get("left_column")), _column(pair.get("right_column"))
        if left == right:
            raise ValueError("Comparison columns must be distinct")
        pairs.append((pair["label"], left, right))
    if len({p[0] for p in pairs}) != len(pairs):
        raise ValueError("Comparison labels must be unique")
    if any(not min_col <= c <= max_col for c in [group_col] + [c for _, l, r in pairs for c in (l, r)]):
        raise ValueError("All comparison/group columns must fall within data_range")
    source = Path(source).resolve()
    if source.suffix.lower() not in (".xlsx", ".xlsm"):
        raise ValueError("Expected an xlsx or xlsm workbook")
    source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    formulas = openpyxl.load_workbook(source, data_only=False)
    data = None
    try:
        if sheet not in formulas.sheetnames:
            return ToolResult.error("sheet_not_found", "工作表不存在", data={"available_sheets": formulas.sheetnames})
        worksheet = formulas[sheet]
        value_unit, unit_sources = declared_amount_unit(formulas, sheet, header_row)
        if end > worksheet.max_row:
            raise ValueError("Requested data range extends past existing rows")
        rows = list(range(start, end+1))
        has_formula = any(worksheet.cell(r, c).data_type == "f" for r in rows
                          for c in [group_col] + [c for _, l, q in pairs for c in (l, q)])
        read_path, calculation = (get_or_create_recalculated_workbook(source) if has_formula
                                 else (source, {"status": "literal_values"}))
        if has_formula:
            provenance = calculation.get("provenance", {})
            if calculation.get("status") not in ("recalculated_verified", "verified"):
                return ToolResult.error("calculation_unverified", "目标数值未通过重算验证", error_details={"calculation": calculation})
            if (provenance.get("workbook_id") != source_hash
                    or provenance.get("output_hash") != hashlib.sha256(Path(read_path).read_bytes()).hexdigest()):
                return ToolResult.error("calculation_provenance_mismatch", "重算凭据与原件/副本不一致")
        _deadline(deadline)
        data = openpyxl.load_workbook(read_path, data_only=True)
        values = data[sheet]
        groups = [_group(values.cell(r, group_col).value) for r in rows]
        if not all(groups) or len(set(groups)) != len(groups):
            raise ValueError("Group labels must be nonempty and unique; aggregate duplicate groups first")
        if expected_groups is not None:
            if (not isinstance(expected_groups, list) or not expected_groups
                    or any(not isinstance(v, str) or not v.strip() for v in expected_groups)
                    or len(set(expected_groups)) != len(expected_groups)
                    or set(expected_groups) != set(groups)):
                return ToolResult.error("incomplete_group_coverage", "实际分组与声明的完整任务范围不一致", data={"groups": groups})
        evaluated = []
        results = []
        for label, left_col, right_col in pairs:
            entries = []
            for row, group in zip(rows, groups):
                _deadline(deadline)
                left = _number(values.cell(row, left_col).value)
                right = _number(values.cell(row, right_col).value)
                difference = left-right
                relative = difference / abs(right) if right else None
                if any(not math.isfinite(float(value)) for value in (difference, relative) if value is not None):
                    raise ValueError("Derived comparison values exceed supported numeric range")
                if rank_by == "absolute_relative_difference" and relative is None:
                    return ToolResult.error("undefined_relative_difference", "基准为零，无法对声明范围内的全部相对差额排序")
                score = abs(relative) if rank_by == "absolute_relative_difference" else (abs(difference) if rank_by == "absolute_difference" else difference)
                entries.append({"group": group, "left": float(left), "right": float(right),
                                "difference": float(difference), "relative_difference": float(relative) if relative is not None else None,
                                "exact": {"left": str(left), "right": str(right), "difference": str(difference),
                                          "relative_difference": str(relative) if relative is not None else None},
                                "score": str(score), "source": {
                                    "group": f"{sheet}!{get_column_letter(group_col)}{row}",
                                    "left": f"{sheet}!{get_column_letter(left_col)}{row}",
                                    "right": f"{sheet}!{get_column_letter(right_col)}{row}"}})
            optimum = (max if extreme == "max" else min)(Decimal(e["score"]) for e in entries)
            matches = [e for e in entries if Decimal(e["score"]) == optimum]
            results.append({"label": label, "left_headers": _source_headers(formulas, worksheet, left_col, header_row, rows),
                            "right_headers": _source_headers(formulas, worksheet, right_col, header_row, rows), "matches": matches})
            evaluated.append({"label": label, "rows": entries})
        if hashlib.sha256(source.read_bytes()).hexdigest() != source_hash:
            return ToolResult.error("source_changed", "计算期间原件发生变化，结果已失效")
        plan = {"sheet": sheet, "data_range": data_range, "header_row": header_row,
                "group_column": group_column, "comparisons": comparisons, "basis": basis,
                "rank_by": rank_by, "extreme": extreme, "expected_groups": expected_groups}
        plan_hash = hashlib.sha256(json.dumps(plan, sort_keys=True, ensure_ascii=False).encode()).hexdigest()
        evidence_id = hashlib.sha256((source_hash+plan_hash).encode()).hexdigest()[:24]
        return ToolResult.success(data={"kind": "column_comparison", "basis": basis,
                                       "plan": plan, "plan_hash": plan_hash, "evidence_id": evidence_id,
                                       "value_unit": value_unit,
                                       "difference_formula": "left-right", "relative_formula": "(left-right)/abs(right)",
                                       "rank_by": rank_by, "extreme": extreme,
                                       "scope": {"data_range": data_range, "groups": groups, "row_count": len(rows),
                                                 "expected_groups": expected_groups, "coverage_verified": expected_groups is not None},
                                       "results": results, "evaluated": evaluated},
                                  provenance={"file_path": str(source), "workbook_id": source_hash,
                                              "sheet": sheet, "calculation": calculation,
                                              "unit_sources": unit_sources})
    finally:
        formulas.close()
        if data is not None:
            data.close()


def execute_column_comparison(params, deadline=None):
    """Sandbox adapter: full typed evidence stays in runtime; compact JSON enters model context."""
    source, error = resolve_sandboxed_path(params.get("file_path", ""))
    if error or source is None:
        result = ToolResult.error("invalid_file_path", str(error))
    else:
        arguments = {k: v for k, v in params.items() if k != "file_path"}
        try:
            result = calculate_column_comparison(source, deadline=deadline, **arguments)
        except TimeoutError:
            raise
        except Exception as exc:
            result = ToolResult.error("invalid_comparison", str(exc), data={"requested_plan": dict(params)})
    if result.is_error:
        result.data = {**(result.data if isinstance(result.data,dict) else {}), 'requested_plan':dict(params)}
    compact = result.to_dict()
    if result.is_success:
        compact["data"] = {k: v for k, v in result.data.items() if k != "evaluated"}
    compact.pop("text_representation", None)
    import json
    return TextToolOutput(json.dumps(compact, ensure_ascii=False), result)
