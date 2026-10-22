"""Validation and semantic totals for generated workbooks, independent of model text."""

from decimal import Decimal, InvalidOperation
from pathlib import Path
import re
import openpyxl
from openpyxl.formula.tokenizer import Tokenizer
from .workbook_calculation import check_workbook_formulas


def inspect_generated_workbook(path):
    issues = []
    summary = check_workbook_formulas(path)
    if summary.get('error'):
        issues.append(str(summary['error']))
    if summary.get('total_errors', 0):
        issues.append(f"公式错误 {summary['total_errors']} 处")
    if summary.get('uncalculated_formulas', 0):
        issues.append(f"缺少公式缓存 {summary['uncalculated_formulas']} 处")
    book = None
    try:
        book = openpyxl.load_workbook(path, read_only=True, data_only=False)
        visited = 0
        for sheet in book:
            sheet.reset_dimensions()
            for row in sheet:
                for cell in row:
                    visited += 1
                    if visited > 200000:
                        raise ValueError('交付工作簿检查超出单元格预算')
                    if cell.data_type != 'f':
                        continue
                    for token in Tokenizer(cell.value).items:
                        if token.subtype != 'RANGE' or '!' not in token.value:
                            continue
                        name = token.value.rsplit('!', 1)[0].strip("'").replace("''", "'")
                        refs = name.split(':')
                        if '[' in name or any(ref not in book.sheetnames for ref in refs):
                            issues.append(f'{sheet.title}!{cell.coordinate} 缺少内部公式依赖：{name}')
    except Exception as error:
        issues.append(str(error))
    finally:
        if book:
            book.close()
    return {'valid': not issues, 'issues': list(dict.fromkeys(issues)), 'formula_summary': summary}


def decimal_value(value):
    if value is None or isinstance(value, bool):
        return None
    try:
        number = Decimal(str(value).replace(',', '').replace('，', '').strip())
        return number if number.is_finite() else None
    except InvalidOperation:
        return None


def generated_workbook_totals(path):
    """Bind totals by header column, including integer/zero amounts and reordered columns."""
    if not inspect_generated_workbook(path)['valid']:
        return None
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        candidates = []
        for sheet in book:
            rows = list(sheet.values)
            headers = []
            for index, row in enumerate(rows[:20]):
                amounts = [i for i, c in enumerate(row) if isinstance(c, str) and re.search(r'金额|amount', c, re.I)]
                counts = [i for i, c in enumerate(row) if isinstance(c, str) and re.search(r'条数|记录数|笔数|count', c, re.I)]
                if len(amounts) == 1 and len(counts) <= 1:
                    headers.append((index, amounts[0], counts[0] if counts else None, row[amounts[0]]))
            if len(headers) != 1:
                continue
            header, amount_col, count_col, caption = headers[0]
            for row in reversed(rows[header + 1:]):
                if not any(isinstance(c, str) and re.match(r'^\s*(?:总计|合计|total|grand total|summary)(?:\s|$|[（(])', c, re.I) for c in row):
                    continue
                amount = decimal_value(row[amount_col])
                count = decimal_value(row[count_col]) if count_col is not None else None
                if amount is None or (count is not None and (count < 0 or count != count.to_integral_value())):
                    continue
                match = re.search(r'[（(]([^()（）]+)[)）]', caption)
                candidates.append({'amount': amount, 'count': int(count) if count is not None else None,
                                   'amount_unit': match.group(1).strip() if match else None,
                                   'file_path': Path(path), 'sheet_name': sheet.title})
                break
        return candidates[0] if len(candidates) == 1 else None
    finally:
        book.close()


def deliverable_is_valid(path):
    path = Path(path)
    return path.suffix.lower() not in ('.xlsx', '.xlsm') or inspect_generated_workbook(path)['valid']


def quarantine_invalid_deliverables(workspace, turn_start_time, text, outbounds):
    """Keep invalid new outputs recoverable but outside root artifact auto-discovery."""
    import json
    import uuid
    from urllib.parse import quote, unquote
    from .file_tools import is_protected_session_file
    root = Path(workspace)
    rejected = []
    for path in root.iterdir() if root.exists() else []:
        if (path.is_symlink() or not path.is_file() or path.name.startswith('.')
                or path.suffix.lower() not in ('.xlsx', '.xlsm')
                or path.stat().st_mtime < turn_start_time - 2
                or is_protected_session_file(path)):
            continue
        validation = inspect_generated_workbook(path)
        if validation['valid']:
            continue
        folder = root / '.dsh' / 'rejected' / uuid.uuid4().hex
        folder.mkdir(parents=True, exist_ok=True)
        path.replace(folder / path.name)
        rejected.append({'fileName': path.name, 'issues': validation['issues'], 'preserved_path': str(folder / path.name)})
    if not rejected:
        return text, outbounds, []
    names = {r['fileName'] for r in rejected}
    kept = []
    for item in outbounds:
        try:
            payload = json.loads(item) if isinstance(item, str) else item
            name = payload.get('fileName') or Path(payload.get('filePath', '')).name
        except (ValueError, TypeError, AttributeError):
            name = Path(str(item)).name
        if unquote(name) not in names:
            kept.append(item)
    lines = [line for line in text.splitlines() if not any(name in unquote(line) or quote(name) in line for name in names)]
    notice = '\n\nExcel 交付校验未通过，未提供该文件的下载链接；错误文件已保留供排查：' + '；'.join(
        r['fileName'] + '（' + '；'.join(r['issues'][:3]) + '）' for r in rejected)
    return '\n'.join(lines) + notice, kept, rejected
