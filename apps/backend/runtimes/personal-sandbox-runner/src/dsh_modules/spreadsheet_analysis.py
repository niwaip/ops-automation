"""Execute generic analysis plans against verified source values, without code generation."""

import hashlib
import json
import os
import re
import time
from zipfile import BadZipFile
from pathlib import Path
from decimal import Decimal
import openpyxl
from openpyxl.utils import get_column_letter
from .file_tools import resolve_sandboxed_path
from .tool_result import ToolResult, TextToolOutput
from .spreadsheet_context import resolve_session_attachment, resolve_sheet_name
from .workbook_calculation import get_or_create_recalculated_workbook
from .spreadsheet_units import declared_amount_unit
from .spreadsheet_expression import ExpressionEvaluator, bounded_range, source_column, numeric, present_quantity
from .spreadsheet_quality import profiles, evaluate_rule
from .spreadsheet_scope import table_months, validate_caption
from .spreadsheet_checks import validate_source_identity
from .spreadsheet_plan import normalize_plan
from .analysis_semantics import domain_pack, runtime_domain, validate_domain_calculation, validate_metric_result
from .spreadsheet_semantic_bindings import semantic_bindings
from .query_executor import execute_query_plan


def compile_legacy_filters_to_predicate(filters):
    if not filters:
        return None
    predicates = []
    for f in filters:
        col = f.get("column")
        op = f.get("op", "in")
        if op == "in":
            predicates.append({"op": "in", "field": col, "values": f.get("values", [])})
        elif op == "not_in":
            predicates.append({"op": "not_in", "field": col, "values": f.get("values", [])})
        elif op == "prefix":
            predicates.append({"op": "starts_with", "field": col, "prefixes": f.get("prefixes", [])})
        elif op == "exclude_prefix":
            predicates.append({"op": "not_starts_with", "field": col, "prefixes": f.get("exclude_prefixes", f.get("prefixes", []))})
    if not predicates:
        return None
    if len(predicates) == 1:
        return predicates[0]
    return {"and": predicates}


def deadline_check(deadline):
    if deadline is not None and time.monotonic() >= deadline:
        raise TimeoutError("Spreadsheet analysis deadline exceeded")


def prepare_tables(plan, workbook, check, apply_filters=True):
    result = {}; cell_budget = 0
    for item in plan:
        check()
        allowed_tbl_keys = {"id", "sheet", "data_range", "header_row", "filters", "where", "group_by", "having"}
        if not isinstance(item, dict) or set(item) - allowed_tbl_keys:
            raise ValueError("非法数据表计划")
        if not {"id", "sheet", "data_range", "header_row"} <= set(item):
            raise ValueError("数据表计划缺少必要字段")
        sheet_target = resolve_sheet_name(item["sheet"], workbook.sheetnames)
        if sheet_target not in workbook.sheetnames:
            raise KeyError(f"Worksheet {item['sheet']} does not exist.")
        item["sheet"] = sheet_target
        ws = workbook[sheet_target]; bounds = bounded_range(ws, item["data_range"])
        a, start, b, end = bounds; header = item["header_row"]
        if isinstance(header, bool) or not isinstance(header, int) or not 1 <= header <= start:
            raise ValueError("表头必须在业务范围之前")
        header_values=[ws.cell(header,col).value for col in range(a,b+1)]
        if any(v is not None and (not isinstance(v,str) or v.startswith('=')) for v in header_values):
            raise ValueError("声明表头含数据或公式，需要明确的单层文本字段名")
        cell_budget += (end-start+1)*(b-a+1)
        if cell_budget > 200000:
            raise ValueError("合计业务范围超出预算")
        # Explicit header metadata is authoritative. Normalize a range starting
        # on that header exactly; do not guess a different header or period.
        excluded=[]
        if start==header:
            excluded.append(start);start+=1

        # 闭环范围约束：自动修剪表尾非业务行（查询条件、报表类别、系统附注、总计行等）
        # 必须带完整词边界或明确分隔符，严禁将 "TotalEnergies" 等正常业务名称误当成 "total" 删行
        footer_pat = re.compile(
            r'^\s*(?:'
            r'[\*＊#※]'
            r'|(?:合\s*计|总\s*计|小\s*计|grand\s+total|subtotal|total|summary)(?:[:：\s]|$)'
            r'|(?:查询条件|报表(?:类别|说明)|填报说明|编制(?:说明|单位)|制表人|制表时间|打印时间|导出时间|生成时间|温馨提示|特别提示|数据来源)(?:[:：\s]|$)'
            r'|(?:备注|注|说明|附注|声明)[:：]'
            r'|(?:第\s*\d+\s*页|page\s*\d+)'
            r')',
            re.I
        )
        col_count = b - a + 1
        while end >= start:
            row_cells = [ws.cell(end, c) for c in range(a, b + 1) if ws.cell(end, c).value is not None]
            if not row_cells:
                excluded.append(end)
                end -= 1
                continue
            first_val = str(row_cells[0].value or "").strip()
            is_footer = False

            # 结构密度与度量连续性校验：
            # 若行填充率较高且后续列包含有效数值度量，严禁因第一列带有敏感词前缀而误删正常业务明细
            # （如 'Total Energy Services'、'*Star Consulting'、'Page 3 Solutions' 均为合法业务记录）
            populated_ratio = len(row_cells) / max(col_count, 1)
            has_measure_numbers = any(
                isinstance(c.value, (int, float, Decimal)) and not isinstance(c.value, bool)
                for c in row_cells[1:]
            )

            if populated_ratio >= 0.5 and has_measure_numbers:
                if re.fullmatch(r'(?:合\s*计|总\s*计|小\s*计|grand\s+total|total|subtotal|summary)[:：]?', first_val, re.I):
                    is_footer = True
            else:
                if footer_pat.search(first_val):
                    is_footer = True
                elif col_count >= 4 and len(row_cells) <= max(2, col_count // 4):
                    if any(footer_pat.search(str(c.value or "").strip()) for c in row_cells):
                        is_footer = True
            if is_footer:
                excluded.append(end)
                end -= 1
            else:
                break
        bounds = (a, start, b, end)
        normalized_data_range = f"{get_column_letter(a)}{start}:{get_column_letter(b)}{end}"
        rows = list(range(start, end+1)); filters = item.get("filters", [])
        if not isinstance(filters, list) or len(filters) > 8:
            raise ValueError("过滤条件超出预算")
        for cond in filters:
            if not isinstance(cond, dict) or "column" not in cond:
                raise ValueError("过滤需要明确column")
            allowed_filter_keys = {"column", "values", "exclude_values", "prefixes", "exclude_prefixes", "op", "case_sensitive"}
            if set(cond) - allowed_filter_keys:
                raise ValueError("过滤条件包含不支持的字段")
            col = source_column(workbook, {**item, "bounds": bounds}, cond["column"])
            if not a <= col <= b:
                raise ValueError("过滤列不在业务范围")
            if apply_filters:
                op = cond.get("op", "")
                case_sensitive = cond.get("case_sensitive", False)
                if op in ("exclude_prefix", "exclude_prefixes", "not_prefix", "prefix_exclude") or "exclude_prefixes" in cond:
                    prefixes = cond.get("exclude_prefixes") or cond.get("prefixes") or cond.get("values") or []
                    if not prefixes or not isinstance(prefixes, list):
                        raise ValueError("前缀排除过滤需要非空的前缀列表")
                    ex_prefixes = [str(p) if case_sensitive else str(p).lower() for p in prefixes]
                    def matches_ex_prefix(val):
                        v = str(val or "").strip()
                        if not case_sensitive: v = v.lower()
                        return not any(v.startswith(p) for p in ex_prefixes)
                    rows = [r for r in rows if matches_ex_prefix(ws.cell(r, col).value)]
                elif op in ("prefix", "prefixes", "starts_with", "startswith") or "prefixes" in cond:
                    prefixes = cond.get("prefixes") or cond.get("values") or []
                    if not prefixes or not isinstance(prefixes, list):
                        raise ValueError("前缀包含过滤需要非空的前缀列表")
                    req_prefixes = [str(p) if case_sensitive else str(p).lower() for p in prefixes]
                    def matches_prefix(val):
                        v = str(val or "").strip()
                        if not case_sensitive: v = v.lower()
                        return any(v.startswith(p) for p in req_prefixes)
                    rows = [r for r in rows if matches_prefix(ws.cell(r, col).value)]
                elif op in ("not_in", "exclude", "not_equal", "!=") or "exclude_values" in cond:
                    ex_vals = cond.get("exclude_values") or cond.get("values")
                    if not ex_vals or not isinstance(ex_vals, list):
                        raise ValueError("过滤需要非空values数组")
                    def matches_not_in(val):
                        if val in ex_vals:
                            return False
                        v_str = str(val or "").strip()
                        if not case_sensitive:
                            v_str = v_str.casefold()
                            return not any(v_str == str(x).strip().casefold() for x in ex_vals if x is not None)
                        return not any(v_str == str(x).strip() for x in ex_vals if x is not None)
                    rows = [r for r in rows if matches_not_in(ws.cell(r, col).value)]
                elif op in ("in", "equals", "==") or (not op and "values" in cond):
                    if "values" not in cond or not isinstance(cond["values"], list) or not cond["values"]:
                        raise ValueError("过滤需要非空values数组")
                    in_vals = cond["values"]
                    def matches_in(val):
                        if val in in_vals:
                            return True
                        v_str = str(val or "").strip()
                        if not case_sensitive:
                            v_str = v_str.casefold()
                            return any(v_str == str(x).strip().casefold() for x in in_vals if x is not None)
                        return any(v_str == str(x).strip() for x in in_vals if x is not None)
                    rows = [r for r in rows if matches_in(ws.cell(r, col).value)]
                else:
                    raise ValueError(f"不支持的过滤操作符: {op}")
        if not rows and not filters:
            raise ValueError("声明范围没有业务记录")
        result[item["id"]] = {**item, "data_range": normalized_data_range, "bounds": bounds, "selected_rows": rows,"excluded_header_rows":excluded,
                               "headers":{get_column_letter(c):str(ws.cell(header,c).value or '') for c in range(a,b+1)},
                               "input_records": end-start+1, "evaluated_records": len(rows)}
    return result


def calculate_spreadsheet_analysis(source, plan, deadline=None, domain="generic"):
    original = values = None
    requested_plan=plan
    error_path="plan"
    try:
        check = lambda: deadline_check(deadline)
        check(); source = Path(source).resolve()
        if source.suffix.lower() == '.csv':
            return ToolResult.error(
                "unsupported_source_format",
                f"原件 '{source.name}' 为纯文本 CSV 格式，openpyxl 表格聚合引擎仅支持 .xlsx 工作簿。如需汇总 CSV 数据，请使用 read_file/bash 执行 Python 脚本（如 pandas/csv 模块）完成计算与导出，或先转换为 xlsx 文件。",
                data={"source_format": "csv", "supported_formats": [".xlsx", ".xlsm"]}
            )
        pack = domain_pack(domain)
        allowed_plan_keys = {
            "tables", "calculations", "checks", "rules", "profile_tables",
            "query", "where", "group_by", "having", "order_by", "limit",
            "offset", "distinct", "distinct_fields", "fields", "metrics", "source"
        }
        if set(plan) - allowed_plan_keys:
            raise ValueError("分析计划包含未知字段")

        source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
        original = openpyxl.load_workbook(source, data_only=False)

        t0 = plan.get("tables", [{}])[0] if isinstance(plan.get("tables"), list) and plan.get("tables") else {}
        has_query_features = (
            "query" in plan
            or "metrics" in plan
            or "source" in plan
            or "where" in plan
            or "where" in t0
            or "having" in plan
            or "having" in t0
            or "order_by" in plan
            or "limit" in plan
            or "distinct" in plan
            or "select" in plan
            or "fields" in plan
            or "mode" in plan
            or (isinstance(plan.get("group_by"), list) and len(plan.get("group_by")) > 1)
            or any(isinstance(c.get("expression"), dict) and (
                c["expression"].get("aggregate") in ("avg", "dec_avg", "count_distinct")
                or c["expression"].get("ref") in ("task_rows", "rows", "records")
                or c["expression"].get("mode") == "records"
            ) for c in plan.get("calculations", []))
        )

        has_formula = any(c.data_type == "f" for ws in original for row in ws for c in row)
        read_path, calculation = get_or_create_recalculated_workbook(source) if has_formula else (source, {"status":"literal_values"})
        is_partial_recalc = False
        if has_formula:
            provenance = calculation.get("provenance", {})
            calc_status = calculation.get("status")
            if (calc_status not in ("verified", "recalculated_verified", "partial")
                    or provenance.get("workbook_id") != source_hash
                    or provenance.get("output_hash") != hashlib.sha256(Path(read_path).read_bytes()).hexdigest()):
                return ToolResult.error("calculation_unverified", "重算或来源验证未通过", data=calculation)
            if calc_status == "partial":
                is_partial_recalc = True

        def _check_partial_error_intersection(visited_cells_set):
            if not is_partial_recalc:
                return None
            err_set = set()
            for err_list in (calculation.get("error_summary") or {}).values():
                for loc in err_list:
                    err_set.add(str(loc).upper())
                    err_set.add(str(loc).split("!")[-1].upper())
            for loc in (calculation.get("uncalculated_locations") or []):
                err_set.add(str(loc).upper())
                err_set.add(str(loc).split("!")[-1].upper())
            for v in visited_cells_set:
                clean_v = str(v).split("!")[-1].upper()
                if clean_v in err_set:
                    return clean_v
            return None
        check(); values = openpyxl.load_workbook(read_path, data_only=True)

        units = [declared_amount_unit(original, ws.title, min(ws.max_row,20)) for ws in original]
        declared = {u for u, _ in units if u}
        amount_unit = next(iter(declared)) if len(declared)==1 else "undeclared_amount"

        if has_query_features:
            legacy_pred = compile_legacy_filters_to_predicate(t0.get("filters", []))
            if "query" in plan:
                q_spec = dict(plan["query"])
                if legacy_pred:
                    q_where = q_spec.get("where")
                    q_spec["where"] = {"and": [legacy_pred, q_where]} if q_where else legacy_pred
            elif "metrics" in plan or (isinstance(plan.get("where"), dict) and not plan.get("calculations")):
                q_spec = dict(plan)
                if legacy_pred:
                    q_where = q_spec.get("where")
                    q_spec["where"] = {"and": [legacy_pred, q_where]} if q_where else legacy_pred
            else:
                metrics_spec = []
                projection_fields = []
                for c in plan.get("calculations", []):
                    expr = c.get("expression", {})
                    if isinstance(expr, dict) and (expr.get("ref") in ("task_rows", "rows", "records") or expr.get("mode") == "records"):
                        if expr.get("fields"):
                            projection_fields.extend(expr.get("fields"))
                        continue
                    col_target = expr.get("column") or expr.get("field")
                    if not col_target and expr.get("aggregate") not in ("count_rows", "count"):
                        continue
                    metrics_spec.append({
                        "id": c["id"],
                        "label": c.get("label", c["id"]),
                        "op": expr.get("aggregate", "sum"),
                        "field": col_target,
                        "unit": c.get("output_unit") or expr.get("unit"),
                    })
                grp = plan.get("group_by") or t0.get("group_by")
                if not grp:
                    declared_groups = [c.get('expression', {}).get('group_by') for c in plan.get('calculations', [])]
                    if declared_groups and all(g == declared_groups[0] for g in declared_groups):
                        grp = declared_groups[0]
                if grp and not isinstance(grp, list):
                    grp = [grp]

                explicit_where = t0.get("where") or plan.get("where")
                if legacy_pred and explicit_where:
                    combined_where = {"and": [legacy_pred, explicit_where]}
                else:
                    combined_where = explicit_where or legacy_pred

                all_fields = projection_fields or plan.get("select") or plan.get("fields") or plan.get("distinct_fields")
                q_spec = {
                    "source": {
                        "sheet": t0.get("sheet"),
                        "header_row": t0.get("header_row", 1),
                        "data_range": t0.get("data_range"),
                        "table_id": t0.get("id", "detail"),
                    },
                    "where": combined_where,
                    "group_by": grp,
                    "metrics": metrics_spec,
                    "having": plan.get("having") or t0.get("having"),
                    "order_by": plan.get("order_by"),
                    "limit": plan.get("limit"),
                    "offset": plan.get("offset"),
                    "distinct": plan.get("distinct", False),
                    "distinct_fields": plan.get("distinct_fields") or plan.get("fields"),
                    "fields": all_fields,
                    "select": all_fields,
                }

            query_res = execute_query_plan(values, q_spec, check, amount_unit)
            if is_partial_recalc:
                visited = getattr(query_res, "visited_cells", []) or (query_res.data.get("visited_cells") if isinstance(getattr(query_res, "data", None), dict) else [])
                err_cell = _check_partial_error_intersection(visited)
                if err_cell:
                    return ToolResult.error("calculation_unverified", f"查询涉及的单元格 {err_cell} 包含公式计算错误", data=calculation)
            canonical = json.dumps(plan, sort_keys=True, ensure_ascii=False, separators=(",", ":"))
            plan_hash = hashlib.sha256(canonical.encode()).hexdigest()
            return query_res.to_tool_result(source, source_hash, requested_plan, q_spec, plan_hash, amount_unit)

        for key, limit in (("tables", 16), ("calculations", 64), ("checks", 32), ("rules", 32)):
            if not isinstance(plan.get(key, []), list) or len(plan.get(key, [])) > limit:
                raise ValueError("分析项超出预算")
        prepare_tables(plan.get("tables", []), original, check, apply_filters=False)
        tasks = [item for key in ("calculations", "checks", "rules") for item in plan.get(key, [])]
        if not tasks and not (plan.get('profile_tables') is True and plan.get('tables')):
            raise ValueError("未制定分析计划，不能判定任务完成")
        ids = [item.get("id") for item in plan.get("tables", [])+tasks]
        if any(not isinstance(i,str) or not 1<=len(i)<=64 or not i.isprintable() or i.strip()!=i for i in ids) or len(set(ids)) != len(ids):
            raise ValueError("计划ID必须非空、唯一、可打印且不超过64字符")
        if any(not isinstance(t.get("label"), str) or not t["label"].strip() for t in tasks):
            raise ValueError("分析项需要明确名称")
        tables = prepare_tables(plan.get("tables", []), values, check)
        plan,normalizations=normalize_plan(plan,values,tables)
        units = [declared_amount_unit(original, ws.title, min(ws.max_row,20)) for ws in original]
        declared = {u for u, _ in units if u}
        amount_unit = next(iter(declared)) if len(declared)==1 else "undeclared_amount"
        evaluator = ExpressionEvaluator(values, tables, amount_unit, check)
        for table in tables.values():
            table["source_months"] = table_months(values, table)
        facts = []
        for item in plan.get("calculations", []):
            error_path=f"calculations.{item.get('id')}（{item.get('label')}）"
            if not {"id", "label", "expression"} <= set(item) or set(item)-{"id", "label", "expression", "output_unit", "metric_id"}:
                raise ValueError("非法计算计划")
            validate_caption(item, values, tables)
            validate_domain_calculation(item, semantic_bindings(item['expression'], values, tables), domain)
            expr = item["expression"]
            group_by_col = expr.get("group_by") if isinstance(expr, dict) else None
            if group_by_col:
                grouped_res = evaluator.evaluate_grouped(expr)
                for item_res in grouped_res:
                    if len(item_res) == 6:
                        g_val, q, g_sources, g_count, g_id_part, g_fields = item_res
                    else:
                        g_val, q, g_sources, g_count = item_res[:4]
                        g_id_part = g_val
                        g_fields = {}
                    sources = sorted(g_sources)
                    evaluator.visited_cells |= g_sources
                    if not sources:
                        raise ValueError("计算必须引用真实来源，不能用常量伪造业务事实")
                    if q.unit == "undeclared_amount":
                        raise ValueError("金额单位未声明或有冲突，请读取并澄清来源单位")
                    presented = present_quantity(q, item.get('output_unit', 'source'))
                    validate_metric_result(item.get('metric_id'), presented['unit'], domain)
                    facts.append({"id": f"{item['id']}__{g_id_part}", "label": f"{g_val} {item['label']}",
                                  "group_key": str(g_val), "group_dimension": group_by_col,
                                  "source_cells": sources, "matched_records": g_count,
                                  **({"metric_id": item["metric_id"]} if "metric_id" in item else {}),
                                  **({"group_fields": g_fields} if g_fields else {}),
                                  **presented})
            else:
                before = evaluator.visited_cells
                evaluator.visited_cells = set()
                q = evaluator.evaluate(item["expression"])
                sources = sorted(evaluator.visited_cells)
                evaluator.visited_cells |= before
                if not sources:
                    raise ValueError("计算必须引用真实来源，不能用常量伪造业务事实")
                if q.unit == "undeclared_amount":
                    raise ValueError("金额单位未声明或有冲突，请读取并澄清来源单位")
                presented=present_quantity(q,item.get('output_unit','source'))
                validate_metric_result(item.get('metric_id'),presented['unit'],domain)
                table_id = item["expression"].get("table") if isinstance(item.get("expression"), dict) else None
                matched_records = tables[table_id]["evaluated_records"] if table_id in tables else None
                facts.append({"id": item["id"], "label": item["label"], "source_cells":sources,
                              **({"metric_id":item["metric_id"]} if "metric_id" in item else {}),
                              **({"matched_records": matched_records} if matched_records is not None else {}),
                              **presented})
        checks = []
        for item in plan.get("checks", []):
            error_path=f"checks.{item.get('id')}（{item.get('label')}）"
            if not {"id", "label", "left", "right", "tolerance", "assumptions"} <= set(item) or set(item)-{"id", "label", "left", "right", "tolerance", "assumptions", "relation_type"} or not isinstance(item["assumptions"],str) or not item["assumptions"].strip():
                raise ValueError("等式检查必须包含明确假设与容差")
            validate_source_identity(item, original)
            before = evaluator.visited_cells
            evaluator.visited_cells = set()
            left, right = evaluator.evaluate(item["left"]), evaluator.evaluate(item["right"])
            sources = sorted(evaluator.visited_cells)
            evaluator.visited_cells |= before
            if not sources:
                raise ValueError("等式检查必须引用真实来源")
            if left.unit != right.unit or left.unit == "undeclared_amount":
                raise ValueError("等式两侧维度不同或单位未声明")
            tolerance = numeric(item["tolerance"])
            if tolerance < 0:
                raise ValueError("容差不能为负")
            checks.append({"id":item["id"],"label":item["label"],"left":str(left.value),"right":str(right.value),
                           "residual":str(left.value-right.value),"unit":left.unit,"tolerance":str(tolerance),
                           "passed":abs(left.value-right.value)<=tolerance,"assumptions":item["assumptions"],
                           "source_cells":sources})
        if is_partial_recalc:
            err_cell = _check_partial_error_intersection(evaluator.visited_cells)
            if err_cell:
                return ToolResult.error("calculation_unverified", f"计算引用的单元格 {err_cell} 包含公式计算错误", data=calculation)

        rules = [evaluate_rule(r, values, tables, check) for r in plan.get("rules", [])]
        profiles_result = profiles(values, tables, check) if rules or plan.get('profile_tables') is True else []
        if hashlib.sha256(source.read_bytes()).hexdigest() != source_hash:
            return ToolResult.error("source_changed", "分析期间原件已变化")
        coverage = [{"id":t["id"],"sheet":t["sheet"],"data_range":t["data_range"],"header_row":t["header_row"],
                     "input_records":t["input_records"],"evaluated_records":t["evaluated_records"],
                     "excluded_header_rows":t["excluded_header_rows"],
                     "headers":t["headers"],"source_months":t.get("source_months"),
                     "selected_rows":t["selected_rows"],"filters":t.get("filters",[]),"coverage_complete":True} for t in tables.values()]
        from openpyxl.utils import column_index_from_string
        extracted_rows = []
        for t in tables.values():
            has_explicit_filter = bool(t.get("filters"))
            is_records_mode = plan.get("mode") == "records" or plan.get("extract_rows") is True
            # 仅当存在明确筛选条件（如筛选特定员工/项目）或显式声明了记录抽取模式时才提取明细行
            if not has_explicit_filter and not is_records_mode:
                continue
            ws_sheet = values[t["sheet"]]
            col_headers = {column_index_from_string(c): h for c, h in t.get("headers", {}).items() if h}
            for r_idx in t.get("selected_rows", [])[:100]:
                r_item = {}
                for c_idx, h_name in col_headers.items():
                    val = ws_sheet.cell(r_idx, c_idx).value
                    r_item[h_name] = str(val) if val is not None else ""
                if r_item:
                    extracted_rows.append(r_item)

        canonical = json.dumps(plan, sort_keys=True, ensure_ascii=False,separators=(",",":"))
        plan_hash = hashlib.sha256(canonical.encode()).hexdigest()
        data = {"kind":"spreadsheet_analysis","plan":plan,"requested_plan":requested_plan,
                "semantic_domain":{"id":domain,"version":pack.VERSION if pack else "1"},
                "normalizations":normalizations,"plan_hash":plan_hash,"facts":facts,"checks":checks,
                "rules":rules,"profiles":profiles_result,"coverage":coverage,"rows":extracted_rows,
                "visited_cells":sorted(evaluator.visited_cells),"amount_unit":amount_unit,
                "unit_sources":sorted({ref for _, refs in units for ref in refs})}
        data["evidence_id"] = hashlib.sha256((source_hash+json.dumps(data,sort_keys=True,ensure_ascii=False,default=str)).encode()).hexdigest()[:24]
        provenance = {"file_path":str(source),"workbook_id":source_hash,"calculation":calculation,"plan_hash":plan_hash}
        # Compact model view, full evidence remains available to the verifier.
        view = {k:data[k] for k in ("evidence_id","rules","profiles","amount_unit","semantic_domain")}
        if normalizations:view["protocol_normalizations"]=normalizations[:20]
        if extracted_rows:
            view["rows"] = extracted_rows[:50]
            view["total_rows"] = len(extracted_rows)
        for key in ("facts", "checks"):
            view[key]=[{**{k:v for k,v in item.items() if k!="source_cells"},
                        "source_cell_count":len(item["source_cells"])} for item in data[key]]
        view["coverage"]=[{k:v for k,v in c.items() if k!="selected_rows"} for c in coverage]
        return ToolResult.success(data=data,provenance=provenance,text=json.dumps(view,ensure_ascii=False,default=str))
    except TimeoutError:
        raise
    except Exception as e:
        err_msg = str(e)
        code = "unsupported_source_format" if ("csv" in err_msg.lower() or "not support .csv" in err_msg) else "invalid_analysis_plan"
        return ToolResult.error(code, f"{error_path}：{err_msg}", data={"requested_plan":requested_plan},
                                provenance={"file_path":str(source),"workbook_id":locals().get("source_hash"),"plan_hash":hashlib.sha256(json.dumps(plan,sort_keys=True,ensure_ascii=False,default=str).encode()).hexdigest()})
    finally:
        if original:original.close()
        if values:values.close()


def execute_spreadsheet_analysis(params, deadline=None):
    # Decode JSON-encoded structured fields at the protocol boundary only.
    # No inference or repair of business ranges/expressions happens here.
    params = dict(params)
    if "plan" in params and isinstance(params["plan"], dict):
        plan_obj = params.pop("plan")
        params.update(plan_obj)
    for key in ("tables", "calculations", "checks", "rules"):
        if isinstance(params.get(key),str):
            try:
                params[key] = json.loads(params[key])
            except ValueError:
                result = ToolResult.error("invalid_analysis_plan",f"{key}必须是JSON数组")
                return TextToolOutput(result.render_text(),result)
    source, err_res = resolve_session_attachment(params.get("file_path", ""), "当前原生分析")
    if err_res:
        return TextToolOutput(err_res.render_text(), err_res)
    result = calculate_spreadsheet_analysis(source, {k: v for k, v in params.items() if k != "file_path"}, deadline,
                                            domain=runtime_domain())
    return TextToolOutput(result.render_text(), result)
