"""Compile request-bound workbook observations into exact verified writes."""
import datetime
import hashlib
from dataclasses import replace
from pathlib import Path
import openpyxl
from openpyxl.utils import get_column_letter
from .bindings import LogicalKey
from .compiler import compile_update_plan
from .discovery import extract_workbook_ir
from .layout_binding import locate,verify_unit
from .result_catalog import bind_result,convert_value,fingerprint,normalize_unit,select_result,serialize_facts
from .sources import FactTable
from ..user_declarations import unit_confirmation

REQUEST_FIELDS={'target_file','result_id','base_sheet','new_sheet','period','period_field','section',
                'target_label','target_kind','target_unit','key_field','metric_id','output_name',
                'target_range','protected_labels','unit_cell'}
REQUIRED_FIELDS={'target_file','base_sheet','new_sheet','period','period_field','section',
                 'target_label','target_kind'}


def build_semantic_plan(request,workspace):
    if request.keys()-REQUEST_FIELDS:
        raise ValueError('不支持的更新参数：'+', '.join(sorted(request.keys()-REQUEST_FIELDS)))
    for key in REQUIRED_FIELDS:
        if not request.get(key):
            raise ValueError('缺少业务参数：'+key)
    if request['target_kind'] not in ('forecast','actual'):
        raise ValueError('仅允许更新预测或实绩输入列')
    period=str(request['period'])
    target=Path(request['target_file'])
    target=(Path(workspace)/target).resolve() if not target.is_absolute() else target.resolve()
    if not target.is_relative_to(Path(workspace).resolve()):
        from ..spreadsheet_context import resolve_session_attachment
        resolved,error=resolve_session_attachment(request['target_file'],'更新目标')
        if error or resolved is None or not resolved.is_relative_to(Path(workspace).resolve()):
            raise ValueError('目标必须是当前会话授权的工作区附件')
        target=resolved
    source_sheet=request['base_sheet'];new_sheet=request['new_sheet']
    if not isinstance(new_sheet,str) or not new_sheet or len(new_sheet)>31 or any(c in new_sheet for c in r'\/*?:[]'):
        raise ValueError('新工作表名不符合 Excel 名称约束')
    unit=normalize_unit(request.get('target_unit'))
    wb=openpyxl.load_workbook(target)
    try:
        if source_sheet not in wb.sheetnames or new_sheet in wb.sheetnames:
            raise ValueError(f'base_sheet 必须是目标工作簿中被复制的快照，候选：{wb.sheetnames}；new_sheet 必须尚不存在')
        ws=wb[source_sheet]
        confirmation=unit_confirmation()
        unit=verify_unit(ws,unit,request.get('unit_cell'),confirmation)
        result=select_result(request.get('result_id'))
        provenance=result['provenance']
        source=provenance.get('file_path')
        source_hash=provenance.get('workbook_id') or result['data'].get('workbook_id')
        if not source or not source_hash or fingerprint(source)!=source_hash:
            raise ValueError('来源指纹缺失或已变化，请重新统计')
        source_facts,metric=bind_result(result,request.get('key_field'),request['period_field'],period,request.get('metric_id'))
        facts=source_facts.all_facts()
        spec,col=locate(ws,request['section'],period,request['target_label'],
                        {f.logical_key.department_code for f in facts},request.get('target_range'),
                        request.get('protected_labels',[]))
    finally:
        wb.close()
    spec.update(metric_id=metric,unit=unit,value_kind=request['target_kind'])
    mapped=FactTable();preview=[]
    ops=[{'op':'copy_sheet','op_id':'copy','source_sheet':source_sheet,'target_sheet':new_sheet}]
    for fact in facts:
        code=fact.logical_key.department_code
        key=LogicalKey(request['section'],code,metric,period,request['target_kind'],unit)
        value=convert_value(fact.numeric_value,fact.unit,unit)
        mapped.add_fact(replace(fact,logical_key=key,numeric_value=value,unit=unit))
        coord=f"{get_column_letter(col)}{spec['department_rows'][code]}"
        ops.append({'op':'fill_from_result','op_id':f'fill_{len(preview)}','target_sheet':new_sheet,
                    'coordinate':coord,'fact_ref':fact.evidence_id,'region_id':key.region_id,
                    'department_code':code,'metric_id':metric,'period':period,'value_kind':key.value_kind,'unit':unit})
        preview.append({'key':code,'coordinate':coord,'value':str(value),'unit':unit})
    output=request.get('output_name') or f'{target.stem}_{new_sheet}.xlsx'
    if not request.get('output_name'):
        base=Path(output).stem
        version=1
        while (Path(workspace)/output).exists():
            version+=1
            output=f'{base}_{version}.xlsx'
    if Path(output).name!=output or not output.endswith('.xlsx') or (Path(workspace)/output).exists():
        raise ValueError('输出必须为工作区内尚不存在的 .xlsx 文件，不能覆盖输入或历史文件')
    payload=serialize_facts(mapped)
    unmatched=spec['department_rows'].keys()-{f.logical_key.department_code for f in facts}
    warnings=['来源无匹配记录，保留原值：'+', '.join(sorted(unmatched))] if unmatched else []
    plan={'schema_version':'SpreadsheetUpdatePlan/v1','task_contract_ref':'current_turn','target_ref':str(target),
          'sources':[],'as_of':datetime.date.today().isoformat(),
          'snapshot':{'selected_sheet_id':source_sheet,'new_sheet_name':new_sheet},
          'semantic_bindings':{'region_specs':[spec],'facts_hash':hashlib.sha256(payload).hexdigest(),
                               'unit_basis': {'source':'user', **confirmation} if confirmation else {'source':'workbook','unit':unit},
                               'input_hashes':{str(target):fingerprint(target),source:source_hash},'warnings':warnings},
          'operations':ops,'preserve':{},'requirements':{},'output':{'mode':'deliverable_file','display_name':output}}
    compile_update_plan(plan,extract_workbook_ir(target,source_sheet,region_specs=[spec]),mapped)
    return plan,mapped,payload,preview,warnings
