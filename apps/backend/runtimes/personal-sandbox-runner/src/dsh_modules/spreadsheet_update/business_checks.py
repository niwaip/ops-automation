"""Read-back and explicitly bound reconciliations, independent of row labels or currencies."""
from decimal import Decimal
import math
import openpyxl


def verify_business_checks(path,target_sheet,compiled,checks):
    workbook=openpyxl.load_workbook(path,data_only=True)
    try:
        ws=workbook[target_sheet]
        for write in compiled.cell_writes:
            if write.value_type!='number':
                continue
            actual=ws[write.coordinate].value
            if not isinstance(actual,(int,float)) or not math.isfinite(actual) or not math.isclose(actual,write.new_value,rel_tol=1e-12,abs_tol=0):
                return [f'Written value does not match locked fact at {write.coordinate}']
        for check in checks:
            name=check.get('check') or check.get('rule') or check.get('check_rule') if isinstance(check,dict) else None
            if name!='reconciliation_total_match':
                return [f"Unknown or unsupported business check '{name}'"]
            if not check.get('total_coordinate') or not check.get('detail_coordinates'):
                return ['Reconciliation check failed: missing control total cell and explicit detail coordinates']
            if 'tolerance' not in check:
                return ['Reconciliation tolerance must be explicitly declared in the target unit']
            tolerance=Decimal(str(check['tolerance']))
            if not tolerance.is_finite() or tolerance<0:
                return ['Invalid reconciliation tolerance']
            actual_total=Decimal(str(ws[check['total_coordinate']].value))
            detail=sum((Decimal(str(ws[coordinate].value)) for coordinate in check['detail_coordinates']),Decimal(0))
            if not actual_total.is_finite() or not detail.is_finite() or abs(actual_total-detail)>tolerance:
                return ['Reconciliation total does not match explicit detail cells']
        return []
    except Exception as exc:
        return ['Business checks failed: '+str(exc)]
    finally:
        workbook.close()
