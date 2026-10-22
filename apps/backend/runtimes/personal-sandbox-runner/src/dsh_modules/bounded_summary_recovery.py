"""Execute an unambiguous sum/count request after a model's incomplete plan.

Only literal fields, categorical predicates and explicit grouping are compiled.
Ambiguous requests decline rather than broadening scope or fabricating values.
"""

import os
import json
import hashlib
from pathlib import Path
from decimal import Decimal
import openpyxl
from openpyxl.utils.cell import range_boundaries, get_column_letter
from .spreadsheet_context import spreadsheet_structure
from .spreadsheet_analysis import calculate_spreadsheet_analysis
from .spreadsheet_request import summary_request
from .analysis_fact_selection import request_fingerprint


def recover_summary(prompt, workspace_dir, deadline=None):
    names = [s.strip() for s in os.getenv('DSH_SESSION_ATTACHMENTS', '').split(',') if s.strip()]
    if len(names) != 1 or Path(names[0]).name != names[0]:
        return None
    workspace = Path(workspace_dir).resolve()
    source = (workspace / names[0]).resolve()
    if source.parent != workspace or source.suffix.lower() not in ('.xlsx', '.xlsm'):
        return None
    if not any(word in prompt for word in ('统计', '汇总', '求和', '合计')):
        return None
    structure = spreadsheet_structure(source, deadline=deadline)
    if not structure.is_success:
        return None
    candidates = []
    book = openpyxl.load_workbook(source, read_only=True, data_only=True)
    try:
        for sheet in structure.data.get('sheets', []):
            for region in sheet.get('data_region_candidates', []):
                headers = next((h['columns'] for h in sheet.get('header_candidates', []) if h['row'] == region['header_row']), {})
                a, start, b, end = range_boundaries(region['data_range'])
                if (end - start + 1) * (b - a + 1) > 200000:
                    continue
                ws = book[sheet['sheet']]
                categories = {}
                for row in ws.iter_rows(min_row=start, max_row=end, min_col=a, max_col=b, values_only=True):
                    for offset, value in enumerate(row):
                        col = get_column_letter(a + offset)
                        bucket = categories.setdefault(col, set())
                        if len(bucket) <= 20 and isinstance(value, str):
                            bucket.add(value)
                categories = {c: list(v) for c, v in categories.items() if len(v) <= 20}
                spec = summary_request(prompt, headers, categories)
                if spec is not None:
                    candidates.append((sheet['sheet'], region, headers, spec))
    finally:
        book.close()
    if len(candidates) != 1:
        return None
    sheet, region, headers, spec = candidates[0]
    calculations = []
    for index, metric in enumerate(spec['metrics']):
        expr = {'aggregate': metric['aggregate'], 'table': 'summary_source', 'column': metric['column'], 'unit': metric['unit']}
        calculations.append({'id': f'total_{index}', 'label': '总' + metric['name'], 'expression': expr})
        if spec['group_by']:
            calculations.append({'id': f'group_{index}', 'label': metric['name'], 'expression': {**expr, 'group_by': spec['group_by']}})
    plan = {'tables': [{'id': 'summary_source', 'sheet': sheet, 'header_row': region['header_row'],
                        'data_range': region['data_range'], 'filters': spec['filters']}], 'calculations': calculations}
    receipt = calculate_spreadsheet_analysis(source, plan, deadline=deadline)
    if not receipt.is_success:
        return receipt
    receipt.data['request_fingerprint'] = request_fingerprint(prompt, receipt.provenance['workbook_id'])
    receipt.data['summary_spec'] = spec
    receipt.provenance['execution_origin'] = 'bounded_summary_recovery'
    receipt.data.pop('evidence_id', None)
    receipt.data['evidence_id'] = hashlib.sha256((receipt.provenance['workbook_id'] + json.dumps(receipt.data, sort_keys=True, ensure_ascii=False, default=str)).encode()).hexdigest()[:24]
    return receipt


class SummaryRecovery:
    def __init__(self, prompt, workspace, deadline=None, enabled=True):
        self.prompt, self.workspace, self.deadline = prompt, workspace, deadline
        self.attempted = not enabled

    def apply(self, evidence, telemetry, outbound_files):
        if self.attempted:
            return False
        self.attempted = True
        from .analysis_task_state import analysis_task_state
        receipt = recover_summary(self.prompt, self.workspace, self.deadline)
        if receipt is None or not receipt.is_success:
            return False
        state = analysis_task_state([receipt], self.prompt)
        if not state['complete']:
            return False
        evidence.append(receipt)
        telemetry.record_tool_call(tool_name='bounded_summary_recovery', params={'file_path': receipt.provenance['file_path'], 'plan_hash': receipt.data['plan_hash']}, status='success')
        telemetry.record_execution_result(receipt)
        telemetry.record_guard('bounded_summary_recovery', 'accept', {'evidence_id': receipt.data['evidence_id']})
        output = export_recovered_summary(receipt, self.prompt, self.workspace)
        if output:
            outbound_files.append(json.dumps({'filePath': output, 'fileName': Path(output).name}, ensure_ascii=False))
        return True


def export_recovered_summary(receipt, prompt, workspace_dir):
    """Atomic, formula-free export of executed facts, with no external sheet dependency."""
    from .intent_contract import extract_source_and_target_artifacts
    from .file_tools import is_protected_session_file
    # The attachment banner is transport metadata, never an output filename request.
    export_prompt = prompt
    if prompt.startswith('【当前轮次用户上传附件】') and '\n用户指令：' in prompt:
        export_prompt = prompt.split('\n用户指令：', 1)[1]
    _, formats, targets = extract_source_and_target_artifacts(export_prompt)
    if not ('.xlsx' in formats or any(str(t).lower().endswith('.xlsx') for t in targets)
            or ('Excel' in prompt and any(w in prompt for w in ('生成', '导出')))):
        return None
    names = [str(t) for t in targets if str(t).lower().endswith('.xlsx')]
    if len(names) > 1:
        return None
    name = names[0] if names else '汇总结果.xlsx'
    if Path(name).name != name:
        return None
    target = Path(workspace_dir) / name
    if target.is_symlink() or is_protected_session_file(target):
        return None
    spec = receipt.data['summary_spec']
    facts = receipt.data['facts']
    book = openpyxl.Workbook()
    ws = book.active
    ws.title = '分组汇总'
    titles = []
    for i, metric in enumerate(spec['metrics']):
        fact = next(f for f in facts if f['id'] == f'total_{i}')
        titles.append(metric['name'] + (f"({fact['unit']})" if fact['unit'] != 'number' else ''))
    ws.append((['分组'] if spec['group_by'] else []) + titles)
    keys = sorted({str(f['group_key']) for f in facts if 'group_key' in f})
    for key in keys:
        ws.append([key] + [Decimal(next(f['value'] for f in facts if f.get('group_key') == key and f['id'].startswith(f'group_{i}__')))
                           for i in range(len(spec['metrics']))])
    ws.append((['总计'] if spec['group_by'] else []) + [Decimal(next(f['value'] for f in facts if f['id'] == f'total_{i}'))
                                                      for i in range(len(spec['metrics']))])
    ws.freeze_panes = 'B2'
    temporary = target.with_name('.' + target.stem + '.tmp.xlsx')
    try:
        book.save(temporary)
        temporary.replace(target)
    finally:
        temporary.unlink(missing_ok=True)
        book.close()
    return str(target)
