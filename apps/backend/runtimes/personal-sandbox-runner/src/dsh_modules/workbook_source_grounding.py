"""Source-to-deliverable numeric grounding for generated Excel workbooks.

Verifies that newly generated `.xlsx`/`.xlsm` deliverables in a spreadsheet session
do not contain fabricated numbers or mismatched source cell bindings when derived
from session source workbooks.
"""

from collections import defaultdict
from decimal import Decimal, InvalidOperation
import os
from pathlib import Path
import re
from typing import Dict, List, Optional, Set, Tuple
import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter


_SINGLE_CELL_RE = re.compile(
    r"^(?:(?:'([^']+)'|([A-Za-z0-9_\u4e00-\u9fa5 Diagnostics]+))!)?\$?([A-Z]{1,3})\$?([1-9]\d{0,4})$"
)
_RANGE_RE = re.compile(
    r"^(?:(?:'([^']+)'|([A-Za-z0-9_\u4e00-\u9fa5]+))!)?\$?([A-Z]{1,3})\$?([1-9]\d{0,4}):\$?([A-Z]{1,3})\$?([1-9]\d{0,4})$"
)
_BARE_COORD_IN_TEXT_RE = re.compile(
    r"(?:(?:'([^']+)'|([A-Za-z0-9_\u4e00-\u9fa5]{1,24}))!)?\$?([A-Z]{1,3})\$?([1-9]\d{0,4})(?![0-9A-Za-z_:])"
)


def _to_decimal(val) -> Optional[Decimal]:
    if val is None or isinstance(val, bool):
        return None
    if isinstance(val, (int, float, Decimal)):
        try:
            d = Decimal(str(val))
            return d if d.is_finite() else None
        except InvalidOperation:
            return None
    if isinstance(val, str):
        cleaned = val.replace(",", "").replace("，", "").replace("¥", "").replace("$", "").strip()
        if not cleaned or cleaned.startswith("="):
            return None
        try:
            d = Decimal(cleaned)
            return d if d.is_finite() else None
        except InvalidOperation:
            return None
    return None


def _norm_label(text: str) -> str:
    if not text:
        return ""
    cleaned = re.sub(r"[（(][^()（）]*[)）]", "", str(text))
    cleaned = re.sub(r"[\s:：_\-—/·•*＊#0-9.、，,]+", "", cleaned)
    return cleaned.strip().lower()


def _format_dec(d: Decimal) -> str:
    normalized = d.quantize(Decimal("0.0001")) if abs(d) < Decimal("1e12") else d
    s = format(normalized.normalize(), "f")
    if "." in s:
        s = s.rstrip("0").rstrip(".")
    return s or "0"


def _discover_source_workbooks(generated_path: Path) -> List[Path]:
    """Finds session input spreadsheet workbooks associated with the workspace."""
    from .file_tools import is_protected_session_file

    ws = generated_path.parent.resolve()
    gen_resolved = generated_path.resolve()
    sources: List[Path] = []
    seen: Set[Path] = set()

    # 1. Prefer pristine backups in .dsh/inputs_backup if available
    backup_dir = ws / ".dsh" / "inputs_backup"
    if backup_dir.is_dir():
        for p in sorted(backup_dir.iterdir()):
            if p.is_file() and p.suffix.lower() in (".xlsx", ".xlsm") and not p.name.startswith("."):
                rp = p.resolve()
                if rp != gen_resolved and rp not in seen:
                    seen.add(rp)
                    sources.append(rp)

    if sources:
        return sources

    # 2. Check DSH_SESSION_ATTACHMENTS environment variable
    env_atts = [v.strip() for v in os.getenv("DSH_SESSION_ATTACHMENTS", "").split(",") if v.strip()]
    for name in env_atts:
        p = (ws / name).resolve()
        if p.is_file() and p.suffix.lower() in (".xlsx", ".xlsm") and p != gen_resolved and p not in seen:
            seen.add(p)
            sources.append(p)

    # 3. Check protected session files in workspace
    if not sources and ws.is_dir():
        for p in sorted(ws.iterdir()):
            if p.is_file() and p.suffix.lower() in (".xlsx", ".xlsm") and not p.name.startswith("."):
                rp = p.resolve()
                if rp != gen_resolved and rp not in seen and is_protected_session_file(rp):
                    seen.add(rp)
                    sources.append(rp)

    return sources


class _SourceWorkbookTruth:
    """Evaluates and indexes ground-truth cell and metric values from source workbooks."""

    def __init__(self, source_paths: List[Path]):
        self.sheet_names: Set[str] = set()
        self.qualified_cells: Dict[Tuple[str, str], Decimal] = {}
        self.bare_cells: Dict[str, Dict[str, Decimal]] = defaultdict(dict)
        self.label_cells: Dict[str, List[Tuple[str, str, Decimal]]] = defaultdict(list)
        self.label_derived: Dict[str, Set[Decimal]] = defaultdict(set)
        self.all_values: Set[Decimal] = set()
        for path in source_paths:
            self._index_workbook(path)

    def _index_workbook(self, path: Path) -> None:
        from .workbook_calculation import get_or_create_recalculated_workbook

        eval_path = path
        try:
            recalc_path, _ = get_or_create_recalculated_workbook(path)
            if recalc_path and Path(recalc_path).exists():
                eval_path = Path(recalc_path)
        except Exception:
            pass

        wb_val = None
        wb_raw = None
        try:
            wb_val = openpyxl.load_workbook(eval_path, data_only=True, read_only=True)
            wb_raw = openpyxl.load_workbook(path, data_only=False, read_only=True)

            raw_cells: Dict[Tuple[str, str], object] = {}
            val_cells: Dict[Tuple[str, str], Decimal] = {}

            for ws in wb_val:
                self.sheet_names.add(ws.title)
                for r_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
                    for c_idx, val in enumerate(row, start=1):
                        d = _to_decimal(val)
                        if d is not None:
                            coord = f"{get_column_letter(c_idx)}{r_idx}"
                            val_cells[(ws.title, coord)] = d

            for ws in wb_raw:
                self.sheet_names.add(ws.title)
                for r_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
                    for c_idx, val in enumerate(row, start=1):
                        if val is not None:
                            coord = f"{get_column_letter(c_idx)}{r_idx}"
                            raw_cells[(ws.title, coord)] = val

            def resolve_cell(sheet: str, coord: str, visiting: Set[Tuple[str, str]]) -> Optional[Decimal]:
                key = (sheet, coord)
                if key in val_cells:
                    return val_cells[key]
                if key in visiting:
                    return None
                raw = raw_cells.get(key)
                if raw is None:
                    return None
                d = _to_decimal(raw)
                if d is not None:
                    val_cells[key] = d
                    return d
                if not isinstance(raw, str) or not raw.startswith("="):
                    return None
                visiting.add(key)
                evaluated = self._eval_formula(raw[1:].strip(), sheet, resolve_cell, visiting)
                visiting.remove(key)
                if evaluated is not None:
                    val_cells[key] = evaluated
                return evaluated

            for (sheet, coord), raw in list(raw_cells.items()):
                if isinstance(raw, str) and raw.startswith("="):
                    resolve_cell(sheet, coord, set())

            for (sheet, coord), d in val_cells.items():
                self.qualified_cells[(sheet, coord)] = d
                self.bare_cells[coord][sheet] = d
                self.all_values.add(d)

            # Index row labels and column headers per sheet
            for ws in wb_raw:
                sheet = ws.title
                rows = list(ws.iter_rows(values_only=True))
                if not rows:
                    continue

                # Group rows by label in first 2 columns to support both financial statements and detail tables
                label_row_nums: Dict[str, Dict[int, List[Decimal]]] = defaultdict(lambda: defaultdict(list))
                col_headers: Dict[int, str] = {}

                for r_idx, row in enumerate(rows, start=1):
                    # Check if row looks like a header row
                    non_empty = [(c_idx, val) for c_idx, val in enumerate(row, start=1) if val is not None]
                    if r_idx <= 10 and len(non_empty) >= 2 and all(
                        isinstance(v, str) and not v.startswith("=") for _, v in non_empty
                    ):
                        for c_idx, val in non_empty:
                            nl = _norm_label(str(val))
                            if nl:
                                col_headers[c_idx] = nl

                    # Extract row label from first 2 columns
                    row_label = ""
                    for c_idx in (1, 2):
                        if c_idx <= len(row) and isinstance(row[c_idx - 1], str):
                            cand = str(row[c_idx - 1]).strip()
                            if cand and not cand.startswith("="):
                                nl = _norm_label(cand)
                                if nl and len(nl) >= 2:
                                    row_label = nl
                                    break

                    row_decimals: List[Tuple[int, str, Decimal]] = []
                    for c_idx in range(1, len(row) + 1):
                        coord = f"{get_column_letter(c_idx)}{r_idx}"
                        d = val_cells.get((sheet, coord))
                        if d is not None:
                            row_decimals.append((c_idx, coord, d))
                            if c_idx in col_headers:
                                ch = col_headers[c_idx]
                                self.label_cells[ch].append((sheet, coord, d))
                                label_row_nums[ch][c_idx].append(d)

                    if row_label and row_decimals:
                        for c_idx, coord, d in row_decimals:
                            self.label_cells[row_label].append((sheet, coord, d))
                            label_row_nums[row_label][c_idx].append(d)
                        if len(row_decimals) >= 2:
                            r_vals = [d for _, _, d in row_decimals]
                            self.label_derived[row_label].add(sum(r_vals, Decimal(0)))
                            self.label_derived[row_label].add(sum(r_vals, Decimal(0)) / Decimal(len(r_vals)))
                            self.label_derived[row_label].add(r_vals[0] - r_vals[1])
                            self.label_derived[row_label].add(r_vals[1] - r_vals[0])

                # Compute sheet-wide column totals for share/ratio derivations
                col_totals: Dict[int, Decimal] = defaultdict(lambda: Decimal(0))
                for r_idx in range(1, len(rows) + 1):
                    for c_idx in range(1, len(rows[r_idx - 1]) + 1):
                        d = val_cells.get((sheet, f"{get_column_letter(c_idx)}{r_idx}"))
                        if d is not None:
                            col_totals[c_idx] += d

                # Compute column/group aggregates (sum, count, avg, min, max, share) for each label
                for nl, col_map in label_row_nums.items():
                    for c_idx, vals in col_map.items():
                        if not vals:
                            continue
                        s = sum(vals, Decimal(0))
                        cnt = Decimal(len(vals))
                        self.label_derived[nl].add(s)
                        self.label_derived[nl].add(cnt)
                        self.label_derived[nl].add(s / cnt)
                        self.label_derived[nl].add(min(vals))
                        self.label_derived[nl].add(max(vals))
                        if col_totals.get(c_idx) and col_totals[c_idx] != Decimal(0):
                            self.label_derived[nl].add(s / col_totals[c_idx])
                        self.all_values.add(s)
                        self.all_values.add(cnt)
        finally:
            if wb_val:
                wb_val.close()
            if wb_raw:
                wb_raw.close()

    def _expand_range(
        self,
        range_str: str,
        default_sheet: str,
        resolve_cell,
        visiting: Set[Tuple[str, str]],
    ) -> List[Decimal]:
        m = _RANGE_RE.match(range_str.strip())
        if not m:
            return []
        sh = m.group(1) or m.group(2) or default_sheet
        c1, r1, c2, r2 = m.group(3), int(m.group(4)), m.group(5), int(m.group(6))
        idx1, idx2 = column_index_from_string(c1), column_index_from_string(c2)
        if idx1 > idx2:
            idx1, idx2 = idx2, idx1
        if r1 > r2:
            r1, r2 = r2, r1
        if (idx2 - idx1 + 1) * (r2 - r1 + 1) > 5000:
            return []
        vals: List[Decimal] = []
        for r in range(r1, r2 + 1):
            for c in range(idx1, idx2 + 1):
                d = resolve_cell(sh, f"{get_column_letter(c)}{r}", visiting)
                if d is not None:
                    vals.append(d)
        return vals

    def _eval_formula(
        self,
        expr: str,
        current_sheet: str,
        resolve_cell,
        visiting: Set[Tuple[str, str]],
    ) -> Optional[Decimal]:
        # 1. Expand ranges (e.g. B4:B15 or Sheet!B4:B15) into comma-separated Decimal literals
        range_in_expr = re.compile(
            r"(?:(?:'([^']+)'|([A-Za-z0-9_\u4e00-\u9fa5]+))!)?\$?([A-Z]{1,3})\$?([1-9]\d{0,4}):\$?([A-Z]{1,3})\$?([1-9]\d{0,4})"
        )

        def _replace_range(m: re.Match) -> str:
            vals = self._expand_range(m.group(0), current_sheet, resolve_cell, visiting)
            if not vals:
                return "Decimal('0')"
            return ", ".join(f"Decimal('{v}')" for v in vals)

        expanded = range_in_expr.sub(_replace_range, expr)

        # 2. Replace single cell references with Decimal('...')
        def _replace_ref(m: re.Match) -> str:
            sh = m.group(1) or m.group(2) or current_sheet
            coord = f"{m.group(3)}{m.group(4)}"
            v = resolve_cell(sh, coord, visiting)
            return f"Decimal('{v}')" if v is not None else "NaN"

        resolved_expr = _BARE_COORD_IN_TEXT_RE.sub(_replace_ref, expanded)
        if "NaN" in resolved_expr:
            return None

        # 3. Wrap remaining standalone numbers in Decimal('...') outside Decimal('...')
        token_re = re.compile(r"Decimal\('[^']*'\)|([A-Za-z_]+)|(\d+(?:\.\d+)?)")

        def _wrap_tokens(m: re.Match) -> str:
            if m.group(1):
                fn = m.group(1).upper()
                if fn in {"SUM", "AVERAGE", "AVG", "MIN", "MAX", "ABS", "ROUND"}:
                    return fn
                return "INVALID"
            if m.group(2):
                return f"Decimal('{m.group(2)}')"
            return m.group(0)

        safe_expr = token_re.sub(_wrap_tokens, resolved_expr)
        if "INVALID" in safe_expr:
            return None

        env = {
            "__builtins__": None,
            "Decimal": Decimal,
            "SUM": lambda *args: sum(args, Decimal(0)),
            "AVERAGE": lambda *args: (sum(args, Decimal(0)) / Decimal(len(args))) if args else Decimal(0),
            "AVG": lambda *args: (sum(args, Decimal(0)) / Decimal(len(args))) if args else Decimal(0),
            "MIN": lambda *args: min(args) if args else Decimal(0),
            "MAX": lambda *args: max(args) if args else Decimal(0),
            "ABS": abs,
            "ROUND": lambda val, n=0: round(val, int(n)),
        }
        try:
            val = eval(safe_expr, env, {})
            return _to_decimal(val)
        except Exception:
            return None


def _values_match(actual: Decimal, expected: Decimal) -> bool:
    """Checks if actual matches expected within standard rounding or unit/percentage scaling."""
    for factor in (Decimal(1), Decimal(100), Decimal("0.01"), Decimal(1000), Decimal("0.001"), Decimal(10000), Decimal("0.0001")):
        scaled = expected * factor
        tol = max(Decimal("0.015"), abs(scaled) * Decimal("0.0005"))
        if abs(actual - scaled) <= tol:
            return True
    return False


def validate_generated_workbook_against_sources(generated_path: Path) -> List[str]:
    """Validates a newly generated workbook against session source spreadsheets.

    Returns a list of human-readable data accuracy issues if hardcoded numbers in the
    generated workbook contradict the ground-truth values in the source workbook(s).
    """
    from .file_tools import is_protected_session_file
    try:
        from .spreadsheet_update.delivery import is_verified_update_file
    except Exception:
        is_verified_update_file = None

    path = Path(generated_path)
    if path.suffix.lower() not in (".xlsx", ".xlsm") or not path.exists():
        return []
    if is_protected_session_file(path):
        return []
    if is_verified_update_file and is_verified_update_file(path):
        return []

    source_paths = _discover_source_workbooks(path)
    if not source_paths:
        return []

    truth = _SourceWorkbookTruth(source_paths)
    if not truth.qualified_cells:
        return []

    issues: List[str] = []
    wb_val = None
    wb_raw = None
    try:
        wb_val = openpyxl.load_workbook(path, data_only=True, read_only=True)
        wb_raw = openpyxl.load_workbook(path, data_only=False, read_only=True)

        for sheet_name in wb_val.sheetnames:
            ws_val = wb_val[sheet_name]
            ws_raw = wb_raw[sheet_name]
            val_rows = list(ws_val.iter_rows(values_only=True))
            raw_rows = list(ws_raw.iter_rows(values_only=True))

            for r_idx, (v_row, f_row) in enumerate(zip(val_rows, raw_rows), start=1):
                if r_idx == 1:
                    # Skip top header row if all cells are strings
                    if all(c is None or isinstance(c, str) for c in v_row):
                        continue

                # Collect literal numbers (not formulas) and text cells in this row
                literal_nums: List[Tuple[str, Decimal]] = []
                all_row_nums: List[Tuple[str, Decimal]] = []
                row_strings: List[Tuple[int, str]] = []

                for c_idx, (v_cell, f_cell) in enumerate(zip(v_row, f_row), start=1):
                    coord = f"{get_column_letter(c_idx)}{r_idx}"
                    is_formula = isinstance(f_cell, str) and f_cell.startswith("=")
                    d = _to_decimal(v_cell)
                    if d is not None:
                        all_row_nums.append((coord, d))
                        # Skip sequential row index in column A (1, 2, 3...)
                        if c_idx == 1 and d == Decimal(r_idx - 1) and len(v_row) >= 3:
                            continue
                        if not is_formula:
                            literal_nums.append((coord, d))
                    elif isinstance(v_cell, str) and v_cell.strip():
                        row_strings.append((c_idx, v_cell.strip()))

                if not all_row_nums:
                    continue

                # Identify any source sheet name mentioned in this row
                mentioned_sheets = [
                    s for _, s in row_strings if s in truth.sheet_names
                ]
                row_label_raw = row_strings[0][1] if row_strings else ""
                row_label_norm = _norm_label(row_label_raw)

                # Check A: Explicit single-cell coordinate cited in the row (e.g., 'B6', '利润表!B6')
                explicit_refs: List[Tuple[Optional[str], str]] = []
                for _, s in row_strings:
                    if ":" in s or "-" in s:
                        continue
                    m = _SINGLE_CELL_RE.match(s)
                    if m:
                        sh = m.group(1) or m.group(2)
                        coord = f"{m.group(3)}{m.group(4)}"
                        explicit_refs.append((sh, coord))

                if explicit_refs:
                    for ref_sh, ref_coord in explicit_refs:
                        candidate_truths: List[Tuple[str, Decimal]] = []
                        if ref_sh and (ref_sh, ref_coord) in truth.qualified_cells:
                            candidate_truths.append((f"{ref_sh}!{ref_coord}", truth.qualified_cells[(ref_sh, ref_coord)]))
                        elif mentioned_sheets:
                            for m_sh in mentioned_sheets:
                                if (m_sh, ref_coord) in truth.qualified_cells:
                                    candidate_truths.append((f"{m_sh}!{ref_coord}", truth.qualified_cells[(m_sh, ref_coord)]))
                        if not candidate_truths and row_label_norm and row_label_norm in truth.label_cells:
                            for l_sh, l_coord, l_val in truth.label_cells[row_label_norm]:
                                if l_coord == ref_coord:
                                    candidate_truths.append((f"{l_sh}!{l_coord}", l_val))
                        if not candidate_truths and ref_coord in truth.bare_cells:
                            for b_sh, b_val in truth.bare_cells[ref_coord].items():
                                candidate_truths.append((f"{b_sh}!{ref_coord}", b_val))

                        if candidate_truths:
                            matched_any = any(
                                _values_match(actual_d, exp_d)
                                for _, actual_d in all_row_nums
                                for _, exp_d in candidate_truths
                            )
                            if not matched_any:
                                num_coord, actual_d = all_row_nums[0]
                                ref_desc, exp_d = candidate_truths[0]
                                label_prefix = f"指标「{row_label_raw}」" if row_label_raw else ""
                                issues.append(
                                    f"{sheet_name}!{num_coord} {label_prefix}数值 {_format_dec(actual_d)} "
                                    f"与所声明的源单元格 {ref_desc} 真实值 {_format_dec(exp_d)} 不一致"
                                )
                    continue

                # Check B: Row metric label grounding for literal hardcoded numbers
                if (
                    literal_nums
                    and row_label_norm
                    and row_label_norm in truth.label_cells
                    and not any(k in row_label_raw for k in ("合计", "总计", "小计", "Total", "Summary"))
                ):
                    label_entries = truth.label_cells[row_label_norm]
                    expected_vals = [d for _, _, d in label_entries]
                    derived_vals = list(truth.label_derived.get(row_label_norm, set()))
                    allowed_for_label = expected_vals + derived_vals

                    for num_coord, actual_d in literal_nums:
                        if any(_values_match(actual_d, exp_d) for exp_d in allowed_for_label):
                            continue
                        sample_truth = " / ".join(
                            f"{sh}!{c}={_format_dec(d)}" for sh, c, d in label_entries[:3]
                        )
                        issues.append(
                            f"{sheet_name}!{num_coord} 指标「{row_label_raw}」硬编码数值 {_format_dec(actual_d)} "
                            f"未经源文件事实核验（源文件真实值：{sample_truth}）"
                        )
    except Exception:
        pass
    finally:
        if wb_val:
            wb_val.close()
        if wb_raw:
            wb_raw.close()

    return list(dict.fromkeys(issues))
