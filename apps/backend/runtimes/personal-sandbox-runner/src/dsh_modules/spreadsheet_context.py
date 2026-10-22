"""Metadata and small coordinate-preserving samples, without eager recalculation."""

import json
import re
import hashlib
import time
from pathlib import Path
from .workbook_reader import inspect_workbook, WorkbookManifest
from .worksheet_metadata import worksheet_dimensions
from typing import Optional, Tuple, List
from .tool_result import ToolResult, TextToolOutput
from .file_tools import resolve_sandboxed_path, match_candidate_name
import openpyxl
from openpyxl.utils import get_column_letter


STRUCTURE_TOOL_NAME = 'inspect_spreadsheet_structure'
STRUCTURE_TOOL = {'type':'function','function':{
    'name':STRUCTURE_TOOL_NAME,
    'description':'按需读取指定工作表结构和带坐标样本，不重算、不证明完整业务数据。名称必须来自工作簿目录。',
    'parameters':{'type':'object','additionalProperties':False,'properties':{
        'file_path':{'type':'string'}, 'sheets':{'type':'array','minItems':1,'maxItems':4,
            'items':{'type':'string'}}},'required':['file_path','sheets']}}}


def resolve_sheet_name(sheet_name: str, available_sheets: List[str]) -> str:
    """Resolves a sheet name against available workbook sheets with fuzzy normalization."""
    if not sheet_name or not available_sheets:
        return sheet_name
    matched = match_candidate_name(sheet_name, available_sheets, min_similarity=0.7, match_extension=False)
    return matched or sheet_name


def resolve_session_attachment(
    file_path: str,
    tool_label: str = "当前工具"
) -> Tuple[Optional[Path], Optional[ToolResult]]:
    """Resolves an input file_path to an authenticated session attachment within sandbox workspace.

    Applies fuzzy normalization against DSH_SESSION_ATTACHMENTS to tolerate CJK variant characters
    without hardcoding. If attachments are specified and the file does not match any attachment,
    returns a clear self-healing ToolResult error specifying the allowed attachments.
    """
    import os
    raw_path = str(file_path or "").strip()
    if not raw_path:
        return None, ToolResult.error("source_unavailable", "未指定有效文件路径")

    from .config import WORKSPACE_DIR
    ws = Path(WORKSPACE_DIR)
    env_atts = [v.strip() for v in os.getenv("DSH_SESSION_ATTACHMENTS", "").split(",") if v.strip()]

    if env_atts:
        allowed = {(ws / name).resolve() for name in env_atts}
        source, error = resolve_sandboxed_path(raw_path)
        if source and source.resolve() in allowed:
            return source.resolve(), None

        # Fuzzy match filename against allowed attachments
        query_filename = Path(raw_path).name
        matched_filename = match_candidate_name(query_filename, env_atts, min_similarity=0.72)
        if matched_filename:
            matched_path = (ws / matched_filename).resolve()
            if matched_path.exists():
                return matched_path, None

        allowed_names_str = ", ".join(env_atts)
        error_msg = f"{tool_label}仅接受当前会话附件原件。当前可用附件: [{allowed_names_str}]。请传入完全匹配的文件名。"
        return None, ToolResult.error("source_outside_session_scope", error_msg)

    source, error = resolve_sandboxed_path(raw_path)
    if error or source is None:
        return None, ToolResult.error("source_unavailable", error or "原件不可用")
    if not source.exists():
        return None, ToolResult.error("source_unavailable", f"文件不存在: {raw_path}")
    return source.resolve(), None


def spreadsheet_structure(source, selected_sheets=None, deadline=None):
    manifest=inspect_workbook(source)
    if not isinstance(manifest,WorkbookManifest):
        return manifest
    resolved_sheets = None
    if selected_sheets is not None:
        if not isinstance(selected_sheets,list) or not 1<=len(selected_sheets)<=4:
            return ToolResult.error('invalid_sheet_selection','请选择目录中真实存在的1至4张工作表',
                                    data={'available_sheets':manifest.sheet_names})
        resolved_sheets = []
        for s in selected_sheets:
            if not isinstance(s, str):
                return ToolResult.error('invalid_sheet_selection','工作表名必须为字符串')
            res_s = resolve_sheet_name(s, manifest.sheet_names)
            if res_s not in manifest.sheet_names:
                return ToolResult.error('invalid_sheet_selection','请选择目录中真实存在的1至4张工作表',
                                        data={'available_sheets':manifest.sheet_names})
            resolved_sheets.append(res_s)
    workbook=openpyxl.load_workbook(source,read_only=True,data_only=False)
    try:
        sheets=[]
        for ws in workbook:
            if resolved_sheets is not None and ws.title not in resolved_sheets:
                continue
            if deadline is not None and time.monotonic()>=deadline:
                raise TimeoutError('Spreadsheet structure deadline exceeded')
            dimensions=worksheet_dimensions(ws)
            sample=[];headers=[];row_labels=[];footers=[];observed={}
            limit=32 if ws.max_row<=32 else 20
            # Read a bounded prefix plus footer, never claim a full nonempty count.
            for row in ws.iter_rows(max_row=min(limit,ws.max_row),max_col=min(ws.max_column,64)):
                filled=[c for c in row if c.value is not None]
                if filled:observed[filled[0].row]=filled
            if ws.max_row>limit:
                for row in ws.iter_rows(min_row=max(limit+1,ws.max_row-25),max_col=min(ws.max_column,64)):
                    filled=[c for c in row if c.value is not None]
                    if filled:observed[filled[0].row]=filled
            footer_pattern = re.compile(
                r"^\s*(?:合计|总计|小计|grand\s*total|total|查询条件|报表类别|报表说明|填报说明|编制说明|编制单位|制表人|打印时间|导出时间|备注|注[:：]|说明[:：]|附注|声明|数据来源|温馨提示|特别提示|[\*＊#])",
                re.I
            )
            for row_index,filled in sorted(observed.items()):
                strings=[c for c in filled if isinstance(c.value,str) and not c.value.startswith("=")]
                if row_index<=20 and len(strings)>=2 and len(strings)==len(filled):
                    headers.append({"row":row_index,"columns":{get_column_letter(c.column):str(c.value) for c in strings}})
                first=next((c for c in filled if c.column==1),None)
                if (first and not (isinstance(first.value,str) and first.value.startswith("="))
                        and (ws.max_row<=32 or len(row_labels)<3 or row_index>=ws.max_row-1)):
                    row_labels.append({"cell":first.coordinate,"value":str(first.value)})
            best=max(headers,key=lambda h:len(h["columns"]),default=None)
            header_col_count = len(best["columns"]) if best else 0
            end = ws.max_row
            while end > (best["row"] if best else 1):
                if end not in observed:
                    break
                row_cells = observed[end]
                first_val = str(row_cells[0].value or '').strip()
                is_footer = False
                if footer_pattern.search(first_val):
                    is_footer = True
                elif header_col_count >= 4 and len(row_cells) <= max(2, header_col_count // 4):
                    if any(footer_pattern.search(str(c.value or '')) for c in row_cells) or ':' in first_val or '：' in first_val:
                        is_footer = True
                if is_footer:
                    footers.append(end)
                    end -= 1
                else:
                    break
            sample_indices=(sorted(observed)[:5]+sorted(observed)[-2:] if selected_sheets is not None
                            else ([r for r in sorted(observed) if r > best['row']][:5] if best else sorted(observed)[:1]))
            for index in dict.fromkeys(sample_indices):
                if index in observed:
                    sample.append({c.coordinate:str(c.value)[:150] for c in observed[index][:20]})
            while end in footers:end-=1
            candidates=([{"header_row":best["row"],"data_range":f"A{best['row']+1}:{get_column_letter(ws.max_column)}{end}",
                           "reason":"候选来自非公式文本表头；排除末尾显式合计行与说明行，请按任务核对"}]
                         if best and best["row"]<end else [])
            sheets.append({"sheet":ws.title,"dimensions":dimensions,"physical_rows":ws.max_row,
                           "coverage_complete":False,"sample_cells":sample,
                           "header_candidates":sorted(headers,key=lambda h:len(h["columns"]),reverse=True)[:3],
                           "data_region_candidates":candidates,"row_labels":row_labels})
        if hashlib.sha256(Path(source).read_bytes()).hexdigest()!=manifest.file_hash:
            return ToolResult.error('source_changed','结构读取期间原件发生变化')
        payload={"kind":"spreadsheet_structure","workbook_id":manifest.file_hash,"extraction_mode":"SCHEMA_SAMPLE",
                 "available_sheets":manifest.sheet_names,
                 "coverage_complete":False,"sheets":sheets}
        return ToolResult.success(data=payload, text=json.dumps(payload,ensure_ascii=False),
                                  provenance={'file_path':str(Path(source).resolve()),'workbook_id':manifest.file_hash})
    finally:
        workbook.close()


def build_spreadsheet_source_context(source, deadline=None, prompt=""):
    from .spreadsheet_preflight import profile_spreadsheet_for_planning
    preflight = profile_spreadsheet_for_planning(source, prompt=prompt)
    preflight_md = preflight.get("compact_markdown", "")

    result = spreadsheet_structure(source, deadline=deadline)
    if not result.is_success:
        return (preflight_md + "\n" + result.render_text()).strip()
    return (
        preflight_md + "\n"
        + '[Workbook Source Structure]:\n' + result.render_text() + (
            '\n这是结构与样本，不是完整业务数据或计算结果。已有明确候选时直接执行分析计划；'
            '需要更多结构时用inspect_spreadsheet_structure指定真实Sheet按需读取。'
        )
    ).strip()


def execute_spreadsheet_structure(params,deadline=None):
    source, err_res = resolve_session_attachment(params.get('file_path',''), '结构工具')
    if err_res:
        return TextToolOutput(err_res.render_text(), err_res)
    if not params.get('sheets'):
        result = ToolResult.error('invalid_sheet_selection', '需要指定真实工作表名')
    else:
        result = spreadsheet_structure(source, params['sheets'], deadline)
    return TextToolOutput(result.render_text(), result)
