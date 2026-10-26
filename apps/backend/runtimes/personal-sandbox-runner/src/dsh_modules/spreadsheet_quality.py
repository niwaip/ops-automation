"""Full declared-range profiles and explicit field-rule checks."""

from collections import Counter
import datetime as dt
import re
import statistics
from decimal import Decimal
from openpyxl.utils import get_column_letter
from .spreadsheet_expression import source_column, numeric


def iqr_fences(nums, multiplier=3.0):
    """Computes standard Tukey outer fences [Q1 - k*IQR, Q3 + k*IQR] when n >= 4 and IQR > 0."""
    if len(nums) < 4:
        return None
    try:
        vals = [float(numeric(v)) for _, v in nums]
        q1, _, q3 = statistics.quantiles(vals, n=4, method="inclusive")
        iqr = q3 - q1
        if iqr <= 0:
            return None
        lower = Decimal(str(round(q1 - multiplier * iqr, 6)))
        upper = Decimal(str(round(q3 + multiplier * iqr, 6)))
        return lower, upper
    except Exception:
        return None


def literal(value):
    if isinstance(value, (dt.date, dt.datetime)):
        return value.isoformat()
    return value


def date_value(value):
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    if isinstance(value, str):
        try:
            return dt.date.fromisoformat(value)
        except ValueError:
            return dt.datetime.fromisoformat(value).date()
    raise ValueError("不是日期")


def profiles(workbook, tables, check_deadline):
    result = []
    for table in tables.values():
        ws = workbook[table["sheet"]]
        rows = table["selected_rows"]
        a, _, b, _ = table["bounds"]
        row_tuples = {}
        for r in rows:
            tup = tuple(literal(ws.cell(r, c).value) for c in range(a, b + 1))
            if any(v is not None and str(v).strip() != "" for v in tup):
                row_tuples.setdefault(tup, []).append(r)
        dup_rows = sorted(r for r_list in row_tuples.values() if len(r_list) > 1 for r in r_list)
        for col in range(a, b+1):
            check_deadline()
            col_letter = get_column_letter(col)
            values = [(r, ws.cell(r, col).value) for r in rows]
            nonblank = [(r, v) for r, v in values if v is not None and str(v).strip()]
            missing_row_list = [r for r, v in values if v is None or not str(v).strip()]
            item = {"table": table["id"], "column": col_letter,
                    "header": str(ws.cell(table["header_row"], col).value or ""),
                    "records": len(rows), "missing": len(rows)-len(nonblank)}
            if col == a:
                item["table_duplicate_row_count"] = len(dup_rows)
                item["table_duplicate_rows"] = dup_rows[:500]
            if missing_row_list:
                item["missing_rows"] = missing_row_list[:500]
            item['missing_examples']=[f"{ws.title}!{col_letter}{r}" for r in missing_row_list[:20]]
            nums = [(r, v) for r, v in nonblank if isinstance(v, (int, float, Decimal)) and not isinstance(v, bool)]
            dates = [(r, v) for r, v in nonblank if isinstance(v, (dt.date, dt.datetime))]
            if nums:
                minimum = min(nums, key=lambda x: numeric(x[1])); maximum = max(nums, key=lambda x: numeric(x[1]))
                num_info = {"count": len(nums), "min": literal(minimum[1]), "min_row": minimum[0],
                            "max": literal(maximum[1]), "max_row": maximum[0]}
                fences = iqr_fences(nums)
                if fences is not None:
                    low_f, high_f = fences
                    outliers = [
                        {"row": r, "cell": f"{ws.title}!{col_letter}{r}", "value": literal(v)}
                        for r, v in nums
                        if numeric(v) < low_f or numeric(v) > high_f
                    ]
                    if outliers:
                        num_info["outlier_count"] = len(outliers)
                        num_info["outliers"] = outliers[:20]
                item["numeric"] = num_info
            if dates:
                first=min(dates,key=lambda x:date_value(x[1]));last=max(dates,key=lambda x:date_value(x[1]))
                item["date"] = {"count": len(dates), "min": str(min(date_value(v) for _, v in dates)),
                                "max": str(max(date_value(v) for _, v in dates)),
                                "min_row":first[0],"max_row":last[0]}
            counts = Counter(str(literal(v)) for _, v in nonblank)
            item["distinct_values"] = len(counts)
            if len(counts) <= 20:
                item["value_counts"] = dict(sorted(counts.items()))
                examples = {}
                for row, value in nonblank:
                    key = str(literal(value))
                    bucket = examples.setdefault(key, [])
                    if len(bucket) < 3: bucket.append(f"{ws.title}!{col_letter}{row}")
                item["value_examples"] = examples
            result.append(item)
    return result


def numeric_bound(value):
    if isinstance(value,str) and re.fullmatch(r"[+-]?(?:[0-9]+(?:\.[0-9]*)?|\.[0-9]+)(?:[eE][+-]?[0-9]+)?",value.strip()):
        return numeric(Decimal(value))
    return numeric(value)


def evaluate_rule(rule, workbook, tables, check_deadline):
    required = {"id", "label", "table", "column", "kind", "assumptions"}
    if not required <= set(rule) or set(rule)-required-{"min", "max", "allowed"}:
        raise ValueError("字段规则缺少必要项或包含未知字段")
    if not isinstance(rule["assumptions"], str) or not rule["assumptions"].strip():
        raise ValueError("规则必须声明依据或待核查假设")
    table = tables[rule["table"]]; ws = workbook[table["sheet"]]
    kind = rule["kind"]
    if kind not in {"unique", "not_empty", "date_range", "allowed_values", "number_range", "outlier"}:
        raise ValueError("未知字段规则")
    col_raw = str(rule["column"] or "").strip()
    is_full_row_unique = (
        kind == "unique"
        and bool(re.fullmatch(r"(?:\*|all|全列|全行|全部列|所有列|所有\d+列|\d+列)", col_raw, re.I))
    )
    if is_full_row_unique:
        a, _, b, _ = table["bounds"]
        col = a
        values = [
            (r, tuple(literal(ws.cell(r, c).value) for c in range(a, b + 1)))
            for r in table["selected_rows"]
        ]
        counts = Counter(v for _, v in values if any(x is not None and str(x).strip() != "" for x in v))
    else:
        col = source_column(workbook,table,rule["column"])
        if not table["bounds"][0] <= col <= table["bounds"][2]:
            raise ValueError("规则列不在声明范围")
        values = [(r, ws.cell(r, col).value) for r in table["selected_rows"]]
        counts = Counter(str(literal(v)) for _, v in values if v is not None)
    if kind == "allowed_values" and (not isinstance(rule.get("allowed"), list) or not rule["allowed"]):
        raise ValueError("枚举规则需要非空allowed")
    if kind in {"date_range", "number_range"} and not ({"min", "max"} & set(rule)):
        raise ValueError("范围规则至少需要min或max，不会推定未声明的边界")
    criteria = {k: rule[k] for k in ("min", "max", "allowed") if k in rule}
    if kind == "date_range":
        lower = date_value(rule["min"]) if 'min' in rule else None
        upper = date_value(rule["max"]) if 'max' in rule else None
    elif kind == "number_range":
        lower = numeric_bound(rule["min"]) if 'min' in rule else None
        upper = numeric_bound(rule["max"]) if 'max' in rule else None
    elif kind == "outlier":
        nums = [(r, v) for r, v in values if isinstance(v, (int, float, Decimal)) and not isinstance(v, bool)]
        fences = iqr_fences(nums)
        lower, upper = fences if fences is not None else (None, None)
        criteria = {"method": "iqr", "multiplier": 3.0}
        if lower is not None and upper is not None:
            criteria["lower_fence"] = str(round(lower, 2))
            criteria["upper_fence"] = str(round(upper, 2))
    else:
        lower = upper = None
    if lower is not None and upper is not None and lower > upper:
        raise ValueError("规则上下界颠倒")
    violations = []
    for row, value in values:
        check_deadline()
        if is_full_row_unique:
            failed = counts.get(value, 0) > 1
        elif kind == "unique":
            failed = value is None or counts[str(literal(value))] > 1
        elif kind == "not_empty":
            failed = value is None or not str(value).strip()
        elif kind == "allowed_values":
            failed = value not in rule["allowed"]
        else:
            try:
                parsed = date_value(value) if kind == "date_range" else numeric(value)
                failed = (lower is not None and parsed < lower) or (upper is not None and parsed > upper)
            except (ValueError, TypeError):
                failed = True
        if failed:
            violations.append({"row": row, "cell": f"{ws.title}!{get_column_letter(col)}{row}",
                               "value": " | ".join(str(x or "") for x in value[:6]) if is_full_row_unique else literal(value)})
    return {"id": rule["id"], "label": rule["label"], "kind": kind,
            "table": table["id"], "column": rule["column"], "assumptions": rule["assumptions"],
            "evaluated_records": len(values), "violation_count": len(violations),
            "violation_rows": [v["row"] for v in violations[:500]],
            "passed": not violations, "examples": violations[:20],
            "examples_complete": len(violations) <= 20,
            "basis_status": "declared_unverified",
            "criteria": criteria}
