#!/usr/bin/env python3
"""Verify immutable office dependencies as the actual non-root sandbox user."""
import importlib
import json
from pathlib import Path
import tempfile
import zipfile
import sys


def main():
    modules = ('pptx', 'docx', 'openpyxl', 'xlsxwriter', 'PIL', 'pypdf',
               'weasyprint', 'markdown', 'jinja2', 'pdfplumber')
    for module in modules:
        importlib.import_module(module)
    from pptx import Presentation
    with tempfile.TemporaryDirectory(prefix='sandbox-runtime-') as folder:
        target = Path(folder) / 'check.pptx'
        deck = Presentation()
        deck.slides.add_slide(deck.slide_layouts[5]).shapes.title.text = '运行时检查'
        deck.save(target)
        assert len(Presentation(target).slides) == 1
        with zipfile.ZipFile(target) as archive:
            assert archive.testzip() is None
        sys.path.insert(0, '/opt/dsh/skills/pdf/scripts')
        from pdf_design import render_report
        from pypdf import PdfReader
        pdf = Path(folder) / 'check.pdf'
        render_report(['# 中文运行时检查', '', '正文与 **粗体**：Hello，世界。'], str(pdf))
        assert '中文运行时检查' in PdfReader(pdf).pages[0].extract_text()
    print(json.dumps({'runtime_contract': 'office-v1', 'modules': list(modules),
                      'pptx_roundtrip': True, 'pdf_roundtrip': True}, ensure_ascii=False))


if __name__ == '__main__':
    main()
