"""Read-only export adapter and independent, exact SQL oracle.

Only this evaluator sees SQL/expected results; no production runtime imports.
Numeric literals are extracted from OOXML and scaled to integers before SQLite.
"""
import argparse
from collections import Counter
from decimal import Decimal
import hashlib
import json
from pathlib import Path
import sqlite3
import xml.etree.ElementTree as ET
import zipfile

import openpyxl

from catalog import CASES, FIELDS

NS = {'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main',
      'r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
NUMERIC = {'amount','hours','unit_price'}


def load_source(source):
    source = Path(source)
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    wb = openpyxl.load_workbook(source, read_only=False, data_only=False)
    try:
        candidates = []
        for ws in wb:
            for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 40)):
                labels = [c.value for c in row]
                if all(v in labels for v in FIELDS.values()):
                    if any(labels.count(v)!=1 for v in FIELDS.values()):
                        raise ValueError('重复字段，需要独立fixture适配器')
                    candidates.append((ws,row[0].row,{k:labels.index(v)+1 for k,v in FIELDS.items()}))
        if len(candidates)!=1:
            raise ValueError('该基线适配器要求唯一的工数实绩表头，不猜其他数据域')
        ws, header, columns = candidates[0]
        # This export marks the footer using full-width merged rows. Do not
        # filter by non-null measures: doing so would conceal missing data.
        footer = min((r.min_row for r in ws.merged_cells.ranges
                      if r.min_row>header and r.min_col==1 and r.max_col>=max(columns.values())),
                     default=ws.max_row+1)
        with zipfile.ZipFile(source) as z:
            book = ET.fromstring(z.read('xl/workbook.xml'))
            sheet = next(s for s in book.find('s:sheets',NS) if s.attrib['name']==ws.title)
            rel_id = sheet.attrib['{'+NS['r']+'}id']
            rels = ET.fromstring(z.read('xl/_rels/workbook.xml.rels'))
            target = next(r.attrib['Target'] for r in rels if r.attrib['Id']==rel_id)
            target = target.lstrip('/') if target.startswith('/') else 'xl/'+target
            root = ET.fromstring(z.read(target))
        numbers = {c.attrib['r']: Decimal(c.find('s:v',NS).text)
                   for c in root.findall('.//s:sheetData/s:row/s:c',NS)
                   if c.attrib.get('t','n')=='n' and c.find('s:v',NS) is not None}
        records = []
        for r in range(header+1,footer):
            record = {'row_id':r}
            for field,col in columns.items():
                c = ws.cell(r,col)
                if c.data_type=='f':
                    raise ValueError('基线有公式：需要明确重算凭据，不能把缓存当原始常量')
                value = c.value
                if field in NUMERIC:
                    value = None if value is None else numbers[c.coordinate]
                elif value is not None:
                    value = str(int(value)) if isinstance(value,(int,float)) and value==int(value) else str(value)
                record[field] = value
            records.append(record)
        if not records:
            raise ValueError('无明细')
        scales = {k:10**max((max(0,-r[k].as_tuple().exponent) for r in records if r[k] is not None),default=0)
                  for k in NUMERIC}
        # SQLite integer arithmetic is exact only within signed 64-bit bounds.
        for field in NUMERIC:
            if sum(abs(int(r[field]*scales[field])) for r in records if r[field] is not None)>=2**63:
                raise ValueError('超出SQL整数预算；需Decimal引擎适配器')
        profile = {'sha256':digest,'bytes':source.stat().st_size,'sheet':ws.title,'header_row':header,
                   'data_range':f'A{header+1}:{openpyxl.utils.get_column_letter(ws.max_column)}{footer-1}',
                   'physical_rows':ws.max_row,'detail_rows':len(records),'footer_start':footer,
                   'headers':{k: {'label':FIELDS[k],'column':openpyxl.utils.get_column_letter(c)} for k,c in columns.items()},
                   'scales':scales,'formula_policy':'literal-only baseline',
                   'missing':{k:sum(r[k] is None for r in records) for k in FIELDS},
                   'cardinality':{k:len({r[k] for r in records}) for k in FIELDS},
                   'categories':{k:dict(Counter(r[k] for r in records)) for k in ('billable','direct','month')},
                   'source_unchanged':digest==hashlib.sha256(source.read_bytes()).hexdigest()}
        if not profile['source_unchanged']:
            raise ValueError('读取期间来源变化')
        return records,profile
    finally:
        wb.close()


class DecimalAverage:
    def __init__(self):
        self.total,self.count=0,0
    def step(self,value):
        if value is not None:
            self.total+=value
            self.count+=1
    def finalize(self):
        return str(Decimal(self.total)/self.count) if self.count else None


def connection(records, scales):
    conn = sqlite3.connect(':memory:')
    conn.row_factory = sqlite3.Row
    conn.create_aggregate('DEC_AVG',1,DecimalAverage)
    fields=['row_id',*FIELDS]
    conn.execute('CREATE TABLE detail ('+','.join(f'"{f}" '+('INTEGER' if f in NUMERIC or f=='row_id' else 'TEXT') for f in fields)+')')
    conn.executemany('INSERT INTO detail VALUES ('+','.join('?' for _ in fields)+')',
                     [[None if r[f] is None else int(r[f]*scales[f]) if f in NUMERIC else r[f] for f in fields] for r in records])
    return conn


def expected_result(conn, c, profile):
    params={k+'_scale':v for k,v in profile['scales'].items()}
    result=[dict(r) for r in conn.execute(c['sql'],params)]
    for row in result:
        for k,v in row.items():
            field='amount' if k=='amount_k' else 'hours' if k in ('hours','min_hours','max_hours','avg_hours') else None
            if field:
                # Empty aggregate SUM is explicitly normalized to zero for these
                # business tasks; AVG/MIN/MAX with no rows remain undefined.
                if v is None and k in ('amount_k','hours'):
                    v=0
                row[k]=None if v is None else str(Decimal(v)/profile['scales'][field])
    return {'rows':result,'selected_row_ids':[r[0] for r in conn.execute(c['selection_sql'],params)],
            'ordered':c['ordered'],'coverage_scope':c['coverage_scope']}


def prepare(source, output):
    output=Path(output)
    output.mkdir(parents=True,exist_ok=False)
    records,profile=load_source(source)
    conn=connection(records,profile['scales'])
    try:
        expected={c['id']:expected_result(conn,c,profile) for c in CASES}
    finally:
        conn.close()
    manifest={'schema_version':1,'fixture_sha256':profile['sha256'],'cases':[
        {'id':c['id'],'family':c['family'],'prompt':c['prompt'],'mode':'chat',
         'assertions':['source_binding','result_semantics','tool_contract','route_evidence','no_source_mutation'],
         'status':'not_run'} for c in CASES]}
    for name,data in [('profile.json',profile),('expected.json',expected),('manifest.json',manifest),
                      ('catalog.json',[{k:v for k,v in c.items() if k!='native'} for c in CASES])]:
        (output/name).write_text(json.dumps(data,ensure_ascii=False,indent=2)+'\n')
    # Relations on the independent oracle catch vacuous/incorrect specifications.
    count=lambda cid: expected[cid]['rows'][0]['records']
    assert count('Q02')+count('Q03')==count('Q01')
    assert count('Q04')+count('Q05')==count('Q01')
    assert count('Q08')+count('Q09')==count('Q01')
    assert count('Q12')+count('Q36')==count('Q01')
    assert count('Q13')+count('Q35')==count('Q01')
    summary={'cases':len(CASES),'native_adapters':sum(c['native'] is not None for c in CASES),
             'detail_rows':len(records),'baseline':expected['Q01']['rows'][0],
             'zero_match_cases':[cid for cid,e in expected.items() if not e['selected_row_ids']],
             'live_e2e_status':'not_run','source_unchanged':profile['source_unchanged']}
    (output/'summary.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(summary,ensure_ascii=False))


if __name__=='__main__':
    p=argparse.ArgumentParser()
    p.add_argument('source',type=Path)
    p.add_argument('output',type=Path)
    args=p.parse_args()
    prepare(args.source,args.output)
