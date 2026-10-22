import os
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    tcPr.append(shd)

def create_styled_document():
    doc = docx.Document()
    
    # Page margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
    return doc

def generate_software_dev_contract(output_path):
    doc = create_styled_document()
    
    # Title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run('软件定制研发与系统集成协议')
    run.font.name = 'SimHei'
    run.font.size = Pt(20)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x1A, 0x2D, 0x42)
    
    # Contract Code
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_sub = sub.add_run('合同编号：CR_DEV_20261017_01')
    r_sub.font.name = 'SimSun'
    r_sub.font.size = Pt(10.5)
    r_sub.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
    
    # Metadata Table
    table = doc.add_table(rows=2, cols=2)
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table.autofit = False
    
    meta_rows = [
        ('甲方（委托方）：', '北京智通数科技术有限公司', '法定地址：', '北京市海淀区中关村软件园二期12号楼'),
        ('乙方（受托方）：', '上海云联软件工程有限公司', '法定地址：', '上海市浦东新区张江高科技园区博云路2号')
    ]
    
    for r_idx, (k1, v1, k2, v2) in enumerate(meta_rows):
        row = table.rows[r_idx]
        row.cells[0].text = f"{k1}{v1}"
        row.cells[1].text = f"{k2}{v2}"
        set_cell_background(row.cells[0], "F8FAFC")
        set_cell_background(row.cells[1], "F8FAFC")
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.space_before = Pt(2)
                for run in p.runs:
                    run.font.name = 'SimSun'
                    run.font.size = Pt(9.5)
                    run.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
    
    doc.add_paragraph()
    
    # Preamble
    p_pre = doc.add_paragraph()
    r_pre = p_pre.add_run('鉴于甲方因企业业务数字化升级需求，委托乙方开展企业级智能协作办公系统定制研发与系统集成服务；乙方具备专业软件技术开发资质与工程团队，双方依据《中华人民共和国民法典》及相关法律法规，在自愿、平等、互利原则基础上，经友好协商订立本协议。')
    r_pre.font.name = 'SimSun'
    r_pre.font.size = Pt(11)
    
    clauses = [
        ("第一条 开发内容与技术规格", 
         "1.1 乙方负责按照双方共同签署的《需求规格说明书》（附件一）开展微服务架构体系下的协同平台开发、数据接口联调及环境部署工作。\n"
         "1.2 乙方承诺交付的系统符合国家软件安全规范及行业通用性能指标要求，提供完备的技术架构文档与接口规范说明。"),
        
        ("第二条 项目工期与研发进度（工期计算口径）", 
         "2.1 项目工期严格限定为自合同签署生效之日起 60 个自然日/日历日完成全部开发及部署上线，遇法定节假日及周末不予顺延。\n"
         "2.2 任何因第三方接口对接、网络联调、数据迁移或系统测试产生的等待时间，均全部计入乙方的 60 个自然日工期内，不构成工期顺延理由。"),
        
        ("第三条 交付源码与知识产权归属", 
         "3.1 本协议项下定制开发产生之全部源代码、设计文档、算法模型及技术成果之全部知识产权，自产生之日起独家归甲方所有。\n"
         "3.2 乙方在签署本协议前已独立拥有的基础支撑框架代码与通用组件，亦在此无偿且永久独占转让归甲方所有，乙方不得在任何其他商业项目中继续复用。"),
        
        ("第四条 系统验收与试运行标准", 
         "4.1 系统功能开发完成后，以甲方业务部门单方主观满意度评价及集团数字化管理委员会最终审批通过作为唯一验收通过条件。\n"
         "4.2 系统进入试运行阶段后，甲方有权在试运行期内随时提出实质性功能变更与整改要求，且本协议不设最终试运行截止期限，直至甲方满意为止。"),
        
        ("第五条 开源组件合规与第三方连带责任", 
         "5.1 乙方交付的软件系统不得包含侵犯任何第三方知识产权的开源组件。\n"
         "5.2 乙方对系统内引用的所有第三方依赖包及开源组件承担无条件的无限连带赔偿责任；如因任何开源许可争议或第三方诉讼导致甲方被索赔，乙方须无条件全额代偿并赔偿甲方全部商誉与间接损失。"),
        
        ("第六条 违约责任与违约金封顶约定", 
         "6.1 乙方逾期交付任何开发阶段里程碑成果的，每逾期一日，应按本协议合同总金额千分之五（0.5%）向甲方支付违约金，且违约金累计金额不设任何封顶上限。\n"
         "6.2 乙方逾期超过 10 个自然日的，甲方有权立即单方解除本协议，乙方应退还甲方已支付的全部款项，并按合同总额的 50% 承担惩罚性赔偿责任。"),
        
        ("第七条 不可抗力与合理免责", 
         "7.1 任何一方因地震、台风、洪水、战争等无法预见、无法避免且无法克服的法定不可抗力事件直接导致延期或不能履行合同的，遭遇不可抗力方应在事件发生后 48 小时内书面通知对方。\n"
         "7.2 在不可抗力影响范围内，遭遇方免除违约责任，项目交付工期按事件持续时间相应对等顺延。"),
        
        ("第八条 争议解决与司法管辖", 
         "8.1 因本协议履行过程中发生的一切争议，双方应首先通过友好协商解决。\n"
         "8.2 协商不成的，任何一方均有权向原告所在地有管辖权的人民法院提起诉讼。"),
        
        ("第九条 协议生效与附则", 
         "9.1 本协议自双方法定代表人或授权委托代理人签字并加盖双方公章之日起生效。\n"
         "9.2 本协议一式肆份，甲乙双方各执贰份，具有同等法律效力。")
    ]
    
    for title_text, content_text in clauses:
        doc.add_paragraph() # spacing
        p_title = doc.add_paragraph()
        p_title.paragraph_format.space_before = Pt(8)
        p_title.paragraph_format.space_after = Pt(4)
        r_t = p_title.add_run(title_text)
        r_t.font.name = 'SimHei'
        r_t.font.size = Pt(12)
        r_t.font.bold = True
        r_t.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
        
        for line in content_text.split('\n'):
            p_c = doc.add_paragraph()
            p_c.paragraph_format.space_after = Pt(3)
            p_c.paragraph_format.line_spacing = 1.3
            r_c = p_c.add_run(line)
            r_c.font.name = 'SimSun'
            r_c.font.size = Pt(10.5)
            r_c.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
            
    doc.save(output_path)
    print(f"Generated DOCX: {output_path}")

def generate_procurement_contract(output_path):
    doc = create_styled_document()
    
    # Title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run('商业采购买卖与设备供货协议')
    run.font.name = 'SimHei'
    run.font.size = Pt(20)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x1A, 0x2D, 0x42)
    
    # Contract Code
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_sub = sub.add_run('合同编号：CR_PROC_20261017_02')
    r_sub.font.name = 'SimSun'
    r_sub.font.size = Pt(10.5)
    r_sub.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
    
    # Metadata Table
    table = doc.add_table(rows=2, cols=2)
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table.autofit = False
    
    meta_rows = [
        ('买方（甲方）：', '星海智能装备制造股份有限公司', '注册地址：', '江苏省苏州市工业园区金鸡湖大道88号'),
        ('卖方（乙方）：', '华东机电精密工业有限公司', '注册地址：', '浙江省宁波市高新区微电子创新园5号')
    ]
    
    for r_idx, (k1, v1, k2, v2) in enumerate(meta_rows):
        row = table.rows[r_idx]
        row.cells[0].text = f"{k1}{v1}"
        row.cells[1].text = f"{k2}{v2}"
        set_cell_background(row.cells[0], "F8FAFC")
        set_cell_background(row.cells[1], "F8FAFC")
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.space_before = Pt(2)
                for run in p.runs:
                    run.font.name = 'SimSun'
                    run.font.size = Pt(9.5)
                    run.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
    
    doc.add_paragraph()
    
    clauses = [
        ("第一条 采购标的与供货范围",
         "1.1 甲方向乙方采购工业级伺服自动化生产线成套设备共计两套，总合同价款为人民币大写壹佰贰拾万元整（¥1,200,000.00）。\n"
         "1.2 乙方负责设备出厂检验、包装运输至甲方指定工厂，并派遣专业工程师提供现场安装调试与技术培训。"),
         
        ("第二条 货款结算与付款账期（背靠背支付条款）",
         "2.1 本合同货款结算严格以买方最终客户向买方足额支付对应期工程款为前提条件（背靠背支付）。\n"
         "2.2 如因买方最终客户验收延迟、预算审批拖延或拒付款项，买方对乙方相应免除付款义务及任何逾期违约利息责任，乙方不得向买方主张索赔或暂停后续售后维护。"),
         
        ("第三条 到货检验与隐蔽瑕疵异议期限",
         "3.1 货物送达甲方指定仓库后，甲方应在到货之日起 3 个自然日内完成外观及运行检验并向乙方提出书面质量异议。\n"
         "3.2 甲方未在 3 个自然日内书面提出异议的，即视为标的物完全合格且符合约定质量，甲方永久放弃对包括隐蔽瑕疵在内的任何质量异议索赔权。"),
         
        ("第四条 违约赔偿与全部间接损失责任",
         "4.1 任何一方违反本合同约定的，应当承担违约责任，赔偿守约方因此遭受的全部经济损失。\n"
         "4.2 违约赔偿范围不仅包括直接实际财产损失，违约方还应赔偿对方因此产生的全部间接损失、利润损失、商业机会损失、预期可得利益及商誉贬损等不确定衍生损失。"),
         
        ("第五条 不可抗力与争议管辖",
         "5.1 因不可抗力导致合同不能履行的，遭遇方应及时书面通知对方，并提供当地商会证明文件，双方协商变更或解除合同。\n"
         "5.2 双方因履行本合同发生任何争议，应友好协商解决；协商不成的，由原告所在地人民法院诉讼管辖。")
    ]
    
    for title_text, content_text in clauses:
        doc.add_paragraph()
        p_title = doc.add_paragraph()
        p_title.paragraph_format.space_before = Pt(8)
        p_title.paragraph_format.space_after = Pt(4)
        r_t = p_title.add_run(title_text)
        r_t.font.name = 'SimHei'
        r_t.font.size = Pt(12)
        r_t.font.bold = True
        r_t.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
        
        for line in content_text.split('\n'):
            p_c = doc.add_paragraph()
            p_c.paragraph_format.space_after = Pt(3)
            p_c.paragraph_format.line_spacing = 1.3
            r_c = p_c.add_run(line)
            r_c.font.name = 'SimSun'
            r_c.font.size = Pt(10.5)
            r_c.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
            
    doc.save(output_path)
    print(f"Generated DOCX: {output_path}")

if __name__ == '__main__':
    base_dir = os.path.dirname(os.path.abspath(__file__))
    generate_software_dev_contract(os.path.join(base_dir, 'software-dev-contract.docx'))
    generate_procurement_contract(os.path.join(base_dir, 'procurement-contract.docx'))
