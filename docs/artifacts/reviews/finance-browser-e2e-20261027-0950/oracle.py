"""Read-only independent expectations for this specific synthetic workbook."""
import hashlib
import json
import re
import sys
from collections import defaultdict
from decimal import Decimal as D
from pathlib import Path

import openpyxl

source = Path(sys.argv[1])
book = openpyxl.load_workbook(source, data_only=False)
cached = openpyxl.load_workbook(source, data_only=True)
sheet = book['月度经营']
revenue = gross = operating = net = D(0)
monthly = []
for row in range(4, 16):
    amount = D(str(sheet.cell(row, 2).value))
    margin = D(re.search(r'1-([0-9.]+)', sheet.cell(row, 3).value)[1])
    sale = D(sheet.cell(row, 6).value.split('*')[1])
    admin = D(sheet.cell(row, 7).value.split('*')[1])
    gp = amount * margin
    op = gp - amount * sale - amount * admin - D(str(sheet.cell(row, 8).value))
    pretax = op + D(str(sheet.cell(row, 11).value))
    np = pretax - max(D(0), pretax * D(str(book['参数']['B8'].value)))
    revenue += amount
    gross += gp
    operating += op
    net += np
    budget = D(str(book['预算对比'].cell(row, 2).value))
    monthly.append({'month': sheet.cell(row, 1).value, 'income_difference': amount-budget})
ids = defaultdict(list)
for row in range(4, 84):
    ids[book['交易明细'].cell(row, 1).value].append(row)
balance = book['资产负债表']
assets = sum(D(str(balance.cell(row, 2).value)) for row in [4, 5, 6, 7, 9, 10])
debts = sum(D(str(balance.cell(row, 2).value)) for row in [12, 13, 14, 15])
cash = D(12100) + net + D(2200)-D(2350)-D(1100)+D(700)-D(1800)-D(2000)-D(3200)
formulas = [(s.title, c.coordinate) for s in book for row in s for c in row if c.data_type == 'f']
result = {
    'sha256': hashlib.sha256(source.read_bytes()).hexdigest(),
    'formula_count': len(formulas),
    'uncached_formula_count': sum(cached[s][c].value is None for s, c in formulas),
    'annual': {'revenue': revenue, 'gross': gross, 'gross_margin': gross/revenue,
               'operating_profit': operating, 'net_profit': net},
    'budget_comparison': {'monthly_income_budget': sum(book['预算对比'].cell(r, 2).value for r in range(4, 16)),
                          'annual_income_budget': book['利润表']['C4'].value,
                          'monthly_profit_budget': sum(book['预算对比'].cell(r, 5).value for r in range(4, 16)),
                          'annual_profit_budget': book['利润表']['C14'].value},
    'largest_income_difference': max(monthly, key=lambda x: abs(x['income_difference'])),
    'duplicates': {key: rows for key, rows in ids.items() if len(rows)>1},
    'balance_gap': assets - debts - D(26000) - net,
    'closing_cash': cash,
    'cash_gap': cash - D(str(balance['B4'].value)),
}
print(json.dumps(result, ensure_ascii=False, default=str, indent=2))
