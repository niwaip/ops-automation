"""Verify live regression evidence against an unchanged, independently calculated oracle."""
import hashlib
import json
import re
from collections import Counter
from decimal import Decimal as D
from pathlib import Path

import openpyxl

BASE = Path(__file__).resolve().parent
OLD = BASE.parent / 'finance-harness-35b-20261025'
oracle = json.loads((OLD/'verification.json').read_text())
source = Path('/Users/chain/Desktop/AI能力测试_财务报表.xlsx')
assert hashlib.sha256(source.read_bytes()).hexdigest() == oracle['source_hash']
golden = {k: D(v) for k, v in oracle['golden'].items()}
golden.update(cost=D('70576.75'), operating_margin=golden['operating_profit']/golden['revenue'])


def load(name):
    return json.loads((BASE/name).read_text())


def evidence(i):
    obj = load(f'F{i:02}-evidence.json')
    assert obj['source_hashes']['AI能力测试_财务报表.xlsx'] == oracle['source_hash']
    return obj


def result(i):
    stream = load(f'F{i:02}-stream.json')
    assert any(e['type'] == 'done' for e in stream)
    return next(e for e in stream if e['type'] == 'result')


def close(actual, expected):
    assert abs(D(str(actual))-expected) < D('0.000001'), (actual, str(expected))


for n in range(1, 8):
    request = load(f'F{n:02}-request.json')
    assert request['modelId'] == 'd13e0d30-87b9-4e87-b96a-d8b1b3cf78af'
    assert request['config']['mode'] == 'chat'
    evidence(n)
    result(n)
prompts = load('prompts.json')
for n in range(1, 7):
    assert load(f'F{n:02}-request.json')['message'] == prompts[n-1]

ratio_unit_errors = []
for n in (1, 2):
    for receipt in evidence(n)['receipts']:
        for fact in (receipt.get('data') or {}).get('facts', []):
            if fact['id'] not in golden:
                continue
            expected = golden[fact['id']]
            is_ratio = 'margin' in fact['id']
            close(fact['value'], expected*100 if is_ratio and fact['unit'] == '%' else expected)
            if is_ratio and fact['unit'] not in ('%', 'ratio', 'rate'):
                ratio_unit_errors.append({'case': n, 'id': fact['id'], 'unit': fact['unit']})
    assert evidence(n)['verified'] is True
    assert '尚未通过执行证据核验' in result(n)['content']

comparison = next(r['data'] for r in evidence(3)['receipts']
                  if (r.get('data') or {}).get('kind') == 'column_comparison')
old_comparison = next(r['data'] for r in json.loads((OLD/'F03-evidence.json').read_text())['receipts']
                      if (r.get('data') or {}).get('kind') == 'column_comparison')
assert comparison['scope']['coverage_verified'] and comparison['scope']['row_count'] == 12
for new_metric, old_metric in zip(comparison['evaluated'], old_comparison['evaluated']):
    assert new_metric['label'] == old_metric['label']
    assert len(new_metric['rows']) == len(old_metric['rows']) == 12
    for new_row, old_row in zip(new_metric['rows'], old_metric['rows']):
        assert new_row['group'] == old_row['group']
        for key, value in old_row['exact'].items():
            close(new_row['exact'][key], D(value))

wb = openpyxl.load_workbook(source, data_only=True)
transactions = [(i, list(row)) for i, row in enumerate(wb['交易明细'].iter_rows(min_row=4, max_row=83, values_only=True), 4)]
counts = Counter(row[0] for _, row in transactions)
anomalies = {
    'duplicate_voucher': [i for i, row in transactions if counts[row[0]] > 1],
    'blank_department': [i for i, row in transactions if row[3] is None],
    'large_amount': [i for i, row in transactions if abs(row[6]) > 10000],
    'out_of_period': [i for i, row in transactions if row[1].year != 2026],
    'unapproved': [i for i, row in transactions if row[9] != '已审批'],
}
assert list(anomalies.values()) == [[12, 13], [28], [43], [58], [73]]
f4 = evidence(4)
receipt = next(r['data'] for r in f4['receipts'] if (r.get('data') or {}).get('facts'))
assert receipt['plan']['group_by'] == ['凭证号']
assert len(receipt['plan']['metrics']) == 1
assert not receipt.get('checks') and not receipt.get('rules')
duplicate = next(f for f in receipt['facts'] if f.get('group_key') == 'V2026-0009')
assert duplicate['source_cells'] == ['交易明细!A12', '交易明细!A13']
answer = result(4)['content'].split('<details>')[0]
zero_claims = re.findall(r'共发现\s*\*\*0\*\*', answer)
assert len(zero_claims) == 4
assert '第 **7** 行、第 **9** 行' in answer
assert f4['verified'] is True

checks = next(r['data']['checks'] for r in evidence(5)['receipts'] if (r.get('data') or {}).get('checks'))
for check, expected in zip(checks, [('60950', '67138.3', '-6188.3'), ('20588.3', '16850', '3738.3')]):
    for field, value in zip(('left', 'right', 'residual'), expected):
        close(check[field], D(value))
    assert check['passed'] is False

assert '尚未生成更新文件' in result(6)['content']
assert not any(t['name'] == 'send_file' for t in result(6)['data']['sandbox']['metrics']['toolCalls'])
out = BASE/'E2E_finance_35B_rerun.xlsx'
output = openpyxl.load_workbook(out, data_only=True)
required = {'营业收入': golden['revenue'], '毛利': golden['gross_profit'],
            '营业利润': golden['operating_profit'], '净利润': golden['net_profit']}
seen = set()
extra_totals = []
for row in output.active.values:
    if row[0] in required:
        close(row[1], required[row[0]])
        assert row[2] == '千元'
        seen.add(row[0])
    elif row[0] == '合计':
        extra_totals.append(row[1])
assert seen == set(required)
assert out.read_bytes() == Path('/Users/chain/Downloads/E2E_finance_35B_rerun (1).xlsx').read_bytes()
assert len(load('history-after-refresh.json')['messages']) == 10
env = load('environment.json')
code_root = Path('/Users/chain/Documents/MyProject/ops-automation/apps/backend/runtimes/personal-sandbox-runner/src/dsh_modules')
assert all(hashlib.sha256((code_root/name).read_bytes()).hexdigest() == digest for name, digest in env['source_files'].items())
assert hashlib.sha256(source.read_bytes()).hexdigest() == oracle['source_hash']
summary = {'verification_completed': True, 'overall_data_acceptance': False,
           'oracle': 'Previous independent Decimal reconstruction, input SHA-256 unchanged',
           'original_six_outcomes': ['FAIL', 'FAIL', 'RECOVERED', 'FAIL_DATA', 'PASS', 'INCOMPLETE'],
           'F01_F02_verified_but_failed_text': True, 'F01_invalid_ratio_units': ratio_unit_errors,
           'F03_verified_month_metric_groups': 24, 'F04_expected_anomaly_rows': anomalies,
           'F04_false_zero_categories': len(zero_claims), 'F04_wrong_duplicate_rows': [7, 9],
           'F05_equations_verified': 2, 'F06_no_new_delivery': True,
           'F07_required_metrics_verified': 4, 'F07_extra_business_undefined_total': extra_totals,
           'download_sha256': hashlib.sha256(out.read_bytes()).hexdigest(),
           'download_bytes': out.stat().st_size, 'refresh_download_identical': True,
           'history_messages_after_refresh': 10, 'code_unchanged_during_test': True,
           'source_unchanged': True}
(BASE/'verification.json').write_text(json.dumps(summary, ensure_ascii=False, indent=2, default=str))
print(json.dumps(summary, ensure_ascii=False, default=str))
