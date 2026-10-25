"""Check downloaded CSV and recovered, undelivered XLSX; never edit workbooks."""
import csv
import hashlib
import json
from decimal import Decimal as D
from pathlib import Path
import openpyxl

BASE = Path(__file__).resolve().parent
oracle = json.loads((BASE/'oracle.json').read_text())
issues = []
checks = {}
def check(condition, text):
    if not condition: issues.append(text)
def same(a,b):
    if a in (None,'') and b in (None,''): return True
    if isinstance(b,(float,int)) and not isinstance(b,bool):
        try: return abs(D(str(a))-D(str(b))) < D('0.0000001')
        except Exception: return False
    return a==b
def numeric(a,b):
    try: return abs(D(str(a))-D(str(b)))<D('0.0000001')
    except Exception: return False

source=Path('/Users/chain/Desktop/工数实绩明细_20261020134124.xlsx')
check(hashlib.sha256(source.read_bytes()).hexdigest()==oracle['source_sha256'],'Source file changed')
path=BASE/'E2E_workhours_r2_filtered.csv'
if path.exists():
    data=list(csv.reader(path.read_text(encoding='utf-8-sig').splitlines()))
    check(len(data)==50,'CSV must contain header + 49 records')
    check(data[0][1:]==oracle['headers'],'CSV source headers changed')
    for actual,record in zip(data[1:],oracle['filter']['records']):
        check(actual[0]==str(record['row']),f'CSV original row/order mismatch at {record["row"]}')
        for j,b in enumerate(record['values']):
            a=actual[j+1] if j+1<len(actual) else None
            check(same(a,b),f'CSV source row {record["row"]} field {j+1}: {a!r} != {b!r}')
    checks['csv_records_checked']=len(data)-1

path=BASE/'E2E_workhours_r2_edit.xlsx'
if path.exists():
    raw=openpyxl.load_workbook(path,data_only=False);val=openpyxl.load_workbook(path,data_only=True)
    ws=raw['明细'];vs=val['明细']
    expected=[r['values'][:] for r in oracle['records']]
    first=expected[0][:];expected[0][14]=10;del expected[1]
    first[4]='E2E_TEST';first[5]='测试记录';first[14]=8;first[15]=0;expected.append(first)
    check(ws.max_row==275 and ws.max_column==19,'Edit copy shape must be 275 x 19')
    check([c.value for c in ws[1]][:18]==oracle['headers'],'Edit original headers changed')
    check(ws['S1'].value=='工数天数','Missing calculated day column')
    for i,row in enumerate(expected,2):
        for j,b in enumerate(row,1):
            check(same(ws.cell(i,j).value,b),f'Edit row {i} col {j}: unauthorized/missing modification')
        formula=str(ws.cell(i,19).value or '')
        check(formula.replace(' ','').upper()==f'=O{i}/8',f'Edit formula incorrect at S{i}: {formula}')
        cached = vs.cell(i,19).value
        check(cached is not None,f'Edit formula cache missing at S{i} (not a wrong computed value)')
        if cached is not None:
            check(numeric(cached,D(str(row[14]))/8),f'Edit formula cached value incorrect at S{i}')
        check(ws.cell(i,15).number_format in ('0.00','#,##0.00'),'Hours format must show 2 decimals')
        check(ws.cell(i,16).number_format in ('0.00000','#,##0.00000'),'Amount format must show 5 decimals')
    check(ws.freeze_panes=='A2','Edit freeze header missing')
    check(ws.auto_filter.ref=='A1:S275','Edit full-range filter missing')
    cf=[{'range':str(k),'rules':[(r.type,r.operator,r.formula) for r in rs]} for k,rs in ws.conditional_formatting._cf_rules.items()]
    check(bool(cf),'Edit conditional formatting missing')
    checks['edit_records_checked']=len(expected);checks['conditional_formatting']=cf

path=BASE/'E2E_workhours_r2_report.xlsx'
if path.exists():
    raw=openpyxl.load_workbook(path,data_only=False);val=openpyxl.load_workbook(path,data_only=True)
    ws=raw['部门汇总'];vs=val['部门汇总'];seen=set();total_row=None
    for i in range(1,ws.max_row+1):
        code=ws.cell(i,1).value
        if code in ('合计','总计'):total_row=i;continue
        match=next((d for d in oracle['departments'] if d['code']==code),None)
        if not match:continue
        seen.add(code);check(ws.cell(i,2).value==match['name'],f'Report department name mismatch {code}')
        paid=next((p for p in oracle['pivot'] if p['department']==code and p['paid']=='有偿'),{'hours':'0'})
        unpaid=next((p for p in oracle['pivot'] if p['department']==code and p['paid']=='无偿'),{'hours':'0'})
        for col,target in [(3,match['hours']),(4,match['amount_K']),(5,paid['hours']),(6,unpaid['hours'])]:
            check(numeric(vs.cell(i,col).value,target),f'Report {code} col{col} value incorrect')
        check(str(ws.cell(i,7).value).replace(' ','').upper()==f'=E{i}/C{i}',f'Report {code} share formula incorrect')
        check(numeric(vs.cell(i,7).value,D(paid['hours'])/D(match['hours'])),f'Report {code} share value incorrect')
    check(seen=={d['code'] for d in oracle['departments']},'Report must contain all 7 departments')
    check(total_row is not None,'Report total row missing')
    if total_row:
        for col,target in [(3,oracle['summary']['hours']),(4,oracle['summary']['amount_K'])]:
            check(numeric(vs.cell(total_row,col).value,target),'Report total incorrect')
            check(str(ws.cell(total_row,col).value).upper().startswith('=SUM('),'Report total must use SUM')
        paid=sum(D(p['hours']) for p in oracle['pivot'] if p['paid']=='有偿')
        check(numeric(vs.cell(total_row,7).value,paid/D(oracle['summary']['hours'])),'Report overall share incorrect')
    check(bool(ws._charts),'Report column chart missing')
    check(ws.freeze_panes=='A2','Report header freeze missing')
    checks['report_departments_checked']=len(seen)
    checks['charts']=[{'type':type(c).__name__,'series':len(c.series),'references':[str(s) for s in c.series]} for c in ws._charts]

checks['output_issues']=issues
checks['inspected_files']=[p.name for p in BASE.iterdir() if p.suffix in ('.xlsx','.csv')]
checks['artifact_origins']={'csv':'downloaded through UI after refresh','edit':'recovered from workspace after timeout; not delivered','report':'recovered from workspace after timeout; not delivered'}
(BASE/'output-verification.json').write_text(json.dumps(checks,ensure_ascii=False,indent=2,default=str))
print(json.dumps(checks,ensure_ascii=False,default=str))
