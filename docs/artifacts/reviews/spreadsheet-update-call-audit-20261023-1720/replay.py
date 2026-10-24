"""Read-only call audit. No model calls, workbook writes, or service changes.

Reconstruct routing from persisted contract and same-hash runtime source.
Replay the recorded partial recalculation receipt; do not invoke LibreOffice.
"""
import copy
import hashlib
import json
from collections import Counter, defaultdict
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch

import openpyxl

from dsh_modules.analysis_contract import build_analysis_contract
from dsh_modules.analysis_request_text import user_instruction
from dsh_modules.agent_loop import _check_no_tool_assertion_guard
from dsh_modules.prompt_builder import build_user_turn
from dsh_modules.runner import select_active_tools
from dsh_modules.skill_admission import general_route, execution_contract_context
from dsh_modules.spreadsheet_analysis_evidence import requires_spreadsheet_analysis
from dsh_modules.spreadsheet_analysis import calculate_spreadsheet_analysis
from dsh_modules.spreadsheet_update.tools import SPREADSHEET_UPDATE_TOOL_DEFINITIONS
from dsh_modules.workbook_calculation import check_workbook_formulas

ROOT = Path.cwd()
OUT = Path(__file__).resolve().parent
WS = ROOT / 'data/users/e7fce333-a8f4-4097-9a53-f0a4c729da46/workspace'
SID = 'chat-session-1792747141627-wheti8di'
BASE = WS / '.dsh/sessions'
EP = next((BASE / (SID + '.evidence')).glob('*.json'))
E = json.loads(EP.read_text())
accepted = next(g['reason'] for g in E['guard_decisions'] if g['guard'] == 'task_intent')
prompt = user_instruction(E['query'])
contract = build_analysis_contract(prompt)
route = general_route(E['query'], [], list(E['source_hashes']), accepted)
tools = select_active_tools(E['query'], route, False, True)
names = [t['function']['name'] for t in tools]
user_turn = build_user_turn(E['query'], session_files=list(E['source_hashes']),
                            is_generate_intent=route.is_generate_intent,
                            is_inspect_intent=route.is_inspect_intent,
                            is_office_intent=route.is_office_intent,
                            available_tool_names=names,
                            timestamp_str='2026-10-23 17:19:16 Asia/Shanghai')
# This candidate is a controlled guard input, not recovered historical model text.
with patch('dsh_modules.agent_loop.detect_missing_claimed_artifacts',
           return_value=['工数3Q预实管理_1023.xlsx']):
    guard_action, guard_hint, _ = _check_no_tool_assertion_guard(
        reply_text='已生成文件 工数3Q预实管理_1023.xlsx。', last_user_prompt=E['query'],
        round_idx=3, max_rounds=10, start_ts=1792747156, is_guide_intent=False,
        expected_deliverables=['.xlsx'], is_generate_intent=True, is_inspect_intent=False,
        tools=tools, hard_ceiling=10)
target = WS / '工数3Q预实管理.xlsx'
source = WS / '工数实绩明细_20261020134124.xlsx'
before = {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in [target, source]}

wb = openpyxl.load_workbook(source, read_only=True, data_only=True)
ss = wb['工数实绩明细']; ss.reset_dimensions()
amounts = defaultdict(Decimal); counts = Counter(); months = Counter()
for row in ss.iter_rows(min_row=5, max_row=278, max_col=18, values_only=True):
    months[str(row[1])] += 1
    if row[12] == '有偿':
        amounts[str(row[2])] += Decimal(str(row[15]))
        counts[str(row[2])] += 1
wb.close()
analysis = next(r['data'] for r in E['receipts'] if (r.get('data') or {}).get('kind') == 'spreadsheet_analysis')
amount_facts = [f for f in analysis['facts'] if f['id'].startswith('total_amount__')]
receipt_amounts = {f['group_fields']['部门编号']: Decimal(f['value']) for f in amount_facts}

wb = openpyxl.load_workbook(target, read_only=True, data_only=False)
sheet_names = wb.sheetnames; ts = wb['1019']
period_headers = {f'{c}3': str(ts[f'{c}3'].value) for c in 'CDEGHI'}
metric_headers = {f'{c}4': str(ts[f'{c}4'].value) for c in 'CDEGHI'}
target_sections = []
dept_rows = defaultdict(list)
for row in ts.iter_rows(min_row=1, max_row=64, max_col=3):
    a, b, c = row
    if a.value and not b.value:
        target_sections.append({'coordinate': a.coordinate, 'label': str(a.value)})
    if isinstance(b.value, str) and b.value.startswith('FC'):
        dept_rows[b.value].append(b.row)
column_b_formula_count = sum(c.data_type == 'f' for (c,) in ts.iter_rows(min_col=2, max_col=2))
wb.close()
baseline = check_workbook_formulas(target)
recorded_partial = copy.deepcopy(E['receipts'][3]['data'])
projection = {'query': {'source': {'sheet': '1019', 'header_row': 4, 'data_range': 'A5:T64'},
                        'distinct': True, 'fields': ['B']}}
with patch('dsh_modules.spreadsheet_analysis.get_or_create_recalculated_workbook',
           return_value=(target, recorded_partial)) as recalc:
    projection_result = calculate_spreadsheet_analysis(target, projection)

plan_schema = next(t for t in SPREADSHEET_UPDATE_TOOL_DEFINITIONS
                   if t['function']['name'] == 'plan_spreadsheet_update')['function']['parameters']
after = {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in [target, source]}
result = {
    'session_id': SID, 'turn_id': E['turn_id'], 'request': prompt,
    'accepted_contract': accepted,
    'parser': {'entities': contract.entity_predicates, 'is_grouped': contract.is_grouped,
               'group_dimension': contract.group_dimension},
    'routing_reconstruction': {
        'flags': {k: getattr(route, k, None) for k in ['is_modify_intent', 'is_generate_intent',
                  'requires_execution', 'is_inspect_intent', 'is_guide_intent', 'is_office_intent']},
        'tools': names, 'requires_spreadsheet_analysis': requires_spreadsheet_analysis(E['query'], route),
        'update_tools_available': all(n in names for n in ['plan_spreadsheet_update', 'execute_spreadsheet_update']),
        'note': 'Reconstructed from saved accepted contract and current same-hash modules; historical tool schema was not persisted.',
        'admission_context': execution_contract_context(route.admission),
        'plan_tool_schema': plan_schema,
    },
    'prompt_and_recovery_reproduction': {
        'user_turn_demands_unavailable_bash': '你必须调用 `bash` 工具' in user_turn and 'bash' not in names,
        'admission_requires_conversation': '并在对话中直接回复结论' in execution_contract_context(route.admission),
        'controlled_candidate_guard_action': guard_action,
        'controlled_candidate_guard_hint': guard_hint,
        'note': 'Same code and recorded task contract. Guard candidate text is synthetic, not recovered historical output. No model calls.',
    },
    'reconciliation': {
        'evidence_id': analysis['evidence_id'], 'input_records': sum(months.values()),
        'paid_records': sum(counts.values()), 'months': dict(months),
        'by_department': [{'department_code': k, 'paid_amount_source_K': str(amounts[k]),
                           'paid_records': counts[k], 'target_rows': dept_rows[k]} for k in sorted(amounts)],
        'total_paid_amount_source_K': str(sum(amounts.values())),
        'matches_receipt': dict(amounts) == receipt_amounts,
    },
    'target': {'sheets': sheet_names, 'period_headers': period_headers, 'metric_headers': metric_headers,
               'section_labels': target_sections, 'column_B_formula_count': column_b_formula_count,
               'baseline_formula_check': baseline},
    'pre_read_gate_reproduction': {
        'query': projection, 'status': projection_result.status,
        'error': projection_result.error,
        'recorded_recalculation_engine': recorded_partial['engine_execution'],
        'recorded_total_errors': recorded_partial['total_errors'],
        'recalculation_receipt_replayed': recalc.call_count,
        'note': 'Injects the historical partial receipt, not a new engine run. Literal B projection is rejected before query compilation.',
    },
    'file_integrity': {'before': before, 'after': after, 'unchanged': before == after,
                       'matches_call': before == E['source_hashes'],
                       'requested_output_exists': (WS / '工数3Q预实管理_1023.xlsx').exists()},
    'model_calls': 0,
}
OUT.mkdir(parents=True, exist_ok=True)
(OUT / 'replay-result.json').write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n')
assert result['reconciliation']['matches_receipt']
assert result['file_integrity']['unchanged'] and result['file_integrity']['matches_call']
assert result['routing_reconstruction']['update_tools_available']
assert projection_result.error['code'] == 'calculation_unverified'
assert result['prompt_and_recovery_reproduction']['user_turn_demands_unavailable_bash']
assert '只读模式' in guard_hint
print(json.dumps({'routing': result['routing_reconstruction']['flags'], 'tools': names,
                  'requires_analysis': result['routing_reconstruction']['requires_spreadsheet_analysis'],
                  'source_matches_receipt': result['reconciliation']['matches_receipt'],
                  'target': result['target'], 'literal_projection_error': projection_result.error,
                  'unchanged': before == after}, ensure_ascii=False, indent=2))
