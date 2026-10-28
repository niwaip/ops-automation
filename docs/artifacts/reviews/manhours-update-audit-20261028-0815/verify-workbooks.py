"""Read-only independent original/result audit. Never save/recalculate workbooks."""
from collections import defaultdict, Counter
from datetime import datetime, date
from decimal import Decimal
import hashlib
import json
from pathlib import Path
import warnings
import openpyxl
from openpyxl.utils.datetime import to_excel
warnings.filterwarnings('ignore',category=UserWarning,module='openpyxl')
root=Path(__file__).resolve().parents[4]
w=root/'data/users/e7fce333-a8f4-4097-9a53-f0a4c729da46/workspace'
p=Path(__file__).parent
names=['工数3Q预实管理.xlsx','工数实绩明细_20261020134124.xlsx','工数3Q预实管理_1028.xlsx']
hashes={n:hashlib.sha256((w/n).read_bytes()).hexdigest() for n in names}
original=openpyxl.load_workbook(w/names[0],data_only=False)
result=openpyxl.load_workbook(w/names[2],data_only=False)
cache=openpyxl.load_workbook(w/names[2],data_only=True)
source=openpyxl.load_workbook(w/names[1],data_only=True).active
sums=defaultdict(Decimal);counts=Counter();unpaid_wbs=Counter()
for row in source.iter_rows(min_row=5,max_row=278,values_only=True):
 if str(row[1])!='202610': continue
 key=(str(row[2]),str(row[12]));sums[key]+=Decimal(str(row[15] or 0));counts[key]+=1
 if row[12]=='无偿':unpaid_wbs[str(row[8])[:2]]+=1
turn=json.loads((p/'turn.json').read_text());plan=json.loads((p/'run-evidence/plan.json').read_text());receipt=json.loads((p/'run-evidence/receipt.json').read_text())
assert hashes[names[2]]==receipt['output_hash']
assert all(hashes[k]==v for k,v in turn['source_hashes'].items())
checks=[]
for r in turn['receipts']:
 data=r.get('data')
 if not isinstance(data,dict) or data.get('kind')!='spreadsheet_analysis': continue
 charge=data['plan']['where']['and'][1]['value'];metric=data['plan']['metrics'][0]['id']
 assert {x['部门编号']:Decimal(x[metric]) for x in data['rows']}=={dep:v for (dep,c),v in sums.items() if c==charge}
for item in turn['receipts'][-2]['data']['preview']:
 charge='有偿' if item['section']=='有偿' else '无偿';expected=sums[item['key'],charge]
 actual=Decimal(str(result['1028'][item['coordinate']].value))
 assert abs(actual-expected)<Decimal('0.000000001')
 checks.append({'department':item['key'],'charge':charge,'cell':item['coordinate'],'source_K':str(expected),'delivered_K':str(actual)})

def equal(a,b):
 if isinstance(a,(datetime,date)): a=to_excel(a)
 if isinstance(b,(datetime,date)): b=to_excel(b)
 if isinstance(a,(int,float)) and isinstance(b,(int,float)):
  return abs(Decimal(str(a))-Decimal(str(b)))<Decimal('0.000000001')
 return a==b

def diffs(a,b):
 return [{'cell':c.coordinate,'old':str(c.value),'new':str(b[c.coordinate].value)} for row in a for c in row if not equal(c.value,b[c.coordinate].value)]
base=original['1019'];out=result['1028'];historic_diff=diffs(base,result['1019']);changes=diffs(base,out)
allowed={i['cell'] for i in checks};assert not historic_diff,historic_diff
assert all(c['cell'] in allowed for c in changes),changes
assert list(result.sheetnames)==['1019','1028']
assert {str(x) for x in base.merged_cells.ranges}=={str(x) for x in out.merged_cells.ranges}
formulas=[(ws.title,c.coordinate,c.value) for ws in result for row in ws for c in row if c.data_type=='f']
uncached=[(s,c) for s,c,f in formulas if cache[s][c].value is None]
errors=[(ws.title,c.coordinate,c.value) for ws in cache for row in ws for c in row if c.data_type=='e']
assert not uncached and not errors
paid=sum(v for (d,c),v in sums.items() if c=='有偿');unpaid=sum(v for (d,c),v in sums.items() if c=='无偿')
report={'hashes':hashes,'source_unchanged':True,'published_hash_matches_receipt':True,'original_sheets':original.sheetnames,'delivered_sheets':result.sheetnames,'original_sheet_semantic_changes':historic_diff,'new_sheet_changed_business_cells':changes,'writes_verified':checks,'write_count':len(checks),'actual_changed_value_count':len(changes),'plan_columns_unchanged':all(equal(base.cell(r,c).value,out.cell(r,c).value) for r in range(1,65) for c in range(7,15)),'other_months_unchanged':all(equal(base.cell(r,c).value,out.cell(r,c).value) for r in range(1,65) for c in [4,5]),'merged_ranges_preserved':True,'formula_count':len(formulas),'uncached_formulas':uncached,'formula_errors':errors,'source_paid_K':str(paid),'source_unpaid_K':str(unpaid),'source_charge_counts':{c:sum(n for (d,k),n in counts.items() if k==c) for c in ['有偿','无偿']},'unpaid_prefix_counts':dict(unpaid_wbs),'cached_totals':{c:cache['1028'][c].value for c in ['C19','C40','C42']},'preserved_unmatched':{'cell':'C37','department':'FC34022','value_K':str(out['C37'].value)},'copied_header_dates':{c:str(out[c].value) for c in ['C2','C23']},'receipt_reconciliation':turn['receipts'][-2]['data']['reconciliations']}
assert all(hashlib.sha256((w/n).read_bytes()).hexdigest()==h for n,h in hashes.items())
(p/'independent-verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
print(json.dumps({k:v for k,v in report.items() if k not in ['writes_verified','receipt_reconciliation']},ensure_ascii=False,indent=2))
