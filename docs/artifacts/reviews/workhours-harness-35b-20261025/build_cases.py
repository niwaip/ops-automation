"""Read-only independent oracle and reusable real-model E2E case definitions."""
import collections
import hashlib
import json
from decimal import Decimal as D
from pathlib import Path
import openpyxl

BASE = Path(__file__).resolve().parent
SOURCE = Path('/Users/chain/Desktop/工数实绩明细_20261020134124.xlsx')
wb = openpyxl.load_workbook(SOURCE, data_only=False)
ws = wb['工数实绩明细']
headers = [c.value for c in ws[4]]
records = [{'row': r, 'values': [c.value for c in ws[r]]} for r in range(5, 279)]
assert len(records) == 274 and len(headers) == 18
def dec(x): return D(str(x))
def totals(rs):
    return {'count': len(rs), 'hours': str(sum((dec(r['values'][14]) for r in rs), D(0))),
            'amount_K': str(sum((dec(r['values'][15]) for r in rs), D(0)))}
groups = collections.defaultdict(list)
employees = collections.defaultdict(list)
pivot = collections.defaultdict(list)
for r in records:
    v = r['values'];groups[(v[2],v[3])].append(r);employees[(v[4],v[5])].append(r);pivot[(v[2],v[12])].append(r)
filtered = [r for r in records if r['values'][2]=='FC45031' and r['values'][12]=='有偿' and dec(r['values'][14])>8]
filtered.sort(key=lambda r:(-dec(r['values'][14]),r['row']))
ranked = sorted([{'employee': k[0], 'name': k[1], **totals(v)} for k,v in employees.items()],
                key=lambda x:(-dec(x['hours']),x['employee']))
oracle = {'source_sha256':hashlib.sha256(SOURCE.read_bytes()).hexdigest(),'sheet':ws.title,
 'header_row':4,'data_rows':[5,278],'headers':headers,'summary':totals(records),
 'unique_employees':len(employees),'unique_WBS':len({r['values'][8] for r in records}),
 'periods':sorted({r['values'][1] for r in records}),
 'departments':[{'code':k[0],'name':k[1],**totals(v)} for k,v in sorted(groups.items())],
 'employee_top5':ranked[:5],
 'filter':{'totals':totals(filtered),'rows':[r['row'] for r in filtered],'records':filtered},
 'pivot':[{'department':k[0],'paid':k[1],**totals(v)} for k,v in sorted(pivot.items())],
 'missing_customer_rows':[r['row'] for r in records if r['values'][11] in ('',None)],
 'full_duplicate_rows':[],
 'lookups':{eid:totals([r for r in records if r['values'][4]==eid]) for eid in ['S22009','S17028','E2E_NOT_FOUND']},
 'missing_daily_dates':True,'records':records}
assert len({tuple(r['values']) for r in records})==274
(BASE/'oracle.json').write_text(json.dumps(oracle,ensure_ascii=False,indent=2))

end=' 只在聊天回复，不生成文件，不联网；无法核算时明确说明或向我澄清，不要猜测。'
prompts=[
('K01','结构、范围与基础汇总','读取附件工数实绩明细，确认真实表头、业务明细范围、业务记录数、去重员工数、去重WBS数、工数合计、金额合计和工数年月。标题和末尾查询条件不是业务数据，单位沿用原列。'+end),
('K02','组合筛选、排序与行定位','筛选部门编号FC45031、有偿无偿为有偿、工数(小时)严格大于8的全部记录。给出匹配条数、总工数、金额(K)合计，并按工数降序、同值按原Excel行号升序列出前5条的原行号、员工编号、WBS编号、工数和金额。'+end),
('K03','去重实体聚合与Top N','按员工编号汇总全部业务明细，给出总工数最高的前5位员工（同值按员工编号升序，固定5位），列出员工编号、姓名、记录数、工数合计、金额(K)合计。不要把员工单条记录排名当成员工汇总排名。'+end),
('K04','分组、交叉汇总与占比','按部门编号及名称、有偿无偿交叉汇总工数和金额(K)，列出全部部门的有偿/无偿工数、部门总工数、部门金额合计、有偿工数占部门工数比例，最后核对部门汇总与全表合计一致。没有匹配记录时可记0，但必须实际核对。'+end),
('K05','空值及重复审计','检查所有18列完全相同的重复业务行、客户名称为空的业务行，分别给条数和全部原Excel行号。客户为空只标记缺失，不推断数据错误；员工编号重复是正常多项目记录，不能按员工编号去重。'+end),
('K06','日期粒度、缺失数据与交互','帮我按每天统计2026年10月工数，如果没有每日日期请明确指出，不能把导出日期当每条记录日期，也不要编造日数据。'+end),
('K07','精确匹配、批量查询与零结果','分别查询员工编号S22009、S17028和E2E_NOT_FOUND，给出每人的记录条数、总工数及金额(K)合计。编号精确匹配，找不到时明确未找到，不要编造姓名或记录。'+end),
('K08','CSV导出、字段与编码','将部门编号FC45031、有偿无偿为有偿、工数(小时)严格大于8的全部业务明细，按工数降序、同值按原行号升序导出为新的UTF-8 CSV文件 E2E_workhours_filtered.csv。首列新增原Excel行号，其后保留原18列及字段值，保留员工编号、年月为文本，不要覆盖附件。请提供下载，不联网。'),
('K09','副本增删改、公式及样式','生成新的Excel E2E_workhours_edit.xlsx，原附件不改。新表“明细”第1行是原18列表头，后面是原业务明细，去掉原标题与查询条件。仅做这些修改：将原Excel第5行业务记录工数改为10，金额原值不变；删除原第6行业务记录；在末尾追加原第5行的副本，但员工编号改为E2E_TEST、员工姓名改为测试记录、工数改为8、金额改为0。追加第19列“工数天数”，对每条记录用公式工数除以8（仅本测试按8小时一天），保留可重算公式。冻结首行、启用全表筛选，工数两位小数、金额五位小数，工数大于8的单元格设条件格式。只生成这个副本并提供下载，不联网。'),
('K10','汇总文件、公式与图表','基于原附件业务明细生成新的Excel E2E_workhours_report.xlsx。一个“部门汇总”表，列为部门编号、部门名称、工数合计、金额(K)合计、有偿工数、无偿工数、有偿占比；每部门一行，有偿占比用公式有偿工数除以工数合计；末尾合计行工数和金额用SUM公式，占比按整体有偿工数除以整体工数计算，不要平均各部门比例。加一张按部门名称展示工数合计的柱形图，数值单位沿用原表，冻结表头，提供下载。保留原附件，不联网。')]
cases={'model':'qwen36-35b-a3b','mode':'chat','source':str(SOURCE),'policy':{'allowed':['CLARIFICATION','SAFE_RETRY','PARTIAL_VERIFIED','DELIVERED_VERIFIED'],'blocking':['WRONG_DATA','FALSE_SUCCESS','SOURCE_MUTATION'],'max_followups_per_case':1,'completion_separate_from_safe_handling':True},'cases':[{'id':i,'capability':c,'prompt':p} for i,c,p in prompts]}
(BASE/'cases.json').write_text(json.dumps(cases,ensure_ascii=False,indent=2))
print(json.dumps({k:oracle[k] for k in ['summary','unique_employees','unique_WBS','employee_top5']},ensure_ascii=False));print('filtered',oracle['filter']['totals'],'missing_customers',len(oracle['missing_customer_rows']))
