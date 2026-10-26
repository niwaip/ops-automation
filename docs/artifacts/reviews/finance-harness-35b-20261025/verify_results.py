"""Independently verify the recorded live runs; does not call a model."""
import hashlib
import json
import re
from decimal import Decimal as D
from pathlib import Path

import openpyxl

BASE = Path(__file__).resolve().parent
SOURCE = Path('/Users/chain/Desktop/AI能力测试_财务报表.xlsx')
SOURCE_HASH = '9d71ca8cc6172add7e558ac6a906f677c784fcc5c9d3444cbb17de2772fd992f'
assert hashlib.sha256(SOURCE.read_bytes()).hexdigest() == SOURCE_HASH
wb = openpyxl.load_workbook(SOURCE, data_only=False)
ws = wb['月度经营']


def number(sheet, cell):
    return D(str(wb[sheet][cell].value))


def close(actual, expected, tolerance='0.000001'):
    assert abs(D(str(actual)) - expected) <= D(tolerance), (actual, str(expected))


months = {}
for r in range(4, 16):
    revenue = number('月度经营', f'B{r}')
    cost_formula = ws[f'C{r}'].value
    match = re.fullmatch(rf'=B{r}\*\(1-([0-9.]+)\)', cost_formula)
    assert match, cost_formula
    gross = revenue * D(match[1])
    expenses = D('0')
    for col in ('F', 'G'):
        formula = ws[f'{col}{r}'].value
        rate = re.fullmatch(rf'=B{r}\*([0-9.]+)', formula)
        assert rate, formula
        expenses += revenue * D(rate[1])
    operating = gross - expenses - number('月度经营', f'H{r}')
    pretax = operating + number('月度经营', f'K{r}')
    net = pretax - max(D('0'), pretax * number('参数', 'B8'))
    months[ws[f'A{r}'].value] = dict(revenue=revenue, gross=gross, operating=operating, net=net, row=r)

revenue = sum(m['revenue'] for m in months.values())
gross = sum(m['gross'] for m in months.values())
operating = sum(m['operating'] for m in months.values())
net = sum(m['net'] for m in months.values())
golden = dict(revenue=revenue, gross_profit=gross, operating_profit=operating, net_profit=net,
              gross_margin=gross/revenue, net_margin=net/revenue)


def evidence(n):
    obj = json.loads((BASE / f'F{n:02}-evidence.json').read_text())
    assert obj['source_hashes']['AI能力测试_财务报表.xlsx'] == SOURCE_HASH
    return obj


f2 = evidence(2)
facts = [f for r in f2['receipts'] for f in (r.get('data') or {}).get('facts', [])]
for fact in facts:
    close(fact['value'], golden[fact['id']])
assert next(f for f in facts if f['id'] == 'gross_margin')['unit'] == 'ratio'
assert not f2['verified']  # Correct values were rejected by the runtime gate.

f3 = evidence(3)
comparison = next(r['data'] for r in f3['receipts'] if (r.get('data') or {}).get('kind') == 'column_comparison')
assert comparison['scope']['coverage_verified']
assert comparison['scope']['row_count'] == 12
for metric in comparison['evaluated']:
    assert len(metric['rows']) == 12
    for row in metric['rows']:
        m = months[row['group']]
        is_revenue = metric['label'] == '营业收入'
        actual = m['revenue' if is_revenue else 'net']
        budget = number('预算对比', f"{'B' if is_revenue else 'E'}{m['row']}")
        for key, expected in dict(left=actual, right=budget, difference=actual-budget,
                                  relative_difference=(actual-budget)/abs(budget)).items():
            close(row['exact'][key], expected)
extremes = {r['label']: r['matches'][0] for r in comparison['results']}
assert extremes['营业收入']['group'] == '12月'
assert extremes['净利润']['group'] == '7月'

f4 = evidence(4)
missing = next(r['data'] for r in f4['receipts'] if (r.get('data') or {}).get('rows', [{}])[0].get('凭证号') == 'V2026-0025')
assert missing['rows'][0]['部门'] is None
assert missing['facts'][0]['source_cells'][0] == '交易明细!A28'
duplicate = next(r['data'] for r in f4['receipts'] if any(f.get('group_key') == 'V2026-0009' for f in (r.get('data') or {}).get('facts', [])))
assert duplicate['facts'][0]['value'] == 2
assert duplicate['facts'][0]['source_cells'] == ['交易明细!A12', '交易明细!A13']

assets = sum(number('资产负债表', f'B{r}') for r in (4, 5, 6, 7, 9, 10))
liabilities_equity = sum(number('资产负债表', f'B{r}') for r in (12, 13, 14, 15, 17)) + number('资产负债表', 'C18') + net
cfo = net + number('现金流量表', 'B5') - (number('资产负债表', 'B5')-number('资产负债表', 'C5')) - (number('资产负债表', 'B6')-number('资产负债表', 'C6')) + (number('资产负债表', 'B13')-number('资产负债表', 'C13'))
debt_delta = sum(number('资产负债表', f'B{r}')-number('资产负债表', f'C{r}') for r in (12, 15))
cash_end = number('资产负债表', 'C4') + cfo - number('参数', 'B11') + debt_delta + number('现金流量表', 'B13')
checks = next(r['data']['checks'] for r in evidence(5)['receipts'] if (r.get('data') or {}).get('checks'))
for check, lhs, rhs in zip(checks, [assets, cash_end], [liabilities_equity, number('资产负债表', 'B4')]):
    close(check['left'], lhs)
    close(check['right'], rhs)
    close(check['residual'], lhs-rhs)
    assert check['passed'] is False

output = BASE / 'E2E_finance_35B.xlsx'
ow = openpyxl.load_workbook(output, data_only=True)
expected_output = {'营业收入': revenue, '毛利': gross, '营业利润': operating, '净利润': net}
seen = set()
for row in ow.active.values:
    if row[0] in expected_output:
        assert isinstance(row[1], (float, int))
        close(row[1], expected_output[row[0]])
        assert row[2] == '千元'
        seen.add(row[0])
assert seen == set(expected_output)
assert output.read_bytes() == Path('/Users/chain/Downloads/E2E_finance_35B (1).xlsx').read_bytes()
assert len(json.loads((BASE/'history-after-refresh.json').read_text())['messages']) == 8
for n in range(1, 7):
    stream = json.loads((BASE / f'F{n:02}-stream.json').read_text())
    assert any(e['type'] == 'done' for e in stream)

summary = {'source_hash': SOURCE_HASH, 'method': 'Decimal arithmetic from original literals and formula rate constants; no cached formula values used',
           'golden': golden, 'bs_delta': assets-liabilities_equity,
           'cash_delta': cash_end-number('资产负债表', 'B4'), 'budget_monthly_rows_verified': 24,
           'output_metrics_verified': 4, 'refresh_history_messages': 8, 'all_six_streams_done': True,
           'source_unchanged': hashlib.sha256(SOURCE.read_bytes()).hexdigest() == SOURCE_HASH}
(BASE / 'verification.json').write_text(json.dumps(summary, ensure_ascii=False, indent=2, default=str))
print(json.dumps(summary, ensure_ascii=False, default=str))
