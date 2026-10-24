"""Test-data transformation only. Reference-file labels are not runtime configuration."""
import copy
import hashlib
import json
import shutil
from pathlib import Path
from decimal import Decimal
import openpyxl
from openpyxl.formula.translate import Translator
from openpyxl.utils.cell import range_boundaries,get_column_letter

workspace=Path('data/users/e7fce333-a8f4-4097-9a53-f0a4c729da46/workspace')
root=workspace/'.validation/generic-binding-20261023'
original=root/'original';changed=root/'changed'
original.mkdir(parents=True,exist_ok=True);changed.mkdir(parents=True,exist_ok=True)
inputs=['工数3Q预实管理.xlsx','工数实绩明细_20261020134124.xlsx']
for name in inputs:
    shutil.copy2(workspace/name,original/name)
source=openpyxl.load_workbook(workspace/inputs[1])
source_sheet=source.active
codes={str(source_sheet.cell(r,3).value) for r in range(5,source_sheet.max_row+1) if str(source_sheet.cell(r,3).value).startswith('FC')}
target=openpyxl.load_workbook(workspace/inputs[0])
codes.update(str(c.value) for row in target.active for c in row if isinstance(c.value,str) and c.value.startswith('FC'))
mixed=['003',0,'研发/甲','R-9','组织_乙','Ω组','七号','LAB::8','Z-零','九部','余项']
code_map={code:(mixed[i] if i<len(mixed) else '组织/'+str(i)) for i,code in enumerate(sorted(codes))}
periods={202610:202811,202611:202812,202612:202901}
labels={'有偿':'收费业务','*P/*Y/EX（不含*L/*B）':'内部业务','预测':'本期估算','计划':'基准预算'}
changed_target=openpyxl.Workbook();changed_target.remove(changed_target.active)
old=target.active
for name in ['快照-2028-10-01','快照-2028-11-01']:
    ws=changed_target.create_sheet(name)
    ws['A1']='金额单位：万元'
    for row in old:
        for cell in row:
            dr=18 if cell.row==1 else 30
            new=ws.cell(cell.row+dr,cell.column+4)
            value=cell.value
            if cell.data_type=='f':
                value=Translator(value,origin=cell.coordinate).translate_formula(new.coordinate)
            elif isinstance(value,str):
                value=code_map.get(value,labels.get(value,value))
            elif value in periods if isinstance(value,(int,float)) else False:
                value=periods[value]
            elif isinstance(value,(int,float)) and 5<=cell.row<=42 and 3<=cell.column<=14:
                value=float(Decimal(str(value))/Decimal(10))
            new.value=value
            if cell.has_style:
                for attribute in ('font','fill','border','alignment','protection'):
                    setattr(new,attribute,copy.copy(getattr(cell,attribute)))
                new.number_format=cell.number_format
    for merged in old.merged_cells.ranges:
        a,r1,b,r2=range_boundaries(str(merged));delta=18 if r1==1 else 30
        ws.merge_cells(start_row=r1+delta,start_column=a+4,end_row=r2+delta,end_column=b+4)
changed_target.save(changed/'滚动管理.xlsx');changed_target.close();target.close()
# Rename fields and data, then shift the source table as well.
field_map={'工数年月':'归属期','部门编号':'组织键','有偿无偿':'结算性质','金额(K)':'收入(K)'}
category={'有偿':'收费','无偿':'内部'}
changed_source=openpyxl.Workbook();out=changed_source.active;out.title='明细记录'
for row in source_sheet:
    for cell in row:
        value=cell.value
        if cell.row==4:
            value=field_map.get(value,value)
        elif cell.column==3 and value is not None:
            value=code_map.get(str(value),value)
        elif cell.column==2 and str(value)=='202610':
            value='202811'
        elif cell.column==13:
            value=category.get(value,value)
        out.cell(cell.row+8,cell.column+3,value)
changed_source.save(changed/'收入记录.xlsx');changed_source.close();source.close()
manifest={'reference_inputs':inputs,'changed_inputs':['滚动管理.xlsx','收入记录.xlsx'],
          'key_mapping':code_map,'source_header_row':12,'target_key_column':'F','target_value_column':'G',
          'target_period_row':33,'target_label_row':34,'section_row':19,'unit_cell':'A1','source_unit':'K','target_unit':'万元',
          'period':'202811','snapshot':'快照-2028-11-01','new_sheet':'复核版',
          'expected_keys':7,'original_input_hashes':{n:hashlib.sha256((workspace/n).read_bytes()).hexdigest() for n in inputs}}
(root/'fixture-manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2))
print(json.dumps(manifest,ensure_ascii=False,indent=2))
