"""Read back model deliverables against independently summed transformed inputs."""
import hashlib,json,sys
from decimal import Decimal
from pathlib import Path
import openpyxl
from dsh_modules.spreadsheet_update.diff import _cells_equivalent

root=Path(sys.argv[1]).resolve()
manifest=json.loads((root.parent/'fixture-manifest.json').read_text())
receipts=[json.loads(path.read_text()) for path in (root/'.dsh/spreadsheet-runs').glob('*/receipt.json')]
assert len(receipts)==1,receipts
receipt=receipts[0]
output=root/Path(receipt['output_path']).name
assert hashlib.sha256(output.read_bytes()).hexdigest()==receipt['output_hash']
wb=openpyxl.load_workbook(root/'收入记录.xlsx',data_only=True)
ws=wb.active
headers={cell.value:cell.column for cell in ws[manifest['source_header_row']] if cell.value is not None}
sums={};count=0
for row in ws.iter_rows(min_row=manifest['source_header_row']+1,values_only=True):
    if str(row[headers['归属期']-1])==manifest['period'] and row[headers['结算性质']-1]=='收费':
        key=str(row[headers['组织键']-1]);sums[key]=sums.get(key,Decimal(0))+Decimal(str(row[headers['收入(K)']-1]));count+=1
wb.close()
reference=root.parent/'original'/'工数3Q预实管理.xlsx'
wb=openpyxl.load_workbook(reference)
expected_coordinates={}
for row in wb.active:
    for cell in row:
        if isinstance(cell.value,str) and cell.value in manifest['key_mapping']:
            key=str(manifest['key_mapping'][cell.value])
            if key in sums:
                expected_coordinates.setdefault(key,'G'+str(cell.row+30))
wb.close()
original=openpyxl.load_workbook(root/'滚动管理.xlsx')
result=openpyxl.load_workbook(output)
cached=openpyxl.load_workbook(output,data_only=True)
assert result.sheetnames==original.sheetnames+[manifest['new_sheet']],result.sheetnames
assert len(sums)==manifest['expected_keys']
for key,value in sums.items():
    actual=result[manifest['new_sheet']][expected_coordinates[key]].value
    assert abs(Decimal(str(actual))-value/10)<Decimal('0.000000001'),(key,actual,value)
formula_count=0;changed=[]
for name in result.sheetnames:
    before=original[manifest['snapshot']] if name==manifest['new_sheet'] else original[name]
    after=result[name]
    assert set(before.merged_cells.ranges)==set(after.merged_cells.ranges)
    for row in before:
        for cell in row:
            current=after[cell.coordinate]
            if cell.data_type=='f':
                formula_count+=1
                assert current.data_type=='f'
                assert cached[name][cell.coordinate].value is not None
            if not _cells_equivalent(cell.value,current.value,original.epoch):
                assert name==manifest['new_sheet'] and cell.coordinate in expected_coordinates.values(),(name,cell.coordinate,cell.value,current.value)
                changed.append(cell.coordinate)
assert receipt['formula_health']['total_errors']==0
assert receipt['formula_health']['uncalculated_formulas']==0
assert formula_count==receipt['formula_health']['cached_formulas']
report={'status':'passed','record_count':count,'key_count':len(sums),'source_unit':'K','target_unit':'万元',
        'source_total':str(sum(sums.values())),'written_total':str(sum(sums.values())/10),
        'coordinates':expected_coordinates,'changed_cells':changed,'formula_count':formula_count,
        'formula_errors':0,'missing_caches':0,'history_budget_other_periods_and_regions_preserved':True,
        'output_file':output.name,'output_hash':receipt['output_hash'],'warnings':receipt['warnings'],
        'input_hashes':{p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in [root/'收入记录.xlsx',root/'滚动管理.xlsx']}}
original.close();result.close();cached.close()
(root/'independent-verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
print(json.dumps(report,ensure_ascii=False,indent=2))
