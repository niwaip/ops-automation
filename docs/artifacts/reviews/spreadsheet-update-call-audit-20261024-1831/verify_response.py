"""Read-only comparison of numeric claims in the actual final response to evidence."""
import json
import re
from decimal import Decimal
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
response = (ROOT / 'final-response.md').read_text()
receipt = json.loads((ROOT / 'run_fdc43e9cb329/receipt.json').read_text())
base = load_workbook(ROOT / '工数3Q预实管理.xlsx', data_only=True)
delivered = load_workbook(ROOT / '工数3Q预实管理_1023_14.xlsx', data_only=True)
checks = []
try:
    for section in ('有偿', '无偿'):
        block = response.split('### ' + section + '区块回填结果')[1].split('### ')[0]
        for line in block.splitlines():
            if not re.match(r'\| FC\d+', line):
                continue
            fields = [part.strip() for part in line.strip('|').split('|')]
            code, _, claimed_amount, claimed_old = fields
            op = next(op for op in receipt['applied_operations'] if op.get('region_id') == section and op.get('department_code') == code)
            coord = op['coordinate']
            actual_old = Decimal(str(base['1019'][coord].value))
            actual_new = Decimal(str(delivered['1023'][coord].value))
            checks.append({'section': section, 'key': code, 'coordinate': coord,
                           'claimed_old': claimed_old, 'actual_old': str(actual_old),
                           'old_claim_matches_2dp': abs(Decimal(claimed_old) - actual_old) <= Decimal('.00501'),
                           'claimed_new': claimed_amount, 'actual_new': str(actual_new),
                           'new_claim_matches_2dp': abs(Decimal(claimed_amount) - actual_new) <= Decimal('.00501')})
    old_claim = re.search(r'`FC34022` 原预测值为 ([\d.]+) K', response).group(1)
    result = {'method': 'Compare final Markdown table amounts and old forecasts to actual workbook cells with 2dp presentation tolerance',
              'table_checks': checks,
              'incorrect_table_old_claim_count': sum(not check['old_claim_matches_2dp'] for check in checks),
              'all_table_new_amounts_match': all(check['new_claim_matches_2dp'] for check in checks),
              'unpaid_warning_old_claim': old_claim,
              'actual_unpaid_warning_old_value': str(base['1019']['C37'].value),
              'unpaid_warning_claim_matches': abs(Decimal(old_claim) - Decimal(str(base['1019']['C37'].value))) <= Decimal('.00501')}
finally:
    base.close()
    delivered.close()
(ROOT / 'response-verification.json').write_text(json.dumps(result, ensure_ascii=False, indent=2))
print(json.dumps({key: value for key, value in result.items() if key != 'table_checks'}, ensure_ascii=False, indent=2))
