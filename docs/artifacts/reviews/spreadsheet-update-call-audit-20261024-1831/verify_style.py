"""Read-only style storage comparison; differences are not automatically visual defects."""
import json
from copy import copy
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
base = load_workbook(ROOT / '工数3Q预实管理.xlsx')
output = load_workbook(ROOT / '工数3Q预实管理_1023_14.xlsx')
result = {}
try:
    for name in output.sheetnames:
        changes = {key: [] for key in ('font', 'fill', 'border', 'alignment', 'number_format', 'protection')}
        for row in base['1019']:
            for cell in row:
                if cell.value is None:
                    continue
                other = output[name][cell.coordinate]
                for key in changes:
                    before, after = copy(getattr(cell, key)), copy(getattr(other, key))
                    if before != after:
                        changes[key].append({'cell': cell.coordinate, 'before': str(before), 'after': str(after)})
        result[name] = {
            'method': 'Compare copied openpyxl style objects, not StyleProxy identity; storage differences are not automatically visual regressions',
            'changed_nonempty_cell_style_components': {key: {'count': len(values), 'sample': values[:3]} for key, values in changes.items()},
            'merged_ranges_preserved': set(map(str, base['1019'].merged_cells.ranges)) == set(map(str, output[name].merged_cells.ranges)),
            'source_dimensions': base['1019'].calculate_dimension(),
            'output_dimensions': output[name].calculate_dimension(),
        }
finally:
    base.close()
    output.close()
(ROOT / 'style-check.json').write_text(json.dumps(result, ensure_ascii=False, indent=2))
for name, data in result.items():
    print(name, {key: value['count'] for key, value in data['changed_nonempty_cell_style_components'].items()})
