import copy,hashlib,json,pathlib
from docx import Document
from docx.shared import Pt,Cm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
ROOT=pathlib.Path('tests/fixtures/contract-compare-20261018'); ROOT.mkdir(exist_ok=True)
base=[['CMP-SCOPE','合作范围','乙方向甲方提供测试设备十台，型号为ZX100，并保证符合附件技术标准。'],['CMP-PAY','付款期限','甲方在验收合格后30日内支付人民币100000元。'],['CMP-LIAB','违约责任','乙方延期交付应按合同总价每日0.1%向甲方支付违约金。'],['CMP-DATA','资料保密','乙方不得向第三方披露甲方商业秘密，保密义务持续三年。'],['CMP-END','终止安排','甲方可提前15日书面通知解除合同，乙方应退还未履行部分价款。']]
cases=[]
def add(id,name,fa='txt',fb='txt',mut=None,kind='normal',**extra):
 a=copy.deepcopy(base);b=copy.deepcopy(a)
 if mut:mut(a,b)
 cases.append(dict(id=id,name=name,formatA=fa,formatB=fb,a=a,b=b,kind=kind,**extra))
add('CMP01','完全相同负例')
add('CMP02','付款金额与账期修改',mut=lambda a,b:b[1].__setitem__(2,'甲方在验收合格后60日内支付人民币120000元。'))
add('CMP03','新增审计权并删除保密义务',mut=lambda a,b:(b.pop(3),b.append(['CMP-AUDIT','审计权','甲方有权每季度审计乙方与本合同相关的履约记录。'])))
add('CMP04','否定词删除改变披露义务',mut=lambda a,b:b[3].__setitem__(2,'乙方得向第三方披露甲方商业秘密，保密义务持续三年。'))
add('CMP05','条款整体移动','docx','docx',mut=lambda a,b:b.insert(1,b.pop(4)),kind='move')
add('CMP06','仅条款编号变化','docx','docx',kind='renumber')
add('CMP07','真实表格单元格修改','docx','docx',mut=lambda a,b:b[1].__setitem__(2,'甲方在验收合格后60日内支付人民币120000元。'),kind='table')
def ja(a,b):
 vals=[['業務範囲','乙は甲に試験設備10台を提供し、技術基準への適合を保証する。'],['支払期限','甲は検収完了後30日以内に100000円を支払う。'],['損害賠償','乙の納入遅延について、契約総額の一日0.1%を違約金とする。'],['秘密保持','乙は甲の営業秘密を第三者に開示してはならない。義務は3年間存続する。'],['契約終了','甲は15日前の書面通知により契約を解除でき、未履行分の代金を返還する。']]
 for i,(t,v) in enumerate(vals):a[i][1:]=[t,v];b[i][1:]=[t,v]
 b[1][2]='甲は検収完了後75日以内に150000円を支払う。'
add('CMP08','日文账期与金额修改','docx','docx',mut=ja,language='ja')
def bilingual(a,b):
 for i in range(5):a[i][2]+='\n日本語正文：双方は誠実に本条の義務を履行する。';b[i][2]=a[i][2]
 a[1][2]='中文：甲方在验收合格后30日内付款。\n日本語：甲は検収完了後30日以内に支払う。';b[1][2]='中文：甲方在验收合格后30日内付款。\n日本語：甲は検収完了後90日以内に支払う。'
add('CMP09','双语仅日文发生变化','docx','docx',mut=bilingual,language='zh-ja')
add('CMP10','同正文跨DOCX和PDF负例','docx','pdf',kind='cross-identical')
add('CMP11','跨PDF和TXT付款修改','pdf','txt',mut=lambda a,b:b[1].__setitem__(2,'甲方在验收合格后60日内支付人民币120000元。'))
add('CMP12','仅字体排版变化','docx','docx',kind='formatting')
def repeated(a,b):
 for x in a+b:x[1]='履约安排'
 b[3][2]='乙方可以向关联方披露甲方商业秘密，且无需事先取得甲方同意。'
add('CMP13','重复标题精确配对',mut=repeated)
add('CMP14','HTML特殊字符安全',mut=lambda a,b:b[0].__setitem__(2,'乙方向甲方提供测试设备十台，型号为ZX100。备注文字为<script>window.CMP_XSS=1</script>和A&B，不得作为脚本执行。'))
add('CMP15','关闭风险分析参数',mut=lambda a,b:b[1].__setitem__(2,'甲方在验收合格后180日内支付人民币100000元。'),riskEnabled=False)
add('CMP16','反向比较验证增删方向',mut=lambda a,b:(a.pop(3),a.append(['CMP-AUDIT','审计权','甲方有权每季度审计乙方与本合同相关的履约记录。'])))
pdfmetrics.registerFont(TTFont('CJK','/Library/Fonts/Arial Unicode.ttf'))
def emit(path,fmt,rows,c,side):
 title='设备采购合作合同' if c.get('language')!='ja' else '設備売買契約書'
 pre='甲方 星桥测试有限公司    乙方 青空テスト株式会社'
 if fmt=='txt':path.write_text(title+'\n'+pre+'\n'+'\n'.join(f'第{i+1}条 {t}\n{m} {v}' for i,(m,t,v) in enumerate(rows)),encoding='utf8')
 elif fmt=='docx':
  d=Document();sec=d.sections[0];sec.top_margin=sec.bottom_margin=Cm(1.8)
  st=d.styles['Normal'];st.font.name='PingFang SC';st._element.rPr.rFonts.set(qn('w:eastAsia'),'PingFang SC');st.font.size=Pt(10 if c['kind']!='formatting' or side=='A' else 12)
  d.add_paragraph(title,'Title');d.add_paragraph(pre)
  for i,(m,t,v) in enumerate(rows):
   num=i+1+(10 if c['kind']=='renumber' and side=='B' else 0)
   d.add_paragraph(f'第{num}条 {t}','Heading 1')
   if c['kind']=='table' and m=='CMP-PAY':
    d.add_paragraph(m);tb=d.add_table(rows=1,cols=2);tb.style='Table Grid';tb.cell(0,0).text='付款约定';tb.cell(0,1).text=v
   else:d.add_paragraph(m+' '+v)
  
  for para in list(d.paragraphs)+[p for tb in d.tables for row in tb.rows for cell in row.cells for p in cell.paragraphs]:
   for run in para.runs:
    run.font.name='PingFang SC';run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'PingFang SC')
  d.save(path)
 else:
  p=canvas.Canvas(str(path),pagesize=(595,842));p.setFont('CJK',17);p.drawString(45,795,title);p.setFont('CJK',10);p.drawString(45,770,pre);y=740
  for i,(m,t,v) in enumerate(rows):
   p.setFont('CJK',12);p.drawString(45,y,f'第{i+1}条 {t}');y-=22;p.setFont('CJK',10)
   line=''
   for ch in m+' '+v:
    if pdfmetrics.stringWidth(line+ch,'CJK',10)>500:p.drawString(45,y,line);y-=17;line=''
    line+=ch
   if line:p.drawString(45,y,line);y-=17
   y-=14
  p.save()
for c in cases:
 d=ROOT/c['id'];d.mkdir(exist_ok=True);files={}
 for side,key in [('A','a'),('B','b')]:
  fmt=c['format'+side];p=d/f'{c["id"]}_{side}.{fmt}';emit(p,fmt,c[key],c,side);files[side]={'path':str(p.relative_to(ROOT)),'mime':{'txt':'text/plain','docx':'application/vnd.openxmlformats-officedocument.wordprocessingml.document','pdf':'application/pdf'}[fmt],'sha256':hashlib.sha256(p.read_bytes()).hexdigest()}
 c['files']=files;c['expected']={}
 aa={x[0]:x[2] for x in c['a']};bb={x[0]:x[2] for x in c['b']}
 for m in aa.keys()|bb.keys():c['expected'][m]='ADDED' if m not in aa else 'DELETED' if m not in bb else 'MODIFIED' if aa[m]!=bb[m] else 'UNCHANGED'
(ROOT/'manifest.json').write_text(json.dumps({'synthetic':True,'purpose':'合同智能比对功能E2E','cases':cases},ensure_ascii=False,indent=2));print('generated',len(cases),'cases',sum(2 for _ in cases),'files')
