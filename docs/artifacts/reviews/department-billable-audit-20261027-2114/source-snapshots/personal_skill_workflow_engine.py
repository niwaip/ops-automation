"""Portable, read-only interpreter for saved spreadsheet aggregation workflows.

This exact module is shipped as scripts/run.py. No generated code is evaluated.
"""

import argparse
import hashlib
import json
import re
import time
from collections import Counter, defaultdict
from decimal import Decimal, InvalidOperation, localcontext
from pathlib import Path

import openpyxl


ENGINE = 'spreadsheet-group-sum-v1'


def digest(value):
    return hashlib.sha256(json.dumps(value, sort_keys=True, ensure_ascii=False,
                                     separators=(',', ':')).encode()).hexdigest()


def validate_workflow(plan):
    allowed = {'schema_version', 'engine', 'group_by', 'measure', 'filters',
               'instructions', 'steps', 'verification'}
    if not isinstance(plan, dict) or set(plan) != allowed:
        raise ValueError('流程字段不完整或包含未知操作')
    if plan['schema_version'] != 1 or plan['engine'] != ENGINE:
        raise ValueError('不支持的核算流程版本')
    for name in ('group_by', 'measure'):
        if not isinstance(plan[name], str) or not plan[name].strip():
            raise ValueError('缺少明确的分组或求和字段')
    if plan['verification'] != {'source_hash': True, 'row_accounting': True, 'group_total': True}:
        raise ValueError('不能关闭核算验收')
    for name in ('instructions', 'steps'):
        if not isinstance(plan[name], list) or not plan[name] or not all(isinstance(v, str) and v.strip() for v in plan[name]):
            raise ValueError('缺少业务规则或执行步骤')
    if not isinstance(plan['filters'], list) or len(plan['filters']) > 16:
        raise ValueError('非法筛选列表')
    for f in plan['filters']:
        if (not isinstance(f, dict) or set(f) != {'field', 'op', 'values'}
                or not isinstance(f['field'], str) or not f['field'].strip()
                or f['op'] not in {'exclude_prefix', 'in', 'not_in'}
                or not isinstance(f['values'], list) or not f['values']
                or not all(isinstance(v, str) and v.strip() for v in f['values'])):
            raise ValueError('筛选规则不完整或操作不受支持')
    return plan


def field_key(value):
    return re.sub(r'\s+', '', str(value or '')).casefold()


def bind_field(headers, field):
    key = field_key(field)
    exact = [i for i, h in enumerate(headers) if field_key(h) == key]
    matches = exact or [i for i, h in enumerate(headers)
                        if re.sub(r'[（(][^()（）]+[）)]$', '', field_key(h)) == key]
    if len(matches) != 1:
        raise ValueError(f'字段缺失或不唯一：{field}')
    return matches[0]


def locate_table(workbook, fields):
    candidates = []
    for ws in workbook:
        if ws.max_row * ws.max_column > 250000:
            raise ValueError('工作表超过可验证核算预算')
        for row in ws.iter_rows(max_row=min(ws.max_row, 50), values_only=True):
            try:
                mapping = {f: bind_field(row, f) for f in fields}
                candidates.append((ws, row, mapping))
            except ValueError:
                continue
    if len(candidates) != 1:
        raise ValueError('无法按必需字段定位唯一表头与工作表；请检查缺失字段、重复字段或多个候选表')
    ws, headers, mapping = candidates[0]
    # Locate the exact header without relying on worksheet dimension metadata.
    header_row = next(r[0].row for r in ws.iter_rows(max_row=min(ws.max_row, 50))
                      if tuple(c.value for c in r) == tuple(headers))
    return ws, header_row, headers, mapping


def matches(value, condition):
    text = str(value if value is not None else '').strip().casefold()
    values = [v.strip().casefold() for v in condition['values']]
    if condition['op'] == 'exclude_prefix':
        return not text.startswith(tuple(values))
    present = text in values
    return present if condition['op'] == 'in' else not present


def decimal_amount(value, coordinate):
    if value is None or isinstance(value, bool):
        raise ValueError(f'金额缺失或不是数值：{coordinate}')
    try:
        amount = Decimal(str(value))
    except InvalidOperation as error:
        raise ValueError(f'金额不是明确数值：{coordinate}') from error
    if not amount.is_finite():
        raise ValueError(f'金额不是有限数值：{coordinate}')
    return amount


def business_end(ws, header_row, mapping):
    """Exclude a standalone trailing report-query section; never skip incomplete detail."""
    for row in ws.iter_rows(min_row=header_row + 1):
        if not re.fullmatch(r'查询条件\s*[:：]\s*', str(row[0].value or '').strip()):
            continue
        start = row[0].row
        for tail in ws.iter_rows(min_row=start):
            filled = [c for c in tail if c.value is not None and str(c.value).strip()]
            if filled and (len(filled) != 1 or filled[0].column != 1
                           or not re.search(r'[:：]', str(filled[0].value))):
                raise ValueError('查询条件说明段与业务字段混杂，不能猜测核算范围')
            if any(tail[i].value is not None for i in mapping.values() if i != 0):
                raise ValueError('表尾说明含业务数据，不能跳过')
        return start - 1, list(range(start, ws.max_row + 1))
    return ws.max_row, []


def execute_workflow(source, plan, deadline=None):
    with localcontext() as context:
        context.prec = 100
        return _execute_workflow(source, plan, deadline)


def _execute_workflow(source, plan, deadline=None):
    validate_workflow(plan)
    source = Path(source).resolve()
    if source.suffix.lower() not in {'.xlsx', '.xlsm'}:
        raise ValueError('当前可验证流程仅支持 XLSX/XLSM 原附件')
    source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    workbook = openpyxl.load_workbook(source, data_only=False)
    cached = None
    try:
        fields = list(dict.fromkeys([plan['group_by'], plan['measure']]
                                   + [f['field'] for f in plan['filters']]))
        ws, header_row, headers, mapping = locate_table(workbook, fields)
        end_row, metadata_rows = business_end(ws, header_row, mapping)
        cached = openpyxl.load_workbook(source, data_only=True)
        value_ws = cached[ws.title]
        sums = defaultdict(Decimal); counts = Counter()
        source_rows = 0; excluded = Counter(); skipped = []; selected = []
        for row in ws.iter_rows(min_row=header_row + 1, max_row=end_row):
            if deadline is not None and time.monotonic() >= deadline:
                raise TimeoutError('核算流程超时')
            if not any(c.value is not None for c in row):
                continue
            r = row[0].row
            group_cell = row[mapping[plan['group_by']]]
            if re.fullmatch(r'(合计|总计|小计|total|grand total)[:：]?', str(group_cell.value or '').strip(), re.I):
                # Only a trailing subtotal may be ignored; embedded totals are ambiguous.
                if any(any(c.value is not None for c in rr) for rr in ws.iter_rows(min_row=r + 1)):
                    raise ValueError('业务区域中存在合计行，不能猜测统计范围')
                skipped.append(r)
                continue
            if group_cell.value is None or group_cell.data_type == 'f':
                raise ValueError(f'分组字段缺失或含公式：{group_cell.coordinate}')
            source_rows += 1
            rejected = False
            for i, f in enumerate(plan['filters']):
                cell = row[mapping[f['field']]]
                if cell.data_type == 'f' or cell.value is None:
                    raise ValueError(f'筛选字段缺失或含公式：{cell.coordinate}')
                if not matches(cell.value, f):
                    excluded[str(i)] += 1
                    rejected = True
                    break
            if rejected:
                continue
            amount_cell = row[mapping[plan['measure']]]
            # Formula caches cannot establish freshness; require a recalculation workflow.
            if amount_cell.data_type == 'f':
                raise ValueError(f'金额含公式，需要显式重算流程：{amount_cell.coordinate}')
            amount = decimal_amount(value_ws[amount_cell.coordinate].value, amount_cell.coordinate)
            group = str(group_cell.value).strip()
            if not group:
                raise ValueError(f'分组字段为空：{group_cell.coordinate}')
            sums[group] += amount; counts[group] += 1; selected.append(r)
        if not source_rows:
            raise ValueError('附件没有可核算业务明细')
        total = sum(sums.values(), Decimal(0))
        amount_header = str(headers[mapping[plan['measure']]])
        unit_match = re.search(r'[（(]([^()（）]+)[）)]$', amount_header)
        data = {
            'kind': 'personal_skill_workflow', 'engine': ENGINE, 'plan_hash': digest(plan),
            'workbook_id': source_hash, 'file_name': source.name, 'sheet': ws.title,
            'header_row': header_row, 'data_range': f'A{header_row + 1}:{openpyxl.utils.get_column_letter(ws.max_column)}{end_row}',
            'field_mapping': {f: openpyxl.utils.get_column_letter(i + 1) for f, i in mapping.items()},
            'group_by': plan['group_by'], 'measure': amount_header,
            'unit': unit_match.group(1) if unit_match else '原表单位',
            'filters': plan['filters'], 'source_rows': source_rows,
            'included_rows': len(selected), 'excluded_rows': sum(excluded.values()),
            'excluded_by_filter': dict(excluded), 'selected_rows': selected, 'skipped_summary_rows': skipped,
            'skipped_metadata_rows': metadata_rows,
            'groups': [{'key': k, 'count': counts[k], 'amount': str(sums[k])} for k in sorted(sums)],
            'total': str(total), 'verified': True,
        }
        data['evidence_id'] = digest(data)[:24]
        verify_receipt(data, plan, source)
        return data
    finally:
        workbook.close()
        if cached:
            cached.close()


def verify_receipt(data, plan, source):
    if data['plan_hash'] != digest(plan) or data['workbook_id'] != hashlib.sha256(Path(source).read_bytes()).hexdigest():
        raise ValueError('流程或输入文件摘要不一致')
    if data['source_rows'] != data['included_rows'] + data['excluded_rows']:
        raise ValueError('筛选前后行数不守恒')
    if sum(g['count'] for g in data['groups']) != data['included_rows']:
        raise ValueError('分组条数不守恒')
    if len(set(data['selected_rows'])) != data['included_rows']:
        raise ValueError('核算范围含重复或缺失行')
    with localcontext() as context:
        context.prec = 100
        if sum((Decimal(g['amount']) for g in data['groups']), Decimal(0)) != Decimal(data['total']):
            raise ValueError('分组金额与总额不一致')
    if data['evidence_id'] != digest({k: v for k, v in data.items() if k != 'evidence_id'})[:24]:
        raise ValueError('核算回执摘要不一致')


def render_receipt(data):
    def cell(value):
        return str(value).replace('|', '\\|').replace('\n', ' ').replace('\r', ' ')
    lines = [f"已完成核算并通过校验：纳入 **{data['included_rows']} 条**，排除 **{data['excluded_rows']} 条**。",
             '', f"| {cell(data['group_by'])} | 条数 | {cell(data['measure'])} |", '|---|---:|---:|']
    lines += [f"| {cell(g['key'])} | {g['count']} | {g['amount']} |" for g in data['groups']]
    lines += [f"| **合计** | **{data['included_rows']}** | **{data['total']}** |", '',
              f"来源：{cell(data['file_name'])}，{cell(data['sheet'])}，{data['data_range']}；单位：{cell(data['unit'])}。"]
    return '\n'.join(lines)


def main():
    parser = argparse.ArgumentParser(description='Replay and verify a saved calculation workflow')
    parser.add_argument('--input', required=True)
    parser.add_argument('--plan', default=str(Path(__file__).resolve().parent.parent / 'references' / 'workflow.json'))
    parser.add_argument('--receipt', help='Optional JSON execution receipt path')
    args = parser.parse_args()
    plan = json.loads(Path(args.plan).read_text(encoding='utf-8'))
    data = execute_workflow(args.input, plan)
    if args.receipt:
        if Path(args.receipt).resolve() == Path(args.input).resolve():
            raise ValueError('回执不能覆盖输入附件')
        Path(args.receipt).write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(data, ensure_ascii=False))


if __name__ == '__main__':
    main()
