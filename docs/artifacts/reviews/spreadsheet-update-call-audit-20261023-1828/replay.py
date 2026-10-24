"""Read-only replay of the 18:28 call; no model or workbook mutations.

Mocks admission model decisions with the exact persisted contract/verdict.
"""
import copy
import hashlib
import json
from collections import Counter, defaultdict
from decimal import Decimal
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import openpyxl

from dsh_modules.agent_loop import (
    detect_missing_claimed_artifacts, detect_missing_requested_deliverable,
    _check_no_tool_assertion_guard, MAX_SPREADSHEET_ROUNDS_CEILING,
)
from dsh_modules.analysis_request_text import user_instruction
from dsh_modules.prompt_builder import build_user_turn
from dsh_modules.runner import select_active_tools
from dsh_modules.runtime_policy import RuntimePolicy
from dsh_modules.skill_admission import admit_skill, execution_contract_context
from dsh_modules.spreadsheet_analysis_evidence import requires_spreadsheet_analysis
from dsh_modules.task_intent import extract_task_intent, has_physical_artifact_request
from dsh_modules.telemetry import TelemetryStats
from dsh_modules.spreadsheet_update.tools import SPREADSHEET_UPDATE_TOOL_DEFINITIONS

ROOT = Path.cwd()
OUT = Path(__file__).resolve().parent
WS = ROOT / 'data/users/e7fce333-a8f4-4097-9a53-f0a4c729da46/workspace'
SID = 'chat-session-1792747141627-wheti8di'
TID = '1792751302313739669-48fdcc65286b485186470ba4fcf23f72'
E = json.loads((WS / '.dsh/sessions' / (SID + '.evidence') / (TID + '.json')).read_text())
A = {'content': json.loads((OUT / 'call-summary.json').read_text())['final_content']}
accepted = next(g['reason'] for g in E['guard_decisions'] if g['guard'] == 'task_intent')
verdict = {'action': 'general', 'skill_id': None, 'use_default_rules': False,
           'rule_basis': '', 'reason': 'scope_mismatch', 'question': ''}
skills = [{'id': 'dept-summary', 'name': 'dept-summary', 'description': 'audit mock candidate',
           'execution_mode': 'verified_workflow'}]
with patch('dsh_modules.skill_admission.extract_task_intent', return_value=copy.deepcopy(accepted)), \
     patch('dsh_modules.skill_admission.structured_decision', return_value=copy.deepcopy(verdict)), \
     patch('dsh_modules.skill_admission.SemanticSkillMatcher') as matcher:
    matcher.return_value.match.return_value = []
    route = admit_skill(E['query'], [], list(E['source_hashes']),
                       SimpleNamespace(skill_id='dept-summary'), 'not-called', RuntimePolicy(),
                       TelemetryStats(), None, available_skills=skills,
                       read_body=lambda _: 'Audit mock: candidate is rejected.')
tools = select_active_tools(E['query'], route, False, True)
names = [t['function']['name'] for t in tools]
analysis_required = requires_spreadsheet_analysis(E['query'], route)
user_turn = build_user_turn(E['query'], session_files=list(E['source_hashes']),
                            is_generate_intent=route.is_generate_intent,
                            is_inspect_intent=route.is_inspect_intent,
                            is_office_intent=route.is_office_intent,
                            available_tool_names=names,
                            timestamp_str='2026-10-23 18:25:51 Asia/Shanghai')
with patch('dsh_modules.task_intent.structured_decision',
           side_effect=[copy.deepcopy(accepted), copy.deepcopy(accepted)]):
    intent_result = extract_task_intent(user_instruction(E['query']), [], 'not-called',
                                       RuntimePolicy(), TelemetryStats(), None, input_count=2)

with patch('dsh_modules.agent_loop.WORKSPACE_DIR', str(WS)):
    missing_claims = detect_missing_claimed_artifacts(A['content'], is_generate_intent=True,
                     verify_claims_only=True, turn_start_time=1792751151.52)
    missing_requested = detect_missing_requested_deliverable(E['query'], 1792751151.52,
                         expected_deliverables=route.deliverables,
                         is_generate_intent=route.is_generate_intent,
                         is_inspect_intent=route.is_inspect_intent)
    guard_action, guard_message, _ = _check_no_tool_assertion_guard(
        reply_text=A['content'], last_user_prompt=E['query'], round_idx=3, max_rounds=4,
        start_ts=1792751151.52, is_guide_intent=False,
        expected_deliverables=route.deliverables, is_generate_intent=route.is_generate_intent,
        is_inspect_intent=route.is_inspect_intent, tools=tools,
        execution_evidence=[], hard_ceiling=4)

files = [WS / n for n in E['source_hashes']]
before = {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in files}
source = WS / '工数实绩明细_20261020134124.xlsx'
target = WS / '工数3Q预实管理.xlsx'
wb = openpyxl.load_workbook(source, read_only=True, data_only=True)
s = wb['工数实绩明细']; s.reset_dimensions()
amounts = defaultdict(Decimal); counts = Counter(); months = Counter(); dept_names = defaultdict(set)
for row in s.iter_rows(min_row=5, max_row=278, max_col=18, values_only=True):
    months[str(row[1])] += 1
    dept_names[str(row[2])].add(str(row[3]))
    if row[12] == '有偿':
        amounts[str(row[2])] += Decimal(str(row[15])); counts[str(row[2])] += 1
wb.close()
analysis = E['receipts'][0]['data']
receipt_amounts = {f['group_fields']['部门编号']: Decimal(f['value']) for f in analysis['facts']}
wb = openpyxl.load_workbook(target, read_only=True, data_only=False)
sheet_names = wb.sheetnames; month_header = wb['1019']['C3'].value
wb.close()
after = {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in files}
name_errors = []
for line in A['content'].splitlines():
    parts = [p.strip().replace('**', '') for p in line.strip().strip('|').split('|')]
    if len(parts) == 5 and parts[0] in dept_names and parts[1] not in dept_names[parts[0]]:
        name_errors.append({'department_code': parts[0], 'reply_name': parts[1],
                            'source_names': sorted(dept_names[parts[0]])})
plan_schema = next(t['function']['parameters'] for t in SPREADSHEET_UPDATE_TOOL_DEFINITIONS
                   if t['function']['name'] == 'plan_spreadsheet_update')
result = {
    'session_id': SID, 'turn_id': TID, 'accepted_contract': accepted,
    'routing_reproduction': {
        'path': 'admit_skill action=general normal return, not general_route exception fallback',
        'flags': {k: getattr(route, k, None) for k in ['is_generate_intent', 'is_inspect_intent',
                  'requires_execution', 'deliverables', 'is_guide_intent']},
        'active_tools': names, 'requires_spreadsheet_analysis': analysis_required,
        'inspection_round_ceiling': MAX_SPREADSHEET_ROUNDS_CEILING,
        'admission_context': execution_contract_context(route.admission),
        'update_execution_guidance_present': 'Excel 表格受限更新执行规范' in user_turn,
        'readonly_guidance_present': '本次任务为只读数据分析与事实核查' in user_turn,
        'note': 'Same-hash current-source reconstruction; historical full schema/prompt not persisted.',
    },
    'intent_reproduction': {'physical_request_regex_matches': has_physical_artifact_request(user_instruction(E['query'])),
                            'after_same_persisted_decisions': intent_result,
                            'note': 'Both decision results mocked with saved accepted contract; no model calls.'},
    'completion_reproduction': {'missing_claimed_files': missing_claims,
        'missing_requested_deliverable': missing_requested, 'no_tool_guard_action': guard_action,
        'no_tool_guard_message': guard_message,
        'input_text': 'Exact persisted final visible reply', 'transaction_receipts': 0},
    'reconciliation': {'evidence_id': analysis['evidence_id'], 'by_department': [
        {'department_code': k, 'paid_amount_source_K': str(amounts[k]), 'paid_records': counts[k]}
        for k in sorted(amounts)], 'total_paid_amount_source_K': str(sum(amounts.values())),
        'paid_records': sum(counts.values()), 'months': dict(months),
        'matches_tool_receipt': dict(amounts) == receipt_amounts,
        'reply_department_name_mismatches': name_errors},
    'file_integrity': {'before': before, 'after': after, 'unchanged': before == after,
        'matches_call': before == E['source_hashes'], 'target_sheet_names': sheet_names,
        'target_C3': month_header,
        'workspace_1023_workbooks': [str(p.relative_to(WS)) for p in WS.rglob('*1023*.xlsx')]},
    'plan_tool_schema_has_properties': bool(plan_schema['properties'].get('target_ref')),
    'model_calls': 0,
}
OUT.mkdir(parents=True, exist_ok=True)
(OUT / 'replay-result.json').write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n')
assert route.requires_execution and route.is_inspect_intent and not route.is_generate_intent
assert not analysis_required
assert 'execute_spreadsheet_update' in names
assert not result['routing_reproduction']['update_execution_guidance_present']
assert not missing_claims and missing_requested is None and guard_action == 'pass'
assert result['file_integrity']['matches_call'] and before == after and sheet_names == ['1019']
assert dict(amounts) == receipt_amounts
print(json.dumps({k: result[k] for k in ['routing_reproduction', 'completion_reproduction',
                   'file_integrity', 'reconciliation']}, ensure_ascii=False, indent=2))
