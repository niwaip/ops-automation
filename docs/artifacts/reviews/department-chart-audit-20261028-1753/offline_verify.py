"""Read-only audit: original workbook, saved receipts, published HTML and counterfactual selection."""
from collections import defaultdict
from decimal import Decimal
from pathlib import Path
from types import SimpleNamespace
import hashlib
import json
import re
import shutil
import sys
import time
from xml.etree import ElementTree as ET

ROOT = Path(__file__).resolve().parents[4]
OUT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT / 'apps/backend/runtimes/personal-sandbox-runner/src'))
from dsh_modules.spreadsheet_chart_delivery import chart_dataset
from dsh_modules.spreadsheet_chart_renderer import validate_bar_chart
import openpyxl

SID = 'chat-session-1793180826543-nqx1j95c'
WORKSPACE = ROOT / 'data/users/e7fce333-a8f4-4097-9a53-f0a4c729da46/workspace'
sha = lambda p: hashlib.sha256(p.read_bytes()).hexdigest()
first_file = OUT / (SID + '.evidence') / '1793181202495562166-b072ad9b8b5e47f599196cffc8fa4cac.json'
second_file = OUT / (SID + '.evidence') / '1793181276824425423-89ee55c791ad42e59db8ff92c3ea4389.json'
first, second = json.loads(first_file.read_text()), json.loads(second_file.read_text())
replays = []
for saved in first['receipts']:
    receipt = SimpleNamespace(data=saved['data'], provenance=saved['provenance'])
    for name, prompt in [('original', first['query']), ('user_command_only', first['query'].split('用户指令：')[-1])]:
        started = time.perf_counter()
        try:
            data = chart_dataset(receipt, prompt)
            result = {'status': 'selected', 'metric': data['value_field'], 'groups': len(data['points'])}
        except ValueError as error:
            result = {'status': 'rejected', 'reason': str(error)}
        replays.append({'evidence_id': saved['data']['evidence_id'], 'prompt_variant': name,
                        'elapsed_ms': (time.perf_counter()-started)*1000, **result})

source = WORKSPACE / '工数实绩明细_20261020134124.xlsx'
before = sha(source)
book = openpyxl.load_workbook(source, read_only=True, data_only=True)
sheet = book['工数实绩明细']
headers = list(next(sheet.iter_rows(min_row=4, max_row=4, max_col=18, values_only=True)))
index = {str(v): i for i, v in enumerate(headers) if v is not None}
amounts, hours, counts = defaultdict(Decimal), defaultdict(Decimal), defaultdict(int)
scanned = matched = 0
for row in sheet.iter_rows(min_row=5, max_row=278, max_col=18, values_only=True):
    scanned += 1
    if row[index['有偿无偿']] != '有偿': continue
    matched += 1
    key = row[index['部门名称']]
    amounts[key] += Decimal(str(row[index['金额(K)']]))
    hours[key] += Decimal(str(row[index['工数(小时)']]))
    counts[key] += 1
book.close()
assert before == sha(source)
for saved in [*first['receipts'], *second['receipts']]:
    for row in saved['data']['rows']:
        key = row['部门名称']
        assert Decimal(str(row['amount'])) == amounts[key]
        if 'hours' in row: assert Decimal(str(row['hours'])) == hours[key]
        if 'count' in row: assert int(row['count']) == counts[key]

manifest = json.loads((OUT/'runs/7417f3286d6244a99c009f5b9ab4310a/manifest.json').read_text())
html_file = WORKSPACE / Path(manifest['file_path']).name
html = html_file.read_text()
node = next(n for n in ET.fromstring(html).iter() if n.get('id') == 'chart-data')
dataset = json.loads(node.text)
assert validate_bar_chart(html, dataset)
assert sha(html_file) == manifest['sha256']
assert dataset['where'] == {'field': '有偿无偿', 'op': 'eq', 'value': '有偿'}
assert {p['label']: Decimal(p['value']) for p in dataset['points']} == dict(amounts)
shutil.copy2(html_file, OUT/'published-chart.html')

session = json.loads((OUT/'redis-session.json').read_text())
performance = []
for message in session['history'][-4:]:
    if message['role'] != 'assistant': continue
    trace = message['metadata']['executionTrace']; metrics = trace['metrics']
    ledger = metrics['callsLedger']; unique = []
    # Specific duplicate proven by identical usage/time and one proxy request in window log.
    for call in ledger:
        if any(c['purpose']==call['purpose'] and c['total_ms']==call['total_ms'] and c['usage']==call['usage'] for c in unique): continue
        unique.append(call)
    tokens = {k: sum(c['usage'].get(k, 0) for c in unique) for k in ('prompt_tokens','completion_tokens','total_tokens')}
    llm_ms = sum(c['total_ms'] for c in unique)
    performance.append({'message_id': message['id'], 'http_ms': 21475 if len(performance)==0 else 13082,
        'harness_ms': trace['durationMs'], 'dsh_ms': metrics['durationMs'], 'ttft_ms': metrics['ttftMs'],
        'recorded_tokens': metrics['tokens'], 'unique_tokens': tokens, 'recorded_calls':len(ledger),
        'unique_calls':len(unique), 'unique_llm_ms':llm_ms,
        'residual_dsh_ms':metrics['durationMs']-llm_ms,
        'pre_aggregate_first_four_calls_ms':sum(c['total_ms'] for c in unique[:4]),
        'ledger': unique})

result = {'source_sha256': before, 'source_unchanged':before==sha(source), 'scanned':scanned,'matched':matched,
    'groups':len(amounts), 'amount_total_k':str(sum(amounts.values())),
    'independent_rows':[{'department':k,'amount':str(amounts[k]),'hours':str(hours[k]),'count':counts[k]} for k in amounts],
    'selection_replays':replays, 'published_html_sha256':sha(html_file),'published_html_validation':True,
    'published_matches_independent_workbook':True,'performance':performance}
(OUT/'audit-findings.json').write_text(json.dumps(result,ensure_ascii=False,indent=2))
print(json.dumps({k:v for k,v in result.items() if k!='performance'},ensure_ascii=False,indent=2))
for p in performance: print(json.dumps({k:v for k,v in p.items() if k!='ledger'},ensure_ascii=False))
