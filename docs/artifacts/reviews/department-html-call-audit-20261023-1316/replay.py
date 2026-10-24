"""Read-only source audit; reproduction writes only to temporary directories. No LLM calls."""
import contextlib
import hashlib
import io
import json
import os
import re
import sys
import tempfile
import time
from collections import defaultdict
from dataclasses import asdict
from decimal import Decimal
from pathlib import Path
from unittest.mock import Mock

ROOT = Path(__file__).resolve().parents[4]
HERE = Path(__file__).resolve().parent
SRC = ROOT / 'apps/backend/runtimes/personal-sandbox-runner/src'
sys.path.insert(0, str(SRC))
from openpyxl import load_workbook
from dsh_modules.analysis_contract import build_analysis_contract
from dsh_modules.analysis_validation import validate_reply_against_contract
from dsh_modules.artifact_exporter import ArtifactExporter
from dsh_modules.final_answer import finalize_agent_text
from dsh_modules.runtime_policy import RuntimePolicy
from dsh_modules.telemetry import TelemetryStats
from dsh_modules.tool_result import ToolResult

WS = ROOT / 'data/users/e7fce333-a8f4-4097-9a53-f0a4c729da46/workspace'
SID = 'chat-session-1792732494415-xbnajy0r'
source = WS / '工数实绩明细_20261020134124.xlsx'
sha = lambda p: hashlib.sha256(Path(p).read_bytes()).hexdigest()
before = {str(p): sha(p) for p in [source, WS/'index.html', WS/'dept_data.json']}
history = json.loads((WS / f'.dsh/sessions/{SID}.json').read_text())
evidence = json.loads((WS / f'.dsh/sessions/{SID}.evidence.json').read_text())
prompt = history[0]['content']
receipts = [ToolResult(**r) for r in evidence['receipts']]

book = load_workbook(source, read_only=True, data_only=True)
sheet = book['工数实绩明细']
declared_dimension = sheet.calculate_dimension()
sheet.reset_dimensions()
groups = defaultdict(lambda: {'amount_k': Decimal(0), 'hours': Decimal(0), 'rows': 0, 'employees': set()})
for row in sheet.iter_rows(min_row=5, max_row=278, max_col=18, values_only=True):
    assert row[2] and row[3] and row[15] is not None
    g = groups[(str(row[2]),str(row[3]))]
    g['amount_k'] += Decimal(str(row[15]))
    g['hours'] += Decimal(str(row[14]))
    g['rows'] += 1
    g['employees'].add(str(row[4]))
book.close()
data = json.loads((WS/'dept_data.json').read_text())
computed = []
for (code, name), g in sorted(groups.items(), key=lambda v:-v[1]['amount_k']):
    idx = next(i for i,c in data['部门编号'].items() if c == code)
    assert Decimal(str(data['金额'][idx])) == g['amount_k']
    computed.append({'code':code,'name':name,'amount_k':str(g['amount_k']),
                     'hours':str(g['hours']),'rows':g['rows'],'employees':len(g['employees'])})

raw_final = '文件已推送，请查收。'
session = json.loads((HERE/'session-sanitized.json').read_text())
responses = session['history'][1]['metadata']['executionTrace']['metrics']['modelResponses']
assert hashlib.sha256(raw_final.encode()).hexdigest() == responses[-1]['rawSha256']
validation = validate_reply_against_contract(build_analysis_contract(prompt),raw_final,10,10,
    messages=[{'role':'tool','content':evidence['receipts'][0]['data']}],execution_evidence=receipts)
old_history = json.loads((WS/'.dsh/sessions/chat-session-1792730309468-pezbjusi.json').read_text())
old_text = next(m['content'] for m in reversed(old_history) if m['role']=='assistant' and '```html' in m['content'])
old_html = re.findall(r'```html\s*\n([\s\S]*?)```',old_text)[0].strip()
variants = {}
for fresh in [False, True]:
    with tempfile.TemporaryDirectory(prefix='department-html-audit-') as td:
        path = Path(td)/'index.html'
        path.write_text(old_html+'\n')
        start = time.time()
        os.utime(path,(start+(1 if fresh else -3600),)*2)
        callback = Mock(side_effect=AssertionError('Unexpected model invocation'))
        with contextlib.redirect_stdout(io.StringIO()):
            final, healed = finalize_agent_text(raw_final,
                [{'role':'user','content':prompt},{'role':'tool','content':evidence['receipts'][0]['data']}],
                model='audit-no-model',policy=RuntimePolicy(),deadline=None,is_guide_intent=False,
                executed_calls_history=['bash:read','bash:wc','send_file:index'],was_token_truncated=False,
                telemetry=TelemetryStats(),last_user_prompt=prompt,expected_deliverables=['.html'],
                execution_evidence=receipts,turn_start_time=start,call_model=callback,workspace_dir=td,
                max_agent_rounds=10,max_spreadsheet_rounds=4)
            output, paths = ArtifactExporter.export_html(final,False,td,start,True,prompt,[])
        variants['current_turn_file' if fresh else 'stale_file'] = {
            'model_calls':callback.call_count,'finalizer_text':final,
            'success_banner':output.startswith('✨ **交互式页面已生成完毕'),
            'output_bytes':path.stat().st_size,'output_sha256':sha(path),
            'matches_actual_file':sha(path)==sha(WS/'index.html'),
            'has_svg_or_canvas':bool(re.search(r'<(?:svg|canvas)\b',path.read_text(),re.I)),
            'original_preserved':path.read_text().strip()==old_html,
            'matches_actual_answer':output.strip()==history[1]['content'].strip()}

all_employees = set().union(*(g['employees'] for g in groups.values()))
total = sum((g['amount_k'] for g in groups.values()),Decimal(0))
result = {'source_sha256':sha(source),'declared_dimension':declared_dimension,'range':'工数实绩明细!A5:R278',
          'groups':computed,'total_amount_k':str(total),'total_hours':str(sum((g['hours'] for g in groups.values()),Decimal(0))),
          'rows':sum(g['rows'] for g in groups.values()),'unique_employees':len(all_employees),
          'amount_per_employee_k':str(total/len(all_employees)),
          'dept_data_amounts_match':True,'last_model_response_hash_matches':True,
          'validator_with_successful_computation':asdict(validation),
          'prior_html':{'session':'chat-session-1792730309468-pezbjusi','lines':len(old_html.splitlines()),
                        'bytes_without_trailing_newline':len(old_html.encode()),'sha256':hashlib.sha256(old_html.encode()).hexdigest()},
          'reproduction':variants,'original_files_unchanged':all(sha(p)==v for p,v in before.items()),
          'local_runtime_hashes':{n:sha(SRC/'dsh_modules'/n) for n in ['artifact_exporter.py','final_answer.py','html_report_fallback.py','agent_loop.py','analysis_validation.py']}}
assert result['original_files_unchanged']
assert variants['stale_file']['matches_actual_file']
assert variants['stale_file']['matches_actual_answer']
assert variants['current_turn_file']['original_preserved']
print(json.dumps(result,ensure_ascii=False,indent=2))
