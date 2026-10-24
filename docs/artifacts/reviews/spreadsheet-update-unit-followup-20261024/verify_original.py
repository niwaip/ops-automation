"""Independent check of this audit's original attachments, never production defaults."""
import hashlib
import json
import sys
from decimal import Decimal
from pathlib import Path

import openpyxl
from dsh_modules.spreadsheet_update.diff import _cells_equivalent

root = Path(sys.argv[1]).resolve()
runs = list((root / '.dsh/spreadsheet-runs').glob('*/receipt.json'))
assert len(runs) == 1, runs
receipt = json.loads(runs[0].read_text())
plan = json.loads(runs[0].with_name('plan.json').read_text())
assert plan['semantic_bindings']['unit_basis'] == {'source': 'user', 'unit': 'K', 'quote': 'k单位'}
target = root / '工数3Q预实管理.xlsx'
source = root / '工数实绩明细_20261020134124.xlsx'
output = root / Path(receipt['output_path']).name
assert hashlib.sha256(output.read_bytes()).hexdigest() == receipt['output_hash']
wb = openpyxl.load_workbook(source, data_only=True)
ws = wb['工数实绩明细']
headers = {cell.value: cell.column for cell in ws[4] if cell.value is not None}
sums = {}
records = 0
for row in ws.iter_rows(min_row=5, values_only=True):
    if str(row[headers['工数年月'] - 1]) == '202610' and row[headers['有偿无偿'] - 1] == '有偿':
        key = str(row[headers['部门编号'] - 1])
        sums[key] = sums.get(key, Decimal(0)) + Decimal(str(row[headers['金额(K)'] - 1]))
        records += 1
wb.close()
before = openpyxl.load_workbook(target)
after = openpyxl.load_workbook(output)
cached = openpyxl.load_workbook(output, data_only=True)
assert after.sheetnames == before.sheetnames + ['1023'], after.sheetnames
coordinates = {str(before['1019'].cell(row, 2).value): 'C' + str(row)
               for row in range(5, 22) if str(before['1019'].cell(row, 2).value) in sums}
assert coordinates.keys() == sums.keys()
for key, value in sums.items():
    assert abs(Decimal(str(after['1023'][coordinates[key]].value)) - value) < Decimal('0.000000001')
changed = []
formula_count = 0
for name in after.sheetnames:
    baseline = before['1019'] if name == '1023' else before[name]
    current = after[name]
    assert set(baseline.merged_cells.ranges) == set(current.merged_cells.ranges)
    for row in baseline:
        for cell in row:
            actual = current[cell.coordinate]
            if cell.data_type == 'f':
                formula_count += 1
                assert actual.data_type == 'f'
                assert cached[name][cell.coordinate].value is not None
            if not _cells_equivalent(cell.value, actual.value, before.epoch):
                assert name == '1023' and cell.coordinate in coordinates.values(), (name, cell.coordinate)
                changed.append(cell.coordinate)
assert receipt['formula_health']['total_errors'] == 0
assert receipt['formula_health']['uncalculated_formulas'] == 0
assert receipt['formula_health']['cached_formulas'] == formula_count
report = {'status': 'passed', 'record_count': records, 'key_count': len(sums), 'unit': 'K',
          'source_total': str(sum(sums.values())), 'coordinates': coordinates,
          'actually_changed': changed, 'formula_count': formula_count,
          'formula_errors': 0, 'missing_caches': 0,
          'history_budget_other_periods_and_regions_preserved': True,
          'unit_basis': plan['semantic_bindings']['unit_basis'],
          'output_file': output.name, 'output_hash': receipt['output_hash']}
for book in (before, after, cached):
    book.close()
(root / 'independent-verification.json').write_text(json.dumps(report, ensure_ascii=False, indent=2))
print(json.dumps(report, ensure_ascii=False, indent=2))
