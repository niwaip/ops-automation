import sys,io
sys.stdin=io.StringIO('{"request": "帮我生成一个小学生 竞选班干部的ppt，3到4页，主要是自我介绍，个人荣誉，竞选的目标和宣言，或者给你的建议，不要太复杂", "contract": {"instruction_relation": "standalone", "task_mode": "execute", "output_mode": "artifact", "artifact_types": [".html"], "unit_confirmation": null, "section_confirmation": null, "skill_invocation": {"mode": "none", "name": "", "replace_defaults": false}, "requirements": [{"id": "1", "quote": "帮我生成一个小学生 竞选班干部的ppt"}, {"id": "2", "quote": "3到4页"}, {"id": "3", "quote": "主要是自我介绍，个人荣誉，竞选的目标和宣言"}, {"id": "4", "quote": "不要太复杂"}]}}')
import copy, hashlib, io, json, sys, tempfile, zipfile
from contextlib import redirect_stdout
from pathlib import Path
sys.path.insert(0, '/usr/local/bin')
import importlib.util, subprocess
from dsh_modules.skills import get_available_skills
from dsh_modules.skill_router import SkillRouter
from dsh_modules.intent_contract import extract_source_and_target_artifacts
from dsh_modules.skill_admission import general_route
from dsh_modules.task_intent import validate_intent
from dsh_modules.planning_phase import build_deliverable_execution_instruction
from dsh_modules.config import SKILL_DIR
payload=json.loads(sys.stdin.read())
request=payload['request'];contract=payload['contract']
result={'python_executable':sys.executable,'pptx_module_available':importlib.util.find_spec('pptx') is not None,'skill_dir':SKILL_DIR}
skills=get_available_skills()
ppt_skill=next((s for s in skills if s['id']=='pptx'),None)
result['pptx_skill']={k:ppt_skill.get(k) for k in ('id','type','path','deliverables','requires_execution','description')} if ppt_skill else None
if ppt_skill:
 path=Path(ppt_skill['path'])/'SKILL.md'
 result['pptx_skill_sha256']=hashlib.sha256(path.read_bytes()).hexdigest()
with redirect_stdout(io.StringIO()):
 candidate=SkillRouter.route(request,[],available_skills=skills,load_skill_context=False)
 general=general_route(request,[],[],copy.deepcopy(contract))
result['deterministic_target_formats']=extract_source_and_target_artifacts(request)[1]
result['initial_route']={k:getattr(candidate,k) for k in ('skill_id','deliverables','is_ppt_intent','requires_execution')}
result['accepted_recorded_contract']=validate_intent(copy.deepcopy(contract),request,[],[],allow_sanitize=True)
result['general_route']={k:getattr(general,k) for k in ('skill_id','deliverables','is_ppt_intent','requires_execution')}
result['execution_instruction']=build_deliverable_execution_instruction(general)
try:
    from pptx import Presentation
    import pptx
    result['pptx_version'] = pptx.__version__
    with tempfile.TemporaryDirectory(prefix='election-ppt-audit-') as folder:
        path=Path(folder)/'smoke.pptx'
        deck=Presentation()
        for title in ['自我介绍','个人荣誉','竞选目标','竞选宣言']:
            slide=deck.slides.add_slide(deck.slide_layouts[5]);slide.shapes.title.text=title
        deck.save(path)
        loaded=Presentation(path)
        with zipfile.ZipFile(path) as archive:
            result['pptx_smoke']={'pages':len(loaded.slides),'zip_valid':archive.testzip() is None,'bytes':path.stat().st_size,'temporary_only':True}
except ImportError as error:
    result['pptx_smoke']={'status':'blocked','error':str(error),'generation_attempted':False}
check = subprocess.run(['bash','-lc','python3 -c "import sys; print(sys.executable); import pptx; print(pptx.__version__)"'],capture_output=True,text=True)
result['bash_python_check']={'exit_code':check.returncode,'stdout':check.stdout,'stderr':check.stderr}
print(json.dumps(result,ensure_ascii=False,indent=2))
