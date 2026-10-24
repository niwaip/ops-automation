"""Read-only reproduction; no LLM, service invocation, skill mutation or source writes."""
import collections
import hashlib
import json
import sys
from decimal import Decimal
from pathlib import Path

sys.path.insert(0, '/usr/local/bin')
from openpyxl import load_workbook
from dsh_modules.skill_router import SkillRouter
from dsh_modules.personal_skill_workflow import requests_workflow_change

source = Path('/workspace/工数实绩明细_20261020134124.xlsx')
before = hashlib.sha256(source.read_bytes()).hexdigest()
history_path = Path('/workspace/.dsh/sessions/chat-session-1792673001244-91mz0cae.json')
history = json.loads(history_path.read_text())
prompt = [h['content'] for h in history if h['role'] == 'user'][-1]
previous = history[:-2]
route = SkillRouter.route(prompt, previous, files=[source.name])
book = load_workbook(source, read_only=True, data_only=True)
sheet = book['工数实绩明细']
dimension = sheet.calculate_dimension()
sheet.reset_dimensions()
headers = next(sheet.iter_rows(min_row=4, max_row=4, max_col=18, values_only=True))
assert headers[2] == '部门编号' and headers[8] == 'WBS编号' and headers[15] == '金额(K)'
rows = list(sheet.iter_rows(min_row=5, max_row=278, max_col=18, values_only=True))
results = {}
for scope in ('all_business_rows', 'saved_wbs_exclusions'):
    amounts = collections.defaultdict(Decimal)
    counts = collections.Counter()
    for row in rows:
        if scope == 'saved_wbs_exclusions' and str(row[8] or '').strip().casefold().startswith(('ex','sy','sp','ip','ib','sz')):
            continue
        assert row[15] is not None
        key = str(row[2])
        amounts[key] += Decimal(str(row[15]))
        counts[key] += 1
    groups = [{'department': k, 'count': counts[k], 'amount_k': str(v)} for k, v in sorted(amounts.items())]
    results[scope] = {'groups': groups, 'total_k': str(sum(amounts.values(), Decimal(0))),
                      'included_rows': sum(counts.values()),
                      'departments_over_500k': [g for g in groups if Decimal(g['amount_k']) > Decimal(500)]}
book.close()
result = {'source_sha256': before, 'declared_dimension': dimension, 'business_range': 'A5:R278',
          'query': prompt, 'routing': {'skill_id': route.skill_id, 'affinity_score': route.affinity_score,
          'matched_reasons': route.matched_reasons, 'requires_execution': route.requires_execution},
          'new_instruction_guard_detected': requests_workflow_change(prompt.split('用户指令：')[-1]),
          'recomputation': results,
          'runtime_hashes': {name: hashlib.sha256(Path('/usr/local/bin/dsh_modules', name).read_bytes()).hexdigest()
                             for name in ('runner.py','skill_router.py','personal_skill_workflow.py',
                                          'personal_skill_workflow_engine.py','personal_skill_presentation.py')},
          'source_unchanged': before == hashlib.sha256(source.read_bytes()).hexdigest()}
assert result['source_unchanged']
assert route.skill_id == 'dept-summary'
assert result['new_instruction_guard_detected'] is False
assert results['saved_wbs_exclusions']['included_rows'] == 253
assert results['saved_wbs_exclusions']['total_k'] == '1637.41990'
assert len(results['saved_wbs_exclusions']['departments_over_500k']) == 1
print('AUDIT_RESULT=' + json.dumps(result, ensure_ascii=False))
