"""Read-only audit: independently aggregate the source and replay routing without an LLM."""
import collections
import copy
import hashlib
import json
import sys
from decimal import Decimal
from pathlib import Path

sys.path.insert(0, '/usr/local/bin')
from openpyxl import load_workbook
from dsh_modules.skills import get_available_skills
from dsh_modules.skill_router import SkillRouter
from dsh_modules.runner import select_active_tools
from dsh_modules.spreadsheet_analysis_evidence import requires_spreadsheet_analysis
from dsh_modules.analysis_contract import build_analysis_contract
from dsh_modules.spreadsheet_preflight import profile_spreadsheet_for_planning

source = Path('/workspace/工数实绩明细_20261020134124.xlsx')
before = hashlib.sha256(source.read_bytes()).hexdigest()
wb = load_workbook(source, read_only=True, data_only=True)
ws = wb['工数实绩明细']
declared_dimension = ws.calculate_dimension()
ws.reset_dimensions()  # The original workbook incorrectly declares A1:A1.
headers = [str(v or '').strip() for v in next(ws.iter_rows(min_row=4, max_row=4, max_col=18, values_only=True))]
rows = list(ws.iter_rows(min_row=5, max_row=278, max_col=18, values_only=True))
prefixes = ('ex', 'sy', 'sp', 'ip', 'ib', 'sz')
states = collections.Counter(str(r[12] or '').strip() for r in rows)

def aggregate(paid_only=False):
    sums = collections.defaultdict(Decimal)
    counts = collections.Counter()
    excluded_prefix = 0
    excluded_paid = 0
    for r in rows:
        if str(r[8] or '').strip().lower().startswith(prefixes):
            excluded_prefix += 1
            continue
        if paid_only and str(r[12] or '').strip() != '有偿':
            excluded_paid += 1
            continue
        if r[15] is None:
            raise AssertionError('Missing cached amount at column P')
        amount = Decimal(str(r[15]))
        sums[str(r[2])] += amount
        counts[str(r[2])] += 1
    return {'groups': {k: str(v) for k, v in sorted(sums.items())},
            'counts': dict(sorted(counts.items())),
            'total': str(sum(sums.values(), Decimal(0))),
            'included_rows': sum(counts.values()), 'excluded_prefix_rows': excluded_prefix,
            'excluded_paid_rows_after_prefix': excluded_paid}

file_name = source.name
prompt = '【当前轮次用户上传附件】: ' + file_name + '（这是用户本轮刚上传的新文件，为本次指令的主要处理对象）\n用户指令：部门别有偿统计'
skills = get_available_skills()
route = SkillRouter.route(prompt, existing_history=[], available_skills=skills, files=[file_name])
tools = select_active_tools(prompt, route, route.is_search_intent, True)
contract = build_analysis_contract(prompt)
baseline = copy.copy(route)
baseline.skill_id = 'xlsx'
base_tools = select_active_tools(prompt, baseline, baseline.is_search_intent, True)
skill_file = Path('/knowledge/skills/dept-summary/SKILL.md')
preflight = profile_spreadsheet_for_planning(source, prompt=prompt)
reported_groups = dict(zip(
    ['FC46005', 'FC34022', 'FC44003', 'FC45031', 'FC44006', 'FC44004', 'FC48005'],
    ['37.80641', '31.79720', '20.50715', '13.11909', '4.64361', '3.00000', '0.00000']))
reported_total = Decimal('114.87346')
result = {
    'source_sha256': before, 'sheet': ws.title, 'source_row_count': len(rows), 'headers': headers,
    'original_declared_dimension': declared_dimension,
    'paid_state_distribution': dict(states),
    'preflight': {k: preflight.get(k) for k in ['sheet', 'header_row', 'data_range', 'records']},
    'assistant_result_check': {
        'reported_groups': reported_groups,
        'reported_total': str(reported_total),
        'sum_of_reported_groups': str(sum(map(Decimal, reported_groups.values()), Decimal(0))),
        'total_minus_reported_group_sum': str(reported_total - sum(map(Decimal, reported_groups.values()), Decimal(0))),
        'correct_total_minus_reported_total': str(Decimal(aggregate(False)['total']) - reported_total),
    },
    'saved_rule': aggregate(False), 'assistant_added_paid_filter': aggregate(True),
    'routing': {'skill_id': route.skill_id, 'requires_execution': route.requires_execution,
        'is_inspect_intent': route.is_inspect_intent, 'is_generate_intent': route.is_generate_intent,
        'contract_has_metrics': contract.has_metrics,
        'native_analysis_required': requires_spreadsheet_analysis(prompt, route),
        'tools': [t.get('function', {}).get('name') for t in tools],
        'xlsx_counterfactual_analysis_required': requires_spreadsheet_analysis(prompt, baseline),
        'xlsx_counterfactual_tools': [t.get('function', {}).get('name') for t in base_tools]},
    'skill_sha256': hashlib.sha256(skill_file.read_bytes()).hexdigest(),
    'runtime_sha256': {name: hashlib.sha256(Path('/usr/local/bin/dsh_modules', name).read_bytes()).hexdigest()
        for name in ['runner.py', 'skill_router.py', 'agent_loop.py', 'planning_phase.py',
                     'prompt_builder.py', 'spreadsheet_analysis_evidence.py', 'spreadsheet_preflight.py']},
}
wb.close()
assert before == hashlib.sha256(source.read_bytes()).hexdigest(), 'source mutated during read-only audit'
assert route.skill_id == 'dept-summary'
assert result['routing']['native_analysis_required'] is False
assert result['routing']['xlsx_counterfactual_analysis_required'] is True
assert result['source_row_count'] == 274
assert result['saved_rule'] == result['assistant_added_paid_filter']
assert result['saved_rule']['total'] == '1637.41990'
assert result['saved_rule']['included_rows'] == 253
assert result['assistant_result_check']['total_minus_reported_group_sum'] == '4.00000'
result['source_sha256_after'] = hashlib.sha256(source.read_bytes()).hexdigest()
print('AUDIT_JSON=' + json.dumps(result, ensure_ascii=False))
