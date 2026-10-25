import openpyxl
from openpyxl.styles import Alignment, Border, Side, Font, PatternFill, numbers

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "2026财务指标"

# 冻结首行
ws.freeze_panes = "A2"

# 表头
headers = ["指标", "数值", "单位", "来源单元格"]
for col, h in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col, value=h)
    cell.font = Font(bold=True)
    cell.alignment = Alignment(horizontal="center")
    cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    cell.font = Font(bold=True, color="FFFFFF")

# 数据行
data = [
    ("营业收入", 124020, "千元", "B4"),
    ("毛利", 50120, "千元", "B6"),
    ("营业利润", 21438.3, "千元", "B10"),
    ("净利润", 16038.3, "千元", "B14"),
]

for i, (metric, value, unit, ref) in enumerate(data, 2):
    ws.cell(row=i, column=1, value=metric).font = Font(bold=True)
    c2 = ws.cell(row=i, column=2, value=value)
    c2.number_format = '#,##0.00'
    ws.cell(row=i, column=3, value=unit).alignment = Alignment(horizontal="center")
    ws.cell(row=i, column=4, value=ref).alignment = Alignment(horizontal="center")

# 列宽
ws.column_dimensions['A'].width = 16
ws.column_dimensions['B'].width = 16
ws.column_dimensions['C'].width = 10
ws.column_dimensions['D'].width = 16

# 边框
thin_border = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)
for row in ws.iter_rows(min_row=1, max_row=5, min_col=1, max_col=4):
    for cell in row:
        cell.border = thin_border

wb.save("/workspace/E2E_finance_interactive_20261025.xlsx")
print("Done")
