"""Independent read-back: source sums, scope, formulas, caches and receipt integrity."""
import hashlib
import json
import sys
from decimal import Decimal
from pathlib import Path
import openpyxl
from dsh_modules.spreadsheet_update.diff import _cells_equivalent

root=Path(sys.argv[1]).resolve()
source=root/'工数实绩明细_20261020134124.xlsx'
target=root/'工数3Q预实管理.xlsx'
output=root/'工数3Q预实管理_1023.xlsx'
wb=openpyxl.load_workbook(source,data_only=True)
sums={};records=0
for row in wb['工数实绩明细'].iter_rows(min_row=5,values_only=True):
    if str(row[1])=='202610' and row[12]=='有偿':
        sums[row[2]]=sums.get(row[2],Decimal(0))+Decimal(str(row[15]));records+=1
wb.close()
original=openpyxl.load_workbook(target)
result=openpyxl.load_workbook(output)
cached=openpyxl.load_workbook(output,data_only=True)
assert result.sheetnames==['1019','1023'],result.sheetnames
expected={'FC48005':'C5','FC46005':'C8','FC45031':'C10','FC44003':'C13',
          'FC44004':'C14','FC44006':'C15','FC34022':'C16'}
assert set(sums)==set(expected)
for code,coord in expected.items():
    assert abs(Decimal(str(result['1023'][coord].value))-sums[code])<Decimal('0.00000001'),(code,coord)
changes=[];formulas=0
for sheet in ('1019','1023'):
    old=original['1019'];new=result[sheet]
    assert set(old.merged_cells.ranges)==set(new.merged_cells.ranges)
    for row in old:
        for cell in row:
            value=new[cell.coordinate].value
            if cell.data_type=='f':
                formulas+=1
                assert new[cell.coordinate].data_type=='f'
                assert cached[sheet][cell.coordinate].value is not None,(sheet,cell.coordinate)
            if not _cells_equivalent(cell.value,value,original.epoch):
                assert sheet=='1023' and cell.coordinate in expected.values(),(sheet,cell.coordinate,cell.value,value)
                changes.append(cell.coordinate)
for coord in ('C6','C9','C11','C18'):
    assert result['1023'][coord].value==original['1019'][coord].value
receipts=[json.loads(p.read_text()) for p in (root/'.dsh/spreadsheet-runs').glob('*/receipt.json')]
assert len(receipts)==1
receipt=receipts[0]
output_hash=hashlib.sha256(output.read_bytes()).hexdigest()
assert output_hash==receipt['output_hash']
assert receipt['status']=='published'
assert receipt['formula_health']['total_errors']==0
assert receipt['formula_health']['uncalculated_formulas']==0
assert receipt['formula_health']['cached_formulas']==formulas
original.close();result.close();cached.close()
report={'status':'passed','input_hashes':{p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in (source,target)},
        'output_hash':output_hash,'paid_records':records,'department_count':len(sums),
        'source_total_K':str(sum(sums.values())),'department_values_K':{k:str(v) for k,v in sorted(sums.items())},
        'changed_cells':sorted(changes),'formula_count':formulas,'formula_errors':0,'missing_caches':0,
        'historical_sheet_preserved':True,'budget_other_months_and_unpaid_preserved':True,
        'unmatched_departments_preserved':4,'receipt_verdict':receipt['validation_verdict']}
(root/'independent-verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
print(json.dumps(report,ensure_ascii=False,indent=2))
