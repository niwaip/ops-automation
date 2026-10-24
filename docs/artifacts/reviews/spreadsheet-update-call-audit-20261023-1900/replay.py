"""Read-only replay of saved plan compilation; never execute or publish updates."""
import copy
import hashlib
import json
from collections import Counter, defaultdict
from decimal import Decimal
from pathlib import Path

import openpyxl

from dsh_modules.spreadsheet_update.bindings import LogicalKey
from dsh_modules.spreadsheet_update.compiler import compile_update_plan, UpdateCompilationError
from dsh_modules.spreadsheet_update.discovery import extract_workbook_ir
from dsh_modules.spreadsheet_update.schema import validate_update_plan
from dsh_modules.spreadsheet_update.sources import FactTable, ingest_facts_from_source_queries

ROOT = Path.cwd()
OUT = Path(__file__).resolve().parent
WS = ROOT / 'data/users/e7fce333-a8f4-4097-9a53-f0a4c729da46/workspace'
SID = 'chat-session-1792753181184-xf6a3kw1'
TID = '1792753256707802129-334251a9cb5d41e0bcc3ce931ae2aec6'
RUN = WS / '.dsh/spreadsheet-runs/run_1e043e19200a'
PLAN = json.loads((RUN / 'plan.json').read_text())
STATE = json.loads((RUN / 'state.json').read_text())
E = json.loads((WS / '.dsh/sessions' / (SID + '.evidence') / (TID + '.json')).read_text())
target = WS / '工数3Q预实管理.xlsx'
source = WS / '工数实绩明细_20261020134124.xlsx'
snapshot = RUN / 'snapshot_工数3Q预实管理.xlsx'
tracked = [target, source, snapshot, RUN / 'plan.json', RUN / 'state.json']
before = {str(p.relative_to(WS)): hashlib.sha256(p.read_bytes()).hexdigest() for p in tracked}

rids = {o['region_id'] for o in PLAN['operations'] if o.get('region_id')}
mids = {o['metric_id'] for o in PLAN['operations'] if o.get('metric_id')}
ir = extract_workbook_ir(snapshot, '1019', declared_region_ids=rids, declared_metric_ids=mids)
first = PLAN['operations'][1]
first_key = LogicalKey(first['region_id'], first['department_code'], first['metric_id'],
                       first['period'], first.get('value_kind', 'actual'), first.get('unit', 'hours'))
facts = FactTable()
required = [LogicalKey(o['region_id'], o['department_code'], o['metric_id'], o['period'],
                      o.get('value_kind', 'actual'), o.get('unit', 'hours'))
            for o in PLAN['operations'] if o['op'] == 'fill_from_result']
cardinality = facts.check_cardinality_against_bindings(required)

def compilation_error(plan):
    try:
        compile_update_plan(plan, ir, facts)
    except UpdateCompilationError as ex:
        return str(ex)
    raise AssertionError('Unexpected successful compilation')

exact_error = compilation_error(PLAN)
diagnostic = copy.deepcopy(PLAN)
diagnostic['operations'] = diagnostic['operations'][:2]
diagnostic['operations'][1]['value_kind'] = 'actual'
next_error = compilation_error(diagnostic)
# Path normalization only for host-side read; source_queries stays exactly empty.
source_descriptors = copy.deepcopy(PLAN['sources'])
for s in source_descriptors:
    s['file_path'] = str(WS / Path(s['file_path']).name)
ingested, ingest_issues = ingest_facts_from_source_queries(source_descriptors,
    PLAN.get('source_queries', []), PLAN.get('mappings', []), workspace_dir=WS,
    default_as_of=PLAN['as_of'])

wb = openpyxl.load_workbook(target, read_only=True, data_only=False)
sheet_names = wb.sheetnames; s = wb['1019']
dept_rows = defaultdict(list)
for row in s.iter_rows(min_row=5, max_row=18, max_col=3):
    if isinstance(row[1].value, str) and row[1].value.startswith('FC'):
        dept_rows[row[1].value].append(row[1].row)
coordinate_checks = []
for op in PLAN['operations']:
    if op['op'] != 'fill_from_result':
        continue
    row = s[op['coordinate']].row
    coordinate_checks.append({'op_id': op['op_id'], 'department_code': op['department_code'],
        'planned_coordinate': op['coordinate'], 'target_department_at_row': s[f'B{row}'].value,
        'expected_paid_amount_coordinates': [f'C{r}' for r in dept_rows[op['department_code']]],
        'is_formula': s[op['coordinate']].data_type == 'f',
        'department_row_matches': row in dept_rows[op['department_code']]})
wb.close()
wb = openpyxl.load_workbook(source, read_only=True, data_only=True)
s = wb['工数实绩明细']; s.reset_dimensions()
amounts = defaultdict(Decimal); counts = Counter(); months = Counter()
for row in s.iter_rows(min_row=5, max_row=278, max_col=18, values_only=True):
    months[str(row[1])] += 1
    if row[12] == '有偿':
        amounts[str(row[2])] += Decimal(str(row[15])); counts[str(row[2])] += 1
wb.close()
analysis = next(r['data'] for r in E['receipts'] if (r.get('data') or {}).get('kind') == 'spreadsheet_analysis')
receipt_amounts_yuan = {f['group_fields']['部门编号']: Decimal(f['value']) for f in analysis['facts']}
bindings_c5 = [{'key': b.logical_key.to_key_string(), 'coordinate': b.coordinate}
               for b in ir.registry.all_bindings()
               if b.logical_key.region_id == first['region_id']
               and b.logical_key.department_code == first['department_code']
               and b.logical_key.metric_id == first['metric_id'] and b.logical_key.period == '202610']
after = {str(p.relative_to(WS)): hashlib.sha256(p.read_bytes()).hexdigest() for p in tracked}
result = {
    'session_id': SID, 'turn_id': TID, 'run_id': STATE['run_id'],
    'schema_validation': validate_update_plan(PLAN),
    'exact_compilation_error': exact_error,
    'matches_real_error': ('Compilation failed: ' + exact_error) in E['receipts'][-1]['error']['message'],
    'requested_first_binding': first_key.to_key_string(), 'discovered_same_cell_bindings': bindings_c5,
    'diagnostic_only_actual_instead_of_forecast': {'error': next_error,
        'note': 'Change made only to an in-memory first-operation diagnostic. Not a fix or executed update.'},
    'fact_ingestion': {'source_queries_count': len(PLAN.get('source_queries', [])),
        'mappings_count': len(PLAN.get('mappings', [])), 'ingested_count': len(ingested.all_facts()),
        'ingest_issues': ingest_issues,
        'empty_fact_cardinality_is_valid': cardinality.is_valid,
        'matched_count': cardinality.matched_count, 'missing_required_count': len(cardinality.missing_keys)},
    'coordinate_checks': coordinate_checks,
    'unit_and_amount_check': {'source_unit': 'K', 'query_metric_unit': analysis['plan']['metrics'][0]['unit'],
        'operation_units': [op.get('unit', '<omitted -> hours>') for op in PLAN['operations'] if op['op']=='fill_from_result'],
        'paid_records': sum(counts.values()), 'months': dict(months),
        'sum_source_K': str(sum(amounts.values())),
        'sum_receipt_yuan': str(sum(receipt_amounts_yuan.values())),
        'receipt_matches_K_times_1000': all(receipt_amounts_yuan[k] == amounts[k]*1000 for k in amounts)},
    'persisted_state': STATE,
    'file_integrity': {'before': before, 'after': after, 'unchanged': before == after,
        'inputs_match_call': all(before[n] == h for n,h in E['source_hashes'].items()),
        'target_sheet_names': sheet_names, 'candidate_exists': (RUN/'candidate.xlsx').exists(),
        'planned_output_exists': (WS/PLAN['output']['display_name']).exists(),
        'notice_output_exists': (WS/'新工作表_1023.xlsx').exists()},
    'model_calls': 0, 'workbook_writes': 0,
}
OUT.mkdir(parents=True, exist_ok=True)
(OUT/'replay-result.json').write_text(json.dumps(result, ensure_ascii=False, indent=2)+'\n')
assert result['matches_real_error'] and result['schema_validation'][0]
assert before == after and result['file_integrity']['inputs_match_call']
assert not result['file_integrity']['candidate_exists'] and sheet_names == ['1019']
assert result['fact_ingestion']['ingested_count'] == 0
assert 'FactTable missing fact' in next_error
assert result['unit_and_amount_check']['receipt_matches_K_times_1000']
print(json.dumps({k:result[k] for k in ['exact_compilation_error', 'discovered_same_cell_bindings',
    'diagnostic_only_actual_instead_of_forecast', 'fact_ingestion','coordinate_checks',
    'unit_and_amount_check','file_integrity']},ensure_ascii=False,indent=2))
