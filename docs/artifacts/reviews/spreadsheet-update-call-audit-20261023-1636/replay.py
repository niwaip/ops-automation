"""Read-only source reconciliation and deterministic routing/gate reproduction.

Run from repository root with runner src on PYTHONPATH. No model calls, service
mutations, or workbook writes. The raw Redis export stays outside the repository.
"""
import copy
import hashlib
import json
import re
from collections import defaultdict
from dataclasses import replace
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch

import openpyxl

from dsh_modules.analysis_contract import build_analysis_contract
from dsh_modules.analysis_request_text import user_instruction
from dsh_modules.analysis_task_state import analysis_task_state
from dsh_modules.intent_contract import is_modify_intent_requested
from dsh_modules.runner import select_active_tools
from dsh_modules.runtime_policy import RuntimePolicy
from dsh_modules.skill_admission import general_route
from dsh_modules.spreadsheet_analysis_evidence import requires_spreadsheet_analysis
from dsh_modules.spreadsheet_analysis_schema import select_analysis_tools
from dsh_modules.task_intent import extract_task_intent, has_physical_artifact_request
from dsh_modules.telemetry import TelemetryStats
from dsh_modules.tool_result import adapt_tool_output

ROOT = Path.cwd()
OUT = Path(__file__).resolve().parent
WORKSPACE = ROOT / 'data/users/e7fce333-a8f4-4097-9a53-f0a4c729da46/workspace'
SESSION_ID = 'chat-session-1792742409266-8tff0l1d'
TURN_ID = '1792744582324015586-1ef76146bd64439690cad10e72d352dc'
EVIDENCE = WORKSPACE / '.dsh/sessions' / (SESSION_ID + '.evidence') / (TURN_ID + '.json')
evidence = json.loads(EVIDENCE.read_text())
prompt = user_instruction(evidence['query'])
accepted_contract = next(g['reason'] for g in evidence['guard_decisions']
                         if g['guard'] == 'task_intent' and g['action'] == 'accept')
contract = build_analysis_contract(prompt)

entity_pattern = r'(?:员工|人员|用户|同事|部门|客户|项目)\s*[:：=]?\s*([^\s，。！？、的]{2,15})'
entity_matches = [{'matched': m.group(0), 'captured_value': m.group(1)}
                  for m in re.finditer(entity_pattern, prompt)]

files = list(evidence['source_hashes'])
route = general_route(evidence['query'], [], files, accepted_contract)
before_tools = select_active_tools(evidence['query'], route, False, True)
requires_analysis = requires_spreadsheet_analysis(evidence['query'], route)
after_tools = select_analysis_tools(before_tools, evidence['query']) if requires_analysis else before_tools

# Match the observed review outcome (artifact) and isolate the subsequent
# deterministic override. The historical review's exact artifact_types are not
# stored in the guard; .xlsx is a controlled reproduction input, not a recovered
# historical model response.
proposed = {**copy.deepcopy(accepted_contract), 'output_mode': 'artifact', 'artifact_types': ['.xlsx']}
telemetry = TelemetryStats()
with patch('dsh_modules.task_intent.structured_decision', side_effect=[copy.deepcopy(proposed), copy.deepcopy(proposed)]):
    overridden = extract_task_intent(prompt, [], 'not-called', RuntimePolicy(), telemetry, None, input_count=2)

receipts = []
for raw in evidence['receipts']:
    item = copy.deepcopy(raw)
    provenance = item.get('provenance') or {}
    if provenance.get('file_path', '').startswith('/workspace/'):
        provenance['file_path'] = str(WORKSPACE / Path(provenance['file_path']).name)
    receipts.append(adapt_tool_output(item))
with patch('dsh_modules.task_completeness_judge.judge_task_completeness', return_value=None):
    task_state = analysis_task_state(receipts, prompt)
    corrected_group_contract = replace(contract, entity_predicates=[], is_grouped=True,
                                       group_dimension='部门编号')
    with patch('dsh_modules.analysis_task_state.build_analysis_contract', return_value=corrected_group_contract):
        calculation_only_state = analysis_task_state(receipts, prompt)

analysis = next(r for r in receipts if r.is_success and isinstance(r.data, dict)
                and r.data.get('kind') == 'spreadsheet_analysis')
source = WORKSPACE / '工数实绩明细_20261020134124.xlsx'
target = WORKSPACE / '工数3Q预实管理.xlsx'
hashes_before = {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in [source, target]}

wb = openpyxl.load_workbook(source, read_only=True, data_only=True)
ws = wb['工数实绩明细']
ws.reset_dimensions()
amounts = defaultdict(Decimal)
counts = defaultdict(int)
source_months = set()
for row in ws.iter_rows(min_row=5, max_row=278, min_col=1, max_col=18, values_only=True):
    source_months.add(str(row[1]))
    if row[12] == '有偿':
        amounts[str(row[2])] += Decimal(str(row[15]))
        counts[str(row[2])] += 1
wb.close()
receipt_amounts = {f['group_key']: Decimal(str(f['value'])) for f in analysis.data['facts']}
wb = openpyxl.load_workbook(target, read_only=True, data_only=False)
target_sheets = wb.sheetnames
wb.close()
hashes_after = {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in [source, target]}

result = {
    'session_id': SESSION_ID, 'turn_id': TURN_ID, 'request': prompt,
    'contract_parser': {'entities': contract.entity_predicates, 'is_grouped': contract.is_grouped,
                        'group_dimension': contract.group_dimension, 'metrics': contract.get_metric_names(),
                        'entity_regex_matches': entity_matches},
    'artifact_intent_reproduction': {
        'physical_artifact_regex_match': has_physical_artifact_request(prompt),
        'reviewed_mode_controlled_input': 'artifact', 'reviewed_types_controlled_input': ['.xlsx'],
        'after_deterministic_override': overridden,
    },
    'routing_reproduction': {
        'is_modify_intent_requested': is_modify_intent_requested(evidence['query']),
        'route_flags': {k: getattr(route, k, None) for k in
                        ['skill_id', 'is_modify_intent', 'is_generate_intent', 'requires_execution', 'is_inspect_intent']},
        'before_analysis_filter': [t['function']['name'] for t in before_tools],
        'requires_spreadsheet_analysis': requires_analysis,
        'after_analysis_filter': [t['function']['name'] for t in after_tools],
        'note': 'Reconstruction from persisted accepted contract and current source; not a historical tool-schema capture.'
    },
    'gate_reproduction': {'complete': task_state['complete'], 'issues': task_state['issues'],
                          'evidence_ids': task_state['evidence_ids'], 'model_calls': 0},
    'counterfactual_analysis_only_gate': {
        'complete': calculation_only_state['complete'], 'issues': calculation_only_state['issues'],
        'note': 'Corrected grouping contract only; no copy/write/file-delivery receipts were added. '
                'An analysis PASS is not completion of the requested update workflow.'
    },
    'reconciliation': {
        'receipt_evidence_id': analysis.data['evidence_id'],
        'input_records': analysis.data['coverage'][0]['input_records'],
        'paid_records': sum(counts.values()), 'source_months': sorted(source_months),
        'by_department': [{'department_code': key, 'paid_amount_source_K': str(amounts[key]),
                           'paid_records': counts[key], 'receipt_value': str(receipt_amounts.get(key))}
                          for key in sorted(amounts)],
        'independent_matches_receipt': dict(amounts) == receipt_amounts,
        'paid_total_source_K': str(sum(amounts.values(), Decimal('0'))),
        'target_sheets_after_call': target_sheets,
        'source_hashes_match_call': hashes_before == evidence['source_hashes'],
        'source_hashes_unchanged_by_audit': hashes_before == hashes_after,
    },
}
(OUT / 'replay-result.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')

# Save only the selected call's operational metadata. Omit thought/reasoning logs
# and workbook row samples containing employee/customer names.
raw_session = Path('/private/tmp/spreadsheet-call-1636-redis.json')
if raw_session.exists():
    session = json.loads(raw_session.read_text())
    message = next(m for m in session['history'] if m.get('role') == 'assistant'
                   and '特定实体【编码】' in m.get('content', ''))
    meta = message.get('metadata') or {}
    trace = meta.get('executionTrace') or {}
    metrics = trace.get('metrics') or {}
    summary = {
        'session_id': SESSION_ID, 'assistant_message_id': message.get('id'),
        'timestamp': message.get('timestamp'), 'mode': meta.get('mode'),
        'model_id': session.get('session', {}).get('modelId'),
        'final_content': message.get('content'), 'usage': meta.get('usage'),
        'execution_trace': {k: trace.get(k) for k in ['harness', 'executed', 'durationMs', 'exitCode']},
        'metrics': {k: metrics.get(k) for k in ['ttftMs', 'durationMs', 'llmDurationMs', 'toolCallsCount',
                                              'toolCalls', 'llmInvocations', 'tokens', 'finishReason',
                                              'roundFinishReasons', 'guardDecisions', 'modelResponses']},
    }
    (OUT / 'call-summary.json').write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding='utf-8')

print(json.dumps({
    'parser': result['contract_parser'], 'routing': result['routing_reproduction'],
    'artifact_mode_after_override': overridden['output_mode'],
    'gate_issue_codes': [i['code'] for i in task_state['issues']],
    'reconciliation': result['reconciliation'],
}, ensure_ascii=False, indent=2))
